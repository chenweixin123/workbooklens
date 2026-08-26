"""Isolated local recalculation used to validate formula-derived repairs."""

from __future__ import annotations

import base64
import contextlib
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from datetime import time as datetime_time
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.constants import MAX_COLUMN, MAX_ROW

from workbooklens import conversion
from workbooklens.exceptions import PatchValidationError, WorkbookLensError
from workbooklens.formulas.error_propagation import EXCEL_ERRORS
from workbooklens.ooxml.safety import PackageLimits, inspect_package
from workbooklens.utils import sha256_file

RecalculationPreference = Literal["auto", "excel", "libreoffice", "none"]

_DEFAULT_TIMEOUT_SECONDS = 180
_DEFAULT_ERROR_SCAN_MAX_CELLS = 1_000_000
_DEFAULT_ERROR_SCAN_TIMEOUT_SECONDS = 15.0
_RECALCULATED_EXCEL_ERRORS = EXCEL_ERRORS | frozenset(
    {
        "#BLOCKED!",
        "#BUSY!",
        "#CALC!",
        "#CONNECT!",
        "#FIELD!",
        "#GETTING_DATA",
        "#PYTHON!",
        "#SPILL!",
        "#UNKNOWN!",
    }
)
_CRITICAL_PART_PREFIXES = (
    "customXml/",
    "xl/charts/",
    "xl/connections.xml",
    "xl/drawings/",
    "xl/externalLinks/",
    "xl/model/",
    "xl/pivotCache/",
    "xl/pivotTables/",
    "xl/queryTables/",
    "xl/slicers/",
    "xl/tables/",
)
_R1C1_LIKE_SHEET_RE = re.compile(r"^R(?:[1-9][0-9]*)?C(?:[1-9][0-9]*)?$", re.IGNORECASE)
_SHA256_RE = re.compile(r"[0-9a-f]{64}")
_WINDOWS_LEGACY_PATH_LIMIT = 259

_EXCEL_RECALCULATION_SCRIPT = r"""
$ErrorActionPreference = 'Stop'
[Console]::OutputEncoding = [System.Text.UTF8Encoding]::new($false)
$inputPath = [Environment]::GetEnvironmentVariable('WORKBOOKLENS_RECALC_INPUT', 'Process')
$outputPath = [Environment]::GetEnvironmentVariable('WORKBOOKLENS_RECALC_OUTPUT', 'Process')
$excelIdentityPath = [Environment]::GetEnvironmentVariable(
    'WORKBOOKLENS_EXCEL_IDENTITY',
    'Process'
)
$excelGoPath = [Environment]::GetEnvironmentVariable('WORKBOOKLENS_EXCEL_GO', 'Process')
if (
    [string]::IsNullOrWhiteSpace($inputPath) -or
    [string]::IsNullOrWhiteSpace($outputPath) -or
    [string]::IsNullOrWhiteSpace($excelIdentityPath) -or
    [string]::IsNullOrWhiteSpace($excelGoPath)
) {
    throw 'WorkbookLens did not provide recalculation paths.'
}

$excelIdentityTempPath = "$excelIdentityPath.tmp"
$excel = $null
$calculationWorkbook = $null
$workbook = $null
$ownsExcelProcess = $false
$existingExcelPids = [System.Collections.Generic.HashSet[int]]::new()
foreach ($runningExcel in @(Get-Process -Name EXCEL -ErrorAction SilentlyContinue)) {
    [void]$existingExcelPids.Add([int]$runningExcel.Id)
}
Add-Type -TypeDefinition @'
using System;
using System.Runtime.InteropServices;

public static class WorkbookLensRecalculationNativeMethods {
    [DllImport("user32.dll", SetLastError = true)]
    public static extern uint GetWindowThreadProcessId(IntPtr hWnd, out uint processId);
}
'@
function Release-ComObject($value) {
    if ($null -ne $value -and [Runtime.InteropServices.Marshal]::IsComObject($value)) {
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($value)
    }
}
try {
    $excel = New-Object -ComObject Excel.Application
    [uint32]$excelPid = 0
    $windowThread = [WorkbookLensRecalculationNativeMethods]::GetWindowThreadProcessId(
        [IntPtr]([long]$excel.Hwnd),
        [ref]$excelPid
    )
    if ($windowThread -eq 0 -or $excelPid -eq 0) {
        throw 'WorkbookLens could not identify the Microsoft Excel process safely.'
    }
    if ($existingExcelPids.Contains([int]$excelPid)) {
        throw 'Microsoft Excel reused an existing user process; recalculation was refused.'
    }
    $ownsExcelProcess = $true
    $ownedExcelProcessInfo = Get-Process -Id ([int]$excelPid) -ErrorAction Stop
    try {
        $startTimeUtc = $ownedExcelProcessInfo.StartTime.ToUniversalTime()
        $normalizedExecutablePath = $null
        try {
            $candidateExecutablePath = [string]$ownedExcelProcessInfo.Path
            if (-not [string]::IsNullOrWhiteSpace($candidateExecutablePath)) {
                $candidateExecutablePath = [IO.Path]::GetFullPath(
                    $candidateExecutablePath
                ).ToLowerInvariant()
                if (
                    [IO.Path]::GetFileName($candidateExecutablePath).ToUpperInvariant() -eq
                        'EXCEL.EXE'
                ) {
                    $normalizedExecutablePath = $candidateExecutablePath
                }
            }
        } catch {
            # Python verifies the executable through the held Windows process handle.
        }
        $identityJson = ConvertTo-Json -InputObject ([ordered]@{
            process_id = [int]$excelPid
            creation_utc = $startTimeUtc.ToString('o')
            creation_filetime = [string]$startTimeUtc.ToFileTimeUtc()
            session_id = [int]$ownedExcelProcessInfo.SessionId
            normalized_executable_path = $normalizedExecutablePath
        }) -Compress
    } finally {
        $ownedExcelProcessInfo.Dispose()
    }
    [IO.File]::WriteAllText(
        $excelIdentityTempPath,
        $identityJson,
        [Text.UTF8Encoding]::new($false)
    )
    [IO.File]::Move($excelIdentityTempPath, $excelIdentityPath)
    $goDeadline = [DateTime]::UtcNow.AddSeconds(35)
    while (-not [IO.File]::Exists($excelGoPath)) {
        if ([DateTime]::UtcNow -ge $goDeadline) {
            throw 'WorkbookLens did not confirm ownership of the Microsoft Excel process.'
        }
        Start-Sleep -Milliseconds 25
    }
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $excel.AskToUpdateLinks = $false
    $excel.EnableEvents = $false
    $excel.AutomationSecurity = 3
    $calculationWorkbook = $excel.Workbooks.Add()
    $excel.Calculation = -4135
    $excel.CalculateBeforeSave = $false
    $calculationWorkbook.Close($false)
    Release-ComObject $calculationWorkbook
    $calculationWorkbook = $null
    $workbook = $excel.Workbooks.Open($inputPath, 0, $true)
    $workbook.CheckCompatibility = $false
    $workbook.ForceFullCalculation = $true
    $excel.CalculateFullRebuild()
    $workbook.SaveCopyAs($outputPath)
} finally {
    if ([IO.File]::Exists($excelIdentityTempPath)) {
        try { [IO.File]::Delete($excelIdentityTempPath) } catch {}
    }
    if ($null -ne $calculationWorkbook) {
        try { $calculationWorkbook.Close($false) } catch {}
        Release-ComObject $calculationWorkbook
    }
    if ($null -ne $workbook) {
        try { $workbook.Close($false) } catch {}
        Release-ComObject $workbook
    }
    if ($null -ne $excel) {
        if ($ownsExcelProcess) {
            try { $excel.Quit() } catch {}
        }
        Release-ComObject $excel
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""


@dataclass(frozen=True, slots=True)
class RecalculationValidation:
    """Formula-error comparison produced by one local recalculation provider."""

    provider: conversion.ConversionProvider
    formula_errors_before: tuple[str, ...]
    formula_errors_after: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class _SheetStructure:
    """Recalculation-sensitive worksheet structure that a provider must preserve."""

    title: str
    kind: str
    state: str
    formulas: tuple[tuple[str, str], ...]
    literal_cells: tuple[tuple[str, str, str, bool], ...]
    merged_ranges: tuple[str, ...]
    tables: tuple[tuple[str, str], ...]
    auto_filter: str | None
    data_validations: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class _WorkbookStructure:
    """Critical workbook structure excluding cached formula values and calculation metadata."""

    sheets: tuple[_SheetStructure, ...]
    defined_names: tuple[tuple[str, ...], ...]
    external_relationships: tuple[str, ...]
    critical_parts: tuple[str, ...]


class _StagedInputHashError(PatchValidationError):
    """Raised when an input changes while a private recalculation copy is staged or used."""


def candidate_providers(
    preference: RecalculationPreference,
) -> tuple[conversion.ConversionProvider, ...]:
    """Resolve installed providers without silently changing an explicit preference."""

    if preference == "none":
        return ()
    providers = conversion.available_conversion_providers()
    if preference == "auto":
        return providers
    return tuple(provider for provider in providers if provider.kind == preference)


def _validated_private_workspace(raw: str | Path, parent: Path, label: str) -> Path:
    """Resolve one tool-created workspace and require it to stay under its parent."""

    resolved_parent = parent.resolve()
    workspace = Path(raw).resolve()
    if workspace.parent != resolved_parent:
        raise PatchValidationError(f"{label} escaped its intended parent directory")
    return workspace


def _validate_excel_tool_path(path: Path, label: str) -> None:
    """Fail clearly before PowerShell/.NET encounters the legacy Windows path limit."""

    if sys.platform == "win32" and len(str(path)) > _WINDOWS_LEGACY_PATH_LIMIT:
        raise PatchValidationError(
            f"{label} exceeds the safe Windows path limit; choose a shorter output directory"
        )


def _normalize_expected_sha256(value: str | None, path: Path, label: str) -> str:
    if value is None:
        return sha256_file(path)
    normalized = value.strip().lower()
    if _SHA256_RE.fullmatch(normalized) is None:
        raise PatchValidationError(f"Expected {label} SHA-256 is invalid")
    return normalized


def _assert_expected_hash(path: Path, expected_sha256: str, label: str) -> None:
    observed = sha256_file(path)
    if observed != expected_sha256:
        raise _StagedInputHashError(
            f"{label} SHA-256 changed during recalculation validation: "
            f"expected {expected_sha256}, observed {observed}"
        )


def _stage_verified_copy(
    source: Path,
    staged: Path,
    *,
    expected_sha256: str,
    label: str,
) -> None:
    """Copy one input exactly once while detecting replacement before or during the copy."""

    _assert_expected_hash(source, expected_sha256, label)
    with source.open("rb") as source_handle, staged.open("xb") as staged_handle:
        shutil.copyfileobj(source_handle, staged_handle, length=1024 * 1024)
        staged_handle.flush()
        os.fsync(staged_handle.fileno())
    try:
        _assert_expected_hash(staged, expected_sha256, f"Staged {label.lower()}")
        _assert_expected_hash(source, expected_sha256, label)
    except Exception:
        with contextlib.suppress(OSError):
            staged.unlink()
        raise


def _check_scan_deadline(deadline: float) -> None:
    if time.monotonic() > deadline:
        raise PatchValidationError("Formula-error scan exceeded its time limit")


def collect_formula_errors(
    path: Path,
    *,
    max_cells: int = _DEFAULT_ERROR_SCAN_MAX_CELLS,
    timeout_seconds: float = _DEFAULT_ERROR_SCAN_TIMEOUT_SECONDS,
    limits: PackageLimits | None = None,
) -> tuple[str, ...]:
    """Return cached errors for formula cells using a sparse, resource-bounded scan."""

    if max_cells <= 0:
        raise ValueError("max_cells must be positive")
    if timeout_seconds <= 0:
        raise ValueError("timeout_seconds must be positive")
    deadline = time.monotonic() + timeout_seconds
    inspect_package(path, limits)
    formula_book = None
    cached_book = None
    errors: list[str] = []
    scanned_cells = 0
    try:
        _check_scan_deadline(deadline)
        formula_book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        _check_scan_deadline(deadline)
        cached_book = load_workbook(path, read_only=False, data_only=True, keep_links=False)
        _check_scan_deadline(deadline)
        for formula_sheet in formula_book.worksheets:
            _check_scan_deadline(deadline)
            cached_sheet = cached_book[formula_sheet.title]
            sparse_cells = tuple(
                cell for cell in formula_sheet._cells.values() if isinstance(cell, Cell)
            )
            scanned_cells += len(sparse_cells)
            if scanned_cells > max_cells:
                raise PatchValidationError(
                    f"Formula-error scan exceeded its {max_cells}-cell limit"
                )
            for formula_cell in sparse_cells:
                _check_scan_deadline(deadline)
                if formula_cell.data_type != "f":
                    continue
                cached_cell = cached_sheet._cells.get((formula_cell.row, formula_cell.column))
                if not isinstance(cached_cell, Cell):
                    continue
                error = str(cached_cell.value).upper()
                if error in _RECALCULATED_EXCEL_ERRORS:
                    errors.append(f"{formula_sheet.title}!{formula_cell.coordinate}:{error}")
    finally:
        if formula_book is not None:
            formula_book.close()
        if cached_book is not None:
            cached_book.close()
    return tuple(sorted(errors))


def _canonicalize_optional_sheet_quotes(value: str) -> str:
    """Normalize only quotes Excel may remove from an ordinary sheet reference."""

    sheet_token, separator, reference = value.rpartition("!")
    if not separator or not reference:
        return value
    if sheet_token.startswith("'") and sheet_token.endswith("'"):
        sheet_name = sheet_token[1:-1].replace("''", "'")
        if "'" + sheet_name.replace("'", "''") + "'" != sheet_token:
            return value
    else:
        sheet_name = sheet_token
    if not sheet_name.isidentifier():
        return value
    try:
        sheet_row, sheet_column = coordinate_to_tuple(sheet_name)
    except (KeyError, ValueError):
        pass
    else:
        if 1 <= sheet_row <= MAX_ROW and 1 <= sheet_column <= MAX_COLUMN:
            return value
    if _R1C1_LIKE_SHEET_RE.fullmatch(sheet_name):
        return value
    components = reference.split(":")
    if not 1 <= len(components) <= 2:
        return value
    for component in components:
        try:
            row, column = coordinate_to_tuple(component.replace("$", ""))
        except (KeyError, ValueError):
            return value
        if not 1 <= row <= MAX_ROW or not 1 <= column <= MAX_COLUMN:
            return value
    return f"{sheet_name}!{reference}"


def _canonical_formula_text(formula: str) -> str:
    """Tokenize a formula while preserving everything except optional sheet quotes."""

    try:
        tokens = Tokenizer(formula).items
    except (IndexError, TokenizerError, ValueError):
        return formula
    signature: list[tuple[str, str, str]] = []
    for token in tokens:
        value = str(token.value)
        if token.type == "OPERAND" and token.subtype == "RANGE":
            value = _canonicalize_optional_sheet_quotes(value)
        signature.append((str(token.type), str(token.subtype), value))
    return repr(tuple(signature))


def _formula_signature(cell: Cell) -> str | None:
    if cell.data_type != "f":
        return None
    if isinstance(cell.value, str):
        return _canonical_formula_text(cell.value)
    if isinstance(cell.value, ArrayFormula):
        return f"array|{cell.value.ref}|{cell.value.text}"
    if isinstance(cell.value, DataTableFormula):
        fields = ("ref", "ca", "dt2D", "dtr", "r1", "r2", "del1", "del2")
        return "dataTable|" + "|".join(str(getattr(cell.value, field)) for field in fields)
    return f"unsupported|{type(cell.value).__qualname__}|{cell.value!r}"


def _defined_name_signatures(workbook: Any) -> tuple[tuple[str, ...], ...]:
    attributes = (
        "name",
        "comment",
        "customMenu",
        "description",
        "help",
        "statusBar",
        "localSheetId",
        "hidden",
        "function",
        "vbProcedure",
        "xlm",
        "functionGroupId",
        "shortcutKey",
        "publishToServer",
        "workbookParameter",
        "attr_text",
    )
    records: list[tuple[str, ...]] = []
    for defined_name in workbook.defined_names.values():
        records.append(("workbook", *(str(getattr(defined_name, key)) for key in attributes)))
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        for defined_name in worksheet.defined_names.values():
            records.append(
                (
                    f"sheet:{sheet_index}",
                    *(str(getattr(defined_name, key)) for key in attributes),
                )
            )
    return tuple(sorted(records))


def _data_validation_signatures(worksheet: Worksheet) -> tuple[str, ...]:
    if worksheet.data_validations is None:
        return ()
    return tuple(
        sorted(
            "|".join(
                (
                    str(validation.sqref),
                    validation.type or "",
                    validation.operator or "",
                    validation.formula1 or "",
                    validation.formula2 or "",
                )
            )
            for validation in worksheet.data_validations.dataValidation
        )
    )


def _canonical_literal_value(value: Any, *, location: str) -> str:
    """Return a stable, type-preserving representation for one non-formula cell."""

    if value is None:
        return "none"
    if isinstance(value, bool):
        return "bool:true" if value else "bool:false"
    if isinstance(value, int):
        return f"int:{value}"
    if isinstance(value, float):
        if not math.isfinite(value):
            raise PatchValidationError(
                f"Workbook structure contains a non-finite numeric value at {location}"
            )
        return f"float:{value.hex()}"
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise PatchValidationError(
                f"Workbook structure contains a non-finite decimal value at {location}"
            )
        return f"decimal:{value.normalize()}"
    if isinstance(value, datetime):
        return f"datetime:{value.isoformat(timespec='microseconds')}"
    if isinstance(value, date):
        return f"date:{value.isoformat()}"
    if isinstance(value, datetime_time):
        return f"time:{value.isoformat(timespec='microseconds')}"
    if isinstance(value, timedelta):
        return f"timedelta:{value.days}:{value.seconds}:{value.microseconds}"
    if isinstance(value, str):
        return f"str:{value}"
    raise PatchValidationError(
        f"Workbook structure contains an unsupported literal value at {location}: "
        f"{type(value).__qualname__}"
    )


def _workbook_structure(
    path: Path,
    limits: PackageLimits | None,
    *,
    max_cells: int = _DEFAULT_ERROR_SCAN_MAX_CELLS,
) -> _WorkbookStructure:
    inspection = inspect_package(path, limits)
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        scanned_cells = 0
        sheets: list[_SheetStructure] = []
        for sheet in getattr(workbook, "_sheets", workbook.worksheets):
            if isinstance(sheet, Worksheet):
                cells = tuple(cell for cell in sheet._cells.values() if isinstance(cell, Cell))
                scanned_cells += len(cells)
                if scanned_cells > max_cells:
                    raise PatchValidationError(
                        f"Workbook structure scan exceeded its {max_cells}-cell limit"
                    )
                formulas = tuple(
                    sorted(
                        (cell.coordinate, signature)
                        for cell in cells
                        if (signature := _formula_signature(cell)) is not None
                    )
                )
                literal_cells = tuple(
                    sorted(
                        (
                            cell.coordinate,
                            str(cell.data_type),
                            _canonical_literal_value(
                                cell.value,
                                location=f"{sheet.title}!{cell.coordinate}",
                            ),
                            bool(cell.quotePrefix),
                        )
                        for cell in cells
                        if cell.data_type != "f"
                    )
                )
                tables = tuple(
                    sorted(
                        (str(table.displayName), str(table.ref)) for table in sheet.tables.values()
                    )
                )
                sheets.append(
                    _SheetStructure(
                        title=sheet.title,
                        kind="worksheet",
                        state=sheet.sheet_state,
                        formulas=formulas,
                        literal_cells=literal_cells,
                        merged_ranges=tuple(
                            sorted(str(item) for item in sheet.merged_cells.ranges)
                        ),
                        tables=tables,
                        auto_filter=(str(sheet.auto_filter.ref) if sheet.auto_filter.ref else None),
                        data_validations=_data_validation_signatures(sheet),
                    )
                )
            else:
                sheets.append(
                    _SheetStructure(
                        title=sheet.title,
                        kind=type(sheet).__qualname__,
                        state=sheet.sheet_state,
                        formulas=(),
                        literal_cells=(),
                        merged_ranges=(),
                        tables=(),
                        auto_filter=None,
                        data_validations=(),
                    )
                )
        critical_parts = tuple(
            sorted(
                part
                for part in inspection.part_names
                if any(part.startswith(prefix) for prefix in _CRITICAL_PART_PREFIXES)
            )
        )
        return _WorkbookStructure(
            sheets=tuple(sheets),
            defined_names=_defined_name_signatures(workbook),
            external_relationships=tuple(sorted(inspection.external_relationships)),
            critical_parts=critical_parts,
        )
    finally:
        workbook.close()


def _validate_recalculated_structure(
    source: Path,
    recalculated: Path,
    limits: PackageLimits | None,
) -> None:
    if _workbook_structure(source, limits) != _workbook_structure(recalculated, limits):
        raise PatchValidationError(
            "The local recalculation provider unexpectedly changed formulas or workbook structure"
        )


def _validate_recalculated_package(
    path: Path,
    limits: PackageLimits | None,
    *,
    expected_input: Path | None = None,
) -> None:
    inspection = inspect_package(path, limits)
    if not inspection.repairable:
        raise PatchValidationError(
            "The local recalculation provider did not produce a verified macro-free .xlsx package"
        )
    if expected_input is not None:
        _validate_recalculated_structure(expected_input, path, limits)


def _best_effort_cleanup(*actions: Any) -> None:
    """Run every cleanup action without allowing one failure to suppress the rest."""

    for action in actions:
        with contextlib.suppress(Exception):
            action()


def _cleanup_excel_startup(
    process: subprocess.Popen[str],
    excel_identity: Path,
    baseline_process_ids: frozenset[int],
) -> None:
    def terminate_late_excel() -> None:
        late_identity = conversion._load_excel_process_identity(excel_identity)
        if late_identity is not None and late_identity.process_id not in baseline_process_ids:
            conversion._terminate_excel_identity_if_present(late_identity)

    _best_effort_cleanup(
        terminate_late_excel,
        lambda: conversion._terminate_process_tree(process),
    )


def _cleanup_owned_excel(
    process: subprocess.Popen[str],
    owned_excel: Any,
) -> None:
    _best_effort_cleanup(
        lambda: conversion._terminate_excel_process(owned_excel),
        lambda: conversion._terminate_process_tree(process),
    )


def _terminate_owned_excel_best_effort(owned_excel: Any) -> None:
    _best_effort_cleanup(lambda: conversion._terminate_excel_process(owned_excel))


def _run_excel_recalculation(
    provider: conversion.ConversionProvider,
    source: Path,
    output: Path,
    timeout_seconds: int,
) -> None:
    if output.exists():
        raise PatchValidationError(f"Recalculation output already exists: {output}")
    encoded_script = base64.b64encode(_EXCEL_RECALCULATION_SCRIPT.encode("utf-16le")).decode(
        "ascii"
    )
    with tempfile.TemporaryDirectory(prefix=".wl-x-", dir=output.parent) as raw:
        workspace = _validated_private_workspace(raw, output.parent, "Excel workspace")
        staged_output = workspace / "o.xlsx"
        excel_identity = workspace / "i.json"
        excel_go = workspace / "g"
        for path, label in (
            (source, "Excel recalculation input path"),
            (staged_output, "Excel recalculation output path"),
            (Path(f"{excel_identity}.tmp"), "Excel process identity path"),
            (excel_go, "Excel ownership signal path"),
        ):
            _validate_excel_tool_path(path, label)
        environment = os.environ.copy()
        environment["WORKBOOKLENS_RECALC_INPUT"] = str(source)
        environment["WORKBOOKLENS_RECALC_OUTPUT"] = str(staged_output)
        environment["WORKBOOKLENS_EXCEL_IDENTITY"] = str(excel_identity)
        environment["WORKBOOKLENS_EXCEL_GO"] = str(excel_go)
        command = [
            str(provider.runner),
            "-NoLogo",
            "-NoProfile",
            "-NonInteractive",
            "-EncodedCommand",
            encoded_script,
        ]
        baseline_process_ids = conversion._snapshot_excel_process_ids(provider.runner)
        if baseline_process_ids is None:
            raise PatchValidationError(
                "Could not establish a safe Microsoft Excel process baseline for recalculation"
            )
        try:
            process = subprocess.Popen(  # noqa: S603
                command,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=environment,
                creationflags=conversion._creation_flags(),
            )
        except OSError as exc:
            raise PatchValidationError(
                f"Could not start PowerShell for recalculation: {exc}"
            ) from exc
        try:
            expected_identity = conversion._wait_for_owned_excel_process_identity(
                process,
                excel_identity,
                command,
            )
        except conversion._ExcelStartupTimeout as exc:
            _cleanup_excel_startup(process, excel_identity, baseline_process_ids)
            raise PatchValidationError("Excel recalculation startup handshake timed out") from exc
        except conversion._ProviderFailure as exc:
            _cleanup_excel_startup(process, excel_identity, baseline_process_ids)
            raise PatchValidationError(str(exc)) from exc
        if expected_identity.process_id in baseline_process_ids:
            _best_effort_cleanup(lambda: conversion._terminate_process_tree(process))
            raise PatchValidationError(
                "Microsoft Excel reused an existing user process; recalculation was refused"
            )
        try:
            owned_excel = conversion._open_verified_excel_process(
                expected_identity,
                missing_is_clean=False,
            )
        except conversion._ProviderFailure as exc:
            _cleanup_excel_startup(process, excel_identity, baseline_process_ids)
            raise PatchValidationError(str(exc)) from exc
        if owned_excel is None:
            _cleanup_excel_startup(process, excel_identity, baseline_process_ids)
            raise PatchValidationError("Excel exited before recalculation could be verified")
        try:
            try:
                conversion._signal_excel_go(excel_go)
                stdout, stderr = process.communicate(timeout=timeout_seconds)
            except subprocess.TimeoutExpired as exc:
                _cleanup_owned_excel(process, owned_excel)
                raise PatchValidationError(
                    f"Excel recalculation timed out after {timeout_seconds} seconds"
                ) from exc
            except (OSError, conversion._ProviderFailure) as exc:
                _cleanup_owned_excel(process, owned_excel)
                raise PatchValidationError(f"Could not monitor Excel recalculation: {exc}") from exc
            completed = subprocess.CompletedProcess(
                command,
                process.returncode if process.returncode is not None else -1,
                stdout or "",
                stderr or "",
            )
            if completed.returncode != 0:
                _terminate_owned_excel_best_effort(owned_excel)
                raise PatchValidationError(conversion._process_detail(completed))
            try:
                if not conversion._wait_for_excel_process_exit(
                    owned_excel,
                    conversion._EXCEL_GRACEFUL_EXIT_TIMEOUT_SECONDS,
                ):
                    _terminate_owned_excel_best_effort(owned_excel)
            except conversion._ProviderFailure as exc:
                _terminate_owned_excel_best_effort(owned_excel)
                raise PatchValidationError(str(exc)) from exc
            if not staged_output.is_file():
                raise PatchValidationError("Excel did not create a recalculated workbook")
            if output.exists():
                raise PatchValidationError(
                    "Recalculation output appeared while Excel was running; refusing to overwrite it"
                )
            staged_output.rename(output)
        finally:
            _best_effort_cleanup(lambda: conversion._close_excel_process(owned_excel))


def _run_libreoffice_recalculation(
    provider: conversion.ConversionProvider,
    source: Path,
    output: Path,
    timeout_seconds: int,
) -> None:
    if output.exists():
        raise PatchValidationError(f"Recalculation output already exists: {output}")
    with tempfile.TemporaryDirectory(
        prefix=".wl-l-",
        dir=output.parent,
    ) as raw:
        workspace = _validated_private_workspace(raw, output.parent, "LibreOffice workspace")
        staged_input = workspace / "source.xlsx"
        profile = workspace / "profile"
        output_dir = workspace / "output"
        profile.mkdir()
        output_dir.mkdir()
        shutil.copyfile(source, staged_input)
        command = [
            str(provider.runner),
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            str(output_dir),
            str(staged_input),
        ]
        popen_options: dict[str, Any] = {
            "stdout": subprocess.PIPE,
            "stderr": subprocess.PIPE,
            "text": True,
            "encoding": "utf-8",
            "errors": "replace",
        }
        if sys.platform == "win32":
            popen_options["creationflags"] = conversion._creation_flags() | int(
                getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
            )
        else:
            popen_options["start_new_session"] = True
        try:
            process = subprocess.Popen(command, **popen_options)  # noqa: S603
        except OSError as exc:
            raise PatchValidationError(f"Could not start LibreOffice: {exc}") from exc
        try:
            stdout, stderr = process.communicate(timeout=timeout_seconds)
        except subprocess.TimeoutExpired as exc:
            _best_effort_cleanup(lambda: conversion._terminate_process_tree(process))
            raise PatchValidationError(
                f"LibreOffice recalculation timed out after {timeout_seconds} seconds"
            ) from exc
        except OSError as exc:
            _best_effort_cleanup(lambda: conversion._terminate_process_tree(process))
            raise PatchValidationError(
                f"Could not monitor LibreOffice recalculation: {exc}"
            ) from exc
        completed = subprocess.CompletedProcess(
            command,
            process.returncode if process.returncode is not None else -1,
            stdout,
            stderr,
        )
        generated = output_dir / staged_input.name
        if completed.returncode != 0:
            raise PatchValidationError(conversion._process_detail(completed))
        if not generated.is_file():
            raise PatchValidationError(
                "LibreOffice reported success but did not create a recalculated workbook"
            )
        if output.exists():
            raise PatchValidationError(
                "Recalculation output appeared while LibreOffice was running; refusing to overwrite it"
            )
        generated.rename(output)


def _run_provider(
    provider: conversion.ConversionProvider,
    source: Path,
    output: Path,
    timeout_seconds: int,
) -> None:
    if provider.kind == "excel":
        _run_excel_recalculation(provider, source, output, timeout_seconds)
    else:
        _run_libreoffice_recalculation(provider, source, output, timeout_seconds)


def _provider_failure(
    message: str,
    provider: conversion.ConversionProvider,
) -> PatchValidationError:
    error = PatchValidationError(message)
    error.__dict__["recalculation_provider"] = provider.kind
    return error


def validate_formula_recalculation(
    source: Path,
    candidate: Path,
    *,
    preference: RecalculationPreference = "auto",
    timeout_seconds: int = _DEFAULT_TIMEOUT_SECONDS,
    limits: PackageLimits | None = None,
    expected_source_sha256: str | None = None,
    expected_candidate_sha256: str | None = None,
) -> RecalculationValidation:
    """Recalculate immutable staged inputs with one provider and reject error regressions."""

    if timeout_seconds <= 0:
        raise ValueError("timeout_seconds must be positive")
    source = source.resolve()
    candidate = candidate.resolve()
    source_sha256 = _normalize_expected_sha256(
        expected_source_sha256,
        source,
        "source workbook",
    )
    candidate_sha256 = _normalize_expected_sha256(
        expected_candidate_sha256,
        candidate,
        "repair candidate",
    )
    providers = candidate_providers(preference)
    if not providers:
        raise PatchValidationError("No requested local recalculation provider is available")
    failures: list[str] = []
    last_provider: conversion.ConversionProvider | None = None
    with tempfile.TemporaryDirectory(
        prefix=".wl-v-",
        dir=candidate.parent,
    ) as raw:
        workspace = _validated_private_workspace(
            raw,
            candidate.parent,
            "Recalculation validation workspace",
        )
        staged_source = workspace / "s.xlsx"
        staged_candidate = workspace / "c.xlsx"
        _stage_verified_copy(
            source,
            staged_source,
            expected_sha256=source_sha256,
            label="Source workbook",
        )
        _stage_verified_copy(
            candidate,
            staged_candidate,
            expected_sha256=candidate_sha256,
            label="Repair candidate",
        )
        _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
        _assert_expected_hash(staged_candidate, candidate_sha256, "Staged repair candidate")
        for index, provider in enumerate(providers):
            last_provider = provider
            provider_workspace = workspace / f"p{index}"
            provider_workspace.mkdir()
            recalculated_source = provider_workspace / "s.xlsx"
            recalculated_candidate = provider_workspace / "c.xlsx"
            try:
                _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
                _assert_expected_hash(
                    staged_candidate,
                    candidate_sha256,
                    "Staged repair candidate",
                )
                _run_provider(provider, staged_source, recalculated_source, timeout_seconds)
                _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
                _assert_expected_hash(
                    staged_candidate,
                    candidate_sha256,
                    "Staged repair candidate",
                )
                _validate_recalculated_package(
                    recalculated_source,
                    limits,
                    expected_input=staged_source,
                )
                _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
                _assert_expected_hash(
                    staged_candidate,
                    candidate_sha256,
                    "Staged repair candidate",
                )
                before = collect_formula_errors(recalculated_source, limits=limits)
                _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
                _assert_expected_hash(
                    staged_candidate,
                    candidate_sha256,
                    "Staged repair candidate",
                )
            except _StagedInputHashError:
                raise
            except (PatchValidationError, WorkbookLensError, OSError, ValueError) as exc:
                failures.append(f"{provider.label}: {exc}")
                continue
            try:
                _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
                _assert_expected_hash(
                    staged_candidate,
                    candidate_sha256,
                    "Staged repair candidate",
                )
                _run_provider(provider, staged_candidate, recalculated_candidate, timeout_seconds)
                _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
                _assert_expected_hash(
                    staged_candidate,
                    candidate_sha256,
                    "Staged repair candidate",
                )
                _validate_recalculated_package(
                    recalculated_candidate,
                    limits,
                    expected_input=staged_candidate,
                )
                _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
                _assert_expected_hash(
                    staged_candidate,
                    candidate_sha256,
                    "Staged repair candidate",
                )
                after = collect_formula_errors(recalculated_candidate, limits=limits)
                _assert_expected_hash(staged_source, source_sha256, "Staged source workbook")
                _assert_expected_hash(
                    staged_candidate,
                    candidate_sha256,
                    "Staged repair candidate",
                )
            except _StagedInputHashError:
                raise
            except (PatchValidationError, WorkbookLensError, OSError, ValueError) as exc:
                raise _provider_failure(
                    f"{provider.label} validated the source workbook but failed on the repair "
                    f"candidate; refusing provider fallback: {exc}",
                    provider,
                ) from exc
            new_errors = sorted(set(after) - set(before))
            if new_errors:
                detail = ", ".join(new_errors)
                raise _provider_failure(
                    f"Formula recalculation introduced new errors: {detail}",
                    provider,
                )
            return RecalculationValidation(provider, before, after)
    detail = "; ".join(failures)
    if last_provider is None:
        raise PatchValidationError("No requested local recalculation provider is available")
    raise _provider_failure(
        f"Every requested recalculation provider failed. {detail}",
        last_provider,
    )


__all__ = [
    "RecalculationPreference",
    "RecalculationValidation",
    "candidate_providers",
    "collect_formula_errors",
    "validate_formula_recalculation",
]
