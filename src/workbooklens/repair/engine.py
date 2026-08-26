"""End-to-end patch application with rescanning and fail-closed cleanup."""

from __future__ import annotations

import contextlib
import json
import os
import re
import secrets
import shutil
import tempfile
import unicodedata
from collections import deque
from collections.abc import Collection
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from workbooklens.exceptions import PatchValidationError, StalePlanError, UsageError
from workbooklens.formulas.error_propagation import _propagation_candidate
from workbooklens.formulas.ir import FormulaIR, FormulaReference, parse_formula_ir
from workbooklens.models import (
    SEVERITY_RANK,
    Finding,
    PatchOperation,
    PatchPlan,
    PatchResult,
    PatchRisk,
    RecalculationProvider,
    Severity,
    ValidationStatus,
    WorkbookSnapshot,
)
from workbooklens.ooxml.safety import PackageLimits
from workbooklens.repair.ooxml_patch import _validate_canonical_plan, patch_ooxml_package
from workbooklens.repair.planning import build_patch_plan, resolve_patch_selection
from workbooklens.repair.recalculation import (
    RecalculationPreference,
    candidate_providers,
    validate_formula_recalculation,
)
from workbooklens.scanner import scan_workbook
from workbooklens.utils import sha256_file

_MAX_DEPENDENCY_RANGE_CHECKS = 1_000_000
_OPAQUE_DEPENDENCY_FUNCTIONS = frozenset({"INDIRECT", "OFFSET"})
_STRICT_A1_CELL = re.compile(r"[A-Za-z]{1,3}[1-9][0-9]*")

_FindingFamily = tuple[str, str | None, str, str]
_FindingIdentity = tuple[str, ...]
_FormulaNode = tuple[str, str]
_RangeDependent = tuple[int, int, int, int, _FormulaNode]


@dataclass(slots=True)
class _FindingChanges:
    before_keys: dict[str, frozenset[_FindingIdentity]]
    after_keys: dict[str, frozenset[_FindingIdentity]]
    after_key_set: frozenset[_FindingIdentity]
    resolved_ids: set[str]
    remaining_ids: set[str]
    new_ids: set[str]
    dangerous_findings: list[Finding]


@dataclass(frozen=True, slots=True)
class _FormulaRepairGroup:
    """One selected formula repair unit and its proven downstream closure."""

    label: str
    member_ids: frozenset[str]
    formula_ids: frozenset[str]
    targets: frozenset[_FormulaNode]


@dataclass(slots=True)
class _FormulaGroupValidation:
    """Per-round formula validation decisions and retained causal proofs."""

    failures: dict[str, tuple[_FormulaRepairGroup, tuple[str, ...]]]
    retained_messages: list[str]


@dataclass(frozen=True, slots=True)
class _FormulaDependencyGraph:
    """Bounded ordinary-A1 forward dependency graph for one workbook snapshot."""

    exact_sheet_names: frozenset[str]
    compatible_sheet_names: dict[str, tuple[str, ...]]
    direct_dependents: dict[_FormulaNode, frozenset[_FormulaNode]]
    range_dependents: dict[str, tuple[_RangeDependent, ...]]
    transparent_direct_dependents: dict[_FormulaNode, frozenset[_FormulaNode]]
    transparent_range_dependents: dict[str, tuple[_RangeDependent, ...]]

    def canonical_node(self, node: _FormulaNode) -> _FormulaNode:
        sheet, coordinate = node
        resolved_sheet, reason = _resolve_worksheet_name(
            sheet,
            self.exact_sheet_names,
            self.compatible_sheet_names,
        )
        if resolved_sheet is None:
            raise _OpaqueFormulaDependencyError(
                "Formula dependency validation encountered an unresolved worksheet: " + str(reason)
            )
        return resolved_sheet, coordinate.upper()

    def closure(
        self,
        starts: Collection[_FormulaNode],
        *,
        blocked_internal_nodes: Collection[_FormulaNode] = (),
        transparent_only: bool = False,
    ) -> set[_FormulaNode]:
        """Reach forward without traversing through modified internal nodes."""

        closure = {self.canonical_node(node) for node in starts}
        blocked = {self.canonical_node(node) for node in blocked_internal_nodes}
        direct_dependents = (
            self.transparent_direct_dependents if transparent_only else self.direct_dependents
        )
        range_dependents = (
            self.transparent_range_dependents if transparent_only else self.range_dependents
        )
        queue = deque(sorted(closure))
        range_checks = 0
        while queue:
            node = queue.popleft()
            dependents = set(direct_dependents.get(node, ()))
            row, column = coordinate_to_tuple(node[1])
            ranges = range_dependents.get(node[0], ())
            range_checks += len(ranges)
            if range_checks > _MAX_DEPENDENCY_RANGE_CHECKS:
                raise _OpaqueFormulaDependencyError(
                    "Formula dependency validation exceeded its bounded range-check budget"
                )
            dependents.update(
                dependent
                for min_row, max_row, min_column, max_column, dependent in ranges
                if min_row <= row <= max_row and min_column <= column <= max_column
            )
            for dependent in sorted(dependents):
                if dependent in closure:
                    continue
                closure.add(dependent)
                if dependent not in blocked:
                    queue.append(dependent)
        return closure


def _finding_family(finding: Finding) -> _FindingFamily:
    return (finding.rule_id, finding.sheet, finding.title, finding.expected)


def _strict_a1_cells(location: str | None) -> tuple[str, ...] | None:
    if location is None:
        return None
    parts = tuple(part.strip().upper() for part in location.split(","))
    if not parts or any(_STRICT_A1_CELL.fullmatch(part) is None for part in parts):
        return None
    cells: list[str] = []
    for part in parts:
        row, column = coordinate_to_tuple(part)
        if row > 1_048_576 or column > 16_384:
            return None
        cell = f"{get_column_letter(column)}{row}"
        if cell not in cells:
            cells.append(cell)
    return tuple(cells)


def _finding_identity_keys(
    finding: Finding,
    grouped_families: Collection[_FindingFamily],
) -> frozenset[_FindingIdentity]:
    family = _finding_family(finding)
    cells = _strict_a1_cells(finding.location)
    if family in grouped_families and cells is not None:
        rule_id, sheet, title, expected = family
        return frozenset(("cell", rule_id, sheet or "", title, expected, cell) for cell in cells)
    return frozenset({("id", finding.id)})


def _compare_findings(
    before_findings: Collection[Finding],
    after_findings: Collection[Finding],
) -> _FindingChanges:
    all_findings = [*before_findings, *after_findings]
    grouped_families = {
        _finding_family(finding)
        for finding in all_findings
        if (cells := _strict_a1_cells(finding.location)) is not None and len(cells) > 1
    }
    before_keys = {
        finding.id: _finding_identity_keys(finding, grouped_families) for finding in before_findings
    }
    after_keys = {
        finding.id: _finding_identity_keys(finding, grouped_families) for finding in after_findings
    }
    before_key_set = {key for finding_keys in before_keys.values() for key in finding_keys}
    after_key_set = frozenset(key for finding_keys in after_keys.values() for key in finding_keys)
    before_by_key: dict[_FindingIdentity, list[Finding]] = {}
    for finding in before_findings:
        for key in before_keys[finding.id]:
            before_by_key.setdefault(key, []).append(finding)

    resolved_ids = {
        finding.id
        for finding in before_findings
        if before_keys[finding.id].isdisjoint(after_key_set)
    }
    new_ids: set[str] = set()
    remaining_ids: set[str] = set()
    dangerous_findings: list[Finding] = []
    for finding in after_findings:
        keys = after_keys[finding.id]
        if keys - before_key_set:
            new_ids.add(finding.id)
        else:
            remaining_ids.add(finding.id)
        if finding.severity not in {Severity.ERROR, Severity.CRITICAL}:
            continue
        after_rank = SEVERITY_RANK[finding.severity]
        for key in keys:
            predecessors = before_by_key.get(key, [])
            if not predecessors or after_rank > max(
                SEVERITY_RANK[predecessor.severity] for predecessor in predecessors
            ):
                dangerous_findings.append(finding)
                break

    return _FindingChanges(
        before_keys=before_keys,
        after_keys=after_keys,
        after_key_set=after_key_set,
        resolved_ids=resolved_ids,
        remaining_ids=remaining_ids,
        new_ids=new_ids,
        dangerous_findings=dangerous_findings,
    )


class _OpaqueFormulaDependencyError(PatchValidationError):
    """Raised when a complete downstream formula closure cannot be proven."""


def _formula_derived_selection(
    plan: PatchPlan,
    resolved_ids: set[str],
) -> tuple[set[str], set[str]]:
    by_id = {patch.id: patch for patch in plan.patches}
    formula_ids = {
        patch_id for patch_id in resolved_ids if by_id[patch_id].risk == PatchRisk.FORMULA_DERIVED
    }
    formula_groups = {
        by_id[patch_id].atomic_group
        for patch_id in formula_ids
        if by_id[patch_id].atomic_group is not None
    }
    skipped_ids = {
        patch_id
        for patch_id in resolved_ids
        if patch_id in formula_ids or by_id[patch_id].atomic_group in formula_groups
    }
    return formula_ids, skipped_ids


def _atomic_members_by_patch(plan: PatchPlan) -> dict[str, frozenset[str]]:
    grouped: dict[str, set[str]] = {}
    for patch in plan.patches:
        if patch.atomic_group is not None:
            grouped.setdefault(patch.atomic_group, set()).add(patch.id)
    return {
        patch.id: (
            frozenset({patch.id})
            if patch.atomic_group is None
            else frozenset(grouped[patch.atomic_group])
        )
        for patch in plan.patches
    }


def _formula_repair_groups(
    plan: PatchPlan,
    selected_ids: Collection[str],
) -> tuple[_FormulaRepairGroup, ...]:
    """Build deterministic formula atomic groups without weakening dependency validation."""

    selected = set(selected_ids)
    by_id = {patch.id: patch for patch in plan.patches}
    members_by_patch = _atomic_members_by_patch(plan)
    grouped_formula_ids: dict[str, set[str]] = {}
    grouped_member_ids: dict[str, frozenset[str]] = {}
    for patch_id in sorted(selected):
        patch = by_id[patch_id]
        if patch.risk != PatchRisk.FORMULA_DERIVED:
            continue
        label = (
            f"atomic:{patch.atomic_group}"
            if patch.atomic_group is not None
            else f"patch:{patch.id}"
        )
        grouped_formula_ids.setdefault(label, set()).add(patch_id)
        grouped_member_ids[label] = frozenset(members_by_patch[patch_id] & selected)

    groups: list[_FormulaRepairGroup] = []
    for label in sorted(grouped_formula_ids):
        formula_ids = frozenset(grouped_formula_ids[label])
        targets = {(by_id[patch_id].sheet, by_id[patch_id].cell) for patch_id in formula_ids}
        groups.append(
            _FormulaRepairGroup(
                label=label,
                member_ids=grouped_member_ids[label],
                formula_ids=formula_ids,
                targets=frozenset((sheet, cell.upper()) for sheet, cell in targets),
            )
        )
    return tuple(groups)


def _unresolved_target_findings(
    before_findings: Collection[Finding],
    after_findings: Collection[Finding],
    patch_ids: Collection[str],
) -> list[Finding]:
    """Return after-findings that still represent findings targeted by these patches."""

    selected = set(patch_ids)
    finding_changes = _compare_findings(before_findings, after_findings)
    targeted_before = [
        finding
        for finding in before_findings
        if finding.patch_ids and set(finding.patch_ids).issubset(selected)
    ]
    targeted_locations = {
        (finding.rule_id, finding.sheet, finding.location) for finding in targeted_before
    }
    targeted_keys = {
        key for finding in targeted_before for key in finding_changes.before_keys[finding.id]
    }
    return [
        finding
        for finding in after_findings
        if (finding.rule_id, finding.sheet, finding.location) in targeted_locations
        or not finding_changes.after_keys[finding.id].isdisjoint(targeted_keys)
    ]


def _validate_formula_groups(
    groups: Collection[_FormulaRepairGroup],
    active_ids: Collection[str],
    before_snapshot: WorkbookSnapshot,
    after_snapshot: WorkbookSnapshot,
    before_findings: Collection[Finding],
    after_findings: Collection[Finding],
    formula_errors_before: Collection[str],
    formula_errors_after: Collection[str],
    modified_nodes: Collection[_FormulaNode],
) -> _FormulaGroupValidation:
    """Validate formula groups using baseline-relative, path-constrained causal evidence."""

    active = set(active_ids)
    modified = {(sheet, cell.upper()) for sheet, cell in modified_nodes}
    before_errors = _formula_error_codes(formula_errors_before)
    after_errors = _formula_error_codes(formula_errors_after)
    failures: dict[str, tuple[_FormulaRepairGroup, tuple[str, ...]]] = {}
    retained_messages: list[str] = []
    for group in groups:
        if not group.formula_ids.issubset(active):
            continue
        unresolved_findings = _unresolved_target_findings(
            before_findings,
            after_findings,
            group.member_ids,
        )
        failure_reasons = [
            f"finding:{finding.rule_id}:{finding.sheet}!{finding.location}"
            for finding in unresolved_findings
        ]
        other_modified_targets = modified - set(group.targets)
        try:
            after_graph = _formula_dependency_graph(
                after_snapshot,
                ignored_formula_nodes=other_modified_targets,
            )
            after_closure = after_graph.closure(group.targets)
        except _OpaqueFormulaDependencyError as error:
            failure_reasons.append(f"dependency-proof:{error}")
            failures[group.label] = (group, tuple(sorted(failure_reasons)))
            continue

        residuals = {node: code for node, code in after_errors.items() if node in after_closure}
        if not residuals:
            if failure_reasons:
                failures[group.label] = (group, tuple(sorted(failure_reasons)))
            continue

        try:
            before_graph = _formula_dependency_graph(
                before_snapshot,
                ignored_formula_nodes=other_modified_targets,
            )
            before_closure = before_graph.closure(group.targets)
            canonical_modified_before = {before_graph.canonical_node(node) for node in modified}
            canonical_modified_after = {after_graph.canonical_node(node) for node in modified}
        except _OpaqueFormulaDependencyError as error:
            failure_reasons.append(f"dependency-proof:{error}")
            failures[group.label] = (group, tuple(sorted(failure_reasons)))
            continue

        common_roots = {
            node
            for node, before_code in before_errors.items()
            if after_errors.get(node) == before_code
            and node not in canonical_modified_before
            and node not in canonical_modified_after
            and node not in before_closure
            and node not in after_closure
        }
        before_reachability: dict[_FormulaNode, set[_FormulaNode]] = {}
        after_reachability: dict[_FormulaNode, set[_FormulaNode]] = {}
        root_proofs: list[str] = []
        try:
            for residual, after_code in sorted(residuals.items()):
                if before_errors.get(residual) != after_code:
                    failure_reasons.append(
                        f"{residual[0]}!{residual[1]}:{after_code} (not the same baseline error)"
                    )
                    continue
                covering_roots: list[_FormulaNode] = []
                for root in sorted(common_roots):
                    if after_errors[root] != after_code:
                        continue
                    if root not in before_reachability:
                        before_reachability[root] = before_graph.closure(
                            {root},
                            blocked_internal_nodes=canonical_modified_before,
                            transparent_only=True,
                        )
                    if residual not in before_reachability[root]:
                        continue
                    if root not in after_reachability:
                        after_reachability[root] = after_graph.closure(
                            {root},
                            blocked_internal_nodes=canonical_modified_after,
                            transparent_only=True,
                        )
                    if residual in after_reachability[root]:
                        covering_roots.append(root)
                if not covering_roots:
                    failure_reasons.append(f"{residual[0]}!{residual[1]}:{after_code}")
                    continue
                root_proofs.append(
                    f"{residual[0]}!{residual[1]}:{after_code} <= "
                    + ", ".join(
                        f"{sheet}!{cell}:{after_errors[(sheet, cell)]}"
                        for sheet, cell in covering_roots
                    )
                )
        except _OpaqueFormulaDependencyError as error:
            failure_reasons.append(f"dependency-proof:{error}")

        if failure_reasons:
            failures[group.label] = (group, tuple(sorted(failure_reasons)))
            continue
        retained_messages.extend(
            [
                f"Formula group {group.label} retained baseline-relative residual errors proven "
                f"to come from unchanged external roots: {'; '.join(root_proofs)}.",
                f"公式组 {group.label} 保留了已证明由未变外部旧根因导致的基线相对残余错误: "
                + "; ".join(root_proofs)
                + "。",
            ]
        )
    return _FormulaGroupValidation(failures=failures, retained_messages=retained_messages)


def _atomic_dependent_removals(
    plan: PatchPlan,
    active_ids: Collection[str],
    rejected_ids: Collection[str],
) -> set[str]:
    """Remove rejected atomic groups and every active group that depends on them."""

    active = set(active_ids)
    members_by_patch = _atomic_members_by_patch(plan)
    removed: set[str] = set()
    for patch_id in rejected_ids:
        if patch_id in active:
            removed.update(members_by_patch[patch_id] & active)

    changed = True
    while changed:
        changed = False
        for patch in plan.patches:
            if patch.id not in active or patch.id in removed:
                continue
            if not (set(patch.prerequisite_patch_ids) & removed):
                continue
            additions = set(members_by_patch[patch.id] & active) - removed
            if additions:
                removed.update(additions)
                changed = True
    return removed


def _canonical_action_key(patch: PatchOperation) -> tuple[str, str, str, str | None, str]:
    return (
        patch.kind.value,
        patch.sheet,
        patch.cell,
        patch.source_cell,
        json.dumps(patch.after, ensure_ascii=False, sort_keys=True, default=str),
    )


def _formula_error_nodes(errors: Collection[str]) -> dict[tuple[str, str], str]:
    nodes: dict[tuple[str, str], str] = {}
    for error in errors:
        location, separator, _error_code = error.rpartition(":")
        if not separator:
            continue
        sheet, separator, coordinate = location.rpartition("!")
        if not separator or not sheet or not coordinate:
            continue
        nodes[(sheet, coordinate.upper())] = error
    return nodes


def _formula_error_codes(errors: Collection[str]) -> dict[_FormulaNode, str]:
    """Index recalculated formula errors by exact cell and error code."""

    codes: dict[_FormulaNode, str] = {}
    for node, error in _formula_error_nodes(errors).items():
        _location, separator, error_code = error.rpartition(":")
        if separator and error_code:
            codes[node] = error_code
    return codes


def _worksheet_name_indexes(
    snapshot: WorkbookSnapshot,
) -> tuple[frozenset[str], dict[str, tuple[str, ...]]]:
    """Index exact names and non-lossy compatibility candidates separately."""

    exact_names = frozenset(sheet.name for sheet in snapshot.sheets)
    compatible: dict[str, list[str]] = {}
    for name in exact_names:
        key = unicodedata.normalize("NFC", name).casefold()
        compatible.setdefault(key, []).append(name)
    return exact_names, {key: tuple(sorted(candidates)) for key, candidates in compatible.items()}


def _resolve_worksheet_name(
    name: str,
    exact_names: Collection[str],
    compatible_names: dict[str, tuple[str, ...]],
) -> tuple[str | None, str | None]:
    """Resolve exact names first and accept compatibility lookup only when unique."""

    if name in exact_names:
        return name, None
    key = unicodedata.normalize("NFC", name).casefold()
    candidates = compatible_names.get(key, ())
    if len(candidates) == 1:
        return candidates[0], None
    if not candidates:
        return None, f"reference to unavailable worksheet {name!r}"
    return None, (
        f"ambiguous worksheet reference {name!r} matches "
        + ", ".join(repr(item) for item in candidates)
    )


def _formula_operand_references(
    formula: str,
    origin_sheet: str,
    coordinate: str,
) -> tuple[FormulaIR, tuple[FormulaReference, ...], tuple[str, ...]]:
    """Parse every raw range operand independently so compatible names never deduplicate."""

    parsed = parse_formula_ir(formula, origin_sheet, coordinate)
    references: list[FormulaReference] = []
    unresolved: list[str] = []
    for operand in parsed.features.references:
        if "#REF!" in operand.upper():
            continue
        operand_ir = parse_formula_ir(f"={operand}", origin_sheet, coordinate)
        if len(operand_ir.references) != 1:
            unresolved.append(operand)
            continue
        references.append(operand_ir.references[0])
    return parsed, tuple(references), tuple(unresolved)


def _opaque_formula_dependencies(
    snapshot: WorkbookSnapshot,
    targets: Collection[tuple[str, str]],
) -> list[str]:
    """Return formulas whose downstream dependency edges cannot be proven complete."""

    opaque: list[str] = []
    exact_sheet_names, compatible_sheet_names = _worksheet_name_indexes(snapshot)
    ignored: set[tuple[str, str]] = set()
    for target_sheet, coordinate in targets:
        resolved_sheet, reason = _resolve_worksheet_name(
            target_sheet,
            exact_sheet_names,
            compatible_sheet_names,
        )
        if resolved_sheet is None:
            opaque.append(f"{target_sheet}!{coordinate.upper()}: {reason}")
            continue
        ignored.add((resolved_sheet, coordinate.upper()))
    for worksheet in snapshot.sheets:
        for coordinate, cell in worksheet.cells.items():
            formula = cell.formula
            if not isinstance(formula, str) or not formula.startswith("="):
                continue
            if (worksheet.name, coordinate.upper()) in ignored:
                continue
            parsed, references, unresolved_operands = _formula_operand_references(
                formula,
                worksheet.name,
                coordinate,
            )
            features = parsed.features
            reasons: list[str] = []
            dynamic_functions = sorted(
                _OPAQUE_DEPENDENCY_FUNCTIONS & set(features.volatile_functions)
            )
            if dynamic_functions:
                reasons.append("dynamic reference function " + ", ".join(dynamic_functions))
            if features.external_references:
                reasons.append("external workbook reference")
            if features.unsupported_reason:
                reasons.append(features.unsupported_reason)
            if features.has_whole_column_reference:
                reasons.append("whole-column reference")
            if unresolved_operands:
                reasons.append("defined name, 3-D, or other opaque range operand")
            for reference in references:
                _resolved_sheet, reason = _resolve_worksheet_name(
                    reference.sheet,
                    exact_sheet_names,
                    compatible_sheet_names,
                )
                if reason is not None:
                    reasons.append(reason)
            if reasons:
                opaque.append(
                    f"{worksheet.name}!{coordinate.upper()}: " + "; ".join(dict.fromkeys(reasons))
                )
    return sorted(opaque)


def _formula_dependency_closure(
    snapshot: WorkbookSnapshot,
    targets: Collection[tuple[str, str]],
) -> set[tuple[str, str]]:
    """Return target formulas and every downstream formula reached by ordinary A1 refs."""

    graph = _formula_dependency_graph(snapshot, ignored_formula_nodes=targets)
    return graph.closure(targets)


def _formula_dependency_graph(
    snapshot: WorkbookSnapshot,
    *,
    ignored_formula_nodes: Collection[_FormulaNode] = (),
) -> _FormulaDependencyGraph:
    """Build a complete ordinary-A1 graph, ignoring only explicitly blocked formulas."""

    opaque = _opaque_formula_dependencies(snapshot, ignored_formula_nodes)
    if opaque:
        raise _OpaqueFormulaDependencyError(
            "Formula dependency validation is incomplete because opaque references exist: "
            + ", ".join(opaque)
        )
    exact_sheet_names, compatible_sheet_names = _worksheet_name_indexes(snapshot)
    ignored: set[_FormulaNode] = set()
    for sheet, coordinate in ignored_formula_nodes:
        resolved_sheet, reason = _resolve_worksheet_name(
            sheet,
            exact_sheet_names,
            compatible_sheet_names,
        )
        if resolved_sheet is None:
            raise _OpaqueFormulaDependencyError(
                "Formula dependency validation encountered an unresolved ignored worksheet: "
                + str(reason)
            )
        ignored.add((resolved_sheet, coordinate.upper()))
    direct_dependents: dict[_FormulaNode, set[_FormulaNode]] = {}
    range_dependents: dict[str, list[_RangeDependent]] = {}
    transparent_direct_dependents: dict[_FormulaNode, set[_FormulaNode]] = {}
    transparent_range_dependents: dict[str, list[_RangeDependent]] = {}
    for worksheet in snapshot.sheets:
        for coordinate, cell in worksheet.cells.items():
            formula = cell.formula
            if not isinstance(formula, str) or not formula.startswith("="):
                continue
            dependent = (worksheet.name, coordinate.upper())
            if dependent in ignored:
                continue
            _parsed, references, unresolved_operands = _formula_operand_references(
                formula,
                worksheet.name,
                coordinate,
            )
            if unresolved_operands:
                raise _OpaqueFormulaDependencyError(
                    "Formula dependency validation encountered an unresolved range operand"
                )
            propagation = _propagation_candidate(formula, worksheet.name, coordinate)
            transparent_references = (
                set(propagation.references) if propagation is not None else set()
            )
            for reference in references:
                referenced_sheet, reason = _resolve_worksheet_name(
                    reference.sheet,
                    exact_sheet_names,
                    compatible_sheet_names,
                )
                if referenced_sheet is None:
                    raise _OpaqueFormulaDependencyError(
                        "Formula dependency validation encountered an unresolved worksheet: "
                        + str(reason)
                    )
                if reference.is_single_cell:
                    referenced_coordinate = (
                        f"{get_column_letter(reference.min_column)}{reference.min_row}"
                    )
                    direct_dependents.setdefault(
                        (referenced_sheet, referenced_coordinate), set()
                    ).add(dependent)
                    if reference in transparent_references:
                        transparent_direct_dependents.setdefault(
                            (referenced_sheet, referenced_coordinate), set()
                        ).add(dependent)
                    continue
                range_edge = (
                    reference.min_row,
                    reference.max_row,
                    reference.min_column,
                    reference.max_column,
                    dependent,
                )
                range_dependents.setdefault(referenced_sheet, []).append(range_edge)
                if reference in transparent_references:
                    transparent_range_dependents.setdefault(referenced_sheet, []).append(range_edge)
    return _FormulaDependencyGraph(
        exact_sheet_names=exact_sheet_names,
        compatible_sheet_names=compatible_sheet_names,
        direct_dependents={
            node: frozenset(dependents) for node, dependents in direct_dependents.items()
        },
        range_dependents={
            sheet: tuple(dependents) for sheet, dependents in range_dependents.items()
        },
        transparent_direct_dependents={
            node: frozenset(dependents)
            for node, dependents in transparent_direct_dependents.items()
        },
        transparent_range_dependents={
            sheet: tuple(dependents) for sheet, dependents in transparent_range_dependents.items()
        },
    )


def _validated_apply_paths(source: Path, output: Path) -> tuple[Path, Path]:
    source = source.expanduser().resolve()
    output = output.expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        raise UsageError("Repair output must use the .xlsx extension")
    if source == output:
        raise UsageError("Output must be a new path; the source workbook is never overwritten")
    if output.exists():
        raise UsageError(f"Output already exists and will not be overwritten: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    return source, output


def _validated_direct_child(path: Path, parent: Path, label: str) -> Path:
    resolved_parent = parent.resolve()
    resolved = path.resolve()
    if resolved.parent != resolved_parent:
        raise PatchValidationError(f"{label} escaped its intended parent directory")
    return resolved


def _create_private_candidate(output: Path) -> tuple[Path, Path]:
    raw_workspace = Path(tempfile.mkdtemp(prefix=".wl-r-", dir=output.parent))
    workspace = _validated_direct_child(
        raw_workspace,
        output.parent,
        "Private repair workspace",
    )
    candidate = _new_private_candidate(workspace)
    return workspace, candidate


def _new_private_candidate(workspace: Path) -> Path:
    return _validated_direct_child(
        workspace / f"c-{secrets.token_hex(8)}.xlsx",
        workspace,
        "Private repair candidate",
    )


def _cleanup_private_candidate(workspace: Path, candidate: Path) -> bool:
    removed = True
    try:
        candidate.unlink(missing_ok=True)
    except OSError:
        removed = False
    try:
        workspace.rmdir()
    except OSError:
        removed = False
    return removed


def _cleanup_failed_publication(
    candidate: Path,
    output: Path,
    *,
    output_existed_before: bool,
) -> bool:
    if output_existed_before:
        return True
    try:
        if output.exists() and candidate.exists() and os.path.samefile(candidate, output):
            output.unlink()
        return True
    except OSError:
        return False


def _publish_candidate(candidate: Path, output: Path, expected_hash: str) -> None:
    if sha256_file(candidate) != expected_hash:
        raise PatchValidationError("Private candidate changed after validation")
    output_existed_before = output.exists()
    try:
        os.link(candidate, output)
    except FileExistsError as exc:
        cleanup_complete = _cleanup_failed_publication(
            candidate,
            output,
            output_existed_before=output_existed_before,
        )
        exists_error = UsageError(f"Output already exists and will not be overwritten: {output}")
        if not cleanup_complete:
            exists_error.__dict__["publication_cleanup_incomplete"] = True
        raise exists_error from exc
    except OSError as exc:
        cleanup_complete = _cleanup_failed_publication(
            candidate,
            output,
            output_existed_before=output_existed_before,
        )
        filesystem_error = PatchValidationError(
            "The destination filesystem does not support atomic no-overwrite publication"
        )
        if not cleanup_complete:
            filesystem_error.__dict__["publication_cleanup_incomplete"] = True
        raise filesystem_error from exc


def _write_unchanged_candidate(
    source: Path,
    candidate: Path,
    *,
    expected_source_sha256: str,
) -> tuple[str, str]:
    source_hash = sha256_file(source)
    if source_hash != expected_source_sha256:
        raise StalePlanError(
            f"Plan source hash {expected_source_sha256} does not match workbook hash {source_hash}"
        )
    with source.open("rb") as source_handle, candidate.open("xb") as output_handle:
        shutil.copyfileobj(source_handle, output_handle, length=1024 * 1024)
        output_handle.flush()
        os.fsync(output_handle.fileno())
    if sha256_file(source) != source_hash:
        raise PatchValidationError("Source workbook hash changed during apply")
    output_hash = sha256_file(candidate)
    if output_hash != source_hash:
        raise PatchValidationError("Unchanged degraded output failed hash verification")
    return source_hash, output_hash


def _failure_result(
    *,
    plan: PatchPlan,
    output: Path,
    before_scan: Any | None,
    attempted_ids: Collection[str],
    downgraded_ids: Collection[str],
    skipped_ids: Collection[str],
    provider: RecalculationProvider,
    formula_errors_before: Collection[str],
    formula_errors_after: Collection[str],
    reason: str,
    rollback_performed: bool,
    extra_messages: Collection[str] = (),
) -> PatchResult:
    findings = getattr(before_scan, "findings", ()) if before_scan is not None else ()
    attempted = ", ".join(sorted(attempted_ids)) or "none"
    return PatchResult(
        source_sha256=plan.source_sha256,
        output_sha256="",
        output_path=str(output),
        applied_patch_ids=[],
        package_changes=[],
        resolved_finding_ids=[],
        remaining_finding_ids=sorted(finding.id for finding in findings),
        new_finding_ids=[],
        validation_messages=[
            *extra_messages,
            f"Validation failed before publication; attempted patches: {attempted}.",
            f"验证在发布前失败; 尝试的补丁: {attempted}。",
            f"Reason: {reason}",
            f"原因: {reason}",
        ],
        recalculation_provider=provider,
        formula_errors_before=list(formula_errors_before),
        formula_errors_after=list(formula_errors_after),
        downgraded_patch_ids=sorted(downgraded_ids),
        skipped_patch_ids=sorted(skipped_ids),
        validation_status=ValidationStatus.FAILED,
        rollback_performed=rollback_performed,
    )


def _write_failure_report(
    output: Path,
    result: PatchResult,
    *,
    attempted_ids: Collection[str],
    reason: str,
) -> Path | None:
    attempted = ", ".join(sorted(attempted_ids)) or "none / 无"
    rollback_message = (
        "Rollback / 回滚: private candidate cleanup completed; intended output not published / "
        "私有候选清理已完成; 目标输出未发布"
        if result.rollback_performed
        else "Rollback / 回滚: private candidate or publication cleanup was incomplete; "
        "the intended output may still exist and must not be trusted / "
        "私有候选或发布未完全清理; 目标输出可能仍然存在且不得信任"
    )
    report = "\n".join(
        [
            "WorkbookLens repair validation failure / WorkbookLens 修复验证失败",
            "",
            "Status / 状态: FAILED / 失败",
            rollback_message,
            f"Intended output / 目标输出: {output}",
            f"Attempted patch IDs / 尝试的补丁 ID: {attempted}",
            f"Recalculation provider / 重算提供者: {result.recalculation_provider.value}",
            f"Reason / 原因: {reason}",
            *(
                [
                    "",
                    "Validation history / 验证历史:",
                    *(f"- {message}" for message in result.validation_messages),
                ]
                if result.validation_messages
                else []
            ),
            "",
            "The original workbook was not overwritten. / 原始工作簿未被覆盖。",
        ]
    )
    try:
        descriptor, report_name = tempfile.mkstemp(
            prefix=".wl-f-",
            suffix=".txt",
            dir=output.parent,
            text=True,
        )
        try:
            report_path = _validated_direct_child(
                Path(report_name),
                output.parent,
                "Failure report",
            )
        except PatchValidationError:
            os.close(descriptor)
            return None
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(report)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        return report_path
    except OSError:
        return None


def _record_validation_failure(
    error: Exception,
    *,
    plan: PatchPlan,
    output: Path,
    before_scan: Any | None,
    attempted_ids: Collection[str],
    downgraded_ids: Collection[str],
    skipped_ids: Collection[str],
    provider: RecalculationProvider,
    formula_errors_before: Collection[str],
    formula_errors_after: Collection[str],
    rollback_performed: bool,
    extra_messages: Collection[str] = (),
) -> None:
    reason = str(error)
    result = _failure_result(
        plan=plan,
        output=output,
        before_scan=before_scan,
        attempted_ids=attempted_ids,
        downgraded_ids=downgraded_ids,
        skipped_ids=skipped_ids,
        provider=provider,
        formula_errors_before=formula_errors_before,
        formula_errors_after=formula_errors_after,
        reason=reason,
        rollback_performed=rollback_performed,
        extra_messages=extra_messages,
    )
    report_path = _write_failure_report(
        output,
        result,
        attempted_ids=attempted_ids,
        reason=reason,
    )
    result.failure_report_path = str(report_path) if report_path is not None else None
    error.__dict__["patch_result"] = result
    error.__dict__["failure_report_path"] = report_path
    if report_path is not None:
        error.args = (f"{reason}\nFailure report / 失败报告: {report_path}",)


def apply_patch_plan(
    source: Path,
    plan: PatchPlan,
    output: Path,
    *,
    selected_ids: Collection[str] | None = None,
    safe_only: bool = False,
    accept_layout_risk: bool = False,
    accept_semantic_risk: bool = False,
    auto_repair: bool = False,
    recalc_provider: RecalculationPreference = "auto",
    recalc_timeout_seconds: int = 180,
    trust_workbook_for_recalculation: bool = False,
    config: dict[str, Any] | None = None,
    limits: PackageLimits | None = None,
) -> PatchResult:
    """Apply into a private candidate, validate it completely, then publish atomically."""

    if recalc_provider not in {"auto", "excel", "libreoffice", "none"}:
        raise UsageError("--recalc-provider must be one of auto, excel, libreoffice, or none")
    if recalc_timeout_seconds <= 0:
        raise UsageError("Recalculation timeout must be positive")
    source, output = _validated_apply_paths(source, output)
    effective_recalc_preference: RecalculationPreference = (
        recalc_provider if trust_workbook_for_recalculation else "none"
    )
    resolved_ids = resolve_patch_selection(
        plan,
        selected_ids,
        safe_only,
        accept_layout_risk=accept_layout_risk,
        accept_formula_derived=not safe_only and effective_recalc_preference != "none",
        accept_semantic_risk=accept_semantic_risk,
        auto_repair=auto_repair,
    )
    before_scan = scan_workbook(source, config=config, limits=limits)
    canonical_plan = build_patch_plan(before_scan)
    formula_derived_ids, formula_group_ids = _formula_derived_selection(plan, resolved_ids)
    patches_by_id = {patch.id: patch for patch in plan.patches}
    downgraded_ids: set[str] = set()
    skipped_ids: set[str] = set()
    effective_ids = set(resolved_ids)
    downgrade_messages: list[str] = []

    formula_targets = {
        (patches_by_id[patch_id].sheet, patches_by_id[patch_id].cell)
        for patch_id in formula_derived_ids
    }
    opaque_dependencies = (
        _opaque_formula_dependencies(before_scan.snapshot, formula_targets)
        if formula_derived_ids
        else []
    )
    if opaque_dependencies:
        if not auto_repair:
            opaque_error = _OpaqueFormulaDependencyError(
                "Formula-derived repair cannot be automatically validated because opaque "
                "references prevent a complete downstream dependency closure: "
                + ", ".join(opaque_dependencies)
            )
            _record_validation_failure(
                opaque_error,
                plan=plan,
                output=output,
                before_scan=before_scan,
                attempted_ids=resolved_ids,
                downgraded_ids=downgraded_ids,
                skipped_ids=skipped_ids,
                provider=RecalculationProvider.NONE,
                formula_errors_before=(),
                formula_errors_after=(),
                rollback_performed=True,
            )
            raise opaque_error
        downgraded_ids.update(formula_derived_ids)
        skipped_ids.update(formula_group_ids)
        effective_ids.difference_update(formula_group_ids)
        downgrade_messages.append(
            "Formula-derived patches were not applied because a complete downstream dependency "
            "closure cannot be proven in the presence of opaque references: "
            + ", ".join(opaque_dependencies)
        )

    recalculation_required_ids = {
        patch_id
        for patch_id in effective_ids
        if patches_by_id[patch_id].derivation.requires_recalculation
    }
    providers = (
        candidate_providers(effective_recalc_preference)
        if recalculation_required_ids and effective_recalc_preference != "none"
        else ()
    )
    remaining_formula_ids, remaining_formula_group_ids = _formula_derived_selection(
        plan, effective_ids
    )
    if remaining_formula_ids and not providers:
        if not auto_repair:
            unavailable_formula_error = UsageError(
                "Selected formula-derived patches require an available local recalculation "
                "provider; no requested provider is installed: "
                + ", ".join(sorted(remaining_formula_ids))
            )
            _record_validation_failure(
                unavailable_formula_error,
                plan=plan,
                output=output,
                before_scan=before_scan,
                attempted_ids=resolved_ids,
                downgraded_ids=downgraded_ids,
                skipped_ids=skipped_ids,
                provider=RecalculationProvider.NONE,
                formula_errors_before=(),
                formula_errors_after=(),
                rollback_performed=True,
            )
            raise unavailable_formula_error
        downgraded_ids.update(remaining_formula_ids)
        skipped_ids.update(remaining_formula_group_ids)
        effective_ids.difference_update(remaining_formula_group_ids)
        if recalc_provider == "none":
            downgrade_messages.append(
                "Formula-derived patches were not applied because recalculation was disabled."
            )
        elif not trust_workbook_for_recalculation:
            downgrade_messages.append(
                "Formula-derived patches were not applied because this workbook was not explicitly "
                "trusted for local application recalculation."
            )
        else:
            downgrade_messages.append(
                "Formula-derived patches were not applied because no requested local recalculation "
                "provider is available."
            )

    remaining_recalculation_ids = {
        patch_id
        for patch_id in effective_ids
        if patches_by_id[patch_id].derivation.requires_recalculation
    }
    unavailable_required_ids = sorted(remaining_recalculation_ids) if not providers else []
    if unavailable_required_ids:
        unavailable_required_error = UsageError(
            "Selected patches require a local recalculation provider and explicit workbook trust: "
            + ", ".join(unavailable_required_ids)
        )
        _record_validation_failure(
            unavailable_required_error,
            plan=plan,
            output=output,
            before_scan=before_scan,
            attempted_ids=effective_ids,
            downgraded_ids=downgraded_ids,
            skipped_ids=skipped_ids,
            provider=RecalculationProvider.NONE,
            formula_errors_before=(),
            formula_errors_after=(),
            rollback_performed=True,
        )
        raise unavailable_required_error

    active_formula_ids = formula_derived_ids & effective_ids
    formula_groups = _formula_repair_groups(plan, effective_ids) if active_formula_ids else ()
    max_formula_downgrade_rounds = len(formula_groups)

    try:
        workspace, candidate = _create_private_candidate(output)
    except Exception as workspace_error:
        _record_validation_failure(
            workspace_error,
            plan=plan,
            output=output,
            before_scan=before_scan,
            attempted_ids=effective_ids,
            downgraded_ids=downgraded_ids,
            skipped_ids=skipped_ids,
            provider=RecalculationProvider.NONE,
            formula_errors_before=(),
            formula_errors_after=(),
            rollback_performed=True,
        )
        raise
    recalculation_provider = RecalculationProvider.NONE
    formula_errors_before: list[str] = []
    formula_errors_after: list[str] = []
    recalculation_preference = effective_recalc_preference
    recalculation_started = False
    final_candidate_recalculated = False
    downgrade_round = 0
    cleanup_attempted = False
    validated_candidate_scan: Any | None = None
    formula_proof_messages: list[str] = []
    try:
        low_level: Any | None = None
        selected: list[PatchOperation] = []
        applied_recalculation_required: list[PatchOperation] = []
        while True:
            validated_candidate_scan = None
            formula_proof_messages = []
            if not effective_ids:
                _validate_canonical_plan(plan, canonical_plan)
                source_hash, output_hash = _write_unchanged_candidate(
                    source,
                    candidate,
                    expected_source_sha256=plan.source_sha256,
                )
                scan_workbook(candidate, config=config, limits=limits)
                if sha256_file(candidate) != output_hash:
                    raise PatchValidationError(
                        "Unchanged private candidate changed during validation"
                    )
                if recalculation_started:
                    formula_errors_after = list(formula_errors_before)
                result = PatchResult(
                    source_sha256=source_hash,
                    output_sha256=output_hash,
                    output_path=str(output),
                    applied_patch_ids=[],
                    package_changes=[],
                    resolved_finding_ids=[],
                    remaining_finding_ids=sorted(finding.id for finding in before_scan.findings),
                    new_finding_ids=[],
                    validation_messages=[
                        *downgrade_messages,
                        "No workbook bytes were changed; the validated output is a byte-identical "
                        "copy and the skipped candidates remain available for review.",
                    ],
                    recalculation_provider=recalculation_provider,
                    formula_errors_before=formula_errors_before,
                    formula_errors_after=formula_errors_after,
                    downgraded_patch_ids=sorted(downgraded_ids),
                    skipped_patch_ids=sorted(skipped_ids),
                    validation_status=ValidationStatus.DEGRADED,
                )
                _publish_candidate(candidate, output, output_hash)
                return result

            low_level, selected = patch_ooxml_package(
                source,
                plan,
                candidate,
                selected_ids=effective_ids,
                safe_only=False,
                accept_layout_risk=accept_layout_risk,
                accept_formula_derived=bool(formula_derived_ids & effective_ids),
                accept_semantic_risk=accept_semantic_risk,
                limits=limits,
                canonical_plan=canonical_plan,
            )
            applied_recalculation_required = [
                patch for patch in selected if patch.derivation.requires_recalculation
            ]
            final_candidate_recalculated = False
            if applied_recalculation_required or recalculation_started:
                validation = validate_formula_recalculation(
                    source,
                    candidate,
                    preference=recalculation_preference,
                    timeout_seconds=recalc_timeout_seconds,
                    limits=limits,
                    expected_source_sha256=plan.source_sha256,
                    expected_candidate_sha256=low_level.output_sha256,
                )
                validated_provider = RecalculationProvider(validation.provider.kind)
                if (
                    recalculation_provider != RecalculationProvider.NONE
                    and validated_provider != recalculation_provider
                ):
                    raise PatchValidationError(
                        "Automatic repair changed recalculation providers between validation rounds"
                    )
                recalculation_provider = validated_provider
                if validated_provider == RecalculationProvider.EXCEL:
                    recalculation_preference = "excel"
                elif validated_provider == RecalculationProvider.LIBREOFFICE:
                    recalculation_preference = "libreoffice"
                else:
                    raise PatchValidationError(
                        "Automatic repair received an unsupported recalculation provider"
                    )
                next_formula_errors_before = list(validation.formula_errors_before)
                if recalculation_started and next_formula_errors_before != formula_errors_before:
                    raise PatchValidationError(
                        "Source formula-error baseline changed between automatic repair rounds"
                    )
                recalculation_started = True
                final_candidate_recalculated = True
                formula_errors_before = next_formula_errors_before
                formula_errors_after = list(validation.formula_errors_after)
                validated_candidate_scan = scan_workbook(
                    candidate,
                    config=config,
                    limits=limits,
                )
                modified_nodes = {
                    (patch.sheet, cell)
                    for patch in selected
                    for cell in (_strict_a1_cells(patch.cell) or ())
                }
                group_validation = _validate_formula_groups(
                    formula_groups,
                    effective_ids,
                    before_scan.snapshot,
                    validated_candidate_scan.snapshot,
                    before_scan.findings,
                    validated_candidate_scan.findings,
                    formula_errors_before,
                    formula_errors_after,
                    modified_nodes,
                )
                formula_proof_messages = group_validation.retained_messages
                failing_groups = group_validation.failures
                if failing_groups:
                    failure_detail = "; ".join(
                        f"{label} [{', '.join(errors)}]"
                        for label, (_group, errors) in sorted(failing_groups.items())
                    )
                    if not auto_repair:
                        raise PatchValidationError(
                            "Formula-derived repair failed target dependency chain validation: "
                            + failure_detail
                        )
                    rejected_ids = {
                        patch_id
                        for group, _errors in failing_groups.values()
                        for patch_id in group.member_ids
                    }
                    removed_ids = _atomic_dependent_removals(
                        plan,
                        effective_ids,
                        rejected_ids,
                    )
                    next_effective_ids = effective_ids - removed_ids
                    if not removed_ids or next_effective_ids == effective_ids:
                        raise PatchValidationError(
                            "Automatic formula-group downgrade made no progress"
                        )
                    downgrade_round += 1
                    if downgrade_round > max_formula_downgrade_rounds:
                        raise PatchValidationError(
                            "Automatic formula-group downgrade exceeded its bounded round limit"
                        )
                    downgraded_ids.update(
                        patch_id
                        for patch_id in removed_ids
                        if patches_by_id[patch_id].risk == PatchRisk.FORMULA_DERIVED
                    )
                    skipped_ids.update(removed_ids)
                    removed_detail = ", ".join(sorted(removed_ids))
                    downgrade_messages.extend(
                        [
                            f"Auto-repair recalculation round {downgrade_round} rejected formula "
                            f"atomic group(s) {failure_detail}; removed patches and reverse "
                            f"dependents: {removed_detail}.",
                            f"自动修复重算第 {downgrade_round} 轮因目标依赖链仍有错误而拒绝公式"
                            f"原子组 {failure_detail}; 已移除补丁及其反向依赖项: {removed_detail}。",
                        ]
                    )
                    effective_ids = next_effective_ids
                    candidate.unlink()
                    candidate = _new_private_candidate(workspace)
                    continue
            break

        if low_level is None:
            raise PatchValidationError("Automatic repair did not produce a final candidate")
        after_scan = validated_candidate_scan or scan_workbook(
            candidate,
            config=config,
            limits=limits,
        )
        finding_changes = _compare_findings(before_scan.findings, after_scan.findings)
        selected_patch_ids = {patch.id for patch in selected}
        targeted_before = [
            finding
            for finding in before_scan.findings
            if finding.patch_ids and set(finding.patch_ids).issubset(selected_patch_ids)
        ]
        targeted_locations = {
            (finding.rule_id, finding.sheet, finding.location) for finding in targeted_before
        }
        targeted_keys = {
            key for finding in targeted_before for key in finding_changes.before_keys[finding.id]
        }
        unresolved_targets = [
            finding
            for finding in after_scan.findings
            if (finding.rule_id, finding.sheet, finding.location) in targeted_locations
            or not finding_changes.after_keys[finding.id].isdisjoint(targeted_keys)
        ]
        if unresolved_targets:
            targets = ", ".join(
                sorted(
                    f"{finding.rule_id}:{finding.sheet}!{finding.location}"
                    for finding in unresolved_targets
                )
            )
            raise PatchValidationError(f"Repair did not resolve targeted findings: {targets}")
        dangerous_new_or_escalated = finding_changes.dangerous_findings
        if dangerous_new_or_escalated:
            identifiers = ", ".join(finding.id for finding in dangerous_new_or_escalated)
            raise PatchValidationError(
                "Repair introduced new or severity-escalated error-level findings: " + identifiers
            )
        selected_action_keys = {_canonical_action_key(patch) for patch in selected}
        repeated_actions = sorted(
            selected_action_keys & {_canonical_action_key(patch) for patch in after_scan.patches}
        )
        if repeated_actions:
            raise PatchValidationError(
                "Repair was not idempotent; the same patch action was proposed again"
            )
        messages = [
            *downgrade_messages,
            *formula_proof_messages,
            "Output reopened through the secure OOXML reader and openpyxl read-only mode.",
            "Every untouched ZIP part has byte-identical uncompressed content.",
        ]
        if final_candidate_recalculated:
            messages.append(
                "Recalculation-required patches were validated on isolated source and candidate "
                f"copies with {recalculation_provider.value}; the final file remains the direct "
                "OOXML output."
            )
        elif low_level.formula_changed:
            messages.append(
                "Formula cached values were removed and full recalculation is requested on next open; "
                "WorkbookLens did not calculate formulas."
            )
        if downgraded_ids:
            messages.append(
                "Downgraded formula-derived patches and their atomic groups were not applied."
            )
        output_hash = sha256_file(candidate)
        if output_hash != low_level.output_sha256:
            raise PatchValidationError("Private candidate changed during validation")
        result = PatchResult(
            source_sha256=low_level.source_sha256,
            output_sha256=output_hash,
            output_path=str(output),
            applied_patch_ids=[patch.id for patch in selected],
            package_changes=low_level.changes,
            resolved_finding_ids=sorted(finding_changes.resolved_ids),
            remaining_finding_ids=sorted(finding_changes.remaining_ids),
            new_finding_ids=sorted(finding_changes.new_ids),
            validation_messages=messages,
            recalculation_provider=recalculation_provider,
            formula_errors_before=formula_errors_before,
            formula_errors_after=formula_errors_after,
            downgraded_patch_ids=sorted(downgraded_ids),
            skipped_patch_ids=sorted(skipped_ids),
            validation_status=(
                ValidationStatus.DEGRADED if downgraded_ids else ValidationStatus.PASSED
            ),
        )
        _publish_candidate(candidate, output, output_hash)
        return result
    except Exception as error:
        attempted_provider = getattr(error, "recalculation_provider", None)
        if attempted_provider is not None:
            with contextlib.suppress(ValueError):
                recalculation_provider = RecalculationProvider(
                    str(getattr(attempted_provider, "value", attempted_provider))
                )
        rollback_performed = _cleanup_private_candidate(workspace, candidate) and not bool(
            getattr(error, "publication_cleanup_incomplete", False)
        )
        cleanup_attempted = True
        _record_validation_failure(
            error,
            plan=plan,
            output=output,
            before_scan=before_scan,
            attempted_ids=effective_ids,
            downgraded_ids=downgraded_ids,
            skipped_ids=skipped_ids,
            provider=recalculation_provider,
            formula_errors_before=formula_errors_before,
            formula_errors_after=formula_errors_after,
            rollback_performed=rollback_performed,
            extra_messages=downgrade_messages,
        )
        raise
    finally:
        if not cleanup_attempted:
            _cleanup_private_candidate(workspace, candidate)
