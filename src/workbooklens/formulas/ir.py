"""Conservative formula IR and dependency inspection without expression evaluation."""

from __future__ import annotations

import re
from bisect import bisect_left, bisect_right
from collections.abc import Callable, Iterable, Iterator
from contextlib import suppress
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any, Literal

from openpyxl.cell.cell import Cell
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token, TokenizerError
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.constants import MAX_COLUMN, MAX_ROW

from workbooklens.formulas.analysis import FormulaFeatures, analyze_formula

AGGREGATE_FUNCTIONS = frozenset({"AVERAGE", "COUNT", "MAX", "MIN", "SUM"})
SHEET_REFERENCE_RE = re.compile(
    r"^(?:(?:'(?P<quoted>(?:[^']|'')+)'|(?P<plain>[^!]+))!)?(?P<body>.+)$"
)


@dataclass(frozen=True, slots=True)
class FormulaReference:
    """One bounded A1 cell or rectangular range resolved to a worksheet name."""

    raw: str
    origin_sheet: str
    sheet: str
    explicit_sheet: bool
    min_row: int
    max_row: int
    min_column: int
    max_column: int

    @property
    def is_cross_sheet(self) -> bool:
        return self.explicit_sheet and self.sheet.casefold() != self.origin_sheet.casefold()

    @property
    def is_single_cell(self) -> bool:
        return self.min_row == self.max_row and self.min_column == self.max_column

    @property
    def is_single_column(self) -> bool:
        return self.min_column == self.max_column

    @property
    def area(self) -> int:
        return (self.max_row - self.min_row + 1) * (self.max_column - self.min_column + 1)

    @property
    def local_range(self) -> str:
        start = f"{get_column_letter(self.min_column)}{self.min_row}"
        end = f"{get_column_letter(self.max_column)}{self.max_row}"
        return start if start == end else f"{start}:{end}"


@dataclass(frozen=True, slots=True)
class AggregateReference:
    """A simple aggregate whose only argument is one ordinary A1 range."""

    function: str
    reference: FormulaReference


@dataclass(frozen=True, slots=True)
class DegenerateExpression:
    """A statically visible top-level algebraic identity or collapse."""

    kind: Literal["self_division", "self_subtraction", "multiply_by_zero"]
    operator: str
    left: str
    right: str


@dataclass(frozen=True, slots=True)
class FormulaIR:
    """Non-evaluating formula representation used by semantic rules."""

    formula: str
    origin_sheet: str
    origin_coordinate: str
    features: FormulaFeatures
    references: tuple[FormulaReference, ...]
    aggregate_references: tuple[AggregateReference, ...]
    degeneracy: DegenerateExpression | None


@dataclass(frozen=True, slots=True)
class ReferenceTarget:
    """Sparse target-state summary for one formula dependency edge."""

    reference: FormulaReference
    sheet_exists: bool
    nonblank_cells: tuple[str, ...]
    content_bounds: tuple[int, int, int, int] | None
    outside_content: bool
    fully_blank: bool


def _cell_row(cell: Cell) -> int:
    return cell.row


def _cell_column(cell: Cell) -> int:
    return cell.column


@dataclass(frozen=True, slots=True)
class _AxisCellIndex:
    """Flat cells grouped by one primary axis and sorted by the other axis."""

    _cells: tuple[Cell, ...]
    _keys: tuple[int, ...]
    _starts: tuple[int, ...]

    @classmethod
    def from_cells(
        cls,
        cells: Iterable[Cell],
        *,
        primary: Callable[[Cell], int],
        secondary: Callable[[Cell], int],
    ) -> _AxisCellIndex:
        ordered = tuple(sorted(cells, key=lambda cell: (primary(cell), secondary(cell))))
        keys: list[int] = []
        starts: list[int] = []
        for offset, cell in enumerate(ordered):
            key = primary(cell)
            if not keys or keys[-1] != key:
                keys.append(key)
                starts.append(offset)
        return cls(_cells=ordered, _keys=tuple(keys), _starts=tuple(starts))

    def iter_keys(self, minimum: int, maximum: int) -> Iterator[int]:
        start = bisect_left(self._keys, minimum)
        end = bisect_right(self._keys, maximum)
        yield from self._keys[start:end]

    def iter_cells(
        self,
        key: int,
        *,
        min_secondary: int,
        max_secondary: int,
        secondary: Callable[[Cell], int],
        limit: int | None = None,
    ) -> Iterator[Cell]:
        if limit is not None and limit <= 0:
            return
        group = bisect_left(self._keys, key)
        if group >= len(self._keys) or self._keys[group] != key:
            return
        start = self._starts[group]
        end = self._starts[group + 1] if group + 1 < len(self._starts) else len(self._cells)
        selected_start = bisect_left(
            self._cells,
            min_secondary,
            start,
            end,
            key=secondary,
        )
        selected_end = bisect_right(
            self._cells,
            max_secondary,
            selected_start,
            end,
            key=secondary,
        )
        if limit is not None:
            selected_end = min(selected_end, selected_start + limit)
        yield from self._cells[selected_start:selected_end]


@dataclass(frozen=True, slots=True)
class WorksheetContentIndex:
    """One scan's flat sparse nonblank-cell indexes for a worksheet."""

    worksheet: Worksheet
    content_bounds: tuple[int, int, int, int] | None
    _rows: _AxisCellIndex
    _columns: _AxisCellIndex

    @classmethod
    def from_cells(
        cls,
        worksheet: Worksheet,
        cells: Iterable[Cell],
    ) -> WorksheetContentIndex:
        retained = [cell for cell in cells if cell.value is not None]
        return cls(
            worksheet=worksheet,
            content_bounds=(
                min(cell.row for cell in retained),
                max(cell.row for cell in retained),
                min(cell.column for cell in retained),
                max(cell.column for cell in retained),
            )
            if retained
            else None,
            _rows=_AxisCellIndex.from_cells(
                retained,
                primary=_cell_row,
                secondary=_cell_column,
            ),
            _columns=_AxisCellIndex.from_cells(
                retained,
                primary=_cell_column,
                secondary=_cell_row,
            ),
        )

    @classmethod
    def from_worksheet(cls, worksheet: Worksheet) -> WorksheetContentIndex:
        return cls.from_cells(
            worksheet,
            (cell for cell in worksheet._cells.values() if isinstance(cell, Cell)),
        )

    def iter_nonblank_row(
        self,
        row: int,
        *,
        min_column: int = 1,
        max_column: int = MAX_COLUMN,
        limit: int | None = None,
    ) -> Iterator[Cell]:
        """Yield nonblank cells in one row, in column order, with an optional cap."""

        yield from self._rows.iter_cells(
            row,
            min_secondary=min_column,
            max_secondary=max_column,
            secondary=_cell_column,
            limit=limit,
        )

    def iter_nonblank_column(
        self,
        column: int,
        *,
        min_row: int = 1,
        max_row: int = MAX_ROW,
        limit: int | None = None,
    ) -> Iterator[Cell]:
        """Yield nonblank cells in one column, in row order, with an optional cap."""

        yield from self._columns.iter_cells(
            column,
            min_secondary=min_row,
            max_secondary=max_row,
            secondary=_cell_row,
            limit=limit,
        )

    def iter_nonblank_rectangle(
        self,
        *,
        min_row: int,
        max_row: int,
        min_column: int,
        max_column: int,
        limit: int | None = None,
    ) -> Iterator[Cell]:
        """Yield bounded rectangle hits in row-major order without expanding blanks."""

        if limit is not None and limit <= 0:
            return
        if min_column == max_column:
            yield from self.iter_nonblank_column(
                min_column,
                min_row=min_row,
                max_row=max_row,
                limit=limit,
            )
            return
        if min_row == max_row:
            yield from self.iter_nonblank_row(
                min_row,
                min_column=min_column,
                max_column=max_column,
                limit=limit,
            )
            return

        yielded = 0
        for row in self._rows.iter_keys(min_row, max_row):
            remaining = None if limit is None else limit - yielded
            for cell in self._rows.iter_cells(
                row,
                min_secondary=min_column,
                max_secondary=max_column,
                secondary=_cell_column,
                limit=remaining,
            ):
                yield cell
                yielded += 1
            if limit is not None and yielded >= limit:
                return


def worksheet_content_index(context: Any, worksheet: Worksheet) -> WorksheetContentIndex:
    """Return the cached sparse index for this worksheet during one scan."""

    cache_key = "formula_ir.worksheet_content_indexes.v1"
    cached = context.analysis_cache.get(cache_key)
    if not isinstance(cached, dict):
        cached = {}
        context.analysis_cache[cache_key] = cached
    index = cached.get(worksheet.title)
    if isinstance(index, WorksheetContentIndex) and index.worksheet is worksheet:
        return index
    index = WorksheetContentIndex.from_worksheet(worksheet)
    cached[worksheet.title] = index
    return index


def _parse_reference(raw: str, origin_sheet: str) -> FormulaReference | None:
    if "[" in raw or "]" in raw:
        return None
    match = SHEET_REFERENCE_RE.fullmatch(raw.strip())
    if match is None:
        return None
    quoted = match.group("quoted")
    plain = match.group("plain")
    if plain is not None and ":" in plain:
        # A three-dimensional sheet span requires workbook-specific semantics.
        return None
    sheet = quoted.replace("''", "'") if quoted is not None else plain
    body = match.group("body").replace("$", "")
    try:
        min_column, min_row, max_column, max_row = range_boundaries(body)
    except (TypeError, ValueError):
        return None
    if min_column is None or min_row is None or max_column is None or max_row is None:
        return None
    if not (1 <= min_row <= max_row <= MAX_ROW and 1 <= min_column <= max_column <= MAX_COLUMN):
        return None
    return FormulaReference(
        raw=raw,
        origin_sheet=origin_sheet,
        sheet=sheet or origin_sheet,
        explicit_sheet=sheet is not None,
        min_row=min_row,
        max_row=max_row,
        min_column=min_column,
        max_column=max_column,
    )


def _tokenize(formula: str) -> tuple[Token, ...] | None:
    try:
        return tuple(token for token in Tokenizer(formula).items if token.type != "WHITE-SPACE")
    except (IndexError, TokenizerError, ValueError):
        return None


def _function_name(token: Token) -> str | None:
    if token.type != "FUNC" or token.subtype != "OPEN":
        return None
    return str(token.value)[:-1].upper().rsplit(".", maxsplit=1)[-1]


def _simple_aggregate(
    tokens: tuple[Token, ...] | None,
    origin_sheet: str,
) -> AggregateReference | None:
    if tokens is None or len(tokens) != 3:
        return None
    function = _function_name(tokens[0])
    operand = tokens[1]
    closing = tokens[2]
    if (
        function not in AGGREGATE_FUNCTIONS
        or operand.type != "OPERAND"
        or operand.subtype != "RANGE"
        or closing.type != "FUNC"
        or closing.subtype != "CLOSE"
    ):
        return None
    reference = _parse_reference(str(operand.value), origin_sheet)
    if reference is None or reference.is_single_cell:
        return None
    return AggregateReference(function=function, reference=reference)


def _strip_outer_parentheses(tokens: tuple[Token, ...]) -> tuple[Token, ...]:
    current = tokens
    while (
        len(current) >= 2
        and current[0].type == "PAREN"
        and current[0].subtype == "OPEN"
        and current[-1].type == "PAREN"
        and current[-1].subtype == "CLOSE"
    ):
        depth = 0
        wraps_all = True
        for index, token in enumerate(current):
            if token.type == "PAREN" and token.subtype == "OPEN":
                depth += 1
            elif token.type == "PAREN" and token.subtype == "CLOSE":
                depth -= 1
                if depth == 0 and index != len(current) - 1:
                    wraps_all = False
                    break
            if depth < 0:
                wraps_all = False
                break
        if not wraps_all or depth != 0:
            break
        current = current[1:-1]
    return current


def _top_level_operator(tokens: tuple[Token, ...]) -> tuple[int, Token] | None:
    depth = 0
    candidates: list[tuple[int, Token]] = []
    for index, token in enumerate(tokens):
        if token.type in {"FUNC", "PAREN"} and token.subtype == "OPEN":
            depth += 1
            continue
        if token.type in {"FUNC", "PAREN"} and token.subtype == "CLOSE":
            depth -= 1
            if depth < 0:
                return None
            continue
        if depth == 0 and token.type == "OPERATOR-INFIX" and token.value in {"/", "-", "*"}:
            candidates.append((index, token))
    return candidates[0] if depth == 0 and len(candidates) == 1 else None


def _canonical_tokens(tokens: tuple[Token, ...]) -> str:
    pieces: list[str] = []
    for token in _strip_outer_parentheses(tokens):
        value = str(token.value)
        if token.type == "OPERAND" and token.subtype == "RANGE":
            value = value.replace("$", "").upper().replace("''", "'")
        elif token.type == "OPERAND" and token.subtype == "NUMBER":
            with suppress(InvalidOperation):
                value = str(Decimal(value).normalize())
        elif token.type != "OPERAND" or token.subtype != "TEXT":
            value = value.upper()
        pieces.append(f"{token.type}:{token.subtype}:{value}")
    return "|".join(pieces)


def _is_zero(tokens: tuple[Token, ...]) -> bool:
    current = _strip_outer_parentheses(tokens)
    if len(current) != 1:
        return False
    token = current[0]
    if token.type != "OPERAND" or token.subtype != "NUMBER":
        return False
    try:
        return Decimal(str(token.value)) == 0
    except InvalidOperation:
        return False


def _degenerate_expression(tokens: tuple[Token, ...] | None) -> DegenerateExpression | None:
    if not tokens:
        return None
    current = _strip_outer_parentheses(tokens)
    candidate = _top_level_operator(current)
    if candidate is None:
        return None
    index, operator = candidate
    left_tokens = current[:index]
    right_tokens = current[index + 1 :]
    if not left_tokens or not right_tokens:
        return None
    left = _canonical_tokens(left_tokens)
    right = _canonical_tokens(right_tokens)
    if operator.value == "/" and left == right:
        kind: Literal["self_division", "self_subtraction", "multiply_by_zero"] = "self_division"
    elif operator.value == "-" and left == right:
        kind = "self_subtraction"
    elif operator.value == "*" and (_is_zero(left_tokens) ^ _is_zero(right_tokens)):
        kind = "multiply_by_zero"
    else:
        return None
    return DegenerateExpression(
        kind=kind,
        operator=str(operator.value),
        left=left,
        right=right,
    )


def parse_formula_ir(formula: str, origin_sheet: str, origin_coordinate: str) -> FormulaIR:
    """Parse safe structural facts from an Excel formula without evaluating it."""

    features = analyze_formula(formula)
    references: list[FormulaReference] = []
    seen: set[tuple[str, int, int, int, int]] = set()
    for raw in features.references:
        reference = _parse_reference(raw, origin_sheet)
        if reference is None:
            continue
        key = (
            reference.sheet.casefold(),
            reference.min_row,
            reference.max_row,
            reference.min_column,
            reference.max_column,
        )
        if key in seen:
            continue
        seen.add(key)
        references.append(reference)
    tokens = _tokenize(formula)
    aggregate = _simple_aggregate(tokens, origin_sheet)
    return FormulaIR(
        formula=formula,
        origin_sheet=origin_sheet,
        origin_coordinate=origin_coordinate,
        features=features,
        references=tuple(references),
        aggregate_references=(aggregate,) if aggregate is not None else (),
        degeneracy=_degenerate_expression(tokens),
    )


def worksheet_by_name(workbook: Workbook, name: str) -> Worksheet | None:
    """Resolve worksheet names case-insensitively, matching Excel semantics."""

    folded = name.casefold()
    return next((sheet for sheet in workbook.worksheets if sheet.title.casefold() == folded), None)


def inspect_reference_target(
    workbook: Workbook,
    reference: FormulaReference,
    *,
    content_index: WorksheetContentIndex | None = None,
) -> ReferenceTarget:
    """Inspect a dependency target through the sparse cell store, never by expanding its area."""

    worksheet = worksheet_by_name(workbook, reference.sheet)
    if worksheet is None:
        return ReferenceTarget(
            reference=reference,
            sheet_exists=False,
            nonblank_cells=(),
            content_bounds=None,
            outside_content=True,
            fully_blank=True,
        )
    index = (
        content_index
        if content_index is not None and content_index.worksheet is worksheet
        else WorksheetContentIndex.from_worksheet(worksheet)
    )
    bounds = index.content_bounds
    if bounds is not None:
        min_row, max_row, min_column, max_column = bounds
        outside = (
            reference.max_row < min_row
            or reference.min_row > max_row
            or reference.max_column < min_column
            or reference.min_column > max_column
        )
    else:
        bounds = None
        outside = True
    hits = [
        cell.coordinate
        for cell in index.iter_nonblank_rectangle(
            min_row=reference.min_row,
            max_row=reference.max_row,
            min_column=reference.min_column,
            max_column=reference.max_column,
            limit=32,
        )
    ]
    return ReferenceTarget(
        reference=reference,
        sheet_exists=True,
        nonblank_cells=tuple(hits[:32]),
        content_bounds=bounds,
        outside_content=outside,
        fully_blank=not hits,
    )


__all__ = [
    "AGGREGATE_FUNCTIONS",
    "AggregateReference",
    "DegenerateExpression",
    "FormulaIR",
    "FormulaReference",
    "ReferenceTarget",
    "WorksheetContentIndex",
    "inspect_reference_target",
    "parse_formula_ir",
    "worksheet_by_name",
    "worksheet_content_index",
]
