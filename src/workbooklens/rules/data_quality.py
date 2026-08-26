"""Conservative, report-only data-quality and workbook-structure rules."""

from __future__ import annotations

import math
import re
import statistics
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import pairwise
from typing import Any, cast

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.constants import MAX_COLUMN, MAX_ROW

from workbooklens.formulas.ir import worksheet_by_name
from workbooklens.layout import column_width, row_height
from workbooklens.models import Confidence, Evidence, Finding, Region, Severity
from workbooklens.rules.base import RuleContext, RuleResult, WorkbookRule
from workbooklens.utils import stable_id

MIN_PROFILE_ROWS = 6
EXCEL_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
SUMMARY_RE = re.compile(
    r"(?<![A-Z0-9_])(?:grand[\s-]+total|subtotal|total|average|summary)(?![A-Z0-9_])"
    r"|(?:合计|总计|小计|汇总|平均)",
    re.IGNORECASE,
)
PLAIN_NUMERIC_RE = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)")
GROUPED_NUMERIC_RE = re.compile(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?")
DIRECT_REFERENCE_RE = re.compile(
    r"^\s*=\s*(?:(?P<sheet>'(?:[^']|'')+'|[^'!=+\-*/(),]+)!)?"
    r"\$?(?P<column>[A-Z]{1,3})\$?(?P<row>\d+)\s*$",
    re.IGNORECASE,
)

IDENTIFIER_TOKENS = {
    "code",
    "id",
    "identifier",
    "key",
    "no",
    "number",
    "sku",
}
IDENTIFIER_MARKERS = (
    "编号",
    "编码",
    "代码",
    "号码",
    "工号",
    "学号",
    "单号",
    "标识",
)
DATE_TOKENS = {"date", "datetime", "day", "joined", "joining", "time", "timestamp"}
DATE_MARKERS = ("日期", "时间", "入职", "出生", "年月日")
PERCENT_TOKENS = {
    "discount",
    "margin",
    "percent",
    "percentage",
    "rate",
    "ratio",
    "tax",
}
PERCENT_MARKERS = ("百分比", "比例", "折扣", "税率", "率")
SIGN_DOMAIN_EXCLUSION_TOKENS = {
    "adjustment",
    "balance",
    "change",
    "delta",
    "growth",
    "margin",
    "net",
    "profit",
    "return",
    "variance",
}
SIGN_DOMAIN_EXCLUSION_MARKERS = ("余额", "利润", "变动", "变化", "调整", "差额", "净额")
STRICT_POSITIVE_TOKENS = {"age", "payroll", "price", "salary", "wage"}
STRICT_POSITIVE_MARKERS = ("年龄", "单价", "价格", "工资", "薪资")
NONNEGATIVE_TOKENS = {"inventory", "outbound", "inbound", "qty", "quantity", "stock", "units"}
NONNEGATIVE_MARKERS = ("数量", "库存", "入库", "出库")
SEMANTIC_GROUPS: dict[str, tuple[str, ...]] = {
    "budget": ("budget", "budgeted", "预算"),
    "spent": (
        "expenditure",
        "expense",
        "spend",
        "spending",
        "spent",
        "已支出",
        "实际支出",
        "支出",
    ),
    "amount": (
        "amount",
        "gross",
        "income",
        "revenue",
        "sales",
        "turnover",
        "金额",
        "营收",
        "收入",
        "销售额",
    ),
    "quantity": ("count", "qty", "quantity", "units", "volume", "件数", "数量"),
    "cost": ("cost", "price", "价格", "单价", "成本"),
    "rate": ("discount", "margin", "percent", "rate", "ratio", "折扣", "比例", "率"),
    "salary": ("payroll", "salary", "wage", "工资", "薪资"),
}
TREND_RE = re.compile(r"(?<![a-z0-9])trends?(?![a-z0-9])|趋势", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class _TableProfile:
    region: Region
    header_row: int
    headers: dict[int, str]
    body_rows: tuple[int, ...]


@dataclass(frozen=True, slots=True)
class _ConditionalRuleEntry:
    target: CellRange
    rule: Any
    priority: int
    sequence: int


def _confidence(value: float) -> Confidence:
    return Confidence(max(0.0, min(1.0, value)))


def _make_finding(
    *,
    context: RuleContext,
    rule_id: str,
    title: str,
    explanation: str,
    severity: Severity,
    confidence: float,
    worksheet: Worksheet,
    location: str,
    evidence: Evidence,
    expected: str,
    suggested_action: str,
    discriminator: Any = None,
) -> Finding:
    return Finding(
        id=stable_id("finding", rule_id, worksheet.title, location, discriminator),
        content_fingerprint=stable_id(
            "finding-content", evidence.model_dump(mode="json"), length=24
        ),
        rule_id=rule_id,
        title=title,
        explanation=explanation,
        severity=severity,
        confidence=_confidence(confidence),
        workbook=context.path.name,
        sheet=worksheet.title,
        location=location,
        evidence=evidence,
        expected=expected,
        suggested_action=suggested_action,
    )


def _value_at(worksheet: Worksheet, row: int, column: int) -> Any:
    cell = worksheet._cells.get((row, column))
    return cell.value if isinstance(cell, (Cell, MergedCell)) else None


def _cell_at(worksheet: Worksheet, row: int, column: int) -> Cell | MergedCell | None:
    cell = worksheet._cells.get((row, column))
    return cell if isinstance(cell, (Cell, MergedCell)) else None


def _is_nonblank(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def _is_formula(cell: Cell | MergedCell | None) -> bool:
    return isinstance(cell, Cell) and cell.data_type == "f"


def _is_error(cell: Cell | MergedCell | None) -> bool:
    return bool(
        isinstance(cell, Cell)
        and (cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in EXCEL_ERRORS))
    )


def _normalize_text(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(unicodedata.normalize("NFKC", value).strip().casefold().split())


def _tokens(value: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]+", value))


def _looks_like_identifier_header(value: Any) -> bool:
    normalized = _normalize_text(value)
    return bool(
        normalized
        and (
            _tokens(normalized) & IDENTIFIER_TOKENS
            or any(marker in normalized for marker in IDENTIFIER_MARKERS)
        )
    )


def _looks_like_date_header(value: Any) -> bool:
    normalized = _normalize_text(value)
    return bool(
        normalized
        and (
            _tokens(normalized) & DATE_TOKENS
            or any(marker in normalized for marker in DATE_MARKERS)
        )
    )


def _looks_like_percentage_header(value: Any) -> bool:
    normalized = _normalize_text(value)
    return bool(
        normalized
        and (
            _tokens(normalized) & PERCENT_TOKENS
            or "%" in normalized
            or any(marker in normalized for marker in PERCENT_MARKERS)
        )
    )


def _sign_domain(value: Any) -> str | None:
    normalized = _normalize_text(value)
    if not normalized:
        return None
    words = _tokens(normalized)
    if words & SIGN_DOMAIN_EXCLUSION_TOKENS or any(
        marker in normalized for marker in SIGN_DOMAIN_EXCLUSION_MARKERS
    ):
        return None
    if words & STRICT_POSITIVE_TOKENS or any(
        marker in normalized for marker in STRICT_POSITIVE_MARKERS
    ):
        return "positive"
    if (
        words & NONNEGATIVE_TOKENS
        or normalized in {"in", "out"}
        or any(marker in normalized for marker in NONNEGATIVE_MARKERS)
    ):
        return "nonnegative"
    return None


def _satisfies_sign_domain(value: float, domain: str) -> bool:
    return value > 0.0 if domain == "positive" else value >= 0.0


def _row_is_summary(worksheet: Worksheet, region: Region, row: int) -> bool:
    for column in range(region.min_column, region.max_column + 1):
        value = _value_at(worksheet, row, column)
        if isinstance(value, str) and SUMMARY_RE.search(value):
            return True
    return False


def _table_profiles(context: RuleContext, worksheet: Worksheet) -> tuple[_TableProfile, ...]:
    if worksheet.sheet_state != "visible":
        return ()
    profiles: list[_TableProfile] = []
    for region in context.data_regions.get(worksheet.title, []):
        width = region.max_column - region.min_column + 1
        if width < 2 or region.max_row - region.min_row < MIN_PROFILE_ROWS:
            continue
        headers = {
            column: _normalize_text(_value_at(worksheet, region.min_row, column))
            for column in range(region.min_column, region.max_column + 1)
        }
        if sum(bool(value) for value in headers.values()) < 2:
            continue
        minimum_populated = max(2, math.ceil(width * 0.35))
        body_rows = tuple(
            row
            for row in range(region.min_row + 1, region.max_row + 1)
            if not _row_is_summary(worksheet, region, row)
            and sum(
                _is_nonblank(_value_at(worksheet, row, column))
                for column in range(region.min_column, region.max_column + 1)
            )
            >= minimum_populated
        )
        if len(body_rows) < MIN_PROFILE_ROWS:
            continue
        profiles.append(
            _TableProfile(
                region=region,
                header_row=region.min_row,
                headers=headers,
                body_rows=body_rows,
            )
        )
    return tuple(profiles)


def _configured_key_columns(context: RuleContext, worksheet: Worksheet) -> set[int]:
    columns: set[int] = set()
    keys = context.config.get("keys", [])
    if not isinstance(keys, list):
        return columns
    for item in keys:
        if not isinstance(item, dict):
            continue
        if str(item.get("sheet", "")).casefold() != worksheet.title.casefold():
            continue
        range_text = item.get("range")
        if not isinstance(range_text, str):
            continue
        try:
            key_range = CellRange(range_text.replace("$", ""))
        except ValueError:
            continue
        columns.update(range(key_range.min_col, key_range.max_col + 1))
    return columns


def _identifier_key(value: Any) -> tuple[str, str] | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return "number", str(value)
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return "number", str(int(value))
    if not isinstance(value, str):
        return None
    normalized = _normalize_text(value)
    if not normalized:
        return None
    if re.fullmatch(r"[+-]?\d+", normalized) and not re.fullmatch(r"[+-]?0\d+", normalized):
        return "number", str(int(normalized))
    return "text", normalized


def _in_merged_range(worksheet: Worksheet, coordinate: str) -> bool:
    return any(coordinate in merged for merged in worksheet.merged_cells.ranges)


def _plain_numeric_text(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    stripped = value.strip()
    if not stripped or re.search(r"[()\-/]", stripped):
        return False
    if (
        PLAIN_NUMERIC_RE.fullmatch(stripped) is None
        and GROUPED_NUMERIC_RE.fullmatch(stripped) is None
    ):
        return False
    try:
        return Decimal(stripped.replace(",", "")).is_finite()
    except InvalidOperation:
        return False


def _numeric_value(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    result = float(value)
    return result if math.isfinite(result) else None


def _conditional_rule_entries(worksheet: Worksheet) -> tuple[_ConditionalRuleEntry, ...]:
    entries: list[_ConditionalRuleEntry] = []
    sequence = 0
    for conditional_formatting in worksheet.conditional_formatting:
        for target in conditional_formatting.sqref.ranges:
            target_range = CellRange(str(target))
            for rule in conditional_formatting.rules:
                raw_priority = getattr(rule, "priority", None)
                priority = (
                    raw_priority
                    if isinstance(raw_priority, int)
                    and not isinstance(raw_priority, bool)
                    and raw_priority > 0
                    else MAX_ROW + 1
                )
                entries.append(
                    _ConditionalRuleEntry(
                        target=target_range,
                        rule=rule,
                        priority=priority,
                        sequence=sequence,
                    )
                )
                sequence += 1
    return tuple(sorted(entries, key=lambda entry: (entry.priority, entry.sequence)))


def _literal_cell_is_operands(rule: Any) -> tuple[float, ...] | None:
    if getattr(rule, "type", None) != "cellIs":
        return None
    operator = getattr(rule, "operator", None)
    arity = 2 if operator in {"between", "notBetween"} else 1
    if operator not in {
        "between",
        "equal",
        "greaterThan",
        "greaterThanOrEqual",
        "lessThan",
        "lessThanOrEqual",
        "notBetween",
        "notEqual",
    }:
        return None
    formulas = getattr(rule, "formula", None)
    if (
        not isinstance(formulas, list)
        or len(formulas) != arity
        or not all(isinstance(formula, str) for formula in formulas)
    ):
        return None
    operands: list[float] = []
    for formula in formulas:
        value = formula.strip().removeprefix("=").strip()
        try:
            operand = float(Decimal(value))
        except (InvalidOperation, ValueError, OverflowError):
            return None
        if not math.isfinite(operand):
            return None
        operands.append(operand)
    return tuple(operands)


def _literal_threshold(rule: Any) -> float | None:
    if getattr(rule, "operator", None) not in {"lessThan", "lessThanOrEqual"}:
        return None
    operands = _literal_cell_is_operands(rule)
    if operands is None or operands[0] > 0.0:
        return None
    return operands[0]


def _matches_literal_cell_is(value: float, rule: Any) -> bool | None:
    operands = _literal_cell_is_operands(rule)
    if operands is None:
        return None
    operator = getattr(rule, "operator", None)
    first = operands[0]
    if operator == "equal":
        return value == first
    if operator == "notEqual":
        return value != first
    if operator == "lessThan":
        return value < first
    if operator == "lessThanOrEqual":
        return value <= first
    if operator == "greaterThan":
        return value > first
    if operator == "greaterThanOrEqual":
        return value >= first
    second = operands[1]
    if operator == "between":
        return first <= value <= second
    if operator == "notBetween":
        return not first <= value <= second
    return None


def _success_green_fill(rule: Any) -> str | None:
    differential = getattr(rule, "dxf", None)
    fill = getattr(differential, "fill", None)
    if fill is None or getattr(fill, "fill_type", None) != "solid":
        return None
    colors = (getattr(fill, "fgColor", None), getattr(fill, "start_color", None))
    for color in colors:
        rgb = getattr(color, "rgb", None)
        if not isinstance(rgb, str) or len(rgb) not in {6, 8}:
            continue
        value = rgb[-6:]
        try:
            red, green, blue = (int(value[index : index + 2], 16) for index in (0, 2, 4))
        except ValueError:
            continue
        if green >= 128 and green >= red + 48 and green >= blue + 24:
            return value.upper()
    return None


def _semantic_groups(value: Any) -> set[str]:
    normalized = _normalize_text(value)
    if not normalized:
        return set()
    words = _tokens(normalized)
    groups: set[str] = set()
    for group, markers in SEMANTIC_GROUPS.items():
        if any(
            marker in normalized if not marker.isascii() else marker in words for marker in markers
        ):
            groups.add(group)
    if groups & {"budget", "spent"}:
        groups.difference_update({"amount", "cost"})
    return groups


class InferredDuplicateIdentifierRule(WorkbookRule):
    rule_id = "WL022_INFERRED_DUPLICATE_IDENTIFIER"
    title = "Duplicate value in inferred identifier column"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        # Local import avoids an import-time cycle: relational_semantics reuses the
        # bounded table-profile helpers in this module.
        from workbooklens.rules.relational_semantics import inferred_foreign_key_columns

        foreign_key_columns = inferred_foreign_key_columns(context)
        for worksheet in context.workbook.worksheets:
            configured_columns = _configured_key_columns(context, worksheet)
            for profile in _table_profiles(context, worksheet):
                for column, header in profile.headers.items():
                    if (
                        column in configured_columns
                        or (worksheet.title, profile.header_row, column) in foreign_key_columns
                        or not _looks_like_identifier_header(header)
                    ):
                        continue
                    groups: dict[tuple[str, str], list[Cell]] = {}
                    for row in profile.body_rows:
                        cell = _cell_at(worksheet, row, column)
                        if not isinstance(cell, Cell) or _is_formula(cell) or _is_error(cell):
                            continue
                        key = _identifier_key(cell.value)
                        if key is not None:
                            groups.setdefault(key, []).append(cell)
                    populated = sum(len(cells) for cells in groups.values())
                    if populated < MIN_PROFILE_ROWS or populated / len(profile.body_rows) < 0.75:
                        continue
                    if len(groups) / populated < 0.7:
                        continue
                    for key, cells in sorted(groups.items(), key=lambda item: item[0]):
                        if len(cells) < 2:
                            continue
                        coordinates = [cell.coordinate for cell in cells]
                        evidence = Evidence(
                            summary=f"Identifier value appears {len(cells)} times in an inferred key column",
                            observed={"normalized_value": key[1], "cells": coordinates},
                            expected={"occurrences": 1},
                            peers=coordinates,
                            details={"header": header, "populated_rows": populated},
                        )
                        result.findings.append(
                            _make_finding(
                                context=context,
                                rule_id=self.rule_id,
                                title=self.title,
                                explanation=(
                                    "A column with an explicit identifier-like header and high uniqueness "
                                    "contains a repeated normalized value."
                                ),
                                severity=Severity.ERROR,
                                confidence=0.98,
                                worksheet=worksheet,
                                location=",".join(coordinates),
                                evidence=evidence,
                                expected="Inferred identifier columns contain one nonblank value per record.",
                                suggested_action=(
                                    "Review the repeated records or configure the intended key explicitly; "
                                    "no record is deleted automatically."
                                ),
                                discriminator=key,
                            )
                        )
        return result


class MissingInferredIdentifierRule(WorkbookRule):
    rule_id = "WL023_MISSING_INFERRED_IDENTIFIER"
    title = "Missing value in inferred identifier column"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            configured_columns = _configured_key_columns(context, worksheet)
            for profile in _table_profiles(context, worksheet):
                for column, header in profile.headers.items():
                    if column in configured_columns or not _looks_like_identifier_header(header):
                        continue
                    keys = [
                        _identifier_key(_value_at(worksheet, row, column))
                        for row in profile.body_rows
                    ]
                    populated = [key for key in keys if key is not None]
                    if len(populated) < MIN_PROFILE_ROWS:
                        continue
                    coverage = len(populated) / len(profile.body_rows)
                    if coverage < 0.75 or len(set(populated)) / len(populated) < 0.7:
                        continue
                    for row, key in zip(profile.body_rows, keys, strict=True):
                        if key is not None:
                            continue
                        coordinate = f"{get_column_letter(column)}{row}"
                        if _in_merged_range(worksheet, coordinate):
                            continue
                        evidence = Evidence(
                            summary="A populated record is missing its inferred identifier",
                            observed=None,
                            expected="nonblank identifier",
                            details={
                                "header": header,
                                "peer_coverage": round(coverage, 4),
                                "record_row": row,
                            },
                        )
                        result.findings.append(
                            _make_finding(
                                context=context,
                                rule_id=self.rule_id,
                                title=self.title,
                                explanation=(
                                    "A populated table row is blank in a column whose header and peer "
                                    "uniqueness strongly indicate record-identifier semantics."
                                ),
                                severity=Severity.ERROR,
                                confidence=0.96,
                                worksheet=worksheet,
                                location=coordinate,
                                evidence=evidence,
                                expected="Inferred identifier columns contain one nonblank value per record.",
                                suggested_action=(
                                    "Recover the identifier from an authoritative source or mark the record "
                                    "for review; WorkbookLens does not invent identifiers."
                                ),
                            )
                        )
        return result


class MixedNumericStorageRule(WorkbookRule):
    rule_id = "WL024_MIXED_NUMERIC_STORAGE"
    title = "Non-numeric text in a numeric-dominant column"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            for profile in _table_profiles(context, worksheet):
                for column, header in profile.headers.items():
                    if _looks_like_identifier_header(header) or _looks_like_date_header(header):
                        continue
                    cells = [
                        cell
                        for row in profile.body_rows
                        if (cell := _cell_at(worksheet, row, column)) is not None
                        and _is_nonblank(cell.value)
                        and not _is_formula(cell)
                        and not _is_error(cell)
                        and not isinstance(cell.value, (date, datetime))
                    ]
                    numeric_cells = [
                        cell for cell in cells if _numeric_value(cell.value) is not None
                    ]
                    if len(numeric_cells) < MIN_PROFILE_ROWS or not cells:
                        continue
                    ratio = len(numeric_cells) / len(cells)
                    if ratio < 0.75:
                        continue
                    anomalies = [
                        cell
                        for cell in cells
                        if isinstance(cell.value, str) and not _plain_numeric_text(cell.value)
                    ]
                    if not anomalies or len(anomalies) > max(3, math.ceil(len(cells) * 0.2)):
                        continue
                    for cell in anomalies:
                        evidence = Evidence(
                            summary=f"{len(numeric_cells)} of {len(cells)} populated peers use numeric storage",
                            observed={"value": cell.value, "type": "text"},
                            expected={"type": "numeric"},
                            peers=[peer.coordinate for peer in numeric_cells[:12]],
                            details={"header": header, "numeric_ratio": round(ratio, 4)},
                        )
                        result.findings.append(
                            _make_finding(
                                context=context,
                                rule_id=self.rule_id,
                                title=self.title,
                                explanation=(
                                    "A rare text value appears in a column whose populated peers are "
                                    "overwhelmingly numeric."
                                ),
                                severity=Severity.WARNING,
                                confidence=0.96 if ratio >= 0.85 else 0.88,
                                worksheet=worksheet,
                                location=cell.coordinate,
                                evidence=evidence,
                                expected="Numeric-dominant columns use numeric storage or a documented exception.",
                                suggested_action=(
                                    "Review the source text and convert it only when its intended numeric value "
                                    "is unambiguous; no value is guessed automatically."
                                ),
                            )
                        )
        return result


def _robust_scores(values: list[float]) -> list[float]:
    median = statistics.median(values)
    deviations = [abs(value - median) for value in values]
    mad = statistics.median(deviations)
    if mad <= 1e-12:
        return [0.0 for _ in values]
    return [0.67448975 * abs(value - median) / mad for value in values]


class RobustNumericOutlierRule(WorkbookRule):
    rule_id = "WL025_ROBUST_NUMERIC_OUTLIER"
    title = "Extreme numeric outlier"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            for profile in _table_profiles(context, worksheet):
                for column, header in profile.headers.items():
                    if (
                        _looks_like_identifier_header(header)
                        or _looks_like_date_header(header)
                        or _looks_like_percentage_header(header)
                    ):
                        continue
                    pairs = [
                        (cell, numeric)
                        for row in profile.body_rows
                        if isinstance((cell := _cell_at(worksheet, row, column)), Cell)
                        and not _is_formula(cell)
                        and not _is_error(cell)
                        and (numeric := _numeric_value(cell.value)) is not None
                    ]
                    if len(pairs) < 8:
                        continue
                    values = [value for _, value in pairs]
                    scores = _robust_scores(values)
                    for (cell, value), score in zip(pairs, scores, strict=True):
                        if score < 8.0:
                            continue
                        evidence = Evidence(
                            summary=f"Robust modified z-score is {score:.1f} across {len(values)} numeric peers",
                            observed=value,
                            expected={
                                "median": statistics.median(values),
                                "maximum_review_score": 8.0,
                            },
                            peers=[
                                peer.coordinate
                                for peer, _ in pairs
                                if peer.coordinate != cell.coordinate
                            ][:12],
                            details={"header": header, "method": "median_absolute_deviation"},
                        )
                        result.findings.append(
                            _make_finding(
                                context=context,
                                rule_id=self.rule_id,
                                title=self.title,
                                explanation=(
                                    "A numeric literal is extremely far from the column median under a "
                                    "median-absolute-deviation check."
                                ),
                                severity=Severity.WARNING,
                                confidence=min(0.99, 0.88 + score / 200.0),
                                worksheet=worksheet,
                                location=cell.coordinate,
                                evidence=evidence,
                                expected=(
                                    "Values remain within the robust peer distribution unless the exception "
                                    "is documented."
                                ),
                                suggested_action=(
                                    "Verify the source value, unit, and decimal placement; statistical "
                                    "outliers are reported but never rewritten automatically."
                                ),
                            )
                        )
        return result


class SignConstrainedMeasureRule(WorkbookRule):
    rule_id = "WL035_SIGN_CONSTRAINED_MEASURE"
    title = "Invalid sign in a constrained measure"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            for profile in _table_profiles(context, worksheet):
                for column, header in profile.headers.items():
                    domain = _sign_domain(header)
                    if domain is None:
                        continue
                    pairs = [
                        (cell, numeric)
                        for row in profile.body_rows
                        if isinstance((cell := _cell_at(worksheet, row, column)), Cell)
                        and not _is_formula(cell)
                        and not _is_error(cell)
                        and (numeric := _numeric_value(cell.value)) is not None
                    ]
                    if len(pairs) < MIN_PROFILE_ROWS:
                        continue
                    satisfies = sum(_satisfies_sign_domain(value, domain) for _, value in pairs)
                    ratio = satisfies / len(pairs)
                    if ratio < 0.85:
                        continue
                    violations = [
                        (cell, value)
                        for cell, value in pairs
                        if not _satisfies_sign_domain(value, domain)
                    ]
                    for cell, value in violations:
                        zero_advisory = domain == "positive" and value == 0.0
                        evidence = Evidence(
                            summary=(
                                "Numeric peers overwhelmingly satisfy the inferred positive domain"
                                if domain == "positive"
                                else "Numeric peers overwhelmingly satisfy the inferred nonnegative domain"
                            ),
                            observed=value,
                            expected={"domain": domain},
                            peers=[
                                peer.coordinate
                                for peer, peer_value in pairs
                                if _satisfies_sign_domain(peer_value, domain)
                            ][:12],
                            details={
                                "header": header,
                                "domain": domain,
                                "candidate_kind": "zero_advisory" if zero_advisory else "negative",
                                "satisfying_peers": satisfies,
                                "numeric_peers": len(pairs),
                            },
                        )
                        result.findings.append(
                            _make_finding(
                                context=context,
                                rule_id=self.rule_id,
                                title=self.title,
                                explanation=(
                                    "A zero literal is unusual in a column whose header and peers suggest a "
                                    "positive measure, but valid business exceptions may exist."
                                    if zero_advisory
                                    else "A negative literal conflicts with a strongly inferred positive or "
                                    "nonnegative measure domain."
                                ),
                                severity=Severity.WARNING if zero_advisory else Severity.ERROR,
                                confidence=(
                                    0.82 if zero_advisory else 0.98 if ratio >= 0.95 else 0.92
                                ),
                                worksheet=worksheet,
                                location=cell.coordinate,
                                evidence=evidence,
                                expected=(
                                    "Review whether zero is valid for this positive-like measure."
                                    if zero_advisory
                                    else "Strongly constrained measures respect their inferred positive or "
                                    "nonnegative domain."
                                ),
                                suggested_action=(
                                    "Verify the source value, sign, and unit; WorkbookLens reports the "
                                    "violation but does not replace semantic values automatically."
                                ),
                            )
                        )
        return result


class ConditionalFormatSemanticConflictRule(WorkbookRule):
    rule_id = "WL051_CONDITIONAL_FORMAT_SEMANTIC_CONFLICT"
    title = "Conditional format conflicts with constrained-value semantics"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        seen_semantics: set[tuple[Any, ...]] = set()
        for worksheet in context.workbook.worksheets:
            conditional_rules = _conditional_rule_entries(worksheet)
            for profile in _table_profiles(context, worksheet):
                for column, header in profile.headers.items():
                    domain = _sign_domain(header)
                    if domain is None:
                        continue
                    pairs = [
                        (cell, numeric)
                        for row in profile.body_rows
                        if isinstance((cell := _cell_at(worksheet, row, column)), Cell)
                        and not _is_formula(cell)
                        and not _is_error(cell)
                        and (numeric := _numeric_value(cell.value)) is not None
                    ]
                    if len(pairs) < MIN_PROFILE_ROWS:
                        continue
                    satisfying = sum(_satisfies_sign_domain(value, domain) for _, value in pairs)
                    ratio = satisfying / len(pairs)
                    if ratio < 0.85:
                        continue
                    violating_pairs = [
                        (cell, value)
                        for cell, value in pairs
                        if not _satisfies_sign_domain(value, domain)
                    ]
                    if not violating_pairs:
                        continue
                    for candidate_index, entry in enumerate(conditional_rules):
                        target_range = entry.target
                        if not (
                            target_range.min_col <= column <= target_range.max_col
                            and target_range.min_row <= profile.body_rows[-1]
                            and target_range.max_row >= profile.body_rows[0]
                        ):
                            continue
                        threshold = _literal_threshold(entry.rule)
                        green = _success_green_fill(entry.rule)
                        if threshold is None or green is None:
                            continue
                        prior_rules = conditional_rules[:candidate_index]
                        # Any matching or unknown higher-priority rule can override part of the
                        # lower-priority fill, so the effective success-green format is not proven.
                        matched_pairs: list[tuple[Cell, float]] = []
                        for cell, value in violating_pairs:
                            if cell.coordinate not in target_range:
                                continue
                            if _matches_literal_cell_is(value, entry.rule) is not True:
                                continue
                            prior_outcomes = [
                                _matches_literal_cell_is(value, prior.rule)
                                for prior in prior_rules
                                if cell.coordinate in prior.target
                            ]
                            if not any(outcome is not False for outcome in prior_outcomes):
                                matched_pairs.append((cell, value))
                        if not matched_pairs:
                            continue
                        operator = str(getattr(entry.rule, "operator", ""))
                        semantic_key = (
                            worksheet.title,
                            column,
                            str(target_range),
                            operator,
                            threshold,
                            green,
                        )
                        if semantic_key in seen_semantics:
                            continue
                        seen_semantics.add(semantic_key)
                        violating_cells = [cell.coordinate for cell, _ in matched_pairs]
                        negative_cells = [
                            cell.coordinate for cell, value in matched_pairs if value < 0.0
                        ]
                        zero_cells = [
                            cell.coordinate for cell, value in matched_pairs if value == 0.0
                        ]
                        evidence = Evidence(
                            summary=(
                                "Conditional formatting uses a success-green fill for "
                                "sign-violating constrained values"
                            ),
                            observed={
                                "target": str(target_range),
                                "operator": operator,
                                "threshold": threshold,
                                "fill_rgb": green,
                                "violating_cells": violating_cells,
                                "negative_cells": negative_cells,
                                "zero_cells": zero_cells,
                            },
                            expected={"warning_or_neutral_format_for_sign_violations": True},
                            peers=[
                                cell.coordinate
                                for cell, value in pairs
                                if _satisfies_sign_domain(value, domain)
                            ][:12],
                            details={
                                "header": header,
                                "domain": domain,
                                "rule_priority": entry.priority,
                                "satisfying_peers": satisfying,
                                "numeric_peers": len(pairs),
                                "automatic_format_change": False,
                            },
                        )
                        result.findings.append(
                            _make_finding(
                                context=context,
                                rule_id=self.rule_id,
                                title=self.title,
                                explanation=(
                                    "A rule that selects sign-violating values in a strongly positive "
                                    "or nonnegative field applies an unmistakably green success fill."
                                ),
                                severity=Severity.WARNING,
                                confidence=0.94 if ratio >= 0.95 else 0.9,
                                worksheet=worksheet,
                                location=str(target_range),
                                evidence=evidence,
                                expected=(
                                    "Conditional formatting for sign violations uses warning or neutral "
                                    "semantics unless the exception is documented."
                                ),
                                suggested_action=(
                                    "Review the conditional-format rule and its intended business meaning; "
                                    "WorkbookLens does not rewrite visual semantics automatically."
                                ),
                                discriminator=(column, operator, threshold, green),
                            )
                        )
        return result


class PercentageScaleOutlierRule(WorkbookRule):
    rule_id = "WL026_PERCENTAGE_SCALE_OUTLIER"
    title = "Percentage scale or range anomaly"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            for profile in _table_profiles(context, worksheet):
                for column, header in profile.headers.items():
                    pairs = [
                        (cell, numeric)
                        for row in profile.body_rows
                        if isinstance((cell := _cell_at(worksheet, row, column)), Cell)
                        and not _is_formula(cell)
                        and not _is_error(cell)
                        and (numeric := _numeric_value(cell.value)) is not None
                    ]
                    if len(pairs) < MIN_PROFILE_ROWS:
                        continue
                    formatted_ratio = sum("%" in cell.number_format for cell, _ in pairs) / len(
                        pairs
                    )
                    if not _looks_like_percentage_header(header) and formatted_ratio < 0.75:
                        continue
                    fractional = sum(0.0 <= value <= 1.0 for _, value in pairs)
                    ratio = fractional / len(pairs)
                    if ratio < 0.75:
                        continue
                    for cell, value in pairs:
                        if 0.0 <= value <= 1.0:
                            continue
                        anomaly = (
                            "whole_percent_scale"
                            if 2.0 <= value <= 100.0
                            else "outside_fraction_range"
                        )
                        evidence = Evidence(
                            summary=(
                                f"{fractional} of {len(pairs)} numeric peers use fractional percentage storage"
                            ),
                            observed=value,
                            expected={"minimum": 0.0, "maximum": 1.0},
                            peers=[
                                peer.coordinate
                                for peer, peer_value in pairs
                                if 0 <= peer_value <= 1
                            ][:12],
                            details={
                                "header": header,
                                "anomaly": anomaly,
                                "fractional_ratio": round(ratio, 4),
                            },
                        )
                        result.findings.append(
                            _make_finding(
                                context=context,
                                rule_id=self.rule_id,
                                title=self.title,
                                explanation=(
                                    "A percentage-like column is dominated by fractional values, but one "
                                    "literal uses a different scale or lies outside the peer range."
                                ),
                                severity=Severity.ERROR,
                                confidence=0.98 if ratio >= 0.85 else 0.92,
                                worksheet=worksheet,
                                location=cell.coordinate,
                                evidence=evidence,
                                expected="Percentage-like inputs use one documented scale and valid range.",
                                suggested_action=(
                                    "Confirm whether the literal should be divided by 100 or rejected; "
                                    "WorkbookLens does not rewrite semantic values automatically."
                                ),
                            )
                        )
        return result


DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%d-%m-%Y",
    "%m-%d-%Y",
    "%b %d, %Y",
    "%B %d, %Y",
    "%Y年%m月%d日",
)


def _parse_date_text(value: str) -> bool:
    stripped = value.strip()
    for date_format in DATE_FORMATS:
        try:
            datetime.strptime(stripped, date_format)
        except ValueError:
            continue
        return True
    return False


class DateStorageAnomalyRule(WorkbookRule):
    rule_id = "WL027_DATE_STORAGE_ANOMALY"
    title = "Invalid or inconsistent date storage"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            for profile in _table_profiles(context, worksheet):
                for column, header in profile.headers.items():
                    cells = [
                        cell
                        for row in profile.body_rows
                        if isinstance((cell := _cell_at(worksheet, row, column)), Cell)
                        and _is_nonblank(cell.value)
                        and not _is_formula(cell)
                        and not _is_error(cell)
                    ]
                    date_cells = [
                        cell for cell in cells if isinstance(cell.value, (date, datetime))
                    ]
                    if len(date_cells) < 5:
                        continue
                    if not _looks_like_date_header(header) and len(date_cells) / len(cells) < 0.75:
                        continue
                    for cell in cells:
                        value = cell.value
                        if isinstance(value, (date, datetime)):
                            continue
                        anomaly: str | None = None
                        summary: str | None = None
                        severity = Severity.WARNING
                        if isinstance(value, str):
                            if _parse_date_text(value):
                                anomaly = "text_date"
                                summary = "Date-like text uses a different storage representation from peer dates"
                            else:
                                anomaly = "unparseable_date_text"
                                summary = "Text in a date column cannot be parsed by supported unambiguous date formats"
                                severity = Severity.ERROR
                        else:
                            numeric = _numeric_value(value)
                            if numeric is not None and 1 <= numeric <= 2_958_465:
                                anomaly = "raw_excel_serial"
                                summary = "A numeric Excel serial appears among typed date cells"
                        if anomaly is None or summary is None:
                            continue
                        evidence = Evidence(
                            summary=summary,
                            observed=value,
                            expected={"type": "date"},
                            peers=[peer.coordinate for peer in date_cells[:12]],
                            details={
                                "header": header,
                                "anomaly": anomaly,
                                "typed_date_peers": len(date_cells),
                            },
                        )
                        result.findings.append(
                            _make_finding(
                                context=context,
                                rule_id=self.rule_id,
                                title=self.title,
                                explanation=(
                                    "A date-dominant column contains an invalid value or a storage "
                                    "representation inconsistent with typed date peers."
                                ),
                                severity=severity,
                                confidence=0.98 if anomaly == "unparseable_date_text" else 0.94,
                                worksheet=worksheet,
                                location=cell.coordinate,
                                evidence=evidence,
                                expected="Date columns use valid typed dates with one consistent storage convention.",
                                suggested_action=(
                                    "Review the source date and locale, then convert explicitly; WorkbookLens "
                                    "does not guess ambiguous dates."
                                ),
                            )
                        )
        return result


def _range_intersection_area(left: CellRange, right: Region) -> int:
    min_row = max(left.min_row, right.min_row)
    max_row = min(left.max_row, right.max_row)
    min_col = max(left.min_col, right.min_column)
    max_col = min(left.max_col, right.max_column)
    if min_row > max_row or min_col > max_col:
        return 0
    return (max_row - min_row + 1) * (max_col - min_col + 1)


class AutoFilterCoverageRule(WorkbookRule):
    rule_id = "WL030_AUTOFILTER_COVERAGE"
    title = "AutoFilter excludes part of an inferred table"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            if worksheet.sheet_state != "visible" or not worksheet.auto_filter.ref:
                continue
            try:
                filter_range = CellRange(str(worksheet.auto_filter.ref).replace("$", ""))
            except ValueError:
                continue
            profiles = _table_profiles(context, worksheet)
            if not profiles:
                continue
            profile = max(
                profiles,
                key=lambda item: _range_intersection_area(filter_range, item.region),
            )
            if _range_intersection_area(filter_range, profile.region) == 0:
                continue
            data_columns = [column for column, header in profile.headers.items() if header]
            if not data_columns:
                continue
            min_column = min(data_columns)
            max_column = max(data_columns)
            max_row = max(profile.body_rows)
            missing_columns = [
                column
                for column in range(min_column, max_column + 1)
                if not filter_range.min_col <= column <= filter_range.max_col
                and sum(
                    _is_nonblank(_value_at(worksheet, row, column)) for row in profile.body_rows
                )
                / len(profile.body_rows)
                >= 0.6
            ]
            missing_rows = [
                row
                for row in profile.body_rows
                if not filter_range.min_row <= row <= filter_range.max_row
                and sum(
                    _is_nonblank(_value_at(worksheet, row, column))
                    for column in range(min_column, max_column + 1)
                )
                / (max_column - min_column + 1)
                >= 0.6
            ]
            header_mismatch = filter_range.min_row != profile.header_row
            if not missing_columns and not missing_rows and not header_mismatch:
                continue
            target = (
                f"{get_column_letter(min_column)}{profile.header_row}:"
                f"{get_column_letter(max_column)}{max_row}"
            )
            evidence = Evidence(
                summary=(
                    f"AutoFilter excludes {len(missing_rows)} populated rows and "
                    f"{len(missing_columns)} populated columns"
                ),
                observed=str(worksheet.auto_filter.ref),
                expected=target,
                details={
                    "header_mismatch": header_mismatch,
                    "missing_rows": missing_rows[:25],
                    "missing_columns": [get_column_letter(column) for column in missing_columns],
                },
            )
            result.findings.append(
                _make_finding(
                    context=context,
                    rule_id=self.rule_id,
                    title=self.title,
                    explanation=(
                        "The worksheet AutoFilter does not cover the header and dense populated extent "
                        "of the intersecting inferred table."
                    ),
                    severity=Severity.WARNING,
                    confidence=0.97,
                    worksheet=worksheet,
                    location=str(worksheet.auto_filter.ref),
                    evidence=evidence,
                    expected="AutoFilter ranges cover the intended table header and all populated records.",
                    suggested_action=(
                        "Review the intended table boundary and reset the filter range manually; no "
                        "filter definition is changed automatically."
                    ),
                )
            )
        return result


def _reference_formula(source: Any) -> str | None:
    if source is None:
        return None
    for attribute in ("numRef", "strRef", "multiLvlStrRef"):
        reference = getattr(source, attribute, None)
        formula = getattr(reference, "f", None)
        if isinstance(formula, str) and formula.strip():
            return formula.strip()
    return None


def _resolve_reference(
    context: RuleContext,
    current_sheet: Worksheet,
    formula: str,
) -> tuple[Worksheet, CellRange] | None:
    text = formula.strip().lstrip("=")
    if "[" in text or "]" in text:
        return None
    if "!" in text:
        qualifier, body = text.rsplit("!", 1)
        qualifier = qualifier.strip()
        if qualifier.startswith("'") and qualifier.endswith("'"):
            qualifier = qualifier[1:-1].replace("''", "'")
        if ":" in qualifier:
            return None
        resolved_worksheet = worksheet_by_name(context.workbook, qualifier)
        if resolved_worksheet is None:
            return None
        worksheet = resolved_worksheet
    else:
        worksheet = current_sheet
        body = text
    if "," in body:
        return None
    try:
        cell_range = CellRange(body.replace("$", ""))
    except ValueError:
        return None
    if (
        cell_range.max_col > MAX_COLUMN
        or cell_range.max_row > MAX_ROW
        or _range_area(cell_range) > 100_000
    ):
        return None
    return worksheet, cell_range


def _confirmed_invalid_direct_a1_reference(
    context: RuleContext,
    current_sheet: Worksheet,
    formula: str,
) -> bool:
    text = formula.strip().lstrip("=")
    if "[" in text or "]" in text:
        return False
    if "!" in text:
        qualifier, body = text.rsplit("!", 1)
        qualifier = qualifier.strip()
        if qualifier.startswith("'") and qualifier.endswith("'"):
            qualifier = qualifier[1:-1].replace("''", "'")
    else:
        qualifier = current_sheet.title
        body = text
    if "," in body:
        return False
    try:
        cell_range = CellRange(body.replace("$", ""))
    except ValueError:
        return False
    if ":" in qualifier:
        return True
    if worksheet_by_name(context.workbook, qualifier) is None:
        return True
    return cell_range.max_col > MAX_COLUMN or cell_range.max_row > MAX_ROW


def _range_area(cell_range: CellRange) -> int:
    return (cell_range.max_row - cell_range.min_row + 1) * (
        cell_range.max_col - cell_range.min_col + 1
    )


def _range_length(cell_range: CellRange) -> int:
    if cell_range.min_col == cell_range.max_col:
        return cell_range.max_row - cell_range.min_row + 1
    if cell_range.min_row == cell_range.max_row:
        return cell_range.max_col - cell_range.min_col + 1
    return _range_area(cell_range)


def _range_values(worksheet: Worksheet, cell_range: CellRange) -> list[Any]:
    return [
        _value_at(worksheet, row, column)
        for row in range(cell_range.min_row, cell_range.max_row + 1)
        for column in range(cell_range.min_col, cell_range.max_col + 1)
    ]


def _chart_title_text(chart: Any) -> str:
    title = getattr(chart, "title", None)
    if isinstance(title, str):
        return title.strip()
    rich = getattr(getattr(title, "tx", None), "rich", None)
    paragraphs = getattr(rich, "p", ()) or ()
    parts = [
        str(run.t)
        for paragraph in paragraphs
        for run in (getattr(paragraph, "r", ()) or ())
        if getattr(run, "t", None) is not None
    ]
    return " ".join(parts).strip()


def _chart_location(chart: Any, index: int) -> str:
    anchor = getattr(chart, "anchor", None)
    if isinstance(anchor, str):
        return f"Chart {index} at {anchor}"
    start = getattr(anchor, "_from", None)
    if start is None:
        return f"Chart {index}"
    return f"Chart {index} at {get_column_letter(int(start.col) + 1)}{int(start.row) + 1}"


def _direct_lineage_header(
    context: RuleContext,
    worksheet: Worksheet,
    cell_range: CellRange,
) -> tuple[str, str] | None:
    targets: list[tuple[str, int, int]] = []
    for value in _range_values(worksheet, cell_range):
        if not isinstance(value, str):
            return None
        match = DIRECT_REFERENCE_RE.fullmatch(value)
        if match is None:
            return None
        sheet_name = match.group("sheet")
        if sheet_name is None:
            target_sheet_name = worksheet.title
        else:
            target_sheet_name = sheet_name.strip()
            if target_sheet_name.startswith("'") and target_sheet_name.endswith("'"):
                target_sheet_name = target_sheet_name[1:-1].replace("''", "'")
        target_sheet = worksheet_by_name(context.workbook, target_sheet_name)
        if target_sheet is None:
            return None
        column = CellRange(f"{match.group('column')}1").min_col
        targets.append((target_sheet.title, column, int(match.group("row"))))
    if not targets:
        return None
    sheets = {item[0] for item in targets}
    columns = {item[1] for item in targets}
    if len(sheets) != 1 or len(columns) != 1:
        return None
    target_sheet_name = next(iter(sheets))
    column = next(iter(columns))
    rows = [item[2] for item in targets]
    target_sheet = cast(Worksheet, context.workbook[target_sheet_name])
    for region in context.data_regions.get(target_sheet_name, []):
        if (
            region.min_column <= column <= region.max_column
            and region.min_row < min(rows)
            and max(rows) <= region.max_row
        ):
            header = _value_at(target_sheet, region.min_row, column)
            if isinstance(header, str) and header.strip():
                return target_sheet_name, header
    return None


def _profiles_containing_source_range(
    context: RuleContext,
    worksheet: Worksheet,
    cell_range: CellRange,
) -> tuple[_TableProfile, ...]:
    if cell_range.min_col != cell_range.max_col:
        return ()
    column = cell_range.min_col
    matches = [
        profile
        for profile in _table_profiles(context, worksheet)
        if column in profile.headers
        and profile.header_row < cell_range.min_row
        and cell_range.max_row <= profile.region.max_row
    ]
    return tuple(
        sorted(
            matches,
            key=lambda profile: (
                profile.region.max_row - profile.region.min_row,
                profile.region.max_column - profile.region.min_column,
                profile.header_row,
            ),
        )
    )


def _chart_source_header(
    context: RuleContext,
    worksheet: Worksheet,
    cell_range: CellRange,
) -> tuple[str, str] | None:
    lineage = _direct_lineage_header(context, worksheet, cell_range)
    if lineage is not None:
        return lineage
    profiles = _profiles_containing_source_range(context, worksheet, cell_range)
    if profiles:
        header = profiles[0].headers.get(cell_range.min_col)
        if isinstance(header, str) and header.strip():
            return worksheet.title, header
    if cell_range.min_col == cell_range.max_col and cell_range.min_row > 1:
        header = _value_at(worksheet, cell_range.min_row - 1, cell_range.min_col)
        if isinstance(header, str) and header.strip():
            return worksheet.title, header
    return None


def _date_headers_in_source_table(
    context: RuleContext,
    worksheet: Worksheet,
    cell_range: CellRange,
) -> list[dict[str, Any]]:
    profiles = _profiles_containing_source_range(context, worksheet, cell_range)
    if not profiles:
        return []
    return [
        {"column": get_column_letter(column), "header": header}
        for column, header in profiles[0].headers.items()
        if column != cell_range.min_col and _looks_like_date_header(header)
    ]


class ChartSourceStructureRule(WorkbookRule):
    rule_id = "WL031_CHART_SOURCE_STRUCTURE"
    title = "Chart source range is structurally inconsistent"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            for index, chart in enumerate(getattr(worksheet, "_charts", ()), start=1):
                issues: list[dict[str, Any]] = []
                chart_title = _chart_title_text(chart)
                declared_prefix = re.split(r"[\[(\n]", chart_title, maxsplit=1)[0]
                declared_groups = _semantic_groups(declared_prefix)
                for series_index, series in enumerate(getattr(chart, "ser", ()), start=1):
                    value_formula = (
                        _reference_formula(getattr(series, "val", None))
                        or _reference_formula(getattr(series, "yVal", None))
                        or _reference_formula(getattr(series, "xVal", None))
                    )
                    category_formula = _reference_formula(getattr(series, "cat", None))
                    title_formula = _reference_formula(getattr(series, "tx", None))
                    value_ref = (
                        _resolve_reference(context, worksheet, value_formula)
                        if value_formula is not None
                        else None
                    )
                    category_ref = (
                        _resolve_reference(context, worksheet, category_formula)
                        if category_formula is not None
                        else None
                    )
                    title_ref = (
                        _resolve_reference(context, worksheet, title_formula)
                        if title_formula is not None
                        else None
                    )
                    if (
                        value_formula is not None
                        and value_ref is None
                        and _confirmed_invalid_direct_a1_reference(
                            context, worksheet, value_formula
                        )
                    ):
                        issues.append(
                            {
                                "series": series_index,
                                "kind": "unresolved_value_reference",
                                "reference": value_formula,
                            }
                        )
                        continue
                    if value_ref is None:
                        continue
                    value_sheet, value_range = value_ref
                    if not any(
                        _is_nonblank(value) for value in _range_values(value_sheet, value_range)
                    ):
                        issues.append(
                            {
                                "series": series_index,
                                "kind": "blank_value_range",
                                "reference": value_formula,
                            }
                        )
                    if (
                        category_formula is not None
                        and category_ref is None
                        and _confirmed_invalid_direct_a1_reference(
                            context, worksheet, category_formula
                        )
                    ):
                        issues.append(
                            {
                                "series": series_index,
                                "kind": "unresolved_category_reference",
                                "reference": category_formula,
                            }
                        )
                    elif category_ref is not None:
                        _, category_range = category_ref
                        if _range_length(value_range) != _range_length(category_range):
                            issues.append(
                                {
                                    "series": series_index,
                                    "kind": "category_value_length_mismatch",
                                    "values": _range_length(value_range),
                                    "categories": _range_length(category_range),
                                }
                            )
                    if title_formula is not None:
                        if title_ref is None and _confirmed_invalid_direct_a1_reference(
                            context, worksheet, title_formula
                        ):
                            issues.append(
                                {
                                    "series": series_index,
                                    "kind": "unresolved_series_title_reference",
                                    "reference": title_formula,
                                }
                            )
                        elif title_ref is not None:
                            title_sheet, title_range = title_ref
                            if not any(
                                _is_nonblank(value)
                                for value in _range_values(title_sheet, title_range)
                            ):
                                issues.append(
                                    {
                                        "series": series_index,
                                        "kind": "blank_series_title_reference",
                                        "reference": title_formula,
                                    }
                                )
                    lineage = _chart_source_header(context, value_sheet, value_range)
                    if lineage is not None and len(declared_groups) == 1:
                        source_sheet_name, source_header = lineage
                        source_groups = _semantic_groups(source_header)
                        if len(source_groups) == 1 and declared_groups.isdisjoint(source_groups):
                            issues.append(
                                {
                                    "series": series_index,
                                    "kind": "title_source_semantic_mismatch",
                                    "chart_title": chart_title,
                                    "source_sheet": source_sheet_name,
                                    "source_header": source_header,
                                }
                            )
                    if category_ref is not None and TREND_RE.search(chart_title):
                        category_sheet, category_range = category_ref
                        category_source = _chart_source_header(
                            context, category_sheet, category_range
                        )
                        date_headers = _date_headers_in_source_table(
                            context, category_sheet, category_range
                        )
                        if (
                            category_source is not None
                            and _looks_like_identifier_header(category_source[1])
                            and date_headers
                        ):
                            issues.append(
                                {
                                    "series": series_index,
                                    "kind": "category_source_semantic_mismatch",
                                    "chart_title": chart_title,
                                    "source_sheet": category_source[0],
                                    "category_header": category_source[1],
                                    "available_date_headers": date_headers,
                                }
                            )
                if not issues:
                    continue
                location = _chart_location(chart, index)
                evidence = Evidence(
                    summary=f"Chart {index} has {len(issues)} source-range issues",
                    observed={"title": chart_title, "issues": issues},
                    expected={"consistent_series_sources": True},
                    details={"chart_index": index},
                )
                result.findings.append(
                    _make_finding(
                        context=context,
                        rule_id=self.rule_id,
                        title=self.title,
                        explanation=(
                            "A chart series has an unresolved, blank, length-mismatched, or semantically "
                            "inconsistent source reference."
                        ),
                        severity=Severity.WARNING,
                        confidence=0.96,
                        worksheet=worksheet,
                        location=location,
                        evidence=evidence,
                        expected="Chart titles, categories, series labels, and value ranges resolve consistently.",
                        suggested_action=(
                            "Review the chart's Select Data dialog and source headers; WorkbookLens does "
                            "not rewrite chart definitions automatically."
                        ),
                    )
                )
        return result


class DeepFreezePaneRule(WorkbookRule):
    rule_id = "WL032_DEEP_FREEZE_PANE"
    title = "Freeze pane starts deep inside visible content"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            if worksheet.sheet_state != "visible" or worksheet.freeze_panes is None:
                continue
            pane = worksheet.freeze_panes
            coordinate = pane.coordinate if isinstance(pane, Cell) else str(pane)
            try:
                pane_range = CellRange(coordinate.replace("$", ""))
            except ValueError:
                continue
            pane_row = pane_range.min_row
            pane_column = pane_range.min_col
            populated = [
                (cell.row, cell.column)
                for cell in worksheet._cells.values()
                if isinstance(cell, Cell) and _is_nonblank(cell.value)
            ]
            if not populated:
                continue
            min_row = min(row for row, _ in populated)
            max_row = max(row for row, _ in populated)
            min_column = min(column for _, column in populated)
            max_column = max(column for _, column in populated)
            height = max_row - min_row + 1
            width = max_column - min_column + 1
            deep_row = pane_row - min_row >= max(8, math.ceil(height * 0.5))
            deep_column = pane_column - min_column >= max(5, math.ceil(width * 0.4))
            matched_profile: _TableProfile | None = None
            for profile in _table_profiles(context, worksheet):
                region = profile.region
                if (
                    region.min_row <= pane_row <= region.max_row
                    and region.min_column <= pane_column <= region.max_column
                ):
                    matched_profile = profile
                    break
            if not deep_row and not deep_column:
                continue
            evidence = Evidence(
                summary=f"Freeze pane at {coordinate} is deep inside visible content",
                observed={
                    "freeze_panes": coordinate,
                    "frozen_rows": max(0, pane_row - 1),
                    "frozen_columns": max(0, pane_column - 1),
                },
                expected={"near_header_or_leading_columns": True},
                details={
                    "deep_row": deep_row,
                    "deep_column": deep_column,
                    "visible_row_bounds": [min_row, max_row],
                    "visible_column_bounds": [min_column, max_column],
                    "outside_visible_content": pane_row > max_row + 1
                    or pane_column > max_column + 1,
                    "table_bounds": (
                        None
                        if matched_profile is None
                        else [
                            matched_profile.region.min_row,
                            matched_profile.region.max_row,
                            matched_profile.region.min_column,
                            matched_profile.region.max_column,
                        ]
                    ),
                },
            )
            result.findings.append(
                _make_finding(
                    context=context,
                    rule_id=self.rule_id,
                    title=self.title,
                    explanation=(
                        "The saved freeze pane locks many rows or columns and begins within the main "
                        "visible data instead of near its header or leading identifiers."
                    ),
                    severity=Severity.WARNING,
                    confidence=0.93,
                    worksheet=worksheet,
                    location=coordinate,
                    evidence=evidence,
                    expected="Freeze panes preserve headers or leading identifiers without obscuring data.",
                    suggested_action=(
                        "Review the intended navigation point and reset the pane manually; WorkbookLens "
                        "does not change freeze-pane definitions automatically."
                    ),
                )
            )
        return result


def _split_reference_list(value: str) -> tuple[str, ...]:
    parts: list[str] = []
    current: list[str] = []
    quoted = False
    index = 0
    while index < len(value):
        character = value[index]
        if character == "'":
            if quoted and index + 1 < len(value) and value[index + 1] == "'":
                current.extend((character, character))
                index += 2
                continue
            quoted = not quoted
        if character == "," and not quoted:
            part = "".join(current).strip()
            if part:
                parts.append(part)
            current = []
        else:
            current.append(character)
        index += 1
    part = "".join(current).strip()
    if part:
        parts.append(part)
    return tuple(parts)


def _contains_range(container: CellRange, target: CellRange) -> bool:
    return (
        container.min_row <= target.min_row <= target.max_row <= container.max_row
        and container.min_col <= target.min_col <= target.max_col <= container.max_col
    )


def _ranges_intersect(left: CellRange, right: CellRange) -> bool:
    return not (
        left.max_row < right.min_row
        or right.max_row < left.min_row
        or left.max_col < right.min_col
        or right.max_col < left.min_col
    )


def _ranges_cover_target(ranges: list[CellRange], target: CellRange) -> bool:
    relevant = [item for item in ranges if _ranges_intersect(item, target)]
    if not relevant:
        return False
    row_boundaries = {target.min_row, target.max_row + 1}
    for item in relevant:
        row_boundaries.add(max(target.min_row, item.min_row))
        row_boundaries.add(min(target.max_row + 1, item.max_row + 1))
    ordered = sorted(row_boundaries)
    for start_row, end_row in pairwise(ordered):
        if start_row >= end_row:
            continue
        intervals = sorted(
            (
                max(target.min_col, item.min_col),
                min(target.max_col, item.max_col),
            )
            for item in relevant
            if item.min_row <= start_row and item.max_row + 1 >= end_row
        )
        covered_to = target.min_col - 1
        for start_column, end_column in intervals:
            if start_column > covered_to + 1:
                break
            covered_to = max(covered_to, end_column)
            if covered_to >= target.max_col:
                break
        if covered_to < target.max_col:
            return False
    return True


def _column_width_pixels(worksheet: Worksheet, column: int) -> float:
    width = column_width(worksheet, column)
    return math.floor(((256.0 * width + math.floor(128.0 / 7.0)) / 256.0) * 7.0)


def _chart_occupied_range(worksheet: Worksheet, chart: Any) -> CellRange | None:
    anchor = getattr(chart, "anchor", None)
    if isinstance(anchor, str):
        try:
            return CellRange(anchor.replace("$", ""))
        except ValueError:
            return None
    start = getattr(anchor, "_from", None)
    if start is None:
        return None
    start_column = int(start.col) + 1
    start_row = int(start.row) + 1
    end = getattr(anchor, "to", None)
    if end is not None:
        return CellRange(
            min_col=start_column,
            min_row=start_row,
            max_col=max(start_column, int(end.col) + int(bool(int(end.colOff)))),
            max_row=max(start_row, int(end.row) + int(bool(int(end.rowOff)))),
        )
    extent = getattr(anchor, "ext", None)
    width_emu = getattr(extent, "cx", None)
    height_emu = getattr(extent, "cy", None)
    if not isinstance(width_emu, int) or not isinstance(height_emu, int):
        return None

    remaining_width = max(0.0, width_emu / 9525.0)
    end_column = start_column
    first_column_offset = max(0.0, int(getattr(start, "colOff", 0)) / 9525.0)
    while end_column < MAX_COLUMN:
        available = max(
            0.0,
            _column_width_pixels(worksheet, end_column) - first_column_offset,
        )
        if remaining_width <= available:
            break
        remaining_width -= available
        end_column += 1
        first_column_offset = 0.0

    remaining_height = max(0.0, height_emu / 9525.0)
    end_row = start_row
    first_row_offset = max(0.0, int(getattr(start, "rowOff", 0)) / 9525.0)
    while end_row < MAX_ROW:
        available = max(
            0.0,
            row_height(worksheet, end_row) * 96.0 / 72.0 - first_row_offset,
        )
        if remaining_height <= available:
            break
        remaining_height -= available
        end_row += 1
        first_row_offset = 0.0
    return CellRange(
        min_col=start_column,
        min_row=start_row,
        max_col=end_column,
        max_row=end_row,
    )


def _chart_is_clipped_by_print_range(
    worksheet: Worksheet,
    chart: Any,
    print_range: CellRange,
) -> dict[str, Any] | None:
    anchor = getattr(chart, "anchor", None)
    start = getattr(anchor, "_from", None)
    if start is None:
        return None
    start_column = int(start.col) + 1
    start_row = int(start.row) + 1
    anchor_cell = CellRange(
        min_col=start_column,
        min_row=start_row,
        max_col=start_column,
        max_row=start_row,
    )
    if not _contains_range(print_range, anchor_cell):
        return None
    end = getattr(anchor, "to", None)
    if end is not None:
        # A zero-offset ``to`` marker lies on the leading edge of its marker cell,
        # so the frame ends in the preceding column or row. A nonzero offset
        # consumes space inside the marker cell itself.
        occupied_end_column = int(end.col) + int(bool(int(end.colOff)))
        occupied_end_row = int(end.row) + int(bool(int(end.rowOff)))
        if occupied_end_column <= print_range.max_col and occupied_end_row <= print_range.max_row:
            return None
        return {
            "anchor_start": anchor_cell.coord,
            "anchor_end": f"{get_column_letter(int(end.col) + 1)}{int(end.row) + 1}",
            "print_range": str(print_range),
        }
    extent = getattr(anchor, "ext", None)
    width_emu = getattr(extent, "cx", None)
    height_emu = getattr(extent, "cy", None)
    if not isinstance(width_emu, int) or not isinstance(height_emu, int):
        return None
    required_width_pixels = width_emu / 9525.0
    required_height_pixels = height_emu / 9525.0
    available_width_pixels = sum(
        _column_width_pixels(worksheet, column)
        for column in range(start_column, print_range.max_col + 1)
    )
    available_height_pixels = sum(
        row_height(worksheet, row) * 96.0 / 72.0
        for row in range(start_row, print_range.max_row + 1)
    )
    if (
        required_width_pixels <= available_width_pixels * 1.05
        and required_height_pixels <= available_height_pixels * 1.05
    ):
        return None
    return {
        "anchor_start": anchor_cell.coord,
        "print_range": str(print_range),
        "required_pixels": [round(required_width_pixels, 1), round(required_height_pixels, 1)],
        "available_pixels": [round(available_width_pixels, 1), round(available_height_pixels, 1)],
    }


class PrintAreaCoverageRule(WorkbookRule):
    rule_id = "WL033_PRINT_AREA_COVERAGE"
    title = "Print area truncates meaningful worksheet content"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            if worksheet.sheet_state != "visible" or not worksheet.print_area:
                continue
            print_ranges = [
                resolved_range
                for item in _split_reference_list(str(worksheet.print_area))
                if (resolved := _resolve_reference(context, worksheet, item)) is not None
                and resolved[0].title == worksheet.title
                for resolved_range in [resolved[1]]
            ]
            if not print_ranges:
                continue
            issues: list[dict[str, Any]] = []
            for profile in _table_profiles(context, worksheet):
                columns = [column for column, header in profile.headers.items() if header]
                if not columns:
                    continue
                target = CellRange(
                    min_col=min(columns),
                    min_row=profile.header_row,
                    max_col=max(columns),
                    max_row=max(profile.body_rows),
                )
                if not any(_ranges_intersect(item, target) for item in print_ranges):
                    continue
                if _ranges_cover_target(print_ranges, target):
                    continue
                issues.append(
                    {
                        "kind": "truncated_inferred_table",
                        "table": str(target),
                    }
                )
            for chart_index, chart in enumerate(getattr(worksheet, "_charts", ()), start=1):
                anchor = getattr(chart, "anchor", None)
                start = getattr(anchor, "_from", None)
                if start is None:
                    continue
                anchor_cell = CellRange(
                    min_col=int(start.col) + 1,
                    min_row=int(start.row) + 1,
                    max_col=int(start.col) + 1,
                    max_row=int(start.row) + 1,
                )
                occupied_range = _chart_occupied_range(worksheet, chart)
                if occupied_range is not None and not any(
                    _ranges_intersect(item, occupied_range) for item in print_ranges
                ):
                    issues.append(
                        {
                            "kind": "chart_anchor_excluded_by_print_area",
                            "chart": chart_index,
                            "anchor_range": str(occupied_range),
                            "print_ranges": [str(item) for item in print_ranges],
                        }
                    )
                    continue
                containing_range = next(
                    (item for item in print_ranges if _contains_range(item, anchor_cell)),
                    None,
                )
                if containing_range is None:
                    continue
                clipping = _chart_is_clipped_by_print_range(
                    worksheet,
                    chart,
                    containing_range,
                )
                if clipping is not None:
                    issues.append(
                        {
                            "kind": "chart_anchor_clipped_by_print_area",
                            "chart": chart_index,
                            **clipping,
                        }
                    )
            if not issues:
                continue
            evidence = Evidence(
                summary=f"Print area leaves {len(issues)} meaningful ranges outside its bounds",
                observed={"print_area": str(worksheet.print_area), "issues": issues},
                expected={"meaningful_content_covered": True},
            )
            result.findings.append(
                _make_finding(
                    context=context,
                    rule_id=self.rule_id,
                    title=self.title,
                    explanation=(
                        "The saved print area intersects a dense table without covering it fully, or "
                        "contains a chart anchor whose saved frame extends beyond the printed bounds."
                    ),
                    severity=Severity.WARNING,
                    confidence=0.95,
                    worksheet=worksheet,
                    location=str(worksheet.print_area),
                    evidence=evidence,
                    expected="Print areas include intended dense tables and complete printed chart frames.",
                    suggested_action=(
                        "Review page setup and expand or redefine the print area manually; WorkbookLens "
                        "does not alter print settings automatically."
                    ),
                )
            )
        return result


DATA_QUALITY_RULES: tuple[type[WorkbookRule], ...] = (
    InferredDuplicateIdentifierRule,
    MissingInferredIdentifierRule,
    MixedNumericStorageRule,
    RobustNumericOutlierRule,
    PercentageScaleOutlierRule,
    DateStorageAnomalyRule,
    AutoFilterCoverageRule,
    ChartSourceStructureRule,
    DeepFreezePaneRule,
    PrintAreaCoverageRule,
    SignConstrainedMeasureRule,
    ConditionalFormatSemanticConflictRule,
)


__all__ = ["DATA_QUALITY_RULES"]
