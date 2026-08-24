"""Conservative diagnostics for contradictory saved print settings."""

from __future__ import annotations

import re
from bisect import bisect_left
from collections.abc import Iterable
from dataclasses import dataclass
from typing import Any

from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.print_settings import PrintArea
from openpyxl.worksheet.worksheet import Worksheet

from workbooklens.models import Confidence, Evidence, Finding, Region, Severity
from workbooklens.rules.base import RuleContext, RuleResult, WorkbookRule
from workbooklens.utils import stable_id

_SUMMARY_RE = re.compile(
    r"(?<![A-Z0-9_])(?:grand[\s-]+total|subtotal|total|average|summary)(?![A-Z0-9_])"
    r"|(?:合计|总计|小计|汇总|平均)",
    re.IGNORECASE,
)
_FIXED_TOTAL_PAGE_RE = re.compile(
    r"(?:&P|&\[Page\])(?:\s*页)?\s*(?:of|/|共)\s*(?P<total>\d{1,4})"
    r"(?:\s*(?:pages?|页))?",
    re.IGNORECASE,
)
_DYNAMIC_TOTAL_PAGE_RE = re.compile(r"&(?:N|\[Pages\])", re.IGNORECASE)
_MIN_DENSE_REGION_BODY_ROWS = 12
_MIN_ROWS_AFTER_EARLY_BREAK = 8
_MIN_DENSE_ROW_RATIO = 0.75


@dataclass(frozen=True)
class _PopulatedBounds:
    """Sparse row/column bounds for one saved print area."""

    cell_range: CellRange
    rows: tuple[int, ...]
    columns: tuple[int, ...]

    def has_row_values(self, row_min: int, row_max: int) -> bool:
        index = bisect_left(self.rows, row_min)
        return index < len(self.rows) and self.rows[index] <= row_max

    def has_column_values(self, column_min: int, column_max: int) -> bool:
        index = bisect_left(self.columns, column_min)
        return index < len(self.columns) and self.columns[index] <= column_max


@dataclass(frozen=True)
class _EffectiveBreaks:
    """Effective manual breaks and their saved print-area scope."""

    range_ref: str | None
    rows: tuple[int, ...]
    columns: tuple[int, ...]


def _build_populated_bounds(
    worksheet: Worksheet,
    print_ranges: tuple[CellRange, ...],
) -> tuple[_PopulatedBounds, ...]:
    """Build each print area's populated row/column index in one cell pass."""

    if not print_ranges:
        return ()
    row_sets: list[set[int]] = [set() for _ in print_ranges]
    column_sets: list[set[int]] = [set() for _ in print_ranges]
    # Deliberately consume the sparse cell mapping once. Pagination checks below
    # use the resulting indexes instead of rescanning every cell for each break.
    for cell in worksheet._cells.values():
        if not isinstance(cell, Cell) or cell.value is None:
            continue
        for index, cell_range in enumerate(print_ranges):
            if (
                cell_range.min_row <= cell.row <= cell_range.max_row
                and cell_range.min_col <= cell.column <= cell_range.max_col
            ):
                row_sets[index].add(cell.row)
                column_sets[index].add(cell.column)
    return tuple(
        _PopulatedBounds(
            cell_range=cell_range,
            rows=tuple(sorted(rows)),
            columns=tuple(sorted(columns)),
        )
        for cell_range, rows, columns in zip(print_ranges, row_sets, column_sets, strict=True)
    )


def _confidence(value: float) -> Confidence:
    return Confidence(max(0.0, min(1.0, value)))


def _finding(
    *,
    context: RuleContext,
    worksheet: Worksheet,
    location: str,
    issue_kind: str,
    title: str,
    explanation: str,
    severity: Severity,
    confidence: float,
    evidence: Evidence,
    expected: str,
    suggested_action: str,
) -> Finding:
    return Finding(
        id=stable_id("finding", PrintSettingConsistencyRule.rule_id, worksheet.title, issue_kind),
        content_fingerprint=stable_id(
            "finding-content", evidence.model_dump(mode="json"), length=24
        ),
        rule_id=PrintSettingConsistencyRule.rule_id,
        title=title,
        explanation=explanation,
        severity=severity,
        confidence=_confidence(confidence),
        workbook=context.path.name,
        sheet=worksheet.title,
        location=location,
        evidence=evidence,
        expected=expected,
        suggested_action=suggested_action,
    )


def _print_ranges(worksheet: Worksheet) -> tuple[CellRange, ...]:
    if not worksheet.print_area:
        return ()
    try:
        area = PrintArea.from_string(str(worksheet.print_area))
    except (TypeError, ValueError):
        return ()
    unique = {
        (item.min_row, item.max_row, item.min_col, item.max_col): item for item in area.ranges
    }
    return tuple(
        sorted(
            unique.values(),
            key=lambda item: (item.min_row, item.min_col, item.max_row, item.max_col),
        )
    )


def _cell_ranges_overlap(left: CellRange, right: CellRange) -> bool:
    """Return whether two print rectangles share at least one cell."""

    return not (
        left.max_row < right.min_row
        or right.max_row < left.min_row
        or left.max_col < right.min_col
        or right.max_col < left.min_col
    )


def _region_boxes_overlap(left: Region, right: Region) -> bool:
    return not (
        left.max_row < right.min_row
        or right.max_row < left.min_row
        or left.max_column < right.min_column
        or right.max_column < left.min_column
    )


def _region_range(region: Region) -> CellRange:
    return CellRange(
        f"{get_column_letter(region.min_column)}{region.min_row}:"
        f"{get_column_letter(region.max_column)}{region.max_row}"
    )


def _ranges_intersect(left: CellRange, right: Region) -> bool:
    return not (
        left.max_row < right.min_row
        or right.max_row < left.min_row
        or left.max_col < right.min_column
        or right.max_column < left.min_col
    )


def _row_density(worksheet: Worksheet, region: Region, row: int) -> float:
    row_dimension = worksheet.row_dimensions.get(row)
    if row_dimension is not None and row_dimension.hidden:
        return 0.0
    populated = sum(
        1
        for column in range(region.min_column, region.max_column + 1)
        if isinstance((cell := worksheet._cells.get((row, column))), Cell)
        and cell.value is not None
    )
    return populated / (region.max_column - region.min_column + 1)


def _row_has_summary_label(worksheet: Worksheet, region: Region, row: int) -> bool:
    return any(
        isinstance((cell := worksheet._cells.get((row, column))), Cell)
        and isinstance(cell.value, str)
        and _SUMMARY_RE.search(cell.value) is not None
        for column in range(region.min_column, region.max_column + 1)
    )


def _boundary_has_group_structure(
    worksheet: Worksheet,
    region: Region,
    break_row: int,
) -> bool:
    before_dimension = worksheet.row_dimensions.get(break_row)
    after_dimension = worksheet.row_dimensions.get(break_row + 1)
    before_level = int(before_dimension.outlineLevel or 0) if before_dimension is not None else 0
    after_level = int(after_dimension.outlineLevel or 0) if after_dimension is not None else 0
    if before_level != after_level or before_level > 0 or after_level > 0:
        return True
    if any(
        _row_has_summary_label(worksheet, region, row)
        for row in range(max(region.min_row, break_row - 1), min(region.max_row, break_row + 1) + 1)
    ):
        return True
    return any(
        merged.min_row <= break_row + 1
        and merged.max_row >= break_row
        and merged.max_col >= region.min_column
        and merged.min_col <= region.max_column
        and (merged.max_col - merged.min_col + 1) >= 2
        for merged in worksheet.merged_cells.ranges
    )


def _early_manual_row_breaks(
    worksheet: Worksheet,
    regions: Iterable[Region],
) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    for page_break in worksheet.row_breaks.brk:
        if page_break.man is not True or page_break.min != 0 or page_break.max < 16_383:
            continue
        break_row = int(page_break.id)
        for region in regions:
            body_rows = region.max_row - region.min_row
            rows_before = break_row - region.min_row
            rows_after = region.max_row - break_row
            if (
                float(region.confidence) < 0.75
                or region.max_column - region.min_column + 1 < 3
                or body_rows < _MIN_DENSE_REGION_BODY_ROWS
                or rows_before < 2
                or rows_after < _MIN_ROWS_AFTER_EARLY_BREAK
                or rows_before / body_rows > 0.25
            ):
                continue
            window = (break_row - 1, break_row, break_row + 1, break_row + 2)
            densities = [_row_density(worksheet, region, row) for row in window]
            if min(densities) < _MIN_DENSE_ROW_RATIO:
                continue
            if _boundary_has_group_structure(worksheet, region, break_row):
                continue
            issues.append(
                {
                    "break_after_row": break_row,
                    "region": (
                        f"{region.min_row}:{region.max_row},"
                        f"columns {region.min_column}:{region.max_column}"
                    ),
                    "body_rows_before_break": rows_before,
                    "body_rows_after_break": rows_after,
                    "relative_position": round(rows_before / body_rows, 3),
                    "adjacent_row_densities": [round(item, 3) for item in densities],
                }
            )
            break
    return issues


def _range_has_values(
    worksheet: Worksheet,
    cell_range: CellRange,
    *,
    row_min: int | None = None,
    row_max: int | None = None,
    column_min: int | None = None,
    column_max: int | None = None,
    populated_bounds: _PopulatedBounds | None = None,
) -> bool:
    min_row = max(cell_range.min_row, row_min if row_min is not None else cell_range.min_row)
    max_row = min(cell_range.max_row, row_max if row_max is not None else cell_range.max_row)
    min_column = max(
        cell_range.min_col,
        column_min if column_min is not None else cell_range.min_col,
    )
    max_column = min(
        cell_range.max_col,
        column_max if column_max is not None else cell_range.max_col,
    )
    if min_row > max_row or min_column > max_column:
        return False
    if populated_bounds is not None:
        # Pagination invokes this helper with one restricted dimension at a
        # time. Keep the general two-dimensional fallback below for private
        # callers that may pass both dimensions.
        if (
            (row_min is not None or row_max is not None)
            and column_min is None
            and column_max is None
        ):
            return populated_bounds.has_row_values(min_row, max_row)
        if (
            (column_min is not None or column_max is not None)
            and row_min is None
            and row_max is None
        ):
            return populated_bounds.has_column_values(min_column, max_column)
        if row_min is None and row_max is None and column_min is None and column_max is None:
            return bool(populated_bounds.rows and populated_bounds.columns)
    return any(
        isinstance(cell, Cell)
        and cell.value is not None
        and min_row <= cell.row <= max_row
        and min_column <= cell.column <= max_column
        for cell in worksheet._cells.values()
    )


def _effective_manual_breaks_with_ranges(
    worksheet: Worksheet,
    regions: tuple[Region, ...],
    print_ranges: tuple[CellRange, ...],
) -> tuple[_EffectiveBreaks, ...]:
    """Return effective manual row/column breaks independently per print area."""

    manual_row_breaks = tuple(
        page_break for page_break in worksheet.row_breaks.brk if page_break.man is True
    )
    manual_column_breaks = tuple(
        page_break for page_break in worksheet.col_breaks.brk if page_break.man is True
    )
    inferred_scopes = not print_ranges
    if inferred_scopes:
        ordered_regions = tuple(
            sorted(
                regions,
                key=lambda item: (
                    item.min_row,
                    item.min_column,
                    item.max_row,
                    item.max_column,
                ),
            )
        )
        if any(
            _region_boxes_overlap(left, right)
            for index, left in enumerate(ordered_regions)
            for right in ordered_regions[index + 1 :]
        ):
            return ()
        print_ranges = tuple(_region_range(region) for region in ordered_regions)
    elif any(
        _cell_ranges_overlap(left, right)
        for index, left in enumerate(print_ranges)
        for right in print_ranges[index + 1 :]
    ):
        # Overlapping saved print areas cannot be assigned to independent page
        # scopes without interpreting Excel's print-order semantics.
        return ()

    areas = _build_populated_bounds(worksheet, print_ranges)
    effective_by_area: list[_EffectiveBreaks] = []
    for populated in areas:
        if not populated.rows or not populated.columns:
            continue
        cell_range = populated.cell_range
        row_breaks = sorted(
            {
                int(page_break.id)
                for page_break in manual_row_breaks
                if isinstance(page_break.min, int)
                and not isinstance(page_break.min, bool)
                and isinstance(page_break.max, int)
                and not isinstance(page_break.max, bool)
                and page_break.min <= populated.columns[0] - 1
                and page_break.max >= populated.columns[-1] - 1
            }
        )
        row_breaks = [
            break_row
            for break_row in row_breaks
            if cell_range.min_row <= break_row < cell_range.max_row
            and _range_has_values(
                worksheet,
                cell_range,
                row_max=break_row,
                populated_bounds=populated,
            )
            and _range_has_values(
                worksheet,
                cell_range,
                row_min=break_row + 1,
                populated_bounds=populated,
            )
        ]
        column_breaks = sorted(
            {
                int(page_break.id)
                for page_break in manual_column_breaks
                if isinstance(page_break.min, int)
                and not isinstance(page_break.min, bool)
                and isinstance(page_break.max, int)
                and not isinstance(page_break.max, bool)
                and page_break.min <= populated.rows[0] - 1
                and page_break.max >= populated.rows[-1] - 1
            }
        )
        column_breaks = [
            break_column
            for break_column in column_breaks
            if cell_range.min_col <= break_column < cell_range.max_col
            and _range_has_values(
                worksheet,
                cell_range,
                column_max=break_column,
                populated_bounds=populated,
            )
            and _range_has_values(
                worksheet,
                cell_range,
                column_min=break_column + 1,
                populated_bounds=populated,
            )
        ]
        effective_by_area.append(
            _EffectiveBreaks(
                range_ref=str(cell_range),
                rows=tuple(row_breaks),
                columns=tuple(column_breaks),
            )
        )
    if (
        inferred_scopes
        and len(effective_by_area) > 1
        and not any(area.rows or area.columns for area in effective_by_area)
    ):
        # Disconnected tables without a manual split can still share one
        # physical page. Do not invent one base page per inferred region.
        return (_EffectiveBreaks(range_ref=None, rows=(), columns=()),)
    return tuple(effective_by_area)


def _effective_manual_breaks(
    worksheet: Worksheet,
    regions: tuple[Region, ...],
    print_ranges: tuple[CellRange, ...],
) -> tuple[list[int], list[int]]:
    """Backward-compatible flattened view of effective breaks."""

    by_area = _effective_manual_breaks_by_area(worksheet, regions, print_ranges)
    return (
        sorted({break_row for rows, _ in by_area for break_row in rows}),
        sorted({break_column for _, columns in by_area for break_column in columns}),
    )


def _effective_manual_breaks_by_area(
    worksheet: Worksheet,
    regions: tuple[Region, ...],
    print_ranges: tuple[CellRange, ...],
) -> tuple[tuple[list[int], list[int]], ...]:
    """Backward-compatible per-area view without range metadata."""

    return tuple(
        (list(area.rows), list(area.columns))
        for area in _effective_manual_breaks_with_ranges(worksheet, regions, print_ranges)
    )


def _footer_texts(worksheet: Worksheet) -> dict[str, str]:
    texts: dict[str, str] = {}
    footer_names = ["oddFooter"]
    header_footer = worksheet.HeaderFooter
    if bool(header_footer.differentOddEven):
        footer_names.append("evenFooter")
    if bool(header_footer.differentFirst):
        footer_names.append("firstFooter")
    for footer_name in footer_names:
        footer = getattr(worksheet, footer_name)
        for part_name in ("left", "center", "right"):
            text = getattr(footer, part_name).text
            if isinstance(text, str) and text.strip():
                texts[f"{footer_name}.{part_name}"] = text
    return texts


def _fixed_footer_issues(
    worksheet: Worksheet,
    row_breaks: list[int],
    column_breaks: list[int],
    *,
    breaks_by_area: tuple[_EffectiveBreaks, ...] | None = None,
) -> list[dict[str, Any]]:
    if breaks_by_area:
        minimum_pages = sum(
            (len(area.rows) + 1) * (len(area.columns) + 1) for area in breaks_by_area
        )
    else:
        minimum_pages = (len(row_breaks) + 1) * (len(column_breaks) + 1)
    issues: list[dict[str, Any]] = []
    for location, text in _footer_texts(worksheet).items():
        if _DYNAMIC_TOTAL_PAGE_RE.search(text) is not None:
            continue
        fixed_totals = [int(match.group("total")) for match in _FIXED_TOTAL_PAGE_RE.finditer(text)]
        invalid_totals = [total for total in fixed_totals if total < minimum_pages]
        if not invalid_totals:
            continue
        issues.append(
            {
                "footer_part": location,
                "text": text,
                # Keep the original scalar field for evidence consumers and
                # add the aggregate list when a footer part repeats the field.
                "fixed_total_pages": invalid_totals[0],
                "fixed_total_page_values": invalid_totals,
                "minimum_pages_from_manual_breaks": minimum_pages,
                "manual_row_breaks": row_breaks,
                "manual_column_breaks": column_breaks,
                "manual_breaks_by_print_area": [
                    {
                        "range": area.range_ref,
                        "manual_row_breaks": list(area.rows),
                        "manual_column_breaks": list(area.columns),
                        "minimum_pages": (len(area.rows) + 1) * (len(area.columns) + 1),
                    }
                    for area in (breaks_by_area or ())
                ],
            }
        )
    return issues


def _fit_width_conflicts(
    worksheet: Worksheet,
    regions: tuple[Region, ...],
    print_ranges: tuple[CellRange, ...],
) -> list[dict[str, Any]]:
    setup_properties = worksheet.sheet_properties.pageSetUpPr
    fit_to_page = bool(setup_properties and setup_properties.fitToPage)
    fit_to_width = worksheet.page_setup.fitToWidth
    if not fit_to_page or not isinstance(fit_to_width, int) or fit_to_width <= 1:
        return []
    if not print_ranges:
        return []
    issues: list[dict[str, Any]] = []
    for region in regions:
        intersecting = [item for item in print_ranges if _ranges_intersect(item, region)]
        if not intersecting:
            continue
        covered_columns = {
            column
            for item in intersecting
            for column in range(
                max(item.min_col, region.min_column),
                min(item.max_col, region.max_column) + 1,
            )
        }
        missing_columns = [
            column
            for column in range(region.min_column, region.max_column + 1)
            if column not in covered_columns
        ]
        if len(missing_columns) < 2:
            continue
        issues.append(
            {
                "fit_to_width": fit_to_width,
                "print_area": str(worksheet.print_area),
                "truncated_region": {
                    "rows": [region.min_row, region.max_row],
                    "columns": [region.min_column, region.max_column],
                },
                "missing_region_columns": missing_columns,
            }
        )
    return issues


class PrintSettingConsistencyRule(WorkbookRule):
    """Report print settings only when independent layout evidence contradicts them."""

    rule_id = "WL055_PRINT_SETTING_CONSISTENCY"
    title = "Saved print settings contradict the inferred printed layout"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            regions = tuple(
                region
                for region in context.data_regions.get(worksheet.title, ())
                if region.kind == "data"
            )
            if not regions:
                continue
            print_ranges = _print_ranges(worksheet)
            early_breaks = _early_manual_row_breaks(worksheet, regions)
            if early_breaks:
                evidence = Evidence(
                    summary=(
                        f"{len(early_breaks)} manual row break(s) split the dense leading portion "
                        "of an inferred table"
                    ),
                    observed=early_breaks,
                    expected={
                        "blank_or_summary_boundary": True,
                        "minimum_rows_after_early_break": _MIN_ROWS_AFTER_EARLY_BREAK,
                    },
                )
                result.findings.append(
                    _finding(
                        context=context,
                        worksheet=worksheet,
                        location=",".join(
                            f"{item['break_after_row']}:{item['break_after_row']}"
                            for item in early_breaks
                        ),
                        issue_kind="early_manual_row_break",
                        title=self.title,
                        explanation=(
                            "A full-width manual page break appears in the first quarter of a long, "
                            "continuous dense table. Two populated rows on each side, and the absence "
                            "of blank, summary, merged-section, or outline-group boundaries, make an "
                            "intentional section break unlikely."
                        ),
                        severity=Severity.WARNING,
                        confidence=0.95,
                        evidence=evidence,
                        expected=(
                            "Manual row breaks align with explicit blank, group, or summary boundaries."
                        ),
                        suggested_action=(
                            "Review the saved manual break in Page Break Preview; WorkbookLens does "
                            "not move or remove page breaks automatically."
                        ),
                    )
                )
            breaks_by_area = _effective_manual_breaks_with_ranges(worksheet, regions, print_ranges)
            effective_rows = sorted(
                {break_row for area in breaks_by_area for break_row in area.rows}
            )
            effective_columns = sorted(
                {break_column for area in breaks_by_area for break_column in area.columns}
            )
            footer_issues = _fixed_footer_issues(
                worksheet,
                effective_rows,
                effective_columns,
                breaks_by_area=breaks_by_area,
            )
            for issue in footer_issues:
                evidence = Evidence(
                    summary=(
                        "A fixed footer page total is below the minimum page count proven by "
                        "effective manual breaks"
                    ),
                    observed=issue,
                    expected={"dynamic_total_pages_field": "&N"},
                )
                result.findings.append(
                    _finding(
                        context=context,
                        worksheet=worksheet,
                        location=str(issue["footer_part"]),
                        issue_kind=f"fixed_footer_total:{issue['footer_part']}",
                        title=self.title,
                        explanation=(
                            "The footer uses the current-page field with a literal total, while saved "
                            "manual breaks prove that the printed layout has more pages than that total."
                        ),
                        severity=Severity.WARNING,
                        confidence=0.99,
                        evidence=evidence,
                        expected="Page totals use Excel's dynamic &N field when the layout is multi-page.",
                        suggested_action=(
                            "Replace the literal total with &N or verify the intended pagination "
                            "manually; WorkbookLens does not edit headers or footers."
                        ),
                    )
                )
            fit_conflicts = _fit_width_conflicts(worksheet, regions, print_ranges)
            if fit_conflicts:
                evidence = Evidence(
                    summary=(
                        "Fit-to-width requests multiple pages while the saved print area provably "
                        "omits columns from an intersected dense table"
                    ),
                    observed=fit_conflicts,
                    expected={"print_area_covers_intersected_dense_regions": True},
                )
                result.findings.append(
                    _finding(
                        context=context,
                        worksheet=worksheet,
                        location=str(worksheet.print_area),
                        issue_kind="fit_width_with_truncated_region",
                        title=self.title,
                        explanation=(
                            "Fit-to-page is enabled with a multi-page width, but the saved print area "
                            "already truncates at least two columns of a dense table it intersects. "
                            "The width setting is not treated as anomalous without this independent "
                            "coverage conflict."
                        ),
                        severity=Severity.INFO,
                        confidence=0.93,
                        evidence=evidence,
                        expected=(
                            "Fit-to-width settings are applied to a print area that covers the intended table."
                        ),
                        suggested_action=(
                            "Review the print area and scaling together in Print Preview; WorkbookLens "
                            "does not change page setup automatically."
                        ),
                    )
                )
        return result


PRINT_QUALITY_RULES: tuple[type[WorkbookRule], ...] = (PrintSettingConsistencyRule,)

__all__ = ["PRINT_QUALITY_RULES", "PrintSettingConsistencyRule"]
