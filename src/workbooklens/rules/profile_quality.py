"""Conservative, profile-aware validation for workbook record fields.

The rules in this module never infer replacement business values.  A profile can
opt in to a review-only trailing-whitespace patch, while every other finding is
report-only and carries the configured or structurally inferred expectation in
its evidence.
"""

from __future__ import annotations

import math
import re
import unicodedata
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Any, Literal, cast

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.xml.constants import MAX_COLUMN, MAX_ROW

from workbooklens.models import (
    Confidence,
    Evidence,
    Finding,
    PatchKind,
    PatchOperation,
    PatchPrecondition,
    PatchRisk,
    Region,
    Severity,
)
from workbooklens.rules.base import RuleContext, RuleResult, WorkbookRule
from workbooklens.snapshot import cell_fingerprint
from workbooklens.utils import stable_id

ProfileRole = Literal[
    "identifier",
    "category",
    "email",
    "phone",
    "currency",
    "percentage",
    "date",
    "number",
    "text",
]

_VALID_ROLES = frozenset(
    {
        "identifier",
        "category",
        "email",
        "phone",
        "currency",
        "percentage",
        "date",
        "number",
        "text",
    }
)
_ROLE_ALIASES = {
    "amount": "currency",
    "datetime": "date",
    "id": "identifier",
    "mail": "email",
    "mobile": "phone",
    "percent": "percentage",
    "rate": "percentage",
    "telephone": "phone",
}
_HORIZONTAL_TRAILING_WHITESPACE_RE = re.compile(r"[ \t\u00a0\u3000]+$")
_EMAIL_RE = re.compile(
    r'^[^\s@<>(),;:\\"]+(?:\.[^\s@<>(),;:\\"]+)*@'
    r"(?:[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?\.)+"
    r"[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?$"
)
_PHONE_EXTENSION_RE = re.compile(r"(?:\s*(?:ext\.?|extension|x|转)\s*\d{1,6})$", re.I)
_PHONE_BODY_RE = re.compile(r"^[+()\d.\-\s]+$")
_QUOTED_FORMAT_TEXT_RE = re.compile(r'"(?:[^"]|"")*"')
_ESCAPED_FORMAT_CHARACTER_RE = re.compile(r"\\.")
_FORMAT_SPACING_RE = re.compile(r"_[^;]|\*[^;]")
_CURRENCY_RE = re.compile(
    r"(?<!\[)\$|[€£¥￥₹₩₽]|"
    r"\[\$(?:[$€£¥￥₹₩₽]|USD|EUR|GBP|CNY|RMB|JPY)(?:-[0-9A-F]+)?\]|"
    r"\b(?:USD|EUR|GBP|CNY|RMB|JPY)\b",
    re.I,
)

_IDENTIFIER_TOKENS = {"code", "id", "identifier", "key", "no", "number", "sku"}
_IDENTIFIER_MARKERS = ("编号", "编码", "代码", "号码", "工号", "学号", "单号", "标识")
_EMAIL_TOKENS = {"email", "mail"}
_EMAIL_MARKERS = ("邮箱", "电子邮件")
_PHONE_TOKENS = {"mobile", "phone", "tel", "telephone"}
_PHONE_MARKERS = ("手机号", "手机", "联系电话", "电话")
_AMOUNT_TOKENS = {
    "amount",
    "cost",
    "income",
    "payroll",
    "price",
    "revenue",
    "salary",
    "sales",
    "turnover",
    "wage",
}
_AMOUNT_MARKERS = ("金额", "成本", "工资", "价格", "单价", "收入", "营收", "销售额", "薪资")
_PERCENTAGE_TOKENS = {"discount", "margin", "percent", "percentage", "rate", "ratio", "tax"}
_PERCENTAGE_MARKERS = ("百分比", "比例", "折扣", "税率", "率")
_DATE_TOKENS = {"date", "datetime", "day", "time", "timestamp"}
_DATE_MARKERS = ("日期", "时间", "年月日")
MAX_PROFILE_RANGE_CELLS = 1_000_000


class ProfileConfigurationError(ValueError):
    """An explicit workbook Profile contract is malformed or cannot be resolved."""


@dataclass(frozen=True, slots=True)
class ColumnProfile:
    """One configured field, selected by a header or explicit column."""

    header: str | None = None
    column: int | str | None = None
    required: bool = False
    allowed_values: tuple[Any, ...] = ()
    role: ProfileRole | None = None
    identifier_width: int | None = None
    preserve_leading_zeros: bool = False
    trim_trailing_whitespace: bool = False


@dataclass(frozen=True, slots=True)
class TableProfile:
    """Configured table bounds and field expectations on one worksheet."""

    sheet: str
    range_ref: str | None = None
    header_row: int | None = None
    columns: tuple[ColumnProfile, ...] = ()


@dataclass(frozen=True, slots=True)
class WorkbookProfile:
    """User-editable semantic expectations consumed by profile quality rules."""

    tables: tuple[TableProfile, ...] = ()
    infer_semantics: bool = True
    report_trailing_whitespace: bool = True
    review_trailing_whitespace_patches: bool = False

    @classmethod
    def from_config(cls, config: Mapping[str, Any] | None) -> WorkbookProfile:
        """Parse the stable ``config['profile']['sheets']`` shape fail-closed.

        The older ``workbook_profile.tables`` spelling remains a compatibility
        alias while the public shape is ``profile.sheets``. Explicit malformed
        contracts raise ``ValueError`` instead of silently disabling rules.
        """

        if not isinstance(config, Mapping):
            return cls()
        configured_version = config.get("version")
        if configured_version is not None and (
            not isinstance(configured_version, int)
            or isinstance(configured_version, bool)
            or configured_version not in {1, 2}
        ):
            raise ValueError("configuration version must be the integer 1 or 2")
        profile_present = "profile" in config
        legacy_profile_present = "workbook_profile" in config
        if profile_present and legacy_profile_present:
            raise ValueError("configuration cannot define both profile and workbook_profile")
        if configured_version == 1 and (
            config.get("profile") is not None or config.get("workbook_profile") is not None
        ):
            raise ValueError("workbook profiles require configuration version 2")
        table_key = "sheets"
        if profile_present:
            raw = config.get("profile")
            if raw is None:
                return cls()
        elif legacy_profile_present:
            raw = config.get("workbook_profile")
            table_key = "tables"
            if raw is None:
                return cls()
        else:
            return cls()
        if not isinstance(raw, Mapping):
            raise ValueError("workbook profile must be a mapping")
        _reject_unknown_keys(
            raw,
            {
                table_key,
                "infer_semantics",
                "report_trailing_whitespace",
                "review_trailing_whitespace_patches",
            },
            location="workbook profile",
        )
        tables = _parse_table_profiles(raw.get(table_key))
        return cls(
            tables=tables,
            infer_semantics=_configured_bool(raw.get("infer_semantics"), True),
            report_trailing_whitespace=_configured_bool(
                raw.get("report_trailing_whitespace"), True
            ),
            review_trailing_whitespace_patches=_configured_bool(
                raw.get("review_trailing_whitespace_patches"), False
            ),
        )


@dataclass(frozen=True, slots=True)
class _ResolvedTable:
    worksheet: Worksheet
    min_row: int
    max_row: int
    min_column: int
    max_column: int
    header_row: int
    headers: dict[int, str]
    body_rows: tuple[int, ...]
    columns: dict[int, ColumnProfile] = field(default_factory=dict)
    explicit: bool = False


@dataclass(frozen=True, slots=True)
class _NumberFormatTraits:
    percentage: bool
    date: bool
    currency: bool
    text: bool


def parse_workbook_profile(config: Mapping[str, Any] | None) -> WorkbookProfile:
    """Public convenience API for parsing a scanner configuration mapping."""

    try:
        return WorkbookProfile.from_config(config)
    except ProfileConfigurationError:
        raise
    except ValueError as exc:
        raise ProfileConfigurationError(str(exc)) from exc


def _context_profile(context: RuleContext) -> WorkbookProfile:
    cache_key = "profile_quality.workbook_profile.v1"
    cached = context.analysis_cache.get(cache_key)
    if isinstance(cached, WorkbookProfile):
        return cached
    profile = parse_workbook_profile(context.config)
    context.analysis_cache[cache_key] = profile
    return profile


def _configured_bool(value: Any, default: bool) -> bool:
    if value is None:
        return default
    if not isinstance(value, bool):
        raise ValueError("profile boolean options must be true or false")
    return value


def _reject_unknown_keys(
    value: Mapping[Any, Any],
    allowed: set[str],
    *,
    location: str,
) -> None:
    unknown = sorted(repr(key) for key in value if key not in allowed)
    if unknown:
        raise ValueError(f"unsupported {location} keys: {', '.join(unknown)}")


def normalize_profile_column_selector(value: Any) -> int | str | None:
    """Normalize an explicit Profile column and reject non-Excel coordinates."""

    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("profile column selector cannot be boolean")
    if isinstance(value, int):
        if 1 <= value <= MAX_COLUMN:
            return value
        raise ValueError("profile column number is outside Excel bounds")
    if not isinstance(value, str):
        raise ValueError(f"invalid profile column selector: {value!r}")

    normalized = value.strip().replace("$", "")
    if normalized.isdigit():
        number = int(normalized)
        if 1 <= number <= MAX_COLUMN:
            return number
        raise ValueError("profile column number is outside Excel bounds")
    try:
        number = column_index_from_string(normalized)
    except ValueError as exc:
        raise ValueError(f"invalid profile column selector: {value!r}") from exc
    if number > MAX_COLUMN:
        raise ValueError("profile column selector is outside Excel bounds")
    return normalized.upper()


def profile_range_bounds(value: str) -> tuple[int, int, int, int]:
    """Return row-first Profile bounds after Excel and area-limit validation."""

    raw = value.replace("$", "")
    if "!" in raw:
        raise ValueError("profile sheet ranges must not repeat the sheet name")
    try:
        min_column, min_row, max_column, max_row = range_boundaries(raw)
    except ValueError as exc:
        raise ValueError(f"invalid profile range {value!r}") from exc
    if None in {min_column, min_row, max_column, max_row}:
        raise ValueError(f"profile range must have explicit row and column bounds: {value!r}")

    resolved_min_column = cast(int, min_column)
    resolved_min_row = cast(int, min_row)
    resolved_max_column = cast(int, max_column)
    resolved_max_row = cast(int, max_row)
    if not (
        1 <= resolved_min_column <= resolved_max_column <= MAX_COLUMN
        and 1 <= resolved_min_row <= resolved_max_row <= MAX_ROW
    ):
        raise ValueError(f"profile range is outside Excel worksheet bounds: {value!r}")

    cell_count = (resolved_max_row - resolved_min_row + 1) * (
        resolved_max_column - resolved_min_column + 1
    )
    if cell_count > MAX_PROFILE_RANGE_CELLS:
        raise ValueError(
            f"profile range {value!r} expands to {cell_count} cells; "
            f"limit is {MAX_PROFILE_RANGE_CELLS}"
        )
    return (
        resolved_min_row,
        resolved_max_row,
        resolved_min_column,
        resolved_max_column,
    )


def normalize_profile_range(value: Any) -> str:
    """Normalize a Profile range while applying the shared validation contract."""

    if not isinstance(value, str):
        raise ValueError("profile range must be a string")
    normalized = value.strip().replace("$", "")
    if not normalized:
        raise ValueError("profile range cannot be blank")
    profile_range_bounds(normalized)
    return normalized


def _parse_table_profiles(raw_tables: Any) -> tuple[TableProfile, ...]:
    if raw_tables is None:
        return ()
    if not isinstance(raw_tables, list):
        raise ValueError("profile sheets must be a list")
    parsed: list[TableProfile] = []
    selectors: set[str] = set()
    for index, raw_table in enumerate(raw_tables, start=1):
        if not isinstance(raw_table, Mapping):
            raise ValueError(f"profile sheet entry {index} must be a mapping")
        _reject_unknown_keys(
            raw_table,
            {"sheet", "range", "header_row", "columns"},
            location=f"profile sheet entry {index}",
        )
        sheet = raw_table.get("sheet")
        if not isinstance(sheet, str) or not sheet.strip():
            raise ValueError(f"profile sheet entry {index} requires a non-blank sheet name")
        raw_range = raw_table.get("range")
        range_ref = None if raw_range is None else normalize_profile_range(raw_range)
        range_bounds = profile_range_bounds(range_ref) if range_ref is not None else None
        raw_header_row = raw_table.get("header_row")
        if raw_header_row is None:
            header_row = range_bounds[0] if range_bounds is not None else 1
        elif (
            not isinstance(raw_header_row, int)
            or isinstance(raw_header_row, bool)
            or not 1 <= raw_header_row <= MAX_ROW
        ):
            raise ValueError("profile header_row must be an integer inside Excel bounds")
        else:
            header_row = raw_header_row
        if range_bounds is not None and not range_bounds[0] <= header_row <= range_bounds[1]:
            raise ValueError("profile header_row must be inside the configured range")
        selector = sheet.strip().casefold()
        if selector in selectors:
            raise ValueError("profile contains duplicate sheet table selectors")
        selectors.add(selector)
        parsed.append(
            TableProfile(
                sheet=sheet.strip(),
                range_ref=None if range_ref is None else range_ref.strip(),
                header_row=header_row,
                columns=_parse_column_profiles(raw_table.get("columns")),
            )
        )
    return tuple(parsed)


def _parse_column_profiles(raw_columns: Any) -> tuple[ColumnProfile, ...]:
    entries: list[tuple[str | None, Any]] = []
    if raw_columns is None:
        return ()
    if isinstance(raw_columns, Mapping):
        entries.extend((str(header), raw) for header, raw in raw_columns.items())
    elif isinstance(raw_columns, list):
        entries.extend((None, raw) for raw in raw_columns)
    else:
        raise ValueError("profile columns must be a list or header mapping")

    parsed: list[ColumnProfile] = []
    selectors: set[tuple[str, str]] = set()
    for index, (default_header, raw) in enumerate(entries, start=1):
        if isinstance(raw, str):
            raw = {"role": raw}
        elif isinstance(raw, bool):
            raw = {"required": raw}
        if not isinstance(raw, Mapping):
            raise ValueError(f"profile column entry {index} must be a mapping")
        _reject_unknown_keys(
            raw,
            {
                "header",
                "column",
                "role",
                "type",
                "required",
                "allowed_values",
                "enum",
                "identifier_width",
                "width",
                "preserve_leading_zeros",
                "trim_trailing_whitespace",
            },
            location=f"profile column entry {index}",
        )
        header_value = raw.get("header", default_header)
        if header_value is None:
            header = None
        elif not isinstance(header_value, str) or not header_value.strip():
            raise ValueError(f"profile column entry {index} has an invalid header selector")
        else:
            header = header_value.strip()
        column = normalize_profile_column_selector(raw.get("column"))
        if header is not None and column is not None:
            raise ValueError("profile columns require exactly one header or column selector")
        if header is None and column is None:
            raise ValueError("profile columns require exactly one header or column selector")
        if column is not None:
            resolved_column = _resolve_column(column)
            if resolved_column is None:
                raise ValueError(f"profile column entry {index} has an invalid column selector")
            selector = ("column", str(resolved_column))
        else:
            selector = ("header", cast(str, header).casefold())
        if selector in selectors:
            raise ValueError("profile sheet contains duplicate column selectors")
        selectors.add(selector)
        allowed = raw.get("allowed_values", raw.get("enum", ()))
        if allowed is None:
            allowed_values: tuple[Any, ...] = ()
        elif isinstance(allowed, (list, tuple)):
            allowed_values = tuple(allowed)
        else:
            raise ValueError("profile allowed_values must be a list")
        role = _parse_role(raw.get("role", raw.get("type")))
        identifier_width = raw.get("identifier_width", raw.get("width"))
        if identifier_width is None:
            identifier_width = None
        elif (
            not isinstance(identifier_width, int)
            or isinstance(identifier_width, bool)
            or not 1 <= identifier_width <= 64
        ):
            raise ValueError("profile identifier_width must be an integer from 1 to 64")
        preserve_leading_zeros = _configured_bool(
            raw.get("preserve_leading_zeros"), identifier_width is not None
        )
        if role is None and preserve_leading_zeros:
            role = "identifier"
        parsed.append(
            ColumnProfile(
                header=header,
                column=column,
                required=_configured_bool(raw.get("required"), False),
                allowed_values=allowed_values,
                role=role,
                identifier_width=identifier_width,
                preserve_leading_zeros=preserve_leading_zeros,
                trim_trailing_whitespace=_configured_bool(
                    raw.get("trim_trailing_whitespace"), False
                ),
            )
        )
    return tuple(parsed)


def _parse_role(value: Any) -> ProfileRole | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError("profile role must be a string")
    normalized = value.strip().casefold().replace("-", "_")
    normalized = _ROLE_ALIASES.get(normalized, normalized)
    if normalized not in _VALID_ROLES:
        raise ValueError(f"unsupported profile role: {value!r}")
    return cast(ProfileRole, normalized)


def _confidence(value: float) -> Confidence:
    return Confidence(max(0.0, min(1.0, value)))


def _make_finding(
    *,
    context: RuleContext,
    rule_id: str,
    title: str,
    explanation: str,
    severity: Severity,
    confidence: float,
    worksheet: Worksheet,
    location: str,
    evidence: Evidence,
    expected: str,
    suggested_action: str,
    patches: Sequence[PatchOperation] = (),
    discriminator: Any = None,
) -> Finding:
    return Finding(
        id=stable_id("finding", rule_id, worksheet.title, location, discriminator),
        content_fingerprint=stable_id(
            "finding-content", evidence.model_dump(mode="json"), length=24
        ),
        rule_id=rule_id,
        title=title,
        explanation=explanation,
        severity=severity,
        confidence=_confidence(confidence),
        workbook=context.path.name,
        sheet=worksheet.title,
        location=location,
        evidence=evidence,
        expected=expected,
        suggested_action=suggested_action,
        safe_patch_available=any(patch.safe_only_eligible for patch in patches),
        patch_ids=[patch.id for patch in patches],
    )


def _make_review_text_patch(
    *, worksheet: Worksheet, cell: Cell, after: str, description: str
) -> PatchOperation:
    before = cell.value
    patch_id = stable_id(
        "patch", PatchKind.SET_TEXT.value, worksheet.title, cell.coordinate, before, after, None
    )
    return PatchOperation(
        id=patch_id,
        kind=PatchKind.SET_TEXT,
        sheet=worksheet.title,
        cell=cell.coordinate,
        before=before,
        after=after,
        confidence=_confidence(0.99),
        safe=False,
        risk=PatchRisk.LAYOUT_REVIEW,
        description=description,
        precondition=PatchPrecondition(
            cell_fingerprint=cell_fingerprint(cell),
            expected_value=cell.value,
            expected_style_id=cell.style_id,
        ),
    )


def _normalize_text(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(unicodedata.normalize("NFKC", value).strip().casefold().split())


def _tokens(value: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]+", value))


def _value_at(worksheet: Worksheet, row: int, column: int) -> Any:
    cell = worksheet._cells.get((row, column))
    return cell.value if isinstance(cell, (Cell, MergedCell)) else None


def _cell_at(worksheet: Worksheet, row: int, column: int) -> Cell | MergedCell | None:
    cell = worksheet._cells.get((row, column))
    return cell if isinstance(cell, (Cell, MergedCell)) else None


def _is_nonblank(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def _in_merged_range(worksheet: Worksheet, coordinate: str) -> bool:
    return any(coordinate in merged_range for merged_range in worksheet.merged_cells.ranges)


def _region_bounds(region: Region) -> tuple[int, int, int, int]:
    return region.min_row, region.max_row, region.min_column, region.max_column


def _profile_range(profile: TableProfile) -> tuple[int, int, int, int] | None:
    if profile.range_ref is None:
        return None
    return profile_range_bounds(profile.range_ref)


def _default_profile_bounds(
    worksheet: Worksheet, profile: TableProfile
) -> tuple[int, int, int, int] | None:
    header_row = profile.header_row or 1
    populated_cells = [
        cell
        for cell in worksheet._cells.values()
        if isinstance(cell, Cell) and _is_nonblank(cell.value)
    ]
    header_columns = {cell.column for cell in populated_cells if cell.row == header_row}
    configured_columns = {
        column
        for column_profile in profile.columns
        if (column := _resolve_column(column_profile.column)) is not None
    }
    columns = header_columns | configured_columns
    if not columns:
        return None
    min_column = min(columns)
    max_column = max(columns)
    body_cells = [
        cell
        for cell in populated_cells
        if cell.row >= header_row and min_column <= cell.column <= max_column
    ]
    if not body_cells:
        return None
    return header_row, max(cell.row for cell in body_cells), min_column, max_column


def _resolve_column(selector: int | str | None) -> int | None:
    if isinstance(selector, int):
        return selector
    if isinstance(selector, str):
        try:
            return column_index_from_string(selector)
        except ValueError:
            return None
    return None


def _body_rows(
    worksheet: Worksheet,
    *,
    header_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
) -> tuple[int, ...]:
    rows = {
        cell.row
        for cell in worksheet._cells.values()
        if isinstance(cell, Cell)
        and header_row < cell.row <= max_row
        and min_column <= cell.column <= max_column
        and _is_nonblank(cell.value)
    }
    return tuple(sorted(rows))


def _resolved_table(
    worksheet: Worksheet,
    bounds: tuple[int, int, int, int],
    *,
    table_profile: TableProfile | None,
) -> _ResolvedTable | None:
    min_row, max_row, min_column, max_column = bounds
    header_row = table_profile.header_row if table_profile and table_profile.header_row else min_row
    if not min_row <= header_row <= max_row:
        raise ProfileConfigurationError(
            "profile header_row must be inside the resolved table bounds"
        )
    headers = {
        column: _normalize_text(_value_at(worksheet, header_row, column))
        for column in range(min_column, max_column + 1)
    }
    if sum(bool(header) for header in headers.values()) < 1:
        if table_profile is not None:
            raise ProfileConfigurationError(
                f"profile worksheet {worksheet.title!r} row {header_row} "
                "does not contain a usable header"
            )
        return None
    body_rows = _body_rows(
        worksheet,
        header_row=header_row,
        max_row=max_row,
        min_column=min_column,
        max_column=max_column,
    )
    if not body_rows and table_profile is None:
        return None
    resolved_columns: dict[int, ColumnProfile] = {}
    if table_profile is not None:
        header_columns: dict[str, list[int]] = {}
        for column, header in headers.items():
            if header:
                header_columns.setdefault(header, []).append(column)
        for column_profile in table_profile.columns:
            resolved_column = _resolve_column(column_profile.column)
            if resolved_column is None and column_profile.header is not None:
                matches = header_columns.get(_normalize_text(column_profile.header), [])
                if not matches:
                    raise ProfileConfigurationError(
                        f"profile header {column_profile.header!r} was not found in "
                        f"{worksheet.title!r} row {header_row}"
                    )
                if len(matches) != 1:
                    raise ProfileConfigurationError(
                        f"profile header {column_profile.header!r} is ambiguous in "
                        f"{worksheet.title!r} row {header_row}"
                    )
                resolved_column = matches[0]
            if resolved_column is None:
                raise ProfileConfigurationError("profile column selector could not be resolved")
            if not min_column <= resolved_column <= max_column:
                raise ProfileConfigurationError(
                    f"profile column {resolved_column} is outside the configured table range"
                )
            if resolved_column in resolved_columns:
                raise ProfileConfigurationError(
                    f"profile contains multiple selectors for worksheet column {resolved_column}"
                )
            resolved_columns[resolved_column] = column_profile
    return _ResolvedTable(
        worksheet=worksheet,
        min_row=min_row,
        max_row=max_row,
        min_column=min_column,
        max_column=max_column,
        header_row=header_row,
        headers=headers,
        body_rows=body_rows,
        columns=resolved_columns,
        explicit=table_profile is not None,
    )


def _resolved_tables(context: RuleContext, profile: WorkbookProfile) -> tuple[_ResolvedTable, ...]:
    cache_key = "profile_quality.resolved_tables.v1"
    cached = context.analysis_cache.get(cache_key)
    if isinstance(cached, tuple) and all(isinstance(table, _ResolvedTable) for table in cached):
        return cached
    worksheets = {
        worksheet.title.casefold(): worksheet for worksheet in context.workbook.worksheets
    }
    resolved: list[_ResolvedTable] = []
    explicit_bounds: dict[str, list[tuple[int, int, int, int]]] = {}
    for table_profile in profile.tables:
        worksheet = worksheets.get(table_profile.sheet.casefold())
        if worksheet is None:
            raise ProfileConfigurationError(
                f"profile worksheet {table_profile.sheet!r} does not exist"
            )
        if worksheet.sheet_state != "visible":
            raise ProfileConfigurationError(
                f"profile worksheet {table_profile.sheet!r} is not visible"
            )
        configured_range = _profile_range(table_profile)
        if configured_range is not None:
            bounds_list = [configured_range]
        else:
            default_bounds = _default_profile_bounds(worksheet, table_profile)
            if default_bounds is not None:
                bounds_list = [default_bounds]
            else:
                bounds_list = [
                    _region_bounds(region)
                    for region in context.data_regions.get(worksheet.title, [])
                ]
        if not bounds_list:
            raise ProfileConfigurationError(
                f"profile table bounds could not be resolved for worksheet {worksheet.title!r}"
            )
        for bounds in bounds_list:
            table = _resolved_table(worksheet, bounds, table_profile=table_profile)
            if table is None:
                raise ProfileConfigurationError(
                    f"profile table could not be resolved for worksheet {worksheet.title!r}"
                )
            resolved.append(table)
            explicit_bounds.setdefault(worksheet.title, []).append(bounds)

    if profile.infer_semantics or profile.report_trailing_whitespace:
        for worksheet in context.workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            for region in context.data_regions.get(worksheet.title, []):
                bounds = _region_bounds(region)
                if any(
                    _bounds_overlap(bounds, other)
                    for other in explicit_bounds.get(worksheet.title, [])
                ):
                    continue
                table = _resolved_table(worksheet, bounds, table_profile=None)
                if table is None or len(table.body_rows) < 3:
                    continue
                if sum(bool(header) for header in table.headers.values()) < 2:
                    continue
                resolved.append(table)
    tables = tuple(resolved)
    context.analysis_cache[cache_key] = tables
    return tables


def _bounds_overlap(left: tuple[int, int, int, int], right: tuple[int, int, int, int]) -> bool:
    return not (
        left[1] < right[0] or right[1] < left[0] or left[3] < right[2] or right[3] < left[2]
    )


def _header_role(header: str) -> ProfileRole | None:
    if not header:
        return None
    words = _tokens(header)
    if words & _EMAIL_TOKENS or any(marker in header for marker in _EMAIL_MARKERS):
        return "email"
    if words & _PHONE_TOKENS or any(marker in header for marker in _PHONE_MARKERS):
        return "phone"
    if words & _DATE_TOKENS or any(marker in header for marker in _DATE_MARKERS):
        return "date"
    if (
        words & _PERCENTAGE_TOKENS
        or "%" in header
        or any(marker in header for marker in _PERCENTAGE_MARKERS)
    ):
        return "percentage"
    if words & _AMOUNT_TOKENS or any(marker in header for marker in _AMOUNT_MARKERS):
        return "currency"
    if words & _IDENTIFIER_TOKENS or any(marker in header for marker in _IDENTIFIER_MARKERS):
        return "identifier"
    return None


def _column_role(
    table: _ResolvedTable, column: int, profile: WorkbookProfile
) -> ProfileRole | None:
    configured = table.columns.get(column)
    if configured is not None and configured.role is not None:
        return configured.role
    return _header_role(table.headers.get(column, "")) if profile.infer_semantics else None


def _enum_key(value: Any) -> tuple[str, str]:
    if isinstance(value, str):
        return "text", _normalize_text(value)
    if isinstance(value, bool):
        return "bool", "true" if value else "false"
    if isinstance(value, int):
        return "number", str(value)
    if isinstance(value, float) and math.isfinite(value):
        return "number", format(value, ".15g")
    return type(value).__name__, repr(value)


def _near_enum_candidate(value: Any, allowed_values: tuple[Any, ...]) -> tuple[Any, float] | None:
    if not isinstance(value, str):
        return None
    normalized = _normalize_text(value)
    if len(normalized) < 3:
        return None
    scores = sorted(
        (
            (SequenceMatcher(None, normalized, _normalize_text(candidate)).ratio(), candidate)
            for candidate in allowed_values
            if isinstance(candidate, str) and _normalize_text(candidate)
        ),
        key=lambda item: (-item[0], str(item[1])),
    )
    if not scores or scores[0][0] < 0.8:
        return None
    if len(scores) > 1 and scores[0][0] - scores[1][0] < 0.08:
        return None
    return scores[0][1], scores[0][0]


def _valid_email(value: Any) -> bool:
    return isinstance(value, str) and _EMAIL_RE.fullmatch(value.strip()) is not None


def _valid_phone(value: Any) -> bool:
    if isinstance(value, bool) or value is None:
        return False
    if isinstance(value, int):
        normalized = str(value)
    elif isinstance(value, float) and math.isfinite(value) and value.is_integer():
        normalized = str(int(value))
    elif isinstance(value, str):
        normalized = unicodedata.normalize("NFKC", value).strip()
    else:
        return False
    normalized = _PHONE_EXTENSION_RE.sub("", normalized).strip()
    if not normalized or _PHONE_BODY_RE.fullmatch(normalized) is None:
        return False
    if normalized.count("+") > 1 or ("+" in normalized and not normalized.startswith("+")):
        return False
    if normalized.count("(") != normalized.count(")"):
        return False
    parenthesis_depth = 0
    for character in normalized:
        if character == "(":
            parenthesis_depth += 1
        elif character == ")":
            parenthesis_depth -= 1
            if parenthesis_depth < 0:
                return False
    digits = re.sub(r"\D", "", normalized)
    return 7 <= len(digits) <= 15


def _integer_digits(value: Any) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int) and value >= 0:
        return str(value)
    if isinstance(value, float) and math.isfinite(value) and value >= 0 and value.is_integer():
        return str(int(value))
    if isinstance(value, str):
        normalized = unicodedata.normalize("NFKC", value).strip()
        return normalized if re.fullmatch(r"\d+", normalized) else None
    return None


def _zero_mask_width(number_format: str) -> int | None:
    sanitized = _sanitize_number_format(number_format).split(";", 1)[0]
    if not sanitized or "#" in sanitized or "?" in sanitized or "%" in sanitized:
        return None
    if re.search(r"[dmyhs]", sanitized, re.I):
        return None
    if re.fullmatch(r"[0\-(). /]+", sanitized) is None:
        return None
    count = sanitized.count("0")
    return count or None


def _inferred_identifier_width(table: _ResolvedTable, column: int) -> int | None:
    zero_prefixed_widths: Counter[int] = Counter()
    all_widths: Counter[int] = Counter()
    for row in table.body_rows:
        value = _value_at(table.worksheet, row, column)
        digits = _integer_digits(value)
        if digits is not None and len(digits) >= 2:
            all_widths[len(digits)] += 1
            if isinstance(value, str) and digits.startswith("0"):
                zero_prefixed_widths[len(digits)] += 1
    for widths, minimum, ratio in (
        (zero_prefixed_widths, 4, 0.75),
        (all_widths, 6, 0.8),
    ):
        total = sum(widths.values())
        if total < minimum:
            continue
        width, count = widths.most_common(1)[0]
        if count / total >= ratio:
            return width
    return None


def _sanitize_number_format(number_format: str) -> str:
    without_literals = _QUOTED_FORMAT_TEXT_RE.sub("", number_format)
    without_escapes = _ESCAPED_FORMAT_CHARACTER_RE.sub("", without_literals)
    return _FORMAT_SPACING_RE.sub("", without_escapes)


def _number_format_traits(number_format: str) -> _NumberFormatTraits:
    sanitized = _sanitize_number_format(number_format)
    return _NumberFormatTraits(
        percentage="%" in sanitized,
        date=is_date_format(number_format),
        currency=_CURRENCY_RE.search(sanitized) is not None,
        text=sanitized.strip() == "@",
    )


def _format_conflicts(role: ProfileRole, traits: _NumberFormatTraits) -> tuple[str, ...]:
    conflicts: list[str] = []
    if traits.text:
        conflicts.append("text")
    if role == "currency":
        if traits.percentage:
            conflicts.append("percentage")
        if traits.date:
            conflicts.append("date")
    elif role == "percentage":
        if traits.currency:
            conflicts.append("currency")
        if traits.date:
            conflicts.append("date")
    elif role == "date":
        if traits.percentage:
            conflicts.append("percentage")
        if traits.currency:
            conflicts.append("currency")
    return tuple(dict.fromkeys(conflicts))


class TrailingWhitespaceRule(WorkbookRule):
    rule_id = "WL036_TRAILING_WHITESPACE"
    title = "Trailing whitespace in a record field"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        profile = _context_profile(context)
        if not profile.report_trailing_whitespace:
            return result
        for table in _resolved_tables(context, profile):
            for row in table.body_rows:
                for column in range(table.min_column, table.max_column + 1):
                    cell = _cell_at(table.worksheet, row, column)
                    if not isinstance(cell, Cell) or cell.data_type == "f":
                        continue
                    value = cell.value
                    if not isinstance(value, str) or not value.strip():
                        continue
                    match = _HORIZONTAL_TRAILING_WHITESPACE_RE.search(value)
                    if match is None:
                        continue
                    trimmed = value[: match.start()]
                    if not trimmed:
                        continue
                    merged = _in_merged_range(table.worksheet, cell.coordinate)
                    configured = table.columns.get(column)
                    patch_requested = profile.review_trailing_whitespace_patches or bool(
                        configured and configured.trim_trailing_whitespace
                    )
                    patches: list[PatchOperation] = []
                    if patch_requested and not merged and set(match.group()) <= {" "}:
                        patch = _make_review_text_patch(
                            worksheet=table.worksheet,
                            cell=cell,
                            after=trimmed,
                            description=(
                                "Remove configured ASCII trailing spaces after explicit review; "
                                "the source text and style preconditions must still match."
                            ),
                        )
                        result.patches.append(patch)
                        patches.append(patch)
                    evidence = Evidence(
                        summary="A literal record value ends with horizontal whitespace",
                        observed={
                            "value": value,
                            "trailing_codepoints": [f"U+{ord(char):04X}" for char in match.group()],
                        },
                        expected={"value": trimmed},
                        details={
                            "evidence_level": "PROVEN_STATIC",
                            "explicit_profile": table.explicit,
                            "merged_cell": merged,
                            "review_patch_requested": patch_requested,
                        },
                    )
                    result.findings.append(
                        _make_finding(
                            context=context,
                            rule_id=self.rule_id,
                            title=self.title,
                            explanation=(
                                "A literal text field has horizontal whitespace after its final "
                                "visible character, which can break exact matching and deduplication."
                            ),
                            severity=Severity.WARNING,
                            confidence=0.99 if table.explicit else 0.96,
                            worksheet=table.worksheet,
                            location=cell.coordinate,
                            evidence=evidence,
                            expected="Record text has no unintended horizontal whitespace at the end.",
                            suggested_action=(
                                "Review and apply the proposed text patch; it is never selected by "
                                "safe-only repair."
                                if patches
                                else "Review the source text and remove the trailing whitespace manually if unintended."
                            ),
                            patches=patches,
                        )
                    )
        return result


class RequiredFieldRule(WorkbookRule):
    rule_id = "WL037_REQUIRED_FIELD"
    title = "Configured required field is blank"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        profile = _context_profile(context)
        for table in _resolved_tables(context, profile):
            for column, column_profile in table.columns.items():
                if not column_profile.required:
                    continue
                for row in table.body_rows:
                    coordinate = f"{get_column_letter(column)}{row}"
                    if _is_nonblank(_value_at(table.worksheet, row, column)) or _in_merged_range(
                        table.worksheet, coordinate
                    ):
                        continue
                    evidence = Evidence(
                        summary="A populated record is blank in a configured required field",
                        observed=None,
                        expected="nonblank value",
                        details={
                            "evidence_level": "PROVEN_STATIC",
                            "header": table.headers.get(column, ""),
                        },
                    )
                    result.findings.append(
                        _make_finding(
                            context=context,
                            rule_id=self.rule_id,
                            title=self.title,
                            explanation=(
                                "The workbook profile marks this field as required, and the row "
                                "contains other record data but no value in this field."
                            ),
                            severity=Severity.ERROR,
                            confidence=1.0,
                            worksheet=table.worksheet,
                            location=coordinate,
                            evidence=evidence,
                            expected="Every populated record has a nonblank value in configured required fields.",
                            suggested_action=(
                                "Supply or confirm the business value manually; WorkbookLens does not "
                                "invent required-field contents."
                            ),
                        )
                    )
        return result


class EnumeratedValueRule(WorkbookRule):
    rule_id = "WL038_ENUM_VALUE"
    title = "Value is outside a configured enumeration"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        profile = _context_profile(context)
        for table in _resolved_tables(context, profile):
            for column, column_profile in table.columns.items():
                if not column_profile.allowed_values:
                    continue
                allowed_keys = {_enum_key(value) for value in column_profile.allowed_values}
                for row in table.body_rows:
                    cell = _cell_at(table.worksheet, row, column)
                    if not isinstance(cell, Cell) or cell.data_type in {"e", "f"}:
                        continue
                    if not _is_nonblank(cell.value) or _enum_key(cell.value) in allowed_keys:
                        continue
                    near = _near_enum_candidate(cell.value, column_profile.allowed_values)
                    details: dict[str, Any] = {
                        "evidence_level": "PROVEN_STATIC",
                        "header": table.headers.get(column, ""),
                    }
                    if near is not None:
                        details["near_match"] = near[0]
                        details["similarity"] = round(near[1], 3)
                    evidence = Evidence(
                        summary="A nonblank value is absent from the configured allowed-value set",
                        observed=cell.value,
                        expected=list(column_profile.allowed_values),
                        details=details,
                    )
                    result.findings.append(
                        _make_finding(
                            context=context,
                            rule_id=self.rule_id,
                            title=self.title,
                            explanation=(
                                "The literal value does not match any configured enumeration value "
                                "after Unicode, whitespace, and case normalization."
                            ),
                            severity=Severity.ERROR,
                            confidence=1.0,
                            worksheet=table.worksheet,
                            location=cell.coordinate,
                            evidence=evidence,
                            expected="Enumerated fields use one of the values declared in the workbook profile.",
                            suggested_action=(
                                "Review the possible near match and choose the intended value manually; "
                                "no spelling correction is applied automatically."
                                if near is not None
                                else "Choose an allowed value or update the workbook profile if this is a legitimate category."
                            ),
                            discriminator=_enum_key(cell.value),
                        )
                    )
        return result


class ContactFormatRule(WorkbookRule):
    rule_id = "WL039_CONTACT_FORMAT"
    title = "Contact value has an invalid format"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        profile = _context_profile(context)
        for table in _resolved_tables(context, profile):
            for column in range(table.min_column, table.max_column + 1):
                role = _column_role(table, column, profile)
                if role not in {"email", "phone"}:
                    continue
                validator = _valid_email if role == "email" else _valid_phone
                for row in table.body_rows:
                    cell = _cell_at(table.worksheet, row, column)
                    if not isinstance(cell, Cell) or cell.data_type in {"e", "f"}:
                        continue
                    if not _is_nonblank(cell.value) or validator(cell.value):
                        continue
                    evidence = Evidence(
                        summary=f"A nonblank {role} value fails conservative structural validation",
                        observed=cell.value,
                        expected=role,
                        details={
                            "evidence_level": (
                                "PROVEN_STATIC" if column in table.columns else "STRONG_STRUCTURAL"
                            ),
                            "header": table.headers.get(column, ""),
                            "configured_role": column in table.columns,
                        },
                    )
                    result.findings.append(
                        _make_finding(
                            context=context,
                            rule_id=self.rule_id,
                            title=self.title,
                            explanation=(
                                "The value does not satisfy the conservative structural validator "
                                "for the configured or strongly inferred contact-field role."
                            ),
                            severity=Severity.WARNING,
                            confidence=0.99 if column in table.columns else 0.9,
                            worksheet=table.worksheet,
                            location=cell.coordinate,
                            evidence=evidence,
                            expected="Contact fields contain structurally valid email addresses or phone numbers.",
                            suggested_action=(
                                "Confirm the intended contact value manually; WorkbookLens does not "
                                "rewrite personal contact information."
                            ),
                            discriminator=role,
                        )
                    )
        return result


class LeadingZeroIdentifierRule(WorkbookRule):
    rule_id = "WL040_LEADING_ZERO_IDENTIFIER"
    title = "Identifier width conflicts with leading-zero semantics"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        profile = _context_profile(context)
        for table in _resolved_tables(context, profile):
            for column in range(table.min_column, table.max_column + 1):
                role = _column_role(table, column, profile)
                if role != "identifier":
                    continue
                column_profile = table.columns.get(column)
                width = column_profile.identifier_width if column_profile else None
                explicitly_required = bool(
                    column_profile and (column_profile.preserve_leading_zeros or width is not None)
                )
                if width is None:
                    width = _inferred_identifier_width(table, column)
                if width is None:
                    continue
                for row in table.body_rows:
                    cell = _cell_at(table.worksheet, row, column)
                    if not isinstance(cell, Cell) or cell.data_type in {"e", "f"}:
                        continue
                    digits = _integer_digits(cell.value)
                    if digits is None or len(digits) == width:
                        continue
                    zero_mask_width = _zero_mask_width(cell.number_format)
                    if (
                        len(digits) < width
                        and zero_mask_width is not None
                        and zero_mask_width >= width
                    ):
                        continue
                    if (
                        len(digits) > width
                        and not explicitly_required
                        and not digits.startswith("0")
                    ):
                        continue
                    issue_kind = (
                        "too_short"
                        if len(digits) < width
                        else "extra_leading_zero_padding"
                        if digits.startswith("0")
                        else "too_long"
                    )
                    evidence = Evidence(
                        summary="A fixed-width identifier does not match its configured or inferred width",
                        observed={
                            "value": cell.value,
                            "digit_count": len(digits),
                            "number_format": cell.number_format,
                        },
                        expected={"width": width, "preserve_leading_zeros": True},
                        details={
                            "evidence_level": (
                                "PROVEN_STATIC" if explicitly_required else "STRONG_STRUCTURAL"
                            ),
                            "header": table.headers.get(column, ""),
                            "configured_width": bool(
                                column_profile and column_profile.identifier_width
                            ),
                            "issue_kind": issue_kind,
                        },
                    )
                    result.findings.append(
                        _make_finding(
                            context=context,
                            rule_id=self.rule_id,
                            title=self.title,
                            explanation=(
                                "The identifier digit count conflicts with the configured width or a "
                                "strong fixed-width peer pattern, including possible missing or extra leading zeros."
                            ),
                            severity=Severity.WARNING,
                            confidence=0.99 if explicitly_required else 0.9,
                            worksheet=table.worksheet,
                            location=cell.coordinate,
                            evidence=evidence,
                            expected="Fixed-width identifiers use the declared or strongly inferred digit width consistently.",
                            suggested_action=(
                                "Confirm the identifier width and restore the original identifier manually; "
                                "WorkbookLens does not guess missing or extra digits."
                            ),
                            discriminator=(width, issue_kind),
                        )
                    )
        return result


class NumberFormatRoleConflictRule(WorkbookRule):
    rule_id = "WL041_NUMBER_FORMAT_ROLE_CONFLICT"
    title = "Number format conflicts with the field role"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        profile = _context_profile(context)
        for table in _resolved_tables(context, profile):
            for column in range(table.min_column, table.max_column + 1):
                role = _column_role(table, column, profile)
                if role not in {"currency", "percentage", "date"}:
                    continue
                configured_role = bool(
                    (column_profile := table.columns.get(column))
                    and column_profile.role is not None
                )
                for row in table.body_rows:
                    cell = _cell_at(table.worksheet, row, column)
                    if (
                        not isinstance(cell, Cell)
                        or cell.data_type == "e"
                        or not _is_nonblank(cell.value)
                    ):
                        continue
                    traits = _number_format_traits(cell.number_format)
                    conflicts = _format_conflicts(role, traits)
                    if not conflicts:
                        continue
                    evidence = Evidence(
                        summary=f"A {role} field uses a conflicting number-format role",
                        observed={
                            "number_format": cell.number_format,
                            "conflicting_roles": list(conflicts),
                        },
                        expected={"role": role},
                        details={
                            "evidence_level": (
                                "PROVEN_STATIC" if configured_role else "STRONG_STRUCTURAL"
                            ),
                            "header": table.headers.get(column, ""),
                            "configured_role": configured_role,
                        },
                    )
                    result.findings.append(
                        _make_finding(
                            context=context,
                            rule_id=self.rule_id,
                            title=self.title,
                            explanation=(
                                "The cell number format expresses a semantic role that contradicts "
                                "the configured or strongly inferred column role."
                            ),
                            severity=Severity.WARNING,
                            confidence=0.99 if configured_role else 0.92,
                            worksheet=table.worksheet,
                            location=cell.coordinate,
                            evidence=evidence,
                            expected="Number formats do not contradict amount, percentage, or date field semantics.",
                            suggested_action=(
                                "Review the stored value and intended display role before changing the number format; "
                                "no format is copied automatically."
                            ),
                            discriminator=(role, conflicts),
                        )
                    )
        return result


PROFILE_QUALITY_RULES: tuple[type[WorkbookRule], ...] = (
    TrailingWhitespaceRule,
    RequiredFieldRule,
    EnumeratedValueRule,
    ContactFormatRule,
    LeadingZeroIdentifierRule,
    NumberFormatRoleConflictRule,
)

PROFILE_QUALITY_RULE_IDS = frozenset(rule.rule_id for rule in PROFILE_QUALITY_RULES)

__all__ = [
    "MAX_PROFILE_RANGE_CELLS",
    "PROFILE_QUALITY_RULES",
    "PROFILE_QUALITY_RULE_IDS",
    "ColumnProfile",
    "ContactFormatRule",
    "EnumeratedValueRule",
    "LeadingZeroIdentifierRule",
    "NumberFormatRoleConflictRule",
    "ProfileConfigurationError",
    "RequiredFieldRule",
    "TableProfile",
    "TrailingWhitespaceRule",
    "WorkbookProfile",
    "normalize_profile_column_selector",
    "normalize_profile_range",
    "parse_workbook_profile",
    "profile_range_bounds",
]
