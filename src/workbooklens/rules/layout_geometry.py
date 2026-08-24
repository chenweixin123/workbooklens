"""Conservative worksheet geometry and role-aware layout diagnostics."""

from __future__ import annotations

import math
import re
import statistics
from collections import Counter, defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime, time
from typing import Any, cast

from openpyxl.cell.cell import Cell
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from workbooklens.formulas.ir import worksheet_by_name
from workbooklens.layout import (
    column_width,
    estimated_text_width,
    is_detail_row_perimeter_border_variant,
    measure_text_cell,
    row_height,
)
from workbooklens.models import Confidence, Evidence, Finding, Region, Severity
from workbooklens.rules.base import RuleContext, RuleResult, WorkbookRule
from workbooklens.utils import stable_id
from workbooklens.worksheet_state import is_column_hidden

_DIRECT_RANGE_RE = re.compile(
    r"^\s*(?:(?P<sheet>'(?:[^']|'')+'|[^'!]+)!)?"
    r"(?P<range>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)\s*$",
    re.IGNORECASE,
)
_SUMMARY_RE = re.compile(
    r"(?<![A-Z0-9_])(?:grand[\s-]+total|subtotal|total|average|summary)(?![A-Z0-9_])"
    r"|(?:合计|总计|小计|汇总|平均)",
    re.IGNORECASE,
)
_MAX_ANCHOR_COLUMNS = 512
_MAX_ANCHOR_ROWS = 4096


@dataclass(frozen=True, slots=True)
class _Box:
    min_column: int
    min_row: int
    max_column: int
    max_row: int

    @property
    def coordinate(self) -> str:
        start = f"{get_column_letter(self.min_column)}{self.min_row}"
        end = f"{get_column_letter(self.max_column)}{self.max_row}"
        return start if start == end else f"{start}:{end}"


def _confidence(value: float) -> Confidence:
    return Confidence(max(0.0, min(1.0, value)))


def _finding(
    *,
    context: RuleContext,
    rule_id: str,
    title: str,
    explanation: str,
    severity: Severity,
    confidence: float,
    worksheet: Worksheet,
    location: str,
    evidence: Evidence,
    expected: str,
    suggested_action: str,
    discriminator: Any = None,
) -> Finding:
    return Finding(
        id=stable_id("finding", rule_id, worksheet.title, location, discriminator),
        content_fingerprint=stable_id(
            "finding-content", evidence.model_dump(mode="json"), length=24
        ),
        rule_id=rule_id,
        title=title,
        explanation=explanation,
        severity=severity,
        confidence=_confidence(confidence),
        workbook=context.path.name,
        sheet=worksheet.title,
        location=location,
        evidence=evidence,
        expected=expected,
        suggested_action=suggested_action,
    )


def _visible_cell(worksheet: Worksheet, cell: Cell) -> bool:
    row_dimension = worksheet.row_dimensions.get(cell.row)
    return not (
        (row_dimension is not None and row_dimension.hidden)
        or is_column_hidden(worksheet, cell.column)
    )


def _box_intersects_region(box: _Box, region: Region) -> bool:
    return not (
        box.max_row < region.min_row
        or region.max_row < box.min_row
        or box.max_column < region.min_column
        or region.max_column < box.min_column
    )


def _cell_in_range(cell: Cell, cell_range: CellRange) -> bool:
    return (
        cell_range.min_row <= cell.row <= cell_range.max_row
        and cell_range.min_col <= cell.column <= cell_range.max_col
    )


def _column_pixels(worksheet: Worksheet, column: int) -> float:
    width = column_width(worksheet, column)
    return math.floor(((256.0 * width + math.floor(128.0 / 7.0)) / 256.0) * 7.0)


def _row_pixels(worksheet: Worksheet, row: int) -> float:
    return row_height(worksheet, row) * 96.0 / 72.0


def _advance_axis(
    worksheet: Worksheet,
    *,
    start: int,
    start_offset: float,
    extent: float,
    columns: bool,
) -> int | None:
    if extent <= 0 or start < 1:
        return None
    limit = _MAX_ANCHOR_COLUMNS if columns else _MAX_ANCHOR_ROWS
    current = start
    remaining = extent
    for _ in range(limit):
        size = _column_pixels(worksheet, current) if columns else _row_pixels(worksheet, current)
        available = max(0.1, size - start_offset)
        if remaining <= available + 0.01:
            return current
        remaining -= available
        current += 1
        start_offset = 0.0
    return None


def _drawing_box(worksheet: Worksheet, drawing: Any) -> _Box | None:
    anchor = getattr(drawing, "anchor", None)
    start = getattr(anchor, "_from", None)
    if start is None:
        return None
    min_column = int(start.col) + 1
    min_row = int(start.row) + 1
    end = getattr(anchor, "to", None)
    if end is not None:
        max_column = int(end.col) + (1 if int(getattr(end, "colOff", 0)) > 0 else 0)
        max_row = int(end.row) + (1 if int(getattr(end, "rowOff", 0)) > 0 else 0)
        return _Box(
            min_column=min_column,
            min_row=min_row,
            max_column=max(min_column, max_column),
            max_row=max(min_row, max_row),
        )
    extent = getattr(anchor, "ext", None)
    width_emu = getattr(extent, "cx", None)
    height_emu = getattr(extent, "cy", None)
    if not isinstance(width_emu, int) or not isinstance(height_emu, int):
        return None
    advanced_column = _advance_axis(
        worksheet,
        start=min_column,
        start_offset=int(getattr(start, "colOff", 0)) / 9525.0,
        extent=width_emu / 9525.0,
        columns=True,
    )
    advanced_row = _advance_axis(
        worksheet,
        start=min_row,
        start_offset=int(getattr(start, "rowOff", 0)) / 9525.0,
        extent=height_emu / 9525.0,
        columns=False,
    )
    if advanced_column is None or advanced_row is None:
        return None
    return _Box(min_column, min_row, advanced_column, advanced_row)


def _reference_formulas(source: Any) -> Iterable[str]:
    if source is None:
        return ()
    formulas: list[str] = []
    for attribute in ("numRef", "strRef"):
        reference = getattr(source, attribute, None)
        formula = getattr(reference, "f", None)
        if isinstance(formula, str) and formula.strip():
            formulas.append(formula)
    return formulas


def _chart_reference_formulas(chart: Any) -> tuple[str, ...]:
    formulas: list[str] = []
    for series in getattr(chart, "ser", ()):
        for attribute in ("val", "yVal", "xVal", "cat", "tx"):
            formulas.extend(_reference_formulas(getattr(series, attribute, None)))
    return tuple(dict.fromkeys(formulas))


def _unquote_sheet(value: str) -> str:
    stripped = value.strip()
    if stripped.startswith("'") and stripped.endswith("'"):
        return stripped[1:-1].replace("''", "'")
    return stripped


def _defined_name_target(
    context: RuleContext,
    worksheet: Worksheet,
    candidate: str,
) -> str | None:
    folded_names = {
        name.casefold(): value for name, value in context.snapshot.defined_names.items()
    }
    try:
        local_sheet_id = context.workbook.index(worksheet)
    except ValueError:
        local_sheet_id = None
    if local_sheet_id is not None:
        local_key = f"{candidate}@sheet:{local_sheet_id}".casefold()
        if local_key in folded_names:
            return folded_names[local_key]
    return folded_names.get(candidate.casefold())


def _resolve_range(
    context: RuleContext,
    worksheet: Worksheet,
    formula: str,
) -> tuple[str, CellRange] | None:
    candidate = formula.strip()
    match = _DIRECT_RANGE_RE.fullmatch(candidate)
    if match is None:
        defined_range = _defined_name_target(context, worksheet, candidate)
        if defined_range is not None:
            candidate = defined_range
            match = _DIRECT_RANGE_RE.fullmatch(candidate)
    if match is None:
        return None
    requested_sheet = _unquote_sheet(match.group("sheet") or worksheet.title)
    resolved_worksheet = worksheet_by_name(context.workbook, requested_sheet)
    if resolved_worksheet is None:
        return None
    range_text = match.group("range").replace("$", "")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(range_text)
    except ValueError:
        return None
    if not all(isinstance(item, int) for item in (min_col, min_row, max_col, max_row)):
        return None
    return resolved_worksheet.title, CellRange(
        min_col=cast(int, min_col),
        min_row=cast(int, min_row),
        max_col=cast(int, max_col),
        max_row=cast(int, max_row),
    )


def _chart_sources(
    context: RuleContext,
    worksheet: Worksheet,
    chart: Any,
) -> tuple[CellRange, ...]:
    ranges: list[CellRange] = []
    for formula in _chart_reference_formulas(chart):
        resolved = _resolve_range(context, worksheet, formula)
        if resolved is not None and resolved[0] == worksheet.title:
            ranges.append(resolved[1])
    return tuple(ranges)


def _region_is_kpi(worksheet: Worksheet, region: Region) -> bool:
    width = region.max_column - region.min_column + 1
    height = region.max_row - region.min_row + 1
    if width != 2 or not 2 <= height <= 15:
        return False
    pairs = 0
    for row in range(region.min_row, region.max_row + 1):
        left = worksheet._cells.get((row, region.min_column))
        right = worksheet._cells.get((row, region.max_column))
        if not isinstance(left, Cell) or not isinstance(left.value, str) or not left.value.strip():
            continue
        if not isinstance(right, Cell) or right.value is None:
            continue
        if right.data_type == "f" or (
            isinstance(right.value, (int, float)) and not isinstance(right.value, bool)
        ):
            pairs += 1
    return pairs >= max(2, math.ceil(height * 0.6))


class DrawingContentOverlapRule(WorkbookRule):
    """Report charts or images whose saved frame covers non-source tabular content."""

    rule_id = "WL047_DRAWING_CONTENT_OVERLAP"
    title = "Drawing overlaps non-source worksheet content"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for raw_worksheet in context.workbook.worksheets:
            worksheet = cast(Worksheet, raw_worksheet)
            if worksheet.sheet_state != "visible":
                continue
            drawings = [
                *(
                    ("chart", index, item)
                    for index, item in enumerate(getattr(worksheet, "_charts", ()), start=1)
                ),
                *(
                    ("image", index, item)
                    for index, item in enumerate(getattr(worksheet, "_images", ()), start=1)
                ),
            ]
            for kind, index, drawing in drawings:
                box = _drawing_box(worksheet, drawing)
                if box is None:
                    continue
                regions = [
                    region
                    for region in context.data_regions.get(worksheet.title, ())
                    if _box_intersects_region(box, region)
                ]
                if not regions:
                    continue
                sources = _chart_sources(context, worksheet, drawing) if kind == "chart" else ()
                overlapped: list[Cell] = []
                kpi_overlap = False
                for region in regions:
                    region_cells = [
                        cell
                        for cell in worksheet._cells.values()
                        if isinstance(cell, Cell)
                        and cell.value is not None
                        and _visible_cell(worksheet, cell)
                        and region.min_row <= cell.row <= region.max_row
                        and region.min_column <= cell.column <= region.max_column
                        and box.min_row <= cell.row <= box.max_row
                        and box.min_column <= cell.column <= box.max_column
                        and not any(_cell_in_range(cell, source) for source in sources)
                    ]
                    if region_cells:
                        overlapped.extend(region_cells)
                        kpi_overlap = kpi_overlap or _region_is_kpi(worksheet, region)
                unique_cells = sorted(
                    {cell.coordinate: cell for cell in overlapped}.values(),
                    key=lambda cell: (cell.row, cell.column),
                )
                if len(unique_cells) < 2 and not (kpi_overlap and unique_cells):
                    continue
                coordinates = [cell.coordinate for cell in unique_cells]
                evidence = Evidence(
                    summary=(
                        f"{kind.title()} {index} covers {len(coordinates)} populated non-source cells"
                    ),
                    observed={
                        "drawing_kind": kind,
                        "drawing_index": index,
                        "anchor_range": box.coordinate,
                        "overlapped_cells": coordinates[:50],
                    },
                    expected={"non_source_populated_cells_covered": 0},
                    peers=[str(source) for source in sources],
                    details={"kpi_region_overlap": kpi_overlap},
                )
                result.findings.append(
                    _finding(
                        context=context,
                        rule_id=self.rule_id,
                        title=self.title,
                        explanation=(
                            "The saved drawing frame intersects populated cells in an inferred data or KPI "
                            "region after excluding cells used by the chart itself."
                        ),
                        severity=Severity.WARNING,
                        confidence=0.97 if kpi_overlap else 0.9,
                        worksheet=worksheet,
                        location=box.coordinate,
                        evidence=evidence,
                        expected="Charts and images do not obscure non-source table or KPI content.",
                        suggested_action=(
                            "Review the drawing position and size in Excel; WorkbookLens does not move or "
                            "resize drawings automatically."
                        ),
                        discriminator=(kind, index),
                    )
                )
        return result


def _format_section(number_format: str, value: float) -> str:
    sections = number_format.split(";")
    if value < 0 and len(sections) >= 2:
        return sections[1]
    if value == 0 and len(sections) >= 3:
        return sections[2]
    return sections[0]


def _clean_format(format_code: str) -> str:
    code = re.sub(r"\[(?!h\]|m\]|s\])[^\]]+\]", "", format_code, flags=re.IGNORECASE)
    code = re.sub(r'"([^"]*)"', r"\1", code)
    code = re.sub(r"\\(.)", r"\1", code)
    code = re.sub(r"[_*].", "", code)
    return code


def _date_skeleton(format_code: str) -> str:
    code = _clean_format(format_code)
    replacements = {
        "am/pm": "PM",
        "a/p": "P",
        "yyyy": "2026",
        "yyy": "2026",
        "yy": "26",
        "mmmm": "September",
        "mmm": "Sep",
        "mm": "09",
        "m": "9",
        "dddd": "Wednesday",
        "ddd": "Wed",
        "dd": "28",
        "d": "8",
        "hh": "23",
        "h": "9",
        "ss": "59",
        "s": "9",
    }
    tokens = re.compile(
        r"AM/PM|A/P|yyyy|yyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s",
        re.IGNORECASE,
    )
    return tokens.sub(lambda match: replacements[match.group(0).casefold()], code).strip()


def _numeric_display_text(cell: Cell) -> str | None:
    value = cell.value
    number_format = (cell.number_format or "General").strip()
    if number_format.casefold() in {"general", "@", "text"}:
        return None
    if isinstance(value, (datetime, date, time)):
        return _date_skeleton(number_format) if is_date_format(number_format) else None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        return None
    section = _clean_format(_format_section(number_format, float(value)))
    if is_date_format(number_format):
        return _date_skeleton(section)
    if re.search(r"[Ee][+-]0+", section):
        decimal_match = re.search(r"\.([0#]+)", section)
        decimals = len(decimal_match.group(1)) if decimal_match is not None else 0
        return f"{value:.{decimals}E}"
    percent = "%" in section
    rendered_value = float(value) * 100.0 if percent else float(value)
    decimal_match = re.search(r"\.([0#]+)", section)
    decimals = len(decimal_match.group(1)) if decimal_match else 0
    grouped = "," in section.split(".", maxsplit=1)[0]
    rendered = (
        f"{abs(rendered_value):,.{decimals}f}" if grouped else f"{abs(rendered_value):.{decimals}f}"
    )
    if rendered_value < 0:
        rendered = f"({rendered})" if "(" in section and ")" in section else f"-{rendered}"
    symbols = "".join(character for character in section if character in "$¥￥€£")
    return f"{symbols}{rendered}{'%' if percent else ''}"


class NumericDisplayWidthRiskRule(WorkbookRule):
    """Report fixed-format numeric columns likely to render as hash marks."""

    rule_id = "WL048_NUMERIC_DISPLAY_WIDTH_RISK"
    title = "Fixed-format numeric value may not fit its column"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            for region in context.data_regions.get(worksheet.title, ()):
                if region.max_row - region.min_row < 5:
                    continue
                for column in range(region.min_column, region.max_column + 1):
                    cells = [
                        cell
                        for row in range(region.min_row + 1, region.max_row + 1)
                        if isinstance((cell := worksheet._cells.get((row, column))), Cell)
                        and cell.value is not None
                        and _visible_cell(worksheet, cell)
                    ]
                    numeric_cells = [
                        cell
                        for cell in cells
                        if not isinstance(cell.value, bool)
                        and isinstance(cell.value, (int, float, datetime, date, time))
                    ]
                    if len(numeric_cells) < 5 or len(numeric_cells) / max(1, len(cells)) < 0.7:
                        continue
                    risks: list[tuple[Cell, str, float]] = []
                    for cell in numeric_cells:
                        display = _numeric_display_text(cell)
                        if not display:
                            continue
                        required = estimated_text_width(cell, display)
                        available = column_width(worksheet, column)
                        if required >= available * 1.15 + 0.5:
                            risks.append((cell, display, required))
                    if not risks:
                        continue
                    coordinates = [cell.coordinate for cell, _, _ in risks]
                    evidence = Evidence(
                        summary=(
                            f"{len(risks)} fixed-format numeric values exceed the estimated column width"
                        ),
                        observed={
                            "column": get_column_letter(column),
                            "column_width": column_width(worksheet, column),
                            "cells": [
                                {
                                    "cell": cell.coordinate,
                                    "display_sample": display,
                                    "estimated_width": round(required, 2),
                                    "number_format": cell.number_format,
                                }
                                for cell, display, required in risks[:25]
                            ],
                        },
                        expected={"estimated_display_fits": True},
                        peers=[cell.coordinate for cell in numeric_cells[:12]],
                    )
                    result.findings.append(
                        _finding(
                            context=context,
                            rule_id=self.rule_id,
                            title=self.title,
                            explanation=(
                                "A numeric or date-dominated column uses explicit display formats whose "
                                "estimated rendered text is wider than the saved column. Excel may show "
                                "hash marks or clipped values."
                            ),
                            severity=Severity.WARNING,
                            confidence=0.93,
                            worksheet=worksheet,
                            location=",".join(coordinates[:25]),
                            evidence=evidence,
                            expected="Fixed-format numeric and date values remain legible at the saved width.",
                            suggested_action=(
                                "Review a local column-width adjustment or a deliberate number format; no "
                                "width or format is changed automatically."
                            ),
                            discriminator=(region.min_row, region.min_column, column),
                        )
                    )
        return result


def _row_has_summary(worksheet: Worksheet, region: Region, row: int) -> bool:
    for column in range(region.min_column, region.max_column + 1):
        cell = worksheet._cells.get((row, column))
        if (
            isinstance(cell, Cell)
            and isinstance(cell.value, str)
            and _SUMMARY_RE.search(cell.value)
        ):
            return True
    return False


def _row_needs_height(worksheet: Worksheet, region: Region, row: int, observed: float) -> bool:
    for column in range(region.min_column, region.max_column + 1):
        cell = worksheet._cells.get((row, column))
        if not isinstance(cell, Cell) or not isinstance(cell.value, str) or not cell.value:
            continue
        measurement = measure_text_cell(worksheet, cell)
        if measurement is not None and measurement.required_height >= observed * 0.75:
            return True
    return False


def _row_intersects_drawing(worksheet: Worksheet, row: int) -> bool:
    for drawing in (
        *getattr(worksheet, "_charts", ()),
        *getattr(worksheet, "_images", ()),
    ):
        box = _drawing_box(worksheet, drawing)
        if box is not None and box.min_row <= row <= box.max_row:
            return True
    return False


class DataRegionRowHeightOutlierRule(WorkbookRule):
    """Report dense detail rows with unexplained extreme explicit heights."""

    rule_id = "WL049_DATA_REGION_ROW_HEIGHT_OUTLIER"
    title = "Data-region row height is an unexplained outlier"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            for region in context.data_regions.get(worksheet.title, ()):
                width = region.max_column - region.min_column + 1
                if width < 2 or region.max_row - region.min_row < 5:
                    continue
                detail_rows = [
                    row
                    for row in range(region.min_row + 1, region.max_row + 1)
                    if not _row_has_summary(worksheet, region, row)
                ]
                dense_rows = [
                    row
                    for row in detail_rows
                    if sum(
                        isinstance(worksheet._cells.get((row, column)), Cell)
                        and worksheet._cells[(row, column)].value is not None
                        for column in range(region.min_column, region.max_column + 1)
                    )
                    >= max(2, math.ceil(width * 0.5))
                ]
                if len(dense_rows) < 5:
                    continue
                baseline = statistics.median(row_height(worksheet, row) for row in dense_rows)
                if baseline <= 0:
                    continue
                for row in dense_rows:
                    dimension = worksheet.row_dimensions.get(row)
                    if dimension is None or dimension.height is None or dimension.hidden:
                        continue
                    observed = float(dimension.height)
                    ratio = observed / baseline
                    if 0.65 <= ratio <= 2.5:
                        continue
                    if ratio > 2.5 and (
                        _row_needs_height(worksheet, region, row, observed)
                        or _row_intersects_drawing(worksheet, row)
                        or any(
                            merged.min_row <= row <= merged.max_row
                            for merged in worksheet.merged_cells
                        )
                    ):
                        continue
                    evidence = Evidence(
                        summary=(
                            f"Explicit row height {observed:g} is {ratio:.2f} times the detail-row median"
                        ),
                        observed={"row": row, "height": observed, "ratio": round(ratio, 3)},
                        expected={"peer_median_height": baseline, "accepted_ratio": [0.65, 2.5]},
                        peers=[str(peer) for peer in dense_rows if peer != row][:12],
                    )
                    result.findings.append(
                        _finding(
                            context=context,
                            rule_id=self.rule_id,
                            title=self.title,
                            explanation=(
                                "A densely populated detail row has an extreme explicit height, without "
                                "wrapped text, merged content, or an anchored drawing that explains it."
                            ),
                            severity=Severity.WARNING,
                            confidence=0.94,
                            worksheet=worksheet,
                            location=f"{row}:{row}",
                            evidence=evidence,
                            expected="Comparable detail rows use heights consistent with their visible content.",
                            suggested_action=(
                                "Review the row locally and reset or resize it if the difference is accidental; "
                                "WorkbookLens does not normalize row heights automatically."
                            ),
                            discriminator=(region.min_row, region.min_column, row),
                        )
                    )
        return result


def _color_signature(color: Any) -> tuple[Any, ...]:
    if color is None:
        return (None,)
    return (
        getattr(color, "type", None),
        getattr(color, "rgb", None),
        getattr(color, "indexed", None),
        getattr(color, "theme", None),
        round(float(getattr(color, "tint", 0.0) or 0.0), 4),
    )


def _style_components(cell: Cell) -> dict[str, tuple[Any, ...] | str]:
    return {
        "font": (
            cell.font.name,
            cell.font.sz,
            bool(cell.font.bold),
            bool(cell.font.italic),
            cell.font.underline,
            bool(cell.font.strike),
            _color_signature(cell.font.color),
        ),
        "fill": (
            cell.fill.fill_type,
            _color_signature(cell.fill.fgColor),
            _color_signature(cell.fill.bgColor),
        ),
        "alignment": (
            cell.alignment.horizontal,
            cell.alignment.vertical,
            bool(cell.alignment.wrap_text),
            bool(cell.alignment.shrink_to_fit),
            cell.alignment.text_rotation,
        ),
        "border": tuple(
            (
                getattr(getattr(cell.border, side), "style", None),
                _color_signature(getattr(getattr(cell.border, side), "color", None)),
            )
            for side in ("left", "right", "top", "bottom")
        ),
        "number_format": cell.number_format or "General",
    }


def _mode(values: Iterable[Any]) -> tuple[Any, int, int]:
    counts = Counter(values)
    value, count = counts.most_common(1)[0]
    return value, count, sum(counts.values())


def _row_wide_alternative(
    worksheet: Worksheet,
    region: Region,
    row: int,
    component: str,
    consensus: Any,
) -> bool:
    populated = [
        cell
        for column in range(region.min_column, region.max_column + 1)
        if isinstance((cell := worksheet._cells.get((row, column))), Cell)
        and cell.value is not None
    ]
    if len(populated) < 4:
        return False
    alternatives = [
        _style_components(cell)[component]
        for cell in populated
        if _style_components(cell)[component] != consensus
    ]
    if len(alternatives) < math.ceil(len(populated) * 0.6):
        return False
    _, count, total = _mode(alternatives)
    return count / total >= 0.8


def _summary_rows_near_region(worksheet: Worksheet, region: Region) -> list[int]:
    rows = [
        row
        for row in range(region.min_row + 1, region.max_row + 1)
        if _row_has_summary(worksheet, region, row)
    ]
    for row in range(region.max_row + 1, min(worksheet.max_row, region.max_row + 3) + 1):
        probe = Region(
            sheet=region.sheet,
            min_row=row,
            max_row=row,
            min_column=region.min_column,
            max_column=region.max_column,
            kind="data",
            confidence=region.confidence,
        )
        if _row_has_summary(worksheet, probe, row):
            rows.append(row)
    return sorted(set(rows))


class RoleAwareStyleOutlierRule(WorkbookRule):
    """Report style anomalies using title, header, body, and total-row roles."""

    rule_id = "WL050_ROLE_AWARE_STYLE_OUTLIER"
    title = "Style is inconsistent with its inferred worksheet role"

    def run(self, context: RuleContext) -> RuleResult:
        result = RuleResult()
        for worksheet in context.workbook.worksheets:
            regions = [
                region
                for region in context.data_regions.get(worksheet.title, ())
                if region.max_row - region.min_row >= 5 and region.max_column > region.min_column
            ]
            if worksheet.sheet_state != "visible" or not regions:
                continue
            first_region = min(regions, key=lambda item: (item.min_row, item.min_column))
            result.extend(self._title_findings(context, worksheet, first_region))
            for region in regions:
                result.extend(self._header_findings(context, worksheet, region))
                result.extend(self._body_findings(context, worksheet, region))
                result.extend(self._total_findings(context, worksheet, region))
        return result

    def _title_findings(
        self,
        context: RuleContext,
        worksheet: Worksheet,
        region: Region,
    ) -> RuleResult:
        result = RuleResult()
        if region.min_row <= 1:
            return result
        body_fonts = [
            cell.font.name
            for cell in worksheet._cells.values()
            if isinstance(cell, Cell)
            and region.min_row <= cell.row <= region.max_row
            and region.min_column <= cell.column <= region.max_column
            and cell.font.name
        ]
        body_sizes = [
            float(cell.font.sz)
            for cell in worksheet._cells.values()
            if isinstance(cell, Cell)
            and region.min_row <= cell.row <= region.max_row
            and region.min_column <= cell.column <= region.max_column
            and cell.font.sz is not None
        ]
        baseline_font = Counter(body_fonts).most_common(1)[0][0] if body_fonts else None
        baseline_size = statistics.median(body_sizes) if body_sizes else 11.0
        minimum_span = max(3, math.ceil((region.max_column - region.min_column + 1) * 0.4))
        for merged in worksheet.merged_cells.ranges:
            if merged.max_row >= region.min_row or merged.min_row < max(1, region.min_row - 4):
                continue
            if merged.max_col - merged.min_col + 1 < minimum_span:
                continue
            anchor = worksheet._cells.get((merged.min_row, merged.min_col))
            if (
                not isinstance(anchor, Cell)
                or not isinstance(anchor.value, str)
                or not anchor.value.strip()
            ):
                continue
            anomalies: list[str] = []
            if baseline_font and anchor.font.name and anchor.font.name != baseline_font:
                anomalies.append("font_family")
            if anchor.font.sz is not None and float(anchor.font.sz) >= baseline_size * 1.6:
                anomalies.append("font_size")
            if anchor.alignment.horizontal not in {"center", "centerContinuous"}:
                anomalies.append("horizontal_alignment")
            if anchor.alignment.vertical not in {"center"}:
                anomalies.append("vertical_alignment")
            if anchor.font.italic:
                anomalies.append("italic")
            if anchor.fill.fill_type and _color_signature(anchor.fill.fgColor) not in {
                ("rgb", "00000000", 0, 0, 0.0),
                ("rgb", "00FFFFFF", 0, 0, 0.0),
            }:
                anomalies.append("fill")
            if (
                len(anomalies) < 3
                or not {"horizontal_alignment", "vertical_alignment"}.intersection(anomalies)
                or not {"font_family", "italic"}.intersection(anomalies)
            ):
                continue
            evidence = Evidence(
                summary=f"Merged title combines {len(anomalies)} unusual role-specific style components",
                observed={
                    "range": str(merged),
                    "anomalies": anomalies,
                    "font": anchor.font.name,
                    "font_size": anchor.font.sz,
                    "horizontal": anchor.alignment.horizontal,
                    "vertical": anchor.alignment.vertical,
                },
                expected={
                    "body_font": baseline_font,
                    "body_font_size": baseline_size,
                    "coherent_merged_title_alignment": True,
                },
            )
            result.findings.append(
                _finding(
                    context=context,
                    rule_id=self.rule_id,
                    title=self.title,
                    explanation=(
                        "A wide merged title combines several style choices that are unusual relative "
                        "to the table beneath it; no single choice is treated as an error by itself."
                    ),
                    severity=Severity.INFO,
                    confidence=0.86,
                    worksheet=worksheet,
                    location=str(merged),
                    evidence=evidence,
                    expected="Title styling is internally coherent and intentionally distinct from table data.",
                    suggested_action=(
                        "Review the title's font, fill, and alignment as one unit; no style is copied "
                        "automatically."
                    ),
                    discriminator=("title", str(merged)),
                )
            )
        return result

    def _header_findings(
        self,
        context: RuleContext,
        worksheet: Worksheet,
        region: Region,
    ) -> RuleResult:
        result = RuleResult()
        cells = [
            cell
            for column in range(region.min_column, region.max_column + 1)
            if isinstance((cell := worksheet._cells.get((region.min_row, column))), Cell)
            and cell.value is not None
        ]
        if len(cells) < 4:
            return result
        invalid_formats = [
            cell.coordinate
            for cell in cells
            if isinstance(cell.value, str)
            and (cell.number_format or "General").strip().casefold() not in {"general", "@", "text"}
        ]
        fragmented: dict[str, dict[str, float | int]] = {}
        for component in ("font", "fill", "alignment", "border"):
            signatures = [_style_components(cell)[component] for cell in cells]
            _, mode_count, total = _mode(signatures)
            unique = len(set(signatures))
            if unique >= 3 and mode_count / total < 0.75:
                fragmented[component] = {
                    "unique": unique,
                    "dominant_ratio": round(mode_count / total, 3),
                }
        if not invalid_formats and len(fragmented) < 2:
            return result
        location = (
            f"{get_column_letter(region.min_column)}{region.min_row}:"
            f"{get_column_letter(region.max_column)}{region.min_row}"
        )
        evidence = Evidence(
            summary="Header role has incompatible number formats or multi-component style fragmentation",
            observed={
                "invalid_text_header_formats": invalid_formats,
                "fragmented_components": fragmented,
            },
            expected={"text_header_number_format": "General or Text", "coherent_role_style": True},
            peers=[cell.coordinate for cell in cells],
        )
        result.findings.append(
            _finding(
                context=context,
                rule_id=self.rule_id,
                title=self.title,
                explanation=(
                    "The inferred header row contains text labels with numeric display formats or is "
                    "fragmented across several independent style components."
                ),
                severity=Severity.WARNING,
                confidence=0.95 if invalid_formats else 0.84,
                worksheet=worksheet,
                location=location,
                evidence=evidence,
                expected="Cells serving the same header role use coherent styling and text-appropriate formats.",
                suggested_action="Review the header row as a group; no global header style is imposed.",
                discriminator=("header", region.min_row, region.min_column),
            )
        )
        return result

    def _body_findings(
        self,
        context: RuleContext,
        worksheet: Worksheet,
        region: Region,
    ) -> RuleResult:
        result = RuleResult()
        anomalies: dict[str, set[str]] = defaultdict(set)
        detail_rows = tuple(
            row
            for row in range(region.min_row + 1, region.max_row + 1)
            if not _row_has_summary(worksheet, region, row)
        )
        boundary_borders_by_row: dict[int, tuple[tuple[Any, ...], ...]] = {}
        if detail_rows:
            for row in {detail_rows[0], detail_rows[-1]}:
                boundary_borders_by_row[row] = tuple(
                    cast(tuple[Any, ...], _style_components(boundary_cell)["border"])
                    for column in range(region.min_column, region.max_column + 1)
                    if isinstance((boundary_cell := worksheet._cells.get((row, column))), Cell)
                    and boundary_cell.value is not None
                    and _visible_cell(worksheet, boundary_cell)
                )
        for column in range(region.min_column, region.max_column + 1):
            cells = [
                cell
                for row in detail_rows
                if isinstance((cell := worksheet._cells.get((row, column))), Cell)
                and cell.value is not None
                and _visible_cell(worksheet, cell)
            ]
            if len(cells) < 7:
                continue
            full_styles = [tuple(_style_components(cell).values()) for cell in cells]
            _, full_count, full_total = _mode(full_styles)
            full_consensus = full_count / full_total
            for component in ("font", "fill", "alignment", "border", "number_format"):
                signatures = [_style_components(cell)[component] for cell in cells]
                consensus, count, total = _mode(signatures)
                if count < 6 or count / total < 0.78:
                    continue
                outliers = [
                    cell for cell in cells if _style_components(cell)[component] != consensus
                ]
                if component == "border":
                    outliers = [
                        cell
                        for cell in outliers
                        if not is_detail_row_perimeter_border_variant(
                            cell,
                            detail_rows=detail_rows,
                            observed_border=cast(
                                tuple[Any, ...], _style_components(cell)["border"]
                            ),
                            consensus_border=cast(tuple[Any, ...], consensus),
                            boundary_borders=boundary_borders_by_row.get(cell.row, ()),
                        )
                    ]
                if not outliers:
                    continue
                if component != "number_format" and len(outliers) == 1 and full_consensus >= 0.8:
                    continue
                for cell in outliers:
                    if component != "number_format" and _row_wide_alternative(
                        worksheet, region, cell.row, component, consensus
                    ):
                        continue
                    anomalies[component].add(cell.coordinate)
        if not anomalies:
            return result
        coordinates = sorted(
            {coordinate for values in anomalies.values() for coordinate in values},
            key=coordinate_to_tuple,
        )
        location = ",".join(coordinates[:50])
        evidence = Evidence(
            summary=f"Body-role component consensus identifies {len(coordinates)} anomalous cells",
            observed={component: sorted(values)[:50] for component, values in anomalies.items()},
            expected={"component_consensus_within_column": True},
            peers=[],
            details={
                "region": f"{get_column_letter(region.min_column)}{region.min_row}:"
                f"{get_column_letter(region.max_column)}{region.max_row}"
            },
        )
        result.findings.append(
            _finding(
                context=context,
                rule_id=self.rule_id,
                title=self.title,
                explanation=(
                    "Component-wise comparison finds clustered or diffuse body-cell deviations that a "
                    "single whole-style comparison can miss. Row-wide coherent highlights are excluded."
                ),
                severity=Severity.INFO,
                confidence=0.88,
                worksheet=worksheet,
                location=location,
                evidence=evidence,
                expected="Cells with the same body-column role use stable component-level formatting.",
                suggested_action=(
                    "Review the listed components and preserve intentional exceptions; no body style is "
                    "copied automatically."
                ),
                discriminator=("body", region.min_row, region.min_column),
            )
        )
        return result

    def _total_findings(
        self,
        context: RuleContext,
        worksheet: Worksheet,
        region: Region,
    ) -> RuleResult:
        result = RuleResult()
        summary_rows = _summary_rows_near_region(worksheet, region)
        if not summary_rows:
            return result
        anomalies: list[dict[str, Any]] = []
        for row in summary_rows:
            for column in range(region.min_column, region.max_column + 1):
                total_cell = worksheet._cells.get((row, column))
                if not isinstance(total_cell, Cell) or total_cell.value is None:
                    continue
                if total_cell.data_type != "f" and (
                    isinstance(total_cell.value, bool)
                    or not isinstance(total_cell.value, (int, float, datetime, date, time))
                ):
                    continue
                peers = [
                    cell
                    for peer_row in range(region.min_row + 1, region.max_row + 1)
                    if peer_row not in summary_rows
                    and isinstance((cell := worksheet._cells.get((peer_row, column))), Cell)
                    and cell.value is not None
                ]
                if len(peers) < 6:
                    continue
                consensus, count, total = _mode(cell.number_format or "General" for cell in peers)
                observed = total_cell.number_format or "General"
                if count / total < 0.8 or observed == consensus:
                    continue
                anomalies.append(
                    {
                        "cell": total_cell.coordinate,
                        "observed": observed,
                        "body_consensus": consensus,
                        "support": count,
                        "population": total,
                    }
                )
        if not anomalies:
            return result
        coordinates = [item["cell"] for item in anomalies]
        evidence = Evidence(
            summary=f"{len(anomalies)} total-row cells use formats inconsistent with their body columns",
            observed=anomalies,
            expected={"preserve_column_value_semantics": True},
        )
        result.findings.append(
            _finding(
                context=context,
                rule_id=self.rule_id,
                title=self.title,
                explanation=(
                    "A labeled summary row is allowed to use distinct emphasis, but its numeric cells use "
                    "number formats inconsistent with the value semantics established by their columns."
                ),
                severity=Severity.WARNING,
                confidence=0.93,
                worksheet=worksheet,
                location=",".join(coordinates),
                evidence=evidence,
                expected="Summary emphasis may differ while numeric display semantics remain compatible.",
                suggested_action="Review each total-row format manually; no formula or style is changed.",
                discriminator=("total", region.min_row, region.min_column),
            )
        )
        return result


LAYOUT_GEOMETRY_RULES: tuple[type[WorkbookRule], ...] = (
    DrawingContentOverlapRule,
    NumericDisplayWidthRiskRule,
    DataRegionRowHeightOutlierRule,
    RoleAwareStyleOutlierRule,
)


__all__ = ["LAYOUT_GEOMETRY_RULES"]
