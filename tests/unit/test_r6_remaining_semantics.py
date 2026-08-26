from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

from workbooklens.models import Finding, Severity
from workbooklens.rules import RuleRegistry
from workbooklens.rules.builtin import TextDisplayRiskRule
from workbooklens.rules.data_quality import ConditionalFormatSemanticConflictRule
from workbooklens.rules.formula_semantics import UnlabeledAggregateRoleRule
from workbooklens.scanner import ScanResult, scan_workbook


def _scan(tmp_path: Path, workbook: Workbook, rule: object, name: str) -> ScanResult:
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    return scan_workbook(path, registry=RuleRegistry([rule]))  # type: ignore[list-item]


def _three_column_table() -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Name", "Notes"])
    for row in range(2, 9):
        worksheet.append([f"R{row:03d}", f"Name {row}", "ok"])
    worksheet.column_dimensions["B"].width = 8
    worksheet.column_dimensions["C"].width = 4
    return workbook


def test_text_display_reports_extreme_overflow_past_inferred_table_edge(tmp_path: Path) -> None:
    workbook = _three_column_table()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["C5"] = (
        "A very long unwrapped note that would visibly spill far beyond the table edge. " * 3
    )

    scan = _scan(tmp_path, workbook, TextDisplayRiskRule(), "table-edge-overflow.xlsx")

    findings = [
        finding for finding in scan.findings if finding.rule_id == "WL016_TEXT_DISPLAY_RISK"
    ]
    assert [finding.location for finding in findings] == ["C5"]
    assert findings[0].evidence.details["crosses_inferred_data_region"] is True


def test_text_display_allows_natural_overflow_inside_inferred_table(tmp_path: Path) -> None:
    workbook = _three_column_table()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["B5"] = "Long text may use the blank peer cell inside the inferred table"
    worksheet["C5"] = None

    scan = _scan(tmp_path, workbook, TextDisplayRiskRule(), "inside-table-overflow.xlsx")

    assert not scan.findings


def _five_column_width_book(*, second_risk_column: str = "D") -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Customer", "Region", "Status", "Product"])
    for row in range(2, 9):
        worksheet.append(
            [
                f"R{row:03d}",
                f"Customer name that is clipped {row}",
                "West",
                f"Status description that is clipped {row}"
                if second_risk_column == "D"
                else "Open",
                f"Product description that is clipped {row}"
                if second_risk_column == "E"
                else "Widget",
            ]
        )
    worksheet.column_dimensions["B"].width = 5
    worksheet.column_dimensions["C"].width = 16
    worksheet.column_dimensions["D"].width = 5
    worksheet.column_dimensions["E"].width = 12 if second_risk_column == "D" else 4
    return workbook


def _width_block_advisories(scan: ScanResult) -> list[Finding]:
    return [
        finding
        for finding in scan.findings
        if finding.evidence.details.get("finding_kind") == "multi_column_repeated_clipping"
    ]


def test_text_display_groups_repeated_clipping_across_table_columns(tmp_path: Path) -> None:
    scan = _scan(
        tmp_path,
        _five_column_width_book(),
        TextDisplayRiskRule(),
        "multi-column-clipping.xlsx",
    )

    advisories = _width_block_advisories(scan)
    assert len(advisories) == 1
    advisory = advisories[0]
    assert advisory.location == "A1:E8"
    assert advisory.severity == Severity.INFO
    assert advisory.evidence.observed["affected_columns"] == ["B", "D"]
    assert advisory.patch_ids == []


def test_text_display_does_not_group_one_repeated_clipping_column(tmp_path: Path) -> None:
    workbook = _five_column_width_book(second_risk_column="")

    scan = _scan(
        tmp_path,
        workbook,
        TextDisplayRiskRule(),
        "one-clipping-column.xlsx",
    )

    assert not _width_block_advisories(scan)


def test_text_display_excludes_free_text_notes_from_block_advisory(tmp_path: Path) -> None:
    for index, header in enumerate(("Notes", "Phone", "Email")):
        workbook = _five_column_width_book(second_risk_column="E")
        worksheet = workbook.active
        assert worksheet is not None
        worksheet["E1"] = header

        scan = _scan(
            tmp_path,
            workbook,
            TextDisplayRiskRule(),
            f"excluded-clipping-column-{index}.xlsx",
        )

        assert not _width_block_advisories(scan)


def test_text_display_requires_a_wide_table_for_block_advisory(tmp_path: Path) -> None:
    workbook = _five_column_width_book()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.delete_cols(5)

    scan = _scan(
        tmp_path,
        workbook,
        TextDisplayRiskRule(),
        "four-column-clipping.xlsx",
    )

    assert not _width_block_advisories(scan)


def test_text_display_does_not_add_block_advisory_for_hidden_sheet(tmp_path: Path) -> None:
    workbook = _five_column_width_book()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.sheet_state = "hidden"
    workbook.create_sheet("Visible")

    scan = _scan(
        tmp_path,
        workbook,
        TextDisplayRiskRule(),
        "hidden-table-clipping.xlsx",
    )

    assert not _width_block_advisories(scan)


def _quantity_book(
    fill_rgb: str,
    *,
    operator: str = "lessThan",
    threshold: str = "0",
) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Quantity"])
    values = [2, 4, 6, -1, 8, 10, 12, 14]
    for row, value in enumerate(values, start=2):
        worksheet.append([f"R{row:03d}", value])
    fill = PatternFill(fill_type="solid", fgColor=fill_rgb)
    worksheet.conditional_formatting.add(
        "B2:B9",
        CellIsRule(operator=operator, formula=[threshold], fill=fill),
    )
    return workbook


def test_negative_constrained_values_are_not_colored_success_green(tmp_path: Path) -> None:
    scan = _scan(
        tmp_path,
        _quantity_book("FF00B050"),
        ConditionalFormatSemanticConflictRule(),
        "negative-success-green.xlsx",
    )

    assert [(finding.rule_id, finding.location) for finding in scan.findings] == [
        ("WL051_CONDITIONAL_FORMAT_SEMANTIC_CONFLICT", "B2:B9")
    ]
    assert scan.findings[0].evidence.observed["negative_cells"] == ["B5"]
    assert not scan.patches


def test_negative_warning_red_and_positive_success_green_are_accepted(tmp_path: Path) -> None:
    red = _scan(
        tmp_path,
        _quantity_book("FFFF0000"),
        ConditionalFormatSemanticConflictRule(),
        "negative-warning-red.xlsx",
    )
    positive_green = _scan(
        tmp_path,
        _quantity_book("FF00B050", operator="greaterThan"),
        ConditionalFormatSemanticConflictRule(),
        "positive-success-green.xlsx",
    )

    assert not red.findings
    assert not positive_green.findings


def test_conditional_format_respects_the_actual_negative_threshold(tmp_path: Path) -> None:
    scan = _scan(
        tmp_path,
        _quantity_book("FF00B050", threshold="-5"),
        ConditionalFormatSemanticConflictRule(),
        "negative-threshold.xlsx",
    )

    assert not scan.findings


def _summary_book(*, unlabeled_function: str) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Item", "Quantity", "Category", "Sales", "Gross"])
    for row in range(2, 9):
        worksheet.append([f"Item {row}", row, "A", row * 10, row * 12])
    worksheet["B10"] = f"={unlabeled_function}(B2:B8)"
    worksheet["C10"] = "TOTAL"
    worksheet["D10"] = "=SUM(D2:D8)"
    worksheet["E10"] = "=SUM(E2:E8)"
    return workbook


def test_distinct_aggregate_left_of_summary_label_requires_its_own_label(tmp_path: Path) -> None:
    scan = _scan(
        tmp_path,
        _summary_book(unlabeled_function="AVERAGE"),
        UnlabeledAggregateRoleRule(),
        "unlabeled-average.xlsx",
    )

    assert [(finding.rule_id, finding.location) for finding in scan.findings] == [
        ("WL052_UNLABELED_AGGREGATE_ROLE", "B10")
    ]
    assert scan.findings[0].evidence.details["summary_label_cell"] == "C10"
    assert scan.findings[0].evidence.details["dominant_labelled_function"] == "SUM"
    assert not scan.patches


def test_same_function_aggregate_outside_label_block_is_not_overstated(tmp_path: Path) -> None:
    scan = _scan(
        tmp_path,
        _summary_book(unlabeled_function="SUM"),
        UnlabeledAggregateRoleRule(),
        "unlabeled-sum.xlsx",
    )

    assert not scan.findings
