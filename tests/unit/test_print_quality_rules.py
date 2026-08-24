from __future__ import annotations

from collections.abc import ValuesView
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

from workbooklens.models import Confidence, Region
from workbooklens.rules.print_quality import (
    PrintSettingConsistencyRule,
    _boundary_has_group_structure,
    _effective_manual_breaks,
    _effective_manual_breaks_with_ranges,
    _footer_texts,
    _print_ranges,
    _row_density,
)
from workbooklens.rules.registry import RuleRegistry
from workbooklens.scanner import ScanResult, scan_workbook


def _dense_workbook(*, min_row: int = 1, max_row: int = 24, columns: int = 8) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for column in range(1, columns + 1):
        worksheet.cell(min_row, column, f"Field {column}")
    for row in range(min_row + 1, max_row + 1):
        for column in range(1, columns + 1):
            worksheet.cell(row, column, f"R{row}C{column}")
    return workbook


def _scan(workbook: Workbook, path: Path) -> ScanResult:
    workbook.save(path)
    workbook.close()
    return scan_workbook(
        path,
        registry=RuleRegistry((PrintSettingConsistencyRule(),)),
    )


def _findings(scan: ScanResult) -> list:
    return [
        finding
        for finding in scan.findings
        if finding.rule_id == PrintSettingConsistencyRule.rule_id
    ]


def _enable_fit_to_page(worksheet, *, width: int) -> None:
    worksheet.sheet_properties.pageSetUpPr = PageSetupProperties(
        fitToPage=True,
        autoPageBreaks=False,
    )
    worksheet.page_setup.fitToWidth = width
    worksheet.page_setup.fitToHeight = 0


def test_print_settings_report_only_combined_provable_conflicts(tmp_path: Path) -> None:
    workbook = _dense_workbook(min_row=4, max_row=24, columns=8)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.row_breaks.append(Break(id=7))
    worksheet.oddFooter.right.text = "Page &P of 1"
    worksheet.print_area = "A1:D15"
    _enable_fit_to_page(worksheet, width=3)

    scan = _scan(workbook, tmp_path / "contradictory-print-settings.xlsx")
    findings = _findings(scan)

    assert len(findings) == 3
    summaries = {finding.evidence.summary for finding in findings}
    assert any(summary.startswith("1 manual row break") for summary in summaries)
    assert any(summary.startswith("A fixed footer page total") for summary in summaries)
    assert any(summary.startswith("Fit-to-width requests") for summary in summaries)
    assert {finding.location for finding in findings} == {
        "7:7",
        "oddFooter.right",
        str(worksheet.print_area),
    }
    assert all(not finding.patch_ids for finding in findings)
    assert not scan.patches


def test_summary_boundary_dynamic_footer_and_full_print_area_are_legal(tmp_path: Path) -> None:
    workbook = _dense_workbook(max_row=20, columns=6)
    worksheet = workbook.active
    assert worksheet is not None
    for column in range(2, 6):
        worksheet.cell(5, column).value = None
    worksheet["A5"] = "Subtotal"
    worksheet["F5"] = "=SUM(F2:F4)"
    worksheet.row_breaks.append(Break(id=5))
    worksheet.oddFooter.center.text = "Page &P of &N"
    worksheet.print_area = "A1:F20"
    _enable_fit_to_page(worksheet, width=2)

    scan = _scan(workbook, tmp_path / "summary-page-boundary.xlsx")

    assert not _findings(scan)
    assert not scan.patches


def test_outline_group_boundary_and_matching_multi_page_template_are_legal(
    tmp_path: Path,
) -> None:
    workbook = _dense_workbook(max_row=20, columns=6)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.row_dimensions[6].outlineLevel = 1
    worksheet.row_breaks.append(Break(id=5))
    worksheet.oddFooter.right.text = "Page &P of 2"
    worksheet.print_area = "A1:F20"
    _enable_fit_to_page(worksheet, width=2)

    scan = _scan(workbook, tmp_path / "grouped-multi-page-template.xlsx")

    assert not _findings(scan)
    assert not scan.patches


def test_fit_width_or_fixed_one_page_footer_alone_are_not_anomalies(tmp_path: Path) -> None:
    workbook = _dense_workbook(max_row=20, columns=6)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.oddFooter.right.text = "Page &P of 1"
    worksheet.print_area = "A1:F20"
    _enable_fit_to_page(worksheet, width=4)

    scan = _scan(workbook, tmp_path / "reasonable-fit-width.xlsx")

    assert not _findings(scan)
    assert not scan.patches


def test_helpers_do_not_materialize_row_dimensions_or_disabled_footers() -> None:
    workbook = _dense_workbook(max_row=20, columns=6)
    worksheet = workbook.active
    assert worksheet is not None
    region = Region(
        sheet=worksheet.title,
        min_row=1,
        max_row=20,
        min_column=1,
        max_column=6,
        kind="data",
        confidence=Confidence(1.0),
    )
    before = set(worksheet.row_dimensions)

    assert _row_density(worksheet, region, 100) == 0.0
    assert not _boundary_has_group_structure(worksheet, region, 100)
    assert set(worksheet.row_dimensions) == before

    even_footer = worksheet.evenFooter
    first_footer = worksheet.firstFooter
    assert even_footer is not None
    assert first_footer is not None
    even_footer.right.text = "Page &P of 1"
    first_footer.right.text = "Page &P of 1"
    assert _footer_texts(worksheet) == {}
    worksheet.HeaderFooter.differentOddEven = True
    worksheet.HeaderFooter.differentFirst = True
    assert set(_footer_texts(worksheet)) == {"evenFooter.right", "firstFooter.right"}


def test_row_and_column_breaks_in_one_print_area_prove_cartesian_page_count(
    tmp_path: Path,
) -> None:
    workbook = _dense_workbook(max_row=20, columns=8)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.row_breaks.append(Break(id=10))
    worksheet.col_breaks.append(Break(id=4))
    worksheet.oddFooter.right.text = "Page &P of 3"
    worksheet.print_area = "A1:H20"

    scan = _scan(workbook, tmp_path / "cartesian-page-count.xlsx")
    footer_findings = [finding for finding in _findings(scan) if "Footer" in finding.location]

    assert len(footer_findings) == 1
    observed = footer_findings[0].evidence.observed
    assert observed["minimum_pages_from_manual_breaks"] == 4


def test_partial_row_break_does_not_prove_two_pages_for_the_whole_print_area(
    tmp_path: Path,
) -> None:
    workbook = _dense_workbook(max_row=20, columns=8)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.row_breaks.append(Break(id=10, min=0, max=2))
    worksheet.oddFooter.right.text = "Page &P of 1"
    worksheet.print_area = "A1:H20"

    scan = _scan(workbook, tmp_path / "partial-width-row-break.xlsx")

    assert not [finding for finding in _findings(scan) if "Footer" in finding.location]


def test_partial_break_spans_do_not_prove_pages_without_a_print_area(tmp_path: Path) -> None:
    workbook = _dense_workbook(max_row=20, columns=8)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.row_breaks.append(Break(id=10, min=0, max=2))
    worksheet.col_breaks.append(Break(id=4, min=0, max=5))
    worksheet.oddFooter.right.text = "Page &P of 1"

    scan = _scan(workbook, tmp_path / "partial-breaks-without-print-area.xlsx")

    assert not _findings(scan)


def test_full_break_spans_still_prove_pages_without_a_print_area(tmp_path: Path) -> None:
    workbook = _dense_workbook(max_row=20, columns=8)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.row_breaks.append(Break(id=10))
    worksheet.col_breaks.append(Break(id=4))
    worksheet.oddFooter.right.text = "Page &P of 3"

    scan = _scan(workbook, tmp_path / "full-breaks-without-print-area.xlsx")
    footer_findings = [finding for finding in _findings(scan) if "Footer" in finding.location]

    assert len(footer_findings) == 1
    assert footer_findings[0].evidence.observed["minimum_pages_from_manual_breaks"] == 4


def test_inferred_regions_sum_local_page_counts_without_cross_product(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for row in range(1, 21):
        for column in range(1, 5):
            worksheet.cell(row, column, f"A{row}:{column}")
    for row in range(30, 50):
        for column in range(8, 12):
            worksheet.cell(row, column, f"B{row}:{column}")
    for row in range(60, 70):
        for column in range(1, 5):
            worksheet.cell(row, column, f"C{row}:{column}")
    worksheet.row_breaks.append(Break(id=10, min=0, max=3))
    worksheet.col_breaks.append(Break(id=9, min=29, max=48))
    worksheet.oddFooter.right.text = "Page &P of 4"

    scan = _scan(workbook, tmp_path / "independent-inferred-regions.xlsx")
    footer_findings = [finding for finding in _findings(scan) if "Footer" in finding.location]

    assert len(footer_findings) == 1
    observed = footer_findings[0].evidence.observed
    assert observed["minimum_pages_from_manual_breaks"] == 5
    assert [item["minimum_pages"] for item in observed["manual_breaks_by_print_area"]] == [
        2,
        2,
        1,
    ]


def test_multiple_inferred_regions_without_breaks_do_not_invent_pages(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for row in range(1, 10):
        for column in range(1, 5):
            worksheet.cell(row, column, f"A{row}:{column}")
    for row in range(20, 29):
        for column in range(8, 12):
            worksheet.cell(row, column, f"B{row}:{column}")
    worksheet.oddFooter.right.text = "Page &P of 1"

    scan = _scan(workbook, tmp_path / "unbroken-inferred-regions.xlsx")

    assert not [finding for finding in _findings(scan) if "Footer" in finding.location]


def test_overlapping_inferred_region_boxes_suppress_ambiguous_page_count() -> None:
    workbook = _dense_workbook(max_row=20, columns=8)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.row_breaks.append(Break(id=8))
    overlapping = (
        Region(
            sheet=worksheet.title,
            min_row=1,
            max_row=12,
            min_column=1,
            max_column=6,
            kind="data",
            confidence=Confidence(1.0),
        ),
        Region(
            sheet=worksheet.title,
            min_row=8,
            max_row=20,
            min_column=3,
            max_column=8,
            kind="data",
            confidence=Confidence(1.0),
        ),
    )

    assert _effective_manual_breaks_with_ranges(worksheet, overlapping, ()) == ()


def test_multiple_print_areas_sum_independent_page_counts_without_cross_product(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for row in range(1, 31):
        for column in range(1, 7):
            worksheet.cell(row, column, f"A{row}:{column}")
    for row in range(40, 61):
        for column in range(8, 27):
            worksheet.cell(row, column, f"B{row}:{column}")
    worksheet.row_breaks.append(Break(id=10))
    worksheet.row_breaks.append(Break(id=20))
    worksheet.col_breaks.append(Break(id=13))
    worksheet.col_breaks.append(Break(id=19))
    worksheet.oddFooter.right.text = "Page &P of 5"
    worksheet.print_area = ["A1:F30", "H40:Z60"]

    scan = _scan(workbook, tmp_path / "independent-print-areas.xlsx")
    footer_findings = [finding for finding in _findings(scan) if "Footer" in finding.location]

    assert len(footer_findings) == 1
    observed = footer_findings[0].evidence.observed
    assert observed["minimum_pages_from_manual_breaks"] == 6
    assert [item["range"] for item in observed["manual_breaks_by_print_area"]] == [
        "A1:F30",
        "H40:Z60",
    ]


def test_empty_print_area_does_not_inflate_the_proven_page_count(tmp_path: Path) -> None:
    workbook = _dense_workbook(max_row=20, columns=6)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.oddFooter.right.text = "Page &P of 1"
    worksheet.print_area = ["A1:F20", "H1:M20"]

    scan = _scan(workbook, tmp_path / "empty-secondary-print-area.xlsx")

    assert not [finding for finding in _findings(scan) if "Footer" in finding.location]


def test_multiple_fixed_totals_in_one_footer_part_produce_one_finding(tmp_path: Path) -> None:
    workbook = _dense_workbook(max_row=20, columns=8)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.row_breaks.append(Break(id=10))
    worksheet.col_breaks.append(Break(id=4))
    worksheet.oddFooter.right.text = "Page &P of 1; alternate Page &P of 2"
    worksheet.print_area = "A1:H20"

    scan = _scan(workbook, tmp_path / "multiple-fixed-footer-totals.xlsx")
    footer_findings = [finding for finding in _findings(scan) if "Footer" in finding.location]

    assert len(footer_findings) == 1
    observed = footer_findings[0].evidence.observed
    assert observed["fixed_total_pages"] == 1
    assert observed["fixed_total_page_values"] == [1, 2]
    assert observed["minimum_pages_from_manual_breaks"] == 4


def test_long_page_fields_support_dynamic_and_fixed_totals(tmp_path: Path) -> None:
    dynamic_workbook = _dense_workbook(max_row=20, columns=8)
    dynamic_sheet = dynamic_workbook.active
    assert dynamic_sheet is not None
    dynamic_sheet.row_breaks.append(Break(id=10))
    dynamic_sheet.oddFooter.right.text = "Legacy Page &P of 1; canonical Page &[Page] of &[Pages]"
    dynamic_sheet.print_area = "A1:H20"

    dynamic_scan = _scan(dynamic_workbook, tmp_path / "dynamic-long-page-fields.xlsx")
    assert not [finding for finding in _findings(dynamic_scan) if "Footer" in finding.location]

    fixed_workbook = _dense_workbook(max_row=20, columns=8)
    fixed_sheet = fixed_workbook.active
    assert fixed_sheet is not None
    fixed_sheet.row_breaks.append(Break(id=10))
    fixed_sheet.oddFooter.right.text = "Page &[Page] of 1"
    fixed_sheet.print_area = "A1:H20"

    fixed_scan = _scan(fixed_workbook, tmp_path / "fixed-long-page-field.xlsx")
    footer_findings = [finding for finding in _findings(fixed_scan) if "Footer" in finding.location]
    assert len(footer_findings) == 1


class _CountingCellDict(dict[tuple[int, int], Cell]):
    def __init__(self, cells: dict[tuple[int, int], Cell]) -> None:
        super().__init__(cells)
        self.values_calls = 0

    def values(self) -> ValuesView[Cell]:
        self.values_calls += 1
        return super().values()


def test_manual_break_analysis_indexes_fifty_thousand_cells_once() -> None:
    rows = 250
    columns = 200
    workbook = _dense_workbook(max_row=rows, columns=columns)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.print_area = f"A1:{worksheet.cell(1, columns).column_letter}{rows}"
    for row in range(1, rows):
        worksheet.row_breaks.append(Break(id=row))
    for column in range(1, columns):
        worksheet.col_breaks.append(Break(id=column))

    original = worksheet._cells
    ordered = {key: original[key] for key in ((1, 1), (rows, 1), (1, columns), (rows, columns))}
    ordered.update(original)
    counting_cells = _CountingCellDict(ordered)
    worksheet._cells = counting_cells

    row_breaks, column_breaks = _effective_manual_breaks(worksheet, (), _print_ranges(worksheet))

    assert len(row_breaks) == rows - 1
    assert len(column_breaks) == columns - 1
    assert counting_cells.values_calls == 1
