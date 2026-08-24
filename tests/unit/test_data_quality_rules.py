from __future__ import annotations

import ast
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.workbook.defined_name import DefinedName

from workbooklens.i18n import canonical_text_translation, localize_scan_result
from workbooklens.models import Severity
from workbooklens.scanner import ScanResult, scan_workbook

DATA_QUALITY_RULE_IDS = {
    "WL022_INFERRED_DUPLICATE_IDENTIFIER",
    "WL023_MISSING_INFERRED_IDENTIFIER",
    "WL024_MIXED_NUMERIC_STORAGE",
    "WL025_ROBUST_NUMERIC_OUTLIER",
    "WL026_PERCENTAGE_SCALE_OUTLIER",
    "WL027_DATE_STORAGE_ANOMALY",
    "WL030_AUTOFILTER_COVERAGE",
    "WL031_CHART_SOURCE_STRUCTURE",
    "WL032_DEEP_FREEZE_PANE",
    "WL033_PRINT_AREA_COVERAGE",
    "WL035_SIGN_CONSTRAINED_MEASURE",
}


def _literal_branches(node: ast.expr) -> list[str]:
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        return [node.value]
    if isinstance(node, ast.IfExp):
        return [*_literal_branches(node.body), *_literal_branches(node.orelse)]
    return []


def test_all_static_data_quality_messages_and_dynamic_samples_are_translated() -> None:
    import workbooklens.rules.data_quality as data_quality

    source = Path(data_quality.__file__).read_text(encoding="utf-8")
    tree = ast.parse(source)
    fields = {"explanation", "expected", "suggested_action", "summary"}
    texts = {
        text
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        for keyword in node.keywords
        if keyword.arg in fields
        for text in _literal_branches(keyword.value)
    }
    missing = sorted(text for text in texts if canonical_text_translation(text, "zh-CN") is None)
    assert not missing

    dynamic_samples = (
        "Identifier value appears 2 times in an inferred key column",
        "9 of 10 populated peers use numeric storage",
        "Robust modified z-score is 12.3 across 10 numeric peers",
        "9 of 10 numeric peers use fractional percentage storage",
        "AutoFilter excludes 4 populated rows and 2 populated columns",
        "Chart 1 has 2 source-range issues",
        "Freeze pane at H10 is deep inside visible content",
        "Print area leaves 2 meaningful ranges outside its bounds",
    )
    assert all(canonical_text_translation(text, "zh-CN") for text in dynamic_samples)


def _save_and_scan(workbook: Workbook, path: Path) -> ScanResult:
    workbook.save(path)
    workbook.close()
    return scan_workbook(path)


def _data_quality_findings(scan: ScanResult) -> list:
    return [finding for finding in scan.findings if finding.rule_id in DATA_QUALITY_RULE_IDS]


def test_infers_duplicate_and_missing_identifiers_without_configured_keys(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record ID", "Name", "Amount"])
    identifiers = ["R001", "R002", "R003", "R002", "R005", "R006", None, "R008", "R009", "R010"]
    for index, identifier in enumerate(identifiers, start=1):
        worksheet.append([identifier, f"Record {index}", index * 10])

    scan = _save_and_scan(workbook, tmp_path / "identifier-integrity.xlsx")
    duplicate = [
        finding
        for finding in scan.findings
        if finding.rule_id == "WL022_INFERRED_DUPLICATE_IDENTIFIER"
    ]
    missing = [
        finding
        for finding in scan.findings
        if finding.rule_id == "WL023_MISSING_INFERRED_IDENTIFIER"
    ]

    assert [finding.location for finding in duplicate] == ["A3,A5"]
    assert [finding.location for finding in missing] == ["A8"]
    assert all(not finding.patch_ids for finding in [*duplicate, *missing])


def test_explicit_key_configuration_suppresses_duplicate_auto_inference(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Record ID", "Name", "Amount"])
    for row in range(2, 12):
        worksheet.append(["R002" if row in {3, 5} else f"R{row:03d}", f"Name {row}", row])
    path = tmp_path / "configured-key.xlsx"
    workbook.save(path)
    workbook.close()

    scan = scan_workbook(path, config={"keys": [{"sheet": "Data", "range": "A2:A11"}]})

    assert not any(
        finding.rule_id
        in {
            "WL022_INFERRED_DUPLICATE_IDENTIFIER",
            "WL023_MISSING_INFERRED_IDENTIFIER",
        }
        for finding in scan.findings
    )
    assert any(finding.rule_id == "WL012_DUPLICATE_CONFIGURED_KEY" for finding in scan.findings)


def test_reports_rare_mixed_numeric_text_and_extreme_mad_outlier(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Item ID", "Qty", "Salary"])
    for index in range(1, 13):
        quantity: int | str = "3O" if index == 5 else index
        salary = 8_500_000 if index == 9 else 5_000 + index * 250
        worksheet.append([f"I{index:03d}", quantity, salary])

    scan = _save_and_scan(workbook, tmp_path / "numeric-profile.xlsx")

    assert [
        finding.location
        for finding in scan.findings
        if finding.rule_id == "WL024_MIXED_NUMERIC_STORAGE"
    ] == ["B6"]
    assert [
        finding.location
        for finding in scan.findings
        if finding.rule_id == "WL025_ROBUST_NUMERIC_OUTLIER"
    ] == ["C10"]


def test_sign_constraints_report_only_strong_domains_and_exclude_signed_measures(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(
        [
            "Record ID",
            "Age",
            "Unit Price",
            "Salary",
            "Qty",
            "Stock",
            "Profit",
            "Change",
            "Balance",
            "Adjustment",
        ]
    )
    for index in range(1, 11):
        row: list[object] = [
            f"R{index:03d}",
            20 + index,
            10 + index,
            5_000 + index * 100,
            index,
            index + 5,
            100 + index,
            index,
            1_000 + index,
            index,
        ]
        if index == 4:
            row[1] = 0
            row[2] = 0
        elif index == 5:
            row[3] = -100
        elif index == 6:
            row[4] = -1
        elif index == 7:
            row[5] = -2
        row[6] = -index
        row[7] = -index
        row[8] = -index
        row[9] = -index
        worksheet.append(row)

    scan = _save_and_scan(workbook, tmp_path / "sign-domains.xlsx")
    findings = [
        finding for finding in scan.findings if finding.rule_id == "WL035_SIGN_CONSTRAINED_MEASURE"
    ]

    by_location = {finding.location: finding for finding in findings}
    assert set(by_location) == {"B5", "C5", "D6", "E7", "F8"}
    assert {by_location[location].severity for location in {"B5", "C5"}} == {Severity.WARNING}
    assert all(float(by_location[location].confidence) == 0.82 for location in {"B5", "C5"})
    assert {by_location[location].severity for location in {"D6", "E7", "F8"}} == {Severity.ERROR}
    assert all(not finding.patch_ids and not finding.safe_patch_available for finding in findings)


def test_reports_percentage_scale_and_date_storage_anomalies(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record ID", "Date", "Discount"])
    for index in range(1, 13):
        date_value: object = datetime(2026, 1, index)
        rate: object = 0.05 + (index % 3) * 0.01
        if index == 10:
            date_value = "2026/1/10"
            rate = 15
        elif index == 11:
            date_value = "2026-02-30"
            rate = -0.1
        elif index == 12:
            date_value = 46_000
            rate = 1.2
        worksheet.append([f"R{index:03d}", date_value, rate])

    scan = _save_and_scan(workbook, tmp_path / "date-percent.xlsx")

    percentage_findings = {
        finding.location: finding
        for finding in scan.findings
        if finding.rule_id == "WL026_PERCENTAGE_SCALE_OUTLIER"
    }
    assert set(percentage_findings) == {"C11", "C12", "C13"}
    assert percentage_findings["C11"].evidence.details["anomaly"] == "whole_percent_scale"
    assert percentage_findings["C12"].evidence.details["anomaly"] == "outside_fraction_range"
    assert percentage_findings["C13"].evidence.details["anomaly"] == "outside_fraction_range"
    assert {
        finding.location
        for finding in scan.findings
        if finding.rule_id == "WL027_DATE_STORAGE_ANOMALY"
    } == {"B11", "B12", "B13"}


def test_autofilter_reports_only_meaningfully_populated_omissions(tmp_path: Path) -> None:
    workbook = Workbook()
    bad = workbook.active
    assert bad is not None
    bad.title = "Bad"
    bad.append(["ID", "Name", "Qty", "Amount"])
    for index in range(1, 11):
        bad.append([f"R{index:03d}", f"Name {index}", index, index * 10])
    bad.auto_filter.ref = "A1:B6"

    clean = workbook.create_sheet("Clean")
    clean.append(["ID", "Name", "Qty", "Amount"])
    for index in range(1, 11):
        clean.append([f"C{index:03d}", f"Name {index}", index, index * 10])
    clean.auto_filter.ref = "A1:D11"

    scan = _save_and_scan(workbook, tmp_path / "filters.xlsx")
    findings = [
        finding for finding in scan.findings if finding.rule_id == "WL030_AUTOFILTER_COVERAGE"
    ]

    assert len(findings) == 1
    assert findings[0].sheet == "Bad"
    assert findings[0].location == "A1:B6"
    assert findings[0].evidence.expected == "A1:D11"


def test_chart_rule_reports_blank_title_length_and_direct_lineage_mismatch(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source"
    source.append(["Group", "Qty"])
    for index in range(1, 8):
        source.append([f"G{index}", index])

    dashboard = workbook.create_sheet("Dashboard")
    dashboard.append(["Group", None])
    for index in range(1, 8):
        dashboard.append([f"G{index}", f"='Source'!B{index + 1}"])
    chart = BarChart()
    chart.title = "Sales by group"
    chart.add_data(
        Reference(dashboard, min_col=2, min_row=1, max_row=8),
        titles_from_data=True,
    )
    chart.set_categories(Reference(dashboard, min_col=1, min_row=2, max_row=7))
    dashboard.add_chart(chart, "D2")

    scan = _save_and_scan(workbook, tmp_path / "chart-sources.xlsx")
    findings = [
        finding for finding in scan.findings if finding.rule_id == "WL031_CHART_SOURCE_STRUCTURE"
    ]

    assert len(findings) == 1
    issue_kinds = {item["kind"] for item in findings[0].evidence.observed["issues"]}
    assert issue_kinds == {
        "blank_series_title_reference",
        "category_value_length_mismatch",
        "title_source_semantic_mismatch",
    }


def test_chart_rule_resolves_sheet_names_case_insensitively(tmp_path: Path) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Data"
    source.append(["Group", "Qty"])
    for index in range(1, 8):
        source.append([f"G{index}", index])

    dashboard = workbook.create_sheet("Dashboard")
    chart = BarChart()
    chart.title = "Quantity by group"
    chart.add_data(
        Reference(source, min_col=2, min_row=1, max_row=8),
        titles_from_data=True,
    )
    chart.set_categories(Reference(source, min_col=1, min_row=2, max_row=8))
    assert chart.ser[0].val is not None
    assert chart.ser[0].val.numRef is not None
    chart.ser[0].val.numRef.f = "data!$B$2:$B$8"
    dashboard.add_chart(chart, "D2")

    scan = _save_and_scan(workbook, tmp_path / "case-insensitive-chart-source.xlsx")

    assert not any(finding.rule_id == "WL031_CHART_SOURCE_STRUCTURE" for finding in scan.findings)


def test_deep_freeze_panes_are_reported_without_flagging_moderate_navigation(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    wide = workbook.active
    assert wide is not None
    wide.title = "Wide"
    wide.append([f"Column {column}" for column in range(1, 13)])
    for row in range(1, 21):
        wide.append([row * column for column in range(1, 13)])
    wide.freeze_panes = "H10"

    moderate = workbook.create_sheet("Moderate")
    moderate.append([f"Column {column}" for column in range(1, 11)])
    for row in range(1, 21):
        moderate.append([row * column for column in range(1, 11)])
    moderate.freeze_panes = "E9"

    outside = workbook.create_sheet("Outside")
    outside.append(["ID", "Name", "Qty", "Amount"])
    for row in range(1, 11):
        outside.append([f"R{row:03d}", f"Name {row}", row, row * 10])
    outside.freeze_panes = "Z2"

    tall = workbook.create_sheet("Tall")
    tall.append(["Metric", "Value"])
    for row in range(1, 18):
        tall.append([f"Metric {row}", row])
    tall.freeze_panes = "B14"

    scan = _save_and_scan(workbook, tmp_path / "freeze-panes.xlsx")
    findings = [finding for finding in scan.findings if finding.rule_id == "WL032_DEEP_FREEZE_PANE"]

    assert {(finding.sheet, finding.location) for finding in findings} == {
        ("Outside", "Z2"),
        ("Tall", "B14"),
        ("Wide", "H10"),
    }
    assert all(not finding.patch_ids for finding in findings)


def test_print_area_reports_truncated_table_and_printed_chart_sources(tmp_path: Path) -> None:
    workbook = Workbook()
    table = workbook.active
    assert table is not None
    table.title = "Table"
    table.append(["ID", "Name", "Qty", "Amount"])
    for index in range(1, 11):
        table.append([f"R{index:03d}", f"Name {index}", index, index * 10])
    table.print_area = "A1:B6"

    split = workbook.create_sheet("Split")
    split.append(["ID", "Name", "Qty", "Amount"])
    for index in range(1, 11):
        split.append([f"S{index:03d}", f"Name {index}", index, index * 10])
    split.print_area = ["A1:B11", "C1:D11"]

    dashboard = workbook.create_sheet("Dashboard")
    dashboard.append(["Metric", "Value"])
    for index in range(1, 10):
        dashboard.append([f"Metric {index}", index])
    dashboard.append([None, None])
    dashboard.append(["Group", "Qty"])
    for index in range(1, 7):
        dashboard.append([f"G{index}", index])
    chart = BarChart()
    chart.title = "Quantity by group"
    chart.add_data(
        Reference(dashboard, min_col=2, min_row=12, max_row=18),
        titles_from_data=True,
    )
    chart.set_categories(Reference(dashboard, min_col=1, min_row=13, max_row=18))
    dashboard.add_chart(chart, "A3")
    dashboard.print_area = "A1:B10"

    scan = _save_and_scan(workbook, tmp_path / "print-areas.xlsx")
    findings = [
        finding for finding in scan.findings if finding.rule_id == "WL033_PRINT_AREA_COVERAGE"
    ]

    assert {finding.sheet for finding in findings} == {"Dashboard", "Table"}
    table_finding = next(finding for finding in findings if finding.sheet == "Table")
    assert table_finding.evidence.observed["issues"] == [
        {"kind": "truncated_inferred_table", "table": "A1:D11"}
    ]
    dashboard_finding = next(finding for finding in findings if finding.sheet == "Dashboard")
    assert [issue["kind"] for issue in dashboard_finding.evidence.observed["issues"]] == [
        "chart_anchor_clipped_by_print_area"
    ]


def test_print_area_respects_two_cell_anchor_end_offsets(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    def add_chart_sheet(name: str, *, col_offset: int, row_offset: int) -> None:
        worksheet = workbook.create_sheet(name)
        worksheet.append(["Group", "Qty"])
        for index in range(1, 10):
            worksheet.append([f"G{index}", index])
        chart = BarChart()
        chart.add_data(
            Reference(worksheet, min_col=2, min_row=1, max_row=10),
            titles_from_data=True,
        )
        chart.set_categories(Reference(worksheet, min_col=1, min_row=2, max_row=10))
        chart.anchor = TwoCellAnchor(
            _from=AnchorMarker(col=0, row=0),
            to=AnchorMarker(col=2, colOff=col_offset, row=10, rowOff=row_offset),
        )
        worksheet.add_chart(chart)
        worksheet.print_area = "A1:B10"

    add_chart_sheet("ExactBoundary", col_offset=0, row_offset=0)
    add_chart_sheet("ColumnOverflow", col_offset=1, row_offset=0)
    add_chart_sheet("RowOverflow", col_offset=0, row_offset=1)

    scan = _save_and_scan(workbook, tmp_path / "two-cell-anchor-print-area.xlsx")
    findings = [
        finding for finding in scan.findings if finding.rule_id == "WL033_PRINT_AREA_COVERAGE"
    ]

    assert {finding.sheet for finding in findings} == {"ColumnOverflow", "RowOverflow"}


def test_chart_rule_accepts_comma_sheet_names_and_named_ranges(tmp_path: Path) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Sales, East"
    source.append(["Group", "Qty"])
    for index in range(1, 8):
        source.append([f"G{index}", index])

    direct = BarChart()
    direct.title = "Quantity by group"
    direct.add_data(Reference(source, min_col=2, min_row=1, max_row=8), titles_from_data=True)
    direct.set_categories(Reference(source, min_col=1, min_row=2, max_row=8))
    source.add_chart(direct, "D2")

    named = BarChart()
    named.title = "Quantity by group"
    named.add_data(Reference(source, min_col=2, min_row=1, max_row=8), titles_from_data=True)
    named.ser[0].val.numRef.f = "QuantitySeries"
    workbook.defined_names.add(DefinedName("QuantitySeries", attr_text="'Sales, East'!$B$2:$B$8"))
    source.add_chart(named, "L2")

    large = BarChart()
    large.title = "Quantity by group"
    large.add_data(Reference(source, min_col=2, min_row=1, max_row=8), titles_from_data=True)
    large.ser[0].val.numRef.f = "'Sales, East'!$B$2:$B$200001"
    source.add_chart(large, "T2")

    scan = _save_and_scan(workbook, tmp_path / "supported-chart-references.xlsx")
    assert not any(finding.rule_id == "WL031_CHART_SOURCE_STRUCTURE" for finding in scan.findings)


def test_clean_profiles_remain_unreported_and_new_rules_localize_strictly(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Clean"
    worksheet.append(["Record ID", "Date", "Qty", "Discount"])
    for index in range(1, 11):
        worksheet.append(
            [f"R{index:03d}", datetime(2026, 1, index), index + 10, 0.05 + index / 1000]
        )
    worksheet.auto_filter.ref = "A1:D11"
    chart = BarChart()
    chart.title = "Quantity by record"
    chart.add_data(Reference(worksheet, min_col=3, min_row=1, max_row=11), titles_from_data=True)
    chart.set_categories(Reference(worksheet, min_col=1, min_row=2, max_row=11))
    worksheet.add_chart(chart, "F2")

    clean_scan = _save_and_scan(workbook, tmp_path / "clean-profiles.xlsx")
    assert not _data_quality_findings(clean_scan)

    dirty_workbook = Workbook()
    dirty = dirty_workbook.active
    assert dirty is not None
    dirty.append(["Record ID", "Name", "Amount"])
    identifiers = ["R1", "R2", "R2", "R4", "R5", "R6", None, "R8"]
    for index, identifier in enumerate(identifiers, start=1):
        dirty.append([identifier, f"Name {index}", index])
    dirty_scan = _save_and_scan(dirty_workbook, tmp_path / "localized-dirty.xlsx")
    localized = localize_scan_result(dirty_scan, "zh-CN", strict=True)
    localized_new = _data_quality_findings(localized)
    canonical_by_id = {finding.id: finding for finding in _data_quality_findings(dirty_scan)}
    assert localized_new
    assert all(finding.title != canonical_by_id[finding.id].title for finding in localized_new)
    assert all(
        not finding.patch_ids and not finding.safe_patch_available for finding in localized_new
    )
