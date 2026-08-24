from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

from workbooklens.rules.layout_geometry import (
    LAYOUT_GEOMETRY_RULES,
    DataRegionRowHeightOutlierRule,
    DrawingContentOverlapRule,
    NumericDisplayWidthRiskRule,
    OrphanMicroLabelRule,
    RoleAwareStyleOutlierRule,
    _style_components,
)
from workbooklens.rules.registry import RuleRegistry
from workbooklens.scanner import ScanResult, scan_workbook


def _scan(workbook: Workbook, path: Path, *rules: type) -> ScanResult:
    workbook.save(path)
    workbook.close()
    return scan_workbook(path, registry=RuleRegistry(rule() for rule in rules))


def _findings(scan: ScanResult, rule_id: str) -> list:
    return [finding for finding in scan.findings if finding.rule_id == rule_id]


def _chart_sheet(
    anchor: str,
    *,
    source_only: bool = False,
    sheet_title: str = "Sheet",
) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = sheet_title
    if not source_only:
        for row in range(1, 7):
            worksheet.cell(row, 1, f"KPI {row}")
            worksheet.cell(row, 2, row * 100)
        source_start, source_end = 10, 15
    else:
        source_start, source_end = 1, 9
    for row in range(source_start, source_end + 1):
        worksheet.cell(row, 1, "Group" if row == source_start else f"G{row}")
        worksheet.cell(row, 2, "Value" if row == source_start else row)
    chart = BarChart()
    chart.add_data(
        Reference(worksheet, min_col=2, min_row=source_start, max_row=source_end),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(worksheet, min_col=1, min_row=source_start + 1, max_row=source_end)
    )
    worksheet.add_chart(chart, anchor)
    return workbook


def test_drawing_overlap_positive_legal_and_boundary(tmp_path: Path) -> None:
    positive = _scan(_chart_sheet("A2"), tmp_path / "overlap.xlsx", DrawingContentOverlapRule)
    findings = _findings(positive, DrawingContentOverlapRule.rule_id)
    assert len(findings) == 1
    assert findings[0].evidence.details["kpi_region_overlap"] is True
    assert "A2" in findings[0].evidence.observed["overlapped_cells"]
    assert not findings[0].patch_ids

    source_only = _scan(
        _chart_sheet("A1", source_only=True),
        tmp_path / "source-only.xlsx",
        DrawingContentOverlapRule,
    )
    source_only_findings = _findings(source_only, DrawingContentOverlapRule.rule_id)
    assert len(source_only_findings) == 1
    assert source_only_findings[0].evidence.observed["overlapped_cells"] == ["A1"]

    boundary = _scan(_chart_sheet("D1"), tmp_path / "adjacent.xlsx", DrawingContentOverlapRule)
    assert not _findings(boundary, DrawingContentOverlapRule.rule_id)


def test_drawing_overlap_keeps_non_source_columns_visible_to_the_rule(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Category", "Series A", "Series B", "Notes"])
    for row in range(2, 11):
        worksheet.append([f"G{row}", row * 10, row * 20, f"Review {row}"])
    chart = BarChart()
    chart.add_data(
        Reference(worksheet, min_col=2, max_col=3, min_row=1, max_row=10),
        titles_from_data=True,
    )
    chart.set_categories(Reference(worksheet, min_col=1, min_row=2, max_row=10))
    worksheet.add_chart(chart, "D2")

    scan = _scan(workbook, tmp_path / "non-source-column-overlap.xlsx", DrawingContentOverlapRule)
    findings = _findings(scan, DrawingContentOverlapRule.rule_id)

    assert len(findings) == 1
    assert "D2" in findings[0].evidence.observed["overlapped_cells"]


def test_drawing_overlap_resolves_chart_sheet_names_case_insensitively(tmp_path: Path) -> None:
    workbook = _chart_sheet("A1", source_only=True, sheet_title="DataSheet")
    worksheet = workbook.active
    assert worksheet is not None
    chart = worksheet._charts[0]
    workbook.defined_names.add(DefinedName("QuantitySeries", attr_text="'DataSheet'!$B$2:$B$9"))
    chart.ser[0].val.numRef.f = "quantityseries"
    for series in chart.ser:
        for attribute in ("val", "cat", "tx"):
            source = getattr(series, attribute, None)
            for reference_name in ("numRef", "strRef"):
                reference = getattr(source, reference_name, None)
                if reference is not None and isinstance(reference.f, str):
                    reference.f = reference.f.replace("DataSheet", "datasheet")

    scan = _scan(
        workbook, tmp_path / "case-insensitive-chart-source.xlsx", DrawingContentOverlapRule
    )
    findings = _findings(scan, DrawingContentOverlapRule.rule_id)

    assert len(findings) == 1
    assert findings[0].evidence.observed["overlapped_cells"] == ["A1"]


def test_drawing_overlap_prefers_case_insensitive_local_defined_name(tmp_path: Path) -> None:
    workbook = _chart_sheet("A1", source_only=True, sheet_title="DataSheet")
    worksheet = workbook.active
    assert worksheet is not None
    chart = worksheet._charts[0]
    worksheet.defined_names.add(
        DefinedName(
            "LocalSeries",
            attr_text="'DataSheet'!$B$2:$B$9",
        )
    )
    chart.ser[0].val.numRef.f = "localseries"
    for source in (chart.ser[0].cat, chart.ser[0].tx):
        for reference_name in ("numRef", "strRef"):
            reference = getattr(source, reference_name, None)
            if reference is not None and isinstance(reference.f, str):
                reference.f = reference.f.replace("DataSheet", "datasheet")

    scan = _scan(
        workbook,
        tmp_path / "case-insensitive-local-defined-name.xlsx",
        DrawingContentOverlapRule,
    )
    findings = _findings(scan, DrawingContentOverlapRule.rule_id)

    assert len(findings) == 1
    assert findings[0].evidence.observed["overlapped_cells"] == ["A1"]
    assert "B2:B9" in findings[0].evidence.peers


def _numeric_workbook(width: float, number_format: str) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Amount"])
    for row in range(1, 9):
        worksheet.append([f"R{row:03d}", 1_234_567.89 + row])
        worksheet.cell(row + 1, 2).number_format = number_format
    worksheet.column_dimensions["B"].width = width
    return workbook


def test_numeric_width_positive_legal_and_boundary(tmp_path: Path) -> None:
    positive = _scan(
        _numeric_workbook(5, "$#,##0.00"),
        tmp_path / "narrow.xlsx",
        NumericDisplayWidthRiskRule,
    )
    findings = _findings(positive, NumericDisplayWidthRiskRule.rule_id)
    assert len(findings) == 1
    assert findings[0].evidence.observed["column"] == "B"
    assert not findings[0].patch_ids

    legal = _scan(
        _numeric_workbook(22, "$#,##0.00"),
        tmp_path / "wide.xlsx",
        NumericDisplayWidthRiskRule,
    )
    assert not _findings(legal, NumericDisplayWidthRiskRule.rule_id)

    boundary = _scan(
        _numeric_workbook(3, "General"),
        tmp_path / "general.xlsx",
        NumericDisplayWidthRiskRule,
    )
    assert not _findings(boundary, NumericDisplayWidthRiskRule.rule_id)


def _row_workbook() -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Name", "Value"])
    for row in range(1, 10):
        worksheet.append([f"R{row:03d}", f"Item {row}", row])
    return workbook


def test_row_height_positive_legal_and_boundary(tmp_path: Path) -> None:
    positive_workbook = _row_workbook()
    assert positive_workbook.active is not None
    positive_workbook.active.row_dimensions[5].height = 60
    positive = _scan(
        positive_workbook, tmp_path / "row-outlier.xlsx", DataRegionRowHeightOutlierRule
    )
    findings = _findings(positive, DataRegionRowHeightOutlierRule.rule_id)
    assert [finding.location for finding in findings] == ["5:5"]
    assert not findings[0].patch_ids

    legal_workbook = _row_workbook()
    assert legal_workbook.active is not None
    legal_workbook.active["B5"] = "A deliberately long wrapped description requiring several lines."
    legal_workbook.active["B5"].alignment = Alignment(wrap_text=True)
    legal_workbook.active.column_dimensions["B"].width = 10
    legal_workbook.active.row_dimensions[5].height = 60
    legal = _scan(legal_workbook, tmp_path / "wrapped.xlsx", DataRegionRowHeightOutlierRule)
    assert not _findings(legal, DataRegionRowHeightOutlierRule.rule_id)

    boundary_workbook = _row_workbook()
    assert boundary_workbook.active is not None
    boundary_workbook.active.row_dimensions[5].height = 37.5
    boundary_workbook.active.row_dimensions[6].height = 9.75
    boundary = _scan(
        boundary_workbook, tmp_path / "row-boundary.xlsx", DataRegionRowHeightOutlierRule
    )
    assert not _findings(boundary, DataRegionRowHeightOutlierRule.rule_id)


def _role_workbook(dirty: bool) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.merge_cells("B1:E2")
    worksheet["B1"] = "Quarterly operating report"
    if dirty:
        worksheet["B1"].font = Font(name="Comic Sans MS", size=22, bold=True, italic=True)
        worksheet["B1"].fill = PatternFill("solid", fgColor="FFFF00")
        worksheet["B1"].alignment = Alignment(horizontal="left", vertical="bottom")
    else:
        worksheet["B1"].font = Font(name="Arial", size=20, bold=True)
        worksheet["B1"].fill = PatternFill("solid", fgColor="D9EAF7")
        worksheet["B1"].alignment = Alignment(horizontal="center", vertical="center")

    for column, value in enumerate(["ID", "Name", "Qty", "Amount", "Status"], start=1):
        cell = worksheet.cell(4, column, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    if dirty:
        worksheet["B4"].number_format = "0.0%"
        worksheet["C4"].font = Font(name="Courier New", size=14, italic=True)
        worksheet["D4"].fill = PatternFill("solid", fgColor="FF9900")
        worksheet["E4"].alignment = Alignment(horizontal="right", vertical="bottom")

    for row in range(5, 15):
        worksheet.cell(row, 1, f"R{row:03d}")
        worksheet.cell(row, 2, f"Item {row}")
        worksheet.cell(row, 3, row)
        worksheet.cell(row, 4, row * 100).number_format = "$#,##0.00"
        worksheet.cell(row, 5, "Open")
    if dirty:
        for coordinate in ("B8", "B9"):
            worksheet[coordinate].font = Font(color="FF0000", italic=True)
            worksheet[coordinate].fill = PatternFill("solid", fgColor="FFFF00")
    else:
        for row in range(5, 15, 2):
            for column in range(1, 6):
                worksheet.cell(row, column).fill = PatternFill("solid", fgColor="EEF5FF")

    worksheet["A16"] = "Total"
    worksheet["D16"] = "=SUM(D5:D14)"
    worksheet["D16"].number_format = "0.0%" if dirty else "$#,##0.00"
    return workbook


def test_role_style_positive_and_legal(tmp_path: Path) -> None:
    positive = _scan(_role_workbook(True), tmp_path / "roles-dirty.xlsx", RoleAwareStyleOutlierRule)
    findings = _findings(positive, RoleAwareStyleOutlierRule.rule_id)
    summaries = [finding.evidence.summary for finding in findings]
    assert any(summary.startswith("Merged title") for summary in summaries)
    assert any(summary.startswith("Header role") for summary in summaries)
    assert any(summary.startswith("Body-role") for summary in summaries)
    assert any("total-row" in summary for summary in summaries)
    assert all(not finding.patch_ids for finding in findings)

    legal = _scan(_role_workbook(False), tmp_path / "roles-clean.xlsx", RoleAwareStyleOutlierRule)
    assert not _findings(legal, RoleAwareStyleOutlierRule.rule_id)


def test_role_style_boundary_single_component_and_single_cell(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.merge_cells("A1:D2")
    worksheet["A1"] = "Compact report"
    worksheet["A1"].font = Font(size=14, bold=True)
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    for column, value in enumerate(["ID", "Name", "Qty", "Amount"], start=1):
        worksheet.cell(4, column, value)
    worksheet["B4"].fill = PatternFill("solid", fgColor="D9EAF7")
    for row in range(5, 12):
        worksheet.cell(row, 1, f"R{row:03d}")
        worksheet.cell(row, 2, f"Item {row}")
        worksheet.cell(row, 3, row)
        worksheet.cell(row, 4, row * 10)
    worksheet["B8"].font = Font(italic=True)

    scan = _scan(workbook, tmp_path / "roles-boundary.xlsx", RoleAwareStyleOutlierRule)
    assert not _findings(scan, RoleAwareStyleOutlierRule.rule_id)


@pytest.mark.parametrize("column_count", [2, 3])
def test_role_style_ignores_first_and_last_detail_row_perimeter_borders(
    tmp_path: Path,
    column_count: int,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append([f"Field {column}" for column in range(1, column_count + 1)])
    side = Side(style="thin", color="336699")
    perimeter = Side(style="medium", color="336699")
    interior = Border(left=side, right=side)
    first_detail = Border(left=side, right=side, top=perimeter)
    last_detail = Border(left=side, right=side, bottom=perimeter)
    for row in range(2, 12):
        worksheet.append([f"R{row}", row * 100, *([f"S{row}"] if column_count == 3 else [])])
        border = first_detail if row == 2 else last_detail if row == 11 else interior
        for cell in worksheet[row]:
            cell.border = border

    scan = _scan(
        workbook,
        tmp_path / f"roles-perimeter-{column_count}-columns.xlsx",
        RoleAwareStyleOutlierRule,
    )

    assert not _findings(scan, RoleAwareStyleOutlierRule.rule_id)


def test_role_style_still_reports_interior_border_difference(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record", "Amount", "Status"])
    side = Side(style="thin", color="336699")
    interior = Border(left=side, right=side)
    different = Border(left=side, right=side, bottom=Side(style="medium", color="336699"))
    for row in range(2, 12):
        worksheet.append([f"R{row}", row * 100, "Open"])
        for cell in worksheet[row]:
            cell.border = interior
    worksheet["B6"].border = different
    worksheet["B7"].border = different

    scan = _scan(
        workbook,
        tmp_path / "roles-interior-border-difference.xlsx",
        RoleAwareStyleOutlierRule,
    )
    finding = next(
        finding
        for finding in _findings(scan, RoleAwareStyleOutlierRule.rule_id)
        if finding.evidence.observed.get("border")
    )

    assert finding.evidence.observed["border"] == ["B6", "B7"]


def test_style_components_handles_missing_border_side() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"].border = Border(left=None)

    components = _style_components(worksheet["A1"])

    assert components["border"][0] == (None, (None,))


def test_layout_geometry_rule_export_order_is_stable() -> None:
    assert (
        DrawingContentOverlapRule,
        NumericDisplayWidthRiskRule,
        DataRegionRowHeightOutlierRule,
        RoleAwareStyleOutlierRule,
        OrphanMicroLabelRule,
    ) == LAYOUT_GEOMETRY_RULES
