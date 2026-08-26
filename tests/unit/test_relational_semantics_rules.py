from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from workbooklens.rules.data_quality import InferredDuplicateIdentifierRule
from workbooklens.rules.registry import RuleRegistry
from workbooklens.rules.relational_semantics import CrossSheetOrphanIdentifierRule
from workbooklens.scanner import ScanResult, scan_workbook


def _save_and_scan(workbook: Workbook, path: Path) -> ScanResult:
    workbook.save(path)
    workbook.close()
    return scan_workbook(
        path,
        registry=RuleRegistry(
            (
                InferredDuplicateIdentifierRule(),
                CrossSheetOrphanIdentifierRule(),
            )
        ),
    )


def _relational_book() -> Workbook:
    workbook = Workbook()
    master = workbook.active
    assert master is not None
    master.title = "Projects"
    master.append(["Project ID", "Project"])
    master_ids = [
        "P-001",
        "P-002",
        "P-003",
        "P-005",
        "P-006",
        "P-007",
        "P-009",
        "P-010",
        "P-011",
        "P-012",
    ]
    for index, project_id in enumerate(master_ids, start=1):
        master.append([project_id, f"Project {index}"])

    expenses = workbook.create_sheet("Expenses")
    expenses.append(["Expense ID", "Project ID", "Amount"])
    expense_ids = [
        "P-001",
        "P-002",
        "P-003",
        "P-004",
        "P-005",
        "P-006",
        "P-007",
        "P-008",
        "P-009",
        "P-010",
        "P-999",
        "P-001",
        "P-002",
        "P-003",
        "P-005",
        "P-006",
        "P-007",
        "P-009",
        "P-010",
        "P-011",
    ]
    for index, project_id in enumerate(expense_ids, start=1):
        expenses.append([f"E-{index:03d}", project_id, index * 10])

    hours = workbook.create_sheet("Hours")
    hours.append(["Entry ID", "Project ID", "Hours"])
    hour_ids = [
        "P-001",
        "P-002",
        "P-003",
        "P-004",
        "P-005",
        "P-006",
        "P-007",
        "P-008",
        "P-009",
        "P-010",
        "P-998",
        "P-001",
        "P-002",
        "P-003",
        "P-005",
        "P-006",
        "P-007",
        "P-009",
        "P-010",
        "P-012",
    ]
    for index, project_id in enumerate(hour_ids, start=1):
        hours.append([f"T-{index:03d}", project_id, index])
    return workbook


def _competing_parent_book(*, competitor_missing: int) -> Workbook:
    workbook = Workbook()
    primary = workbook.active
    assert primary is not None
    primary.title = "Primary"
    primary.append(["Project ID", "Project"])

    project_ids = [f"P-{index:03d}" for index in range(1, 15)]
    for index, project_id in enumerate(project_ids, start=1):
        primary.append([project_id, f"Primary project {index}"])

    competitor = workbook.create_sheet("Competitor")
    competitor.append(["Project ID", "Project"])
    retained = project_ids[: len(project_ids) - competitor_missing]
    replacements = [f"ALT-{index:03d}" for index in range(1, competitor_missing + 1)]
    for index, project_id in enumerate([*retained, *replacements], start=1):
        competitor.append([project_id, f"Competing project {index}"])

    transactions = workbook.create_sheet("Transactions")
    transactions.append(["Project ID", "Amount"])
    for index, project_id in enumerate([*project_ids, *project_ids[:6]], start=1):
        transactions.append([project_id, index * 10])
    return workbook


def test_repeated_foreign_keys_are_not_reported_as_duplicate_primary_keys(
    tmp_path: Path,
) -> None:
    scan = _save_and_scan(_relational_book(), tmp_path / "foreign-keys.xlsx")

    duplicate_findings = [
        finding
        for finding in scan.findings
        if finding.rule_id == InferredDuplicateIdentifierRule.rule_id
    ]
    assert not any(
        finding.sheet in {"Expenses", "Hours"} and "B" in (finding.location or "")
        for finding in duplicate_findings
    )


def test_reports_only_unsupported_child_identifier_not_shared_by_another_child(
    tmp_path: Path,
) -> None:
    scan = _save_and_scan(_relational_book(), tmp_path / "orphans.xlsx")

    orphans = {
        (finding.sheet, finding.location, finding.evidence.observed["normalized_value"])
        for finding in scan.findings
        if finding.rule_id == CrossSheetOrphanIdentifierRule.rule_id
    }
    assert orphans == {
        ("Expenses", "B12", "p-999"),
        ("Hours", "B12", "p-998"),
    }
    assert not any(
        finding.evidence.observed["normalized_value"] in {"p-004", "p-008"}
        for finding in scan.findings
        if finding.rule_id == CrossSheetOrphanIdentifierRule.rule_id
    )
    assert all(
        not finding.patch_ids and not finding.safe_patch_available
        for finding in scan.findings
        if finding.rule_id == CrossSheetOrphanIdentifierRule.rule_id
    )


def test_keeps_true_duplicate_in_parent_and_unrelated_identifier_columns(
    tmp_path: Path,
) -> None:
    workbook = _relational_book()
    projects = workbook["Projects"]
    projects["A11"] = "P-001"
    expenses = workbook["Expenses"]
    expenses["A21"] = "E-001"

    scan = _save_and_scan(workbook, tmp_path / "real-duplicates.xlsx")
    duplicates = {
        (finding.sheet, finding.location)
        for finding in scan.findings
        if finding.rule_id == InferredDuplicateIdentifierRule.rule_id
    }
    assert ("Projects", "A2,A11") in duplicates
    assert ("Expenses", "A2,A21") in duplicates


def test_does_not_infer_relationship_from_low_overlap_same_named_columns(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "First"
    first.append(["Project ID", "Value"])
    second = workbook.create_sheet("Second")
    second.append(["Project ID", "Value"])
    for index in range(1, 13):
        first.append([f"A-{index:03d}", index])
        second.append([f"B-{index:03d}", index])
    second["A13"] = "B-001"

    scan = _save_and_scan(workbook, tmp_path / "low-overlap.xlsx")
    assert not any(
        finding.rule_id == CrossSheetOrphanIdentifierRule.rule_id for finding in scan.findings
    )
    assert any(
        finding.rule_id == InferredDuplicateIdentifierRule.rule_id and finding.sheet == "Second"
        for finding in scan.findings
    )


@pytest.mark.parametrize("competitor_missing", [0, 1])
def test_ambiguous_parent_candidates_do_not_suppress_true_child_duplicates(
    tmp_path: Path,
    competitor_missing: int,
) -> None:
    workbook = _competing_parent_book(competitor_missing=competitor_missing)

    scan = _save_and_scan(
        workbook,
        tmp_path / f"ambiguous-parent-{competitor_missing}.xlsx",
    )
    duplicates = [
        finding
        for finding in scan.findings
        if finding.rule_id == InferredDuplicateIdentifierRule.rule_id
        and finding.sheet == "Transactions"
    ]

    assert duplicates
    assert not any(
        finding.rule_id == CrossSheetOrphanIdentifierRule.rule_id
        and finding.sheet == "Transactions"
        for finding in scan.findings
    )


def test_clear_parent_candidate_still_suppresses_child_duplicates(tmp_path: Path) -> None:
    workbook = _competing_parent_book(competitor_missing=3)

    scan = _save_and_scan(workbook, tmp_path / "clear-parent.xlsx")

    assert not any(
        finding.rule_id == InferredDuplicateIdentifierRule.rule_id
        and finding.sheet == "Transactions"
        for finding in scan.findings
    )
