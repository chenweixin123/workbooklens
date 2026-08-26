from __future__ import annotations

import subprocess
import sys
from copy import copy
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils.datetime import to_excel

from workbooklens.exceptions import UsageError
from workbooklens.models import PatchKind, PatchRisk
from workbooklens.repair.ooxml_patch import patch_ooxml_package
from workbooklens.repair.planning import build_patch_plan
from workbooklens.rules.builtin import NumericTextRule, TextDisplayRiskRule
from workbooklens.rules.profile_quality import (
    MAX_PROFILE_RANGE_CELLS,
    PROFILE_QUALITY_RULE_IDS,
    PROFILE_QUALITY_RULES,
    ContactFormatRule,
    EnumeratedValueRule,
    LeadingZeroIdentifierRule,
    LosslessValueNormalizationRule,
    NumberFormatRoleConflictRule,
    RequiredFieldRule,
    TrailingWhitespaceRule,
    _valid_email,
    parse_workbook_profile,
)
from workbooklens.rules.registry import RuleRegistry
from workbooklens.scanner import ScanResult, scan_workbook


def _save_and_scan(
    workbook: Workbook,
    path: Path,
    *,
    config: dict | None = None,
    rules: tuple[type, ...] = PROFILE_QUALITY_RULES,
) -> ScanResult:
    workbook.save(path)
    workbook.close()
    registry = RuleRegistry(rule_type() for rule_type in rules)
    return scan_workbook(path, config=config, registry=registry)


def _profile(columns: list[dict], **options: object) -> dict:
    return {
        "profile": {
            **options,
            "sheets": [
                {
                    "sheet": "Data",
                    "columns": columns,
                }
            ],
        }
    }


def _versioned_profile(
    columns: list[dict],
    *,
    version: int,
    range_ref: str | None = None,
) -> dict:
    sheet: dict[str, object] = {"sheet": "Data", "columns": columns}
    if range_ref is not None:
        sheet["range"] = range_ref
    return {"version": version, "profile": {"sheets": [sheet]}}


def test_profile_parser_accepts_stable_shape() -> None:
    profile = parse_workbook_profile(
        {
            "profile": {
                "infer_semantics": False,
                "review_trailing_whitespace_patches": True,
                "sheets": [
                    {
                        "sheet": "Data",
                        "columns": [
                            {
                                "header": "Region",
                                "role": "category",
                                "required": True,
                                "allowed_values": ["North", "South"],
                            },
                            {
                                "column": "B",
                                "role": "currency",
                                "preserve_leading_zeros": False,
                            },
                        ],
                    },
                ],
            }
        }
    )

    assert not profile.infer_semantics
    assert profile.review_trailing_whitespace_patches
    assert len(profile.tables) == 1
    table = profile.tables[0]
    assert table.header_row == 1
    assert len(table.columns) == 2
    assert table.columns[0].role == "category"
    assert table.columns[0].required
    assert table.columns[0].allowed_values == ("North", "South")
    assert table.columns[1].column == "B"
    assert table.columns[1].role == "currency"
    assert not table.columns[1].preserve_leading_zeros
    assert table.columns[1].repair is None

    assert parse_workbook_profile({"profile": {}}).tables == ()
    assert parse_workbook_profile(None).tables == ()


def test_profile_parser_accepts_v3_column_repair_and_rejects_it_in_v2() -> None:
    profile = parse_workbook_profile(
        {
            "version": 3,
            "profile": {
                "sheets": [
                    {
                        "sheet": "Data",
                        "columns": [
                            {"header": "Amount", "role": "currency", "repair": "auto"},
                            {"header": "Notes", "role": "text", "repair": "report"},
                        ],
                    }
                ]
            },
        }
    )

    assert profile.version == 3
    assert [column.repair for column in profile.tables[0].columns] == ["auto", "report"]
    with pytest.raises(ValueError, match="repair requires configuration version 3"):
        parse_workbook_profile(
            {
                "version": 2,
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "columns": [{"header": "Amount", "role": "currency", "repair": "auto"}],
                        }
                    ]
                },
            }
        )


@pytest.mark.parametrize(
    ("config", "message"),
    [
        ({"profile": []}, "must be a mapping"),
        ({"profile": {"sheets": "Data"}}, "sheets must be a list"),
        ({"profile": {"sheets": ["Data"]}}, "entry 1 must be a mapping"),
        ({"profile": {"sheets": [{"sheet": ""}]}}, "non-blank sheet name"),
        (
            {"profile": {"sheets": [{"sheet": "Data", "columns": [{"role": "email"}]}]}},
            "exactly one header or column selector",
        ),
        (
            {
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "columns": [{"column": "B", "preserve_leading_zeros": "not-a-bool"}],
                        }
                    ]
                }
            },
            "boolean options",
        ),
        (
            {
                "profile": {
                    "sheets": [
                        {"sheet": "Data", "columns": [{"column": "B", "identifier_width": True}]}
                    ]
                }
            },
            "identifier_width",
        ),
        (
            {"version": 1, "profile": {"sheets": [{"sheet": "Data"}]}},
            "require configuration version 2",
        ),
        (
            {
                "version": 1,
                "workbook_profile": {"tables": [{"sheet": "Data"}]},
            },
            "require configuration version 2",
        ),
        (
            {
                "profile": {"sheets": []},
                "workbook_profile": {"tables": []},
            },
            "cannot define both profile and workbook_profile",
        ),
        (
            {"profile": {"infer_semantcs": False}},
            "unsupported workbook profile keys",
        ),
        (
            {"profile": {"sheets": [{"sheet": "Data", "header_rows": 1}]}},
            "unsupported profile sheet entry 1 keys",
        ),
        (
            {
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "columns": [{"header": "ID", "requried": True}],
                        }
                    ]
                }
            },
            "unsupported profile column entry 1 keys",
        ),
    ],
)
def test_direct_profile_parser_rejects_malformed_contracts(
    config: dict,
    message: str,
) -> None:
    with pytest.raises(ValueError, match=message):
        parse_workbook_profile(config)


def test_profile_rule_exports_are_stable_and_complete() -> None:
    assert {
        "WL036_TRAILING_WHITESPACE",
        "WL037_REQUIRED_FIELD",
        "WL038_ENUM_VALUE",
        "WL039_CONTACT_FORMAT",
        "WL040_LEADING_ZERO_IDENTIFIER",
        "WL041_NUMBER_FORMAT_ROLE_CONFLICT",
        "WL058_LOSSLESS_VALUE_NORMALIZATION",
    } == PROFILE_QUALITY_RULE_IDS
    assert {rule_type.rule_id for rule_type in PROFILE_QUALITY_RULES} == PROFILE_QUALITY_RULE_IDS


def test_trailing_whitespace_reports_literal_fields_and_only_proposes_opted_in_ascii_patch(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Record ID", "Name", "Notes"])
    worksheet.append(["R001", "Alice ", "two internal words"])
    worksheet.append(["R002", "Bob", "Line one\nLine two"])
    worksheet.append(["R003", "Cara\t", "tab suffix"])
    worksheet.append(["R004", "   ", "whitespace-only placeholder"])
    worksheet.append(["R005", '="Dan "', "formula is not a literal field"])

    scan = _save_and_scan(
        workbook,
        tmp_path / "trailing-whitespace.xlsx",
        config=_profile([{"header": "Name", "role": "text", "trim_trailing_whitespace": True}]),
        rules=(TrailingWhitespaceRule,),
    )

    assert {finding.location for finding in scan.findings} == {"B2", "B4"}
    assert len(scan.patches) == 1
    patch = scan.patches[0]
    assert patch.cell == "B2"
    assert patch.kind == PatchKind.SET_TEXT
    assert patch.after == "Alice"
    assert not patch.safe
    assert patch.risk == PatchRisk.LAYOUT_REVIEW
    assert not patch.safe_only_eligible
    assert scan.findings[0].safe_patch_available is False


def test_trailing_whitespace_patch_requires_explicit_review_opt_in(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Name"])
    for index, name in enumerate(["Alice ", "Bob", "Cara", "Dan"], start=1):
        worksheet.append([f"R{index:03d}", name])

    scan = _save_and_scan(
        workbook,
        tmp_path / "trailing-no-patch.xlsx",
        config=_profile([{"header": "Name", "role": "text"}]),
        rules=(TrailingWhitespaceRule,),
    )

    assert [finding.location for finding in scan.findings] == ["B2"]
    assert not scan.patches
    assert not scan.findings[0].patch_ids


def test_trailing_whitespace_in_merged_record_anchor_is_report_only(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Name", "Alias"])
    worksheet.append(["R001", "Alice ", None])
    worksheet.merge_cells("B2:C2")

    config = {
        "profile": {
            "review_trailing_whitespace_patches": True,
            "sheets": [
                {
                    "sheet": "Data",
                    "range": "A1:C2",
                    "columns": [
                        {"header": "Name", "role": "text", "trim_trailing_whitespace": True}
                    ],
                }
            ],
        }
    }
    scan = _save_and_scan(
        workbook,
        tmp_path / "merged-trailing-whitespace.xlsx",
        config=config,
        rules=(TrailingWhitespaceRule,),
    )

    assert [finding.location for finding in scan.findings] == ["B2"]
    assert scan.findings[0].evidence.details["merged_cell"] is True
    assert not scan.patches


def test_required_fields_only_flag_populated_records_and_configured_columns(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Name", "Comment"])
    worksheet.append(["R001", None, None])
    worksheet.append(["R002", "Bob", None])
    worksheet.append(["R003", "   ", "Name is whitespace"])
    worksheet.append([None, None, None])

    scan = _save_and_scan(
        workbook,
        tmp_path / "required.xlsx",
        config=_profile(
            [
                {"header": "ID", "role": "identifier"},
                {"header": "Name", "role": "text", "required": True},
            ]
        ),
        rules=(RequiredFieldRule,),
    )

    assert {finding.location for finding in scan.findings} == {"B2", "B4"}
    assert all(
        finding.evidence.details["evidence_level"] == "PROVEN_STATIC" for finding in scan.findings
    )
    assert not scan.patches


def test_explicit_profile_scans_a_small_table_without_inference_thresholds(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Name"])
    worksheet.append(["R001", None])

    scan = _save_and_scan(
        workbook,
        tmp_path / "small-explicit-profile.xlsx",
        config=_profile([{"header": "Name", "role": "text", "required": True}]),
        rules=(RequiredFieldRule,),
    )

    assert [finding.location for finding in scan.findings] == ["B2"]


def test_explicit_profile_sparse_rows_preserve_required_field_semantics(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Name", "Notes"])
    worksheet.append(['="Formula record"', None, None])
    worksheet.append(["Merged record", None, None])
    worksheet.merge_cells("A3:B3")
    worksheet.append(["Single-column record", None, None])
    worksheet.append([None, None, None])

    scan = _save_and_scan(
        workbook,
        tmp_path / "sparse-body-rows.xlsx",
        config={
            "profile": {
                "sheets": [
                    {
                        "sheet": "Data",
                        "range": "A1:C5",
                        "columns": [{"header": "Name", "required": True}],
                    }
                ]
            }
        },
        rules=(RequiredFieldRule,),
    )

    assert [finding.location for finding in scan.findings] == ["B2", "B4"]


def test_direct_profile_config_rejects_ranges_above_shared_limit(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID"])
    worksheet.append(["R001"])

    with pytest.raises(UsageError, match=rf"limit is {MAX_PROFILE_RANGE_CELLS}"):
        _save_and_scan(
            workbook,
            tmp_path / "oversized-profile-range.xlsx",
            config={
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "range": "A1:K100000",
                            "columns": [{"header": "ID", "required": True}],
                        }
                    ]
                }
            },
            rules=(RequiredFieldRule,),
        )


def test_direct_profile_config_accepts_range_at_shared_limit(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID"])
    worksheet.append(["R001"])

    scan = _save_and_scan(
        workbook,
        tmp_path / "maximum-profile-range.xlsx",
        config={
            "profile": {
                "sheets": [
                    {
                        "sheet": "Data",
                        "range": "A1:J100000",
                        "columns": [{"header": "ID", "required": True}],
                    }
                ]
            }
        },
        rules=(RequiredFieldRule,),
    )

    assert scan.findings == []


def test_direct_profile_parser_rejects_huge_numeric_column_immediately() -> None:
    with pytest.raises(ValueError, match="outside Excel bounds"):
        parse_workbook_profile(
            {
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "columns": [
                                {
                                    "column": 1_000_000_000,
                                    "required": True,
                                }
                            ],
                        }
                    ]
                }
            }
        )


@pytest.mark.parametrize(
    ("range_ref", "message"),
    [
        ("not-a-range", "invalid profile range"),
        ("A1:XFE1", "outside Excel worksheet bounds"),
        ("Data!A1:B2", "must not repeat the sheet name"),
    ],
)
def test_direct_profile_config_rejects_invalid_ranges_instead_of_fallback(
    tmp_path: Path,
    range_ref: str,
    message: str,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID"])
    worksheet.append(["R001"])

    with pytest.raises(UsageError, match=message):
        _save_and_scan(
            workbook,
            tmp_path / f"invalid-profile-range-{range_ref.replace(':', '-')}.xlsx",
            config={
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "range": range_ref,
                            "columns": [{"header": "ID", "required": True}],
                        }
                    ]
                }
            },
            rules=(RequiredFieldRule,),
        )


def test_direct_profile_range_defaults_header_row_to_first_range_row(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet["A5"] = "ID"
    worksheet["B5"] = "Name"
    worksheet["A6"] = "R001"
    worksheet["A7"] = "R002"
    worksheet["B7"] = "Alice"

    scan = _save_and_scan(
        workbook,
        tmp_path / "range-default-header.xlsx",
        config={
            "profile": {
                "sheets": [
                    {
                        "sheet": "Data",
                        "range": "A5:B7",
                        "columns": [{"header": "Name", "required": True}],
                    }
                ]
            }
        },
        rules=(RequiredFieldRule,),
    )

    assert [finding.location for finding in scan.findings] == ["B6"]


@pytest.mark.parametrize("header_row", [True, 0, 1_048_577, "5"])
def test_direct_profile_rejects_invalid_header_rows(header_row: object) -> None:
    with pytest.raises(ValueError, match="header_row"):
        parse_workbook_profile(
            {
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "range": "A5:B10",
                            "header_row": header_row,
                        }
                    ]
                }
            }
        )


def test_direct_profile_rejects_header_row_outside_range() -> None:
    with pytest.raises(ValueError, match="inside the configured range"):
        parse_workbook_profile(
            {
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "range": "A5:B10",
                            "header_row": 1,
                        }
                    ]
                }
            }
        )


def test_direct_profile_rejects_duplicate_table_selectors() -> None:
    with pytest.raises(ValueError, match="duplicate sheet table selectors"):
        parse_workbook_profile(
            {
                "profile": {
                    "sheets": [
                        {"sheet": "Data", "range": "A1:B10"},
                        {"sheet": "data", "range": "A20:B30"},
                    ]
                }
            }
        )


def test_direct_profile_rejects_duplicate_resolved_columns(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Name"])
    worksheet.append(["R001", "Alice"])
    worksheet.append(["R002", "Bob"])

    with pytest.raises(UsageError, match="multiple selectors for worksheet column 2"):
        _save_and_scan(
            workbook,
            tmp_path / "duplicate-resolved-column.xlsx",
            config={
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "range": "A1:B3",
                            "columns": [
                                {"header": "Name", "required": True},
                                {"column": "B", "role": "text"},
                            ],
                        }
                    ]
                }
            },
            rules=(RequiredFieldRule,),
        )


def test_direct_profile_rejects_missing_or_ambiguous_headers(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Name", "Name"])
    worksheet.append(["Alice", "Bob"])

    with pytest.raises(UsageError, match="ambiguous"):
        _save_and_scan(
            workbook,
            tmp_path / "ambiguous-profile-header.xlsx",
            config=_profile([{"header": "Name", "required": True}]),
            rules=(RequiredFieldRule,),
        )

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Value"])
    worksheet.append(["R001", 1])

    with pytest.raises(UsageError, match="was not found"):
        _save_and_scan(
            workbook,
            tmp_path / "missing-profile-header.xlsx",
            config=_profile([{"header": "Missing", "required": True}]),
            rules=(RequiredFieldRule,),
        )


def test_direct_profile_rejects_missing_worksheet(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID"])
    worksheet.append(["R001"])

    with pytest.raises(UsageError, match="profile worksheet 'Missing' does not exist"):
        _save_and_scan(
            workbook,
            tmp_path / "missing-profile-sheet.xlsx",
            config={
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Missing",
                            "columns": [{"header": "ID", "required": True}],
                        }
                    ]
                }
            },
            rules=(RequiredFieldRule,),
        )


def test_direct_profile_rejects_unresolved_bounds(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"

    with pytest.raises(UsageError, match="table bounds could not be resolved"):
        _save_and_scan(
            workbook,
            tmp_path / "unresolved-profile-bounds.xlsx",
            config=_profile([{"header": "ID", "required": True}]),
            rules=(RequiredFieldRule,),
        )


def test_direct_profile_rejects_blank_header_row(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet["A2"] = "R001"

    with pytest.raises(UsageError, match="does not contain a usable header"):
        _save_and_scan(
            workbook,
            tmp_path / "blank-profile-header.xlsx",
            config={
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "range": "A1:B2",
                            "columns": [{"column": "A", "required": True}],
                        }
                    ]
                }
            },
            rules=(RequiredFieldRule,),
        )


def test_direct_profile_rejects_column_outside_range(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Name"])
    worksheet.append(["R001", "Alice"])

    with pytest.raises(UsageError, match="column 3 is outside the configured table range"):
        _save_and_scan(
            workbook,
            tmp_path / "outside-profile-column.xlsx",
            config={
                "profile": {
                    "sheets": [
                        {
                            "sheet": "Data",
                            "range": "A1:B2",
                            "columns": [{"column": "C", "required": True}],
                        }
                    ]
                }
            },
            rules=(RequiredFieldRule,),
        )


def test_enumeration_reports_unknown_values_and_only_suggests_unique_near_match(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Region"])
    for row in [
        ["R001", "North"],
        ["R002", "NORTH"],
        ["R003", "Nroth"],
        ["R004", "Other"],
        ["R005", None],
    ]:
        worksheet.append(row)

    scan = _save_and_scan(
        workbook,
        tmp_path / "enumeration.xlsx",
        config=_profile(
            [
                {"header": "ID", "role": "identifier"},
                {
                    "header": "Region",
                    "role": "category",
                    "allowed_values": ["North", "South"],
                },
            ]
        ),
        rules=(EnumeratedValueRule,),
    )

    findings = {finding.location: finding for finding in scan.findings}
    assert set(findings) == {"B4", "B5"}
    assert findings["B4"].evidence.details["near_match"] == "North"
    assert "near_match" not in findings["B5"].evidence.details
    assert not scan.patches


def test_contact_validation_accepts_common_legitimate_forms_and_skips_blanks(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Email", "Phone"])
    worksheet.append(["R001", "alice@example.com", "+86 (10) 1234-5678 ext 9"])
    worksheet.append(["R002", "alice@@example.com", ")1234567("])
    worksheet.append(["R003", None, None])
    worksheet.append(["R004", "bob.smith@sub.example.org", 13800138000])
    worksheet.append(["R005", "first.middle.last+tag@example.co.uk", None])
    worksheet.append(["R006", ".leading@example.com", None])
    worksheet.append(["R007", "trailing.@example.com", None])
    worksheet.append(["R008", "double..dot@example.com", None])

    scan = _save_and_scan(
        workbook,
        tmp_path / "contacts.xlsx",
        config=_profile(
            [
                {"header": "Email", "role": "email"},
                {"header": "Phone", "role": "phone"},
            ]
        ),
        rules=(ContactFormatRule,),
    )

    assert {finding.location for finding in scan.findings} == {
        "B3",
        "C3",
        "B7",
        "B8",
        "B9",
    }
    assert all(float(finding.confidence) == 0.99 for finding in scan.findings)
    assert not scan.patches


def test_email_validation_rejects_ambiguous_dot_input_without_backtracking() -> None:
    assert _valid_email("first.middle.last+tag@example.co.uk")
    assert not _valid_email(".leading@example.com")
    assert not _valid_email("trailing.@example.com")
    assert not _valid_email("double..dot@example.com")
    assert not _valid_email(("!." * 10_000) + "@example.com")

    script = (
        "from workbooklens.rules.profile_quality import _valid_email; "
        "assert not _valid_email(('!.' * 30) + '!')"
    )
    completed = subprocess.run(  # noqa: S603 - fixed interpreter and constant script
        [sys.executable, "-c", script],
        capture_output=True,
        check=False,
        text=True,
        timeout=5,
    )
    assert completed.returncode == 0, completed.stderr


def test_leading_zero_rule_respects_zero_mask_and_non_numeric_identifiers(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Account ID", "Name"])
    worksheet.append([123, "Missing zeros"])
    worksheet.append([456, "Displayed with zeros"])
    worksheet["A3"].number_format = "000000"
    worksheet.append(["000789", "Stored as text"])
    worksheet.append(["000321", "Stored as text"])
    worksheet.append(["000654", "Stored as text"])
    worksheet.append(["000987", "Stored as text"])
    worksheet.append(["0000456", "Extra zero padding"])
    worksheet.append(["A001", "Alphanumeric identifier"])

    scan = _save_and_scan(
        workbook,
        tmp_path / "leading-zero-explicit.xlsx",
        config=_profile(
            [
                {
                    "header": "Account ID",
                    "role": "identifier",
                    "preserve_leading_zeros": True,
                }
            ]
        ),
        rules=(LeadingZeroIdentifierRule,),
    )

    assert [finding.location for finding in scan.findings] == ["A2", "A8"]
    assert {finding.evidence.details["issue_kind"] for finding in scan.findings} == {
        "too_short",
        "extra_leading_zero_padding",
    }
    assert all(finding.evidence.expected["width"] == 6 for finding in scan.findings)
    assert not scan.patches


def test_leading_zero_rule_can_infer_a_dominant_fixed_width_pattern(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Record ID", "Name"])
    for identifier in ["0001", "0002", "0003", "0004", 5, "000006"]:
        worksheet.append([identifier, f"Record {identifier}"])

    scan = _save_and_scan(
        workbook,
        tmp_path / "leading-zero-inferred.xlsx",
        rules=(LeadingZeroIdentifierRule,),
    )

    assert [finding.location for finding in scan.findings] == ["A6", "A7"]
    assert {finding.evidence.details["issue_kind"] for finding in scan.findings} == {
        "too_short",
        "extra_leading_zero_padding",
    }
    assert all(
        finding.evidence.details["evidence_level"] == "STRONG_STRUCTURAL"
        for finding in scan.findings
    )


def test_no_profile_does_not_invent_enum_or_variable_width_identifier_constraints(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Record ID", "Region"])
    for identifier, region in [
        ("1", "North"),
        ("02", "Nroth"),
        ("003", "Custom"),
        ("AB04", "South"),
        ("50000", "Elsewhere"),
    ]:
        worksheet.append([identifier, region])

    scan = _save_and_scan(
        workbook,
        tmp_path / "no-profile-legitimate-variation.xlsx",
        rules=(EnumeratedValueRule, LeadingZeroIdentifierRule),
    )

    assert not scan.findings
    assert not scan.patches


def test_number_format_rule_reports_clear_role_conflicts_but_not_general_or_valid_formats(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["ID", "Amount", "Rate", "Date"])
    worksheet.append(["R001", 2500, 0.2, datetime(2026, 1, 1)])
    worksheet["B2"].number_format = "0.00%"
    worksheet["C2"].number_format = "$#,##0.00"
    worksheet["D2"].number_format = "0.0%"
    worksheet.append(["R002", 3000, 0.25, datetime(2026, 1, 2)])
    worksheet["B3"].number_format = "¥#,##0.00"
    worksheet["C3"].number_format = "0.0%"
    worksheet["D3"].number_format = "[$-409]mmm d, yyyy"
    worksheet.append(["R003", 3500, 0.3, 46025])

    scan = _save_and_scan(
        workbook,
        tmp_path / "format-roles.xlsx",
        config=_profile(
            [
                {"header": "Amount", "role": "currency"},
                {"header": "Rate", "role": "percentage"},
                {"header": "Date", "role": "date"},
            ]
        ),
        rules=(NumberFormatRoleConflictRule,),
    )

    assert {finding.location for finding in scan.findings} == {"B2", "C2", "D2"}
    assert all(
        finding.evidence.details["evidence_level"] == "PROVEN_STATIC" for finding in scan.findings
    )
    assert not scan.patches


def test_wl058_strict_literals_produce_safe_native_storage_patches(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Amount", "Rate", "Date", "Quantity", "Notes"])
    for row in range(2, 6):
        worksheet.append(
            [
                row * 1000,
                row / 100,
                date(2026, 8, row),
                row * 100,
                f"note {row}",
            ]
        )
        worksheet[f"A{row}"].number_format = "$#,##0.00"
        worksheet[f"B{row}"].number_format = "0.0%"
        worksheet[f"C{row}"].number_format = "yyyy-mm-dd"
        worksheet[f"D{row}"].number_format = "#,##0"
    worksheet.append(["$1,234.50", "12.5%", "2026-08-25", "1,234  ", "123  "])

    scan = _save_and_scan(
        workbook,
        tmp_path / "strict-lossless-literals.xlsx",
        config=_versioned_profile(
            [
                {"header": "Amount", "role": "currency", "repair": "auto"},
                {"header": "Rate", "role": "percentage", "repair": "auto"},
                {"header": "Date", "role": "date", "repair": "auto"},
                {"header": "Quantity", "role": "number", "repair": "auto"},
                {"header": "Notes", "role": "text", "repair": "auto"},
            ],
            version=3,
            range_ref="A1:E6",
        ),
        rules=(LosslessValueNormalizationRule,),
    )

    value_patches = {
        patch.cell: patch
        for patch in scan.patches
        if patch.kind in {PatchKind.SET_NUMERIC, PatchKind.NORMALIZE_TEXT}
    }
    assert set(value_patches) == {"A6", "B6", "C6", "D6", "E6"}
    assert value_patches["A6"].after == 1234.5
    assert value_patches["B6"].after == 0.125
    assert value_patches["C6"].after == int(to_excel(date(2026, 8, 25)))
    assert value_patches["D6"].after == 1234
    assert value_patches["E6"].after == "123"
    assert value_patches["E6"].kind == PatchKind.NORMALIZE_TEXT
    assert all(patch.safe and patch.risk == PatchRisk.SAFE for patch in scan.patches)

    format_patches = {
        patch.cell: patch for patch in scan.patches if patch.kind == PatchKind.COPY_NUMBER_FORMAT
    }
    assert set(format_patches) == {"A6", "B6", "C6"}
    for cell in format_patches:
        assert format_patches[cell].atomic_group == value_patches[cell].atomic_group
        assert format_patches[cell].atomic_group is not None
        assert format_patches[cell].source_cell in {
            f"{cell[0]}2",
            f"{cell[0]}3",
            f"{cell[0]}4",
            f"{cell[0]}5",
        }


def test_wl058_rejects_leading_zero_and_malformed_literals(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Quantity", "Amount", "Rate", "Date"])
    for row in range(2, 6):
        worksheet.append([row * 10, row * 1000, row / 100, date(2026, 8, row)])
        worksheet[f"B{row}"].number_format = "$#,##0.00"
        worksheet[f"C{row}"].number_format = "0.0%"
        worksheet[f"D{row}"].number_format = "yyyy-mm-dd"
    worksheet.append(["00123", "$12,34", "12%%", "2026-02-30"])

    scan = _save_and_scan(
        workbook,
        tmp_path / "rejected-lossless-literals.xlsx",
        config=_versioned_profile(
            [
                {"header": "Quantity", "role": "number", "repair": "auto"},
                {"header": "Amount", "role": "currency", "repair": "auto"},
                {"header": "Rate", "role": "percentage", "repair": "auto"},
                {"header": "Date", "role": "date", "repair": "auto"},
            ],
            version=3,
            range_ref="A1:D6",
        ),
        rules=(LosslessValueNormalizationRule,),
    )

    assert not scan.findings
    assert not scan.patches


@pytest.mark.parametrize(
    ("repair_mode", "expected_risk"),
    [
        ("auto", PatchRisk.SAFE),
        ("review", PatchRisk.SEMANTIC_REVIEW),
        ("report", None),
    ],
)
def test_profile_v3_repair_modes_gate_wl058_authority(
    tmp_path: Path,
    repair_mode: str,
    expected_risk: PatchRisk | None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Business Field"])
    for value in (100, 200, 300, 400):
        worksheet.append([value])
    worksheet.append(["1200"])

    scan = _save_and_scan(
        workbook,
        tmp_path / f"profile-v3-{repair_mode}.xlsx",
        config=_versioned_profile(
            [{"header": "Business Field", "role": "number", "repair": repair_mode}],
            version=3,
            range_ref="A1:A6",
        ),
        rules=(LosslessValueNormalizationRule,),
    )

    assert [finding.location for finding in scan.findings] == ["A6"]
    patches = [
        patch
        for patch in scan.patches
        if patch.cell == "A6" and patch.kind == PatchKind.SET_NUMERIC
    ]
    if expected_risk is None:
        assert not patches
        assert not scan.findings[0].patch_ids
        return
    assert len(patches) == 1
    assert patches[0].risk == expected_risk
    assert patches[0].safe is (expected_risk == PatchRisk.SAFE)


def test_v1_and_v2_configs_do_not_gain_implicit_profile_repair_authority(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Business Field"])
    for value in (100, 200, 300, 400):
        worksheet.append([value])
    worksheet.append(["1200"])

    v2_scan = _save_and_scan(
        workbook,
        tmp_path / "profile-v2-no-repair.xlsx",
        config=_versioned_profile(
            [{"header": "Business Field", "role": "number"}],
            version=2,
            range_ref="A1:A6",
        ),
        rules=(LosslessValueNormalizationRule,),
    )
    assert [finding.location for finding in v2_scan.findings] == ["A6"]
    assert not v2_scan.patches

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Quantity", "Context"])
    for index, value in enumerate((100, 200, 300, 400), start=1):
        worksheet.append([value, f"R{index}"])
    worksheet.append(["1200", "R5"])
    v1_scan = _save_and_scan(
        workbook,
        tmp_path / "profile-v1-inferred.xlsx",
        config={"version": 1},
        rules=(LosslessValueNormalizationRule,),
    )
    patch = next(
        patch
        for patch in v1_scan.patches
        if patch.cell == "A6" and patch.kind == PatchKind.SET_NUMERIC
    )
    assert patch.risk == PatchRisk.SAFE
    assert patch.safe


def test_chinese_number_candidates_are_semantic_review_and_require_recalculation(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Amount", "Hours", "Multiplier"])
    for row in range(2, 6):
        worksheet.append([row * 10_000, row * 2, 1 + row / 10])
        worksheet[f"A{row}"].number_format = "$#,##0.00"
        worksheet[f"C{row}"].number_format = "0.0x"
    worksheet.append(["八万九千元", "八小时", "1.5倍"])
    worksheet["A6"].number_format = "$#,##0.00"
    worksheet["C6"].number_format = "0.0x"

    scan = _save_and_scan(
        workbook,
        tmp_path / "chinese-semantic-review.xlsx",
        config=_versioned_profile(
            [
                {"header": "Amount", "role": "currency", "repair": "auto"},
                {"header": "Hours", "role": "number", "repair": "auto"},
                {"header": "Multiplier", "role": "number", "repair": "auto"},
            ],
            version=3,
            range_ref="A1:C6",
        ),
        rules=(LosslessValueNormalizationRule,),
    )

    patches = {patch.cell: patch for patch in scan.patches if patch.kind == PatchKind.SET_NUMERIC}
    assert patches["A6"].after == 89_000
    assert patches["B6"].after == 8
    assert patches["C6"].after == 1.5
    assert all(patch.risk == PatchRisk.SEMANTIC_REVIEW for patch in patches.values())
    assert all(not patch.safe for patch in patches.values())
    assert all(patch.derivation.requires_recalculation for patch in patches.values())
    assert all(patch.derivation.candidate_count == 1 for patch in patches.values())
    assert all(
        patch.derivation.strategy == "semantic_numeric_normalization" for patch in patches.values()
    )
    assert all(
        any(source.startswith("profile_column_role:") for source in patch.derivation.sources)
        for patch in patches.values()
    )
    assert all(
        "explicit_v3_numeric_role" in patch.derivation.invariants for patch in patches.values()
    )


@pytest.mark.parametrize(
    ("repair_mode", "expects_patch"),
    [(None, False), ("report", False), ("review", True), ("auto", True)],
)
def test_v3_semantic_number_requires_explicit_repair_authorization(
    tmp_path: Path,
    repair_mode: str | None,
    expects_patch: bool,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Hours"])
    for value in (2, 4, 6, 10):
        worksheet.append([value])
    worksheet.append(["八小时"])
    column = {"header": "Hours", "role": "number"}
    if repair_mode is not None:
        column["repair"] = repair_mode

    scan = _save_and_scan(
        workbook,
        tmp_path / f"semantic-repair-{repair_mode or 'omitted'}.xlsx",
        config=_versioned_profile(
            [column],
            version=3,
            range_ref="A1:A6",
        ),
        rules=(LosslessValueNormalizationRule,),
    )

    assert [finding.location for finding in scan.findings] == ["A6"]
    patches = [
        patch
        for patch in scan.patches
        if patch.cell == "A6" and patch.kind == PatchKind.SET_NUMERIC
    ]
    assert bool(patches) is expects_patch
    if expects_patch:
        assert len(patches) == 1
        assert patches[0].risk == PatchRisk.SEMANTIC_REVIEW
        assert patches[0].after == 8
    else:
        assert not scan.findings[0].patch_ids
        assert (
            "semantic_repair_authorization_required"
            in scan.findings[0].evidence.details["blocked_reasons"]
        )


def test_same_cell_layout_review_does_not_hide_semantic_review_candidate(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.column_dimensions["A"].width = 2
    worksheet.append(["Hours", "Context"])
    for row in range(2, 6):
        worksheet.append([row * 2, f"R{row}"])
    worksheet.append(["八小时", "R6"])

    scan = _save_and_scan(
        workbook,
        tmp_path / "layout-and-semantic-review.xlsx",
        config=_versioned_profile(
            [{"header": "Hours", "role": "number", "repair": "auto"}],
            version=3,
            range_ref="A1:B6",
        ),
        rules=(TextDisplayRiskRule, LosslessValueNormalizationRule),
    )

    same_cell = [patch for patch in scan.patches if patch.cell == "A6"]
    assert any(patch.risk == PatchRisk.LAYOUT_REVIEW for patch in same_cell)
    semantic = next(
        patch
        for patch in same_cell
        if patch.kind == PatchKind.SET_NUMERIC and patch.risk == PatchRisk.SEMANTIC_REVIEW
    )
    assert semantic.after == 8
    finding = next(
        finding
        for finding in scan.findings
        if finding.rule_id == "WL058_LOSSLESS_VALUE_NORMALIZATION" and finding.location == "A6"
    )
    assert finding.patch_ids == [semantic.id]


def test_inferred_semantic_numbers_are_reported_without_applicable_patches(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["已支出", "工时", "加班系数"])
    for row in range(2, 6):
        worksheet.append([row * 10_000, row * 2, 1 + row / 10])
        worksheet[f"A{row}"].number_format = "¥#,##0"
        worksheet[f"C{row}"].number_format = "0.0x"
    worksheet.append(["八万九千", "八小时", "1.5倍"])
    worksheet["A6"].number_format = "¥#,##0"
    worksheet["C6"].number_format = "0.0x"

    scan = _save_and_scan(
        workbook,
        tmp_path / "inferred-semantic-review.xlsx",
        rules=(LosslessValueNormalizationRule,),
    )

    findings = {finding.location: finding for finding in scan.findings}
    assert set(findings) == {"A6", "B6", "C6"}
    assert {
        location: finding.evidence.expected["value"] for location, finding in findings.items()
    } == {
        "A6": 89_000,
        "B6": 8,
        "C6": 1.5,
    }
    assert all(finding.evidence.details["semantic_candidate"] for finding in findings.values())
    assert all(
        "explicit_v3_numeric_profile_required" in finding.evidence.details["blocked_reasons"]
        for finding in findings.values()
    )
    assert scan.patches == []
    assert not any(patch.risk == PatchRisk.SEMANTIC_REVIEW for patch in scan.patches)


@pytest.mark.parametrize("blocked", ["hidden", "merged", "protected", "text_format"])
def test_semantic_number_patch_requires_a_safe_profile_target(
    tmp_path: Path,
    blocked: str,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Multiplier", "Context"])
    for row in range(2, 6):
        worksheet.append([1 + row / 10, f"R{row}"])
        worksheet[f"A{row}"].number_format = "0.0x"
    worksheet.append(["1.5倍", "R6"])
    worksheet["A6"].number_format = "0.0x"
    if blocked == "hidden":
        worksheet.row_dimensions[6].hidden = True
    elif blocked == "merged":
        worksheet.merge_cells("A6:B6")
    elif blocked == "protected":
        worksheet.protection.sheet = True
    else:
        worksheet["A6"].number_format = "@"

    scan = _save_and_scan(
        workbook,
        tmp_path / f"semantic-{blocked}.xlsx",
        config=_versioned_profile(
            [{"header": "Multiplier", "role": "number", "repair": "auto"}],
            version=3,
            range_ref="A1:B6",
        ),
        rules=(LosslessValueNormalizationRule,),
    )

    assert not any(patch.cell == "A6" for patch in scan.patches)


@pytest.mark.parametrize(
    ("repair_mode", "expected_risk"),
    [
        ("auto", PatchRisk.SAFE),
        ("review", PatchRisk.SEMANTIC_REVIEW),
        ("report", None),
    ],
)
def test_wl006_defers_to_explicit_v3_wl058_repair_mode(
    tmp_path: Path,
    repair_mode: str,
    expected_risk: PatchRisk | None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Quantity", "Context"])
    for row in range(2, 22):
        worksheet.append([row * 100, f"R{row}"])
    worksheet["A10"] = "1200"

    scan = _save_and_scan(
        workbook,
        tmp_path / f"wl006-v3-{repair_mode}.xlsx",
        config=_versioned_profile(
            [{"header": "Quantity", "role": "number", "repair": repair_mode}],
            version=3,
            range_ref="A1:B21",
        ),
        rules=(NumericTextRule, LosslessValueNormalizationRule),
    )

    wl006 = next(
        finding
        for finding in scan.findings
        if finding.rule_id == "WL006_NUMERIC_TEXT" and finding.location == "A10"
    )
    assert not wl006.patch_ids
    numeric_patches = [
        patch
        for patch in scan.patches
        if patch.cell == "A10" and patch.kind == PatchKind.SET_NUMERIC
    ]
    if expected_risk is None:
        assert not numeric_patches
        return
    assert len(numeric_patches) == 1
    assert numeric_patches[0].risk == expected_risk


def test_wl006_keeps_legacy_v2_authority_without_duplicate_wl058_patch(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Quantity", "Context"])
    for row in range(2, 22):
        worksheet.append([row * 100, f"R{row}"])
    worksheet["A10"] = "1200"

    scan = _save_and_scan(
        workbook,
        tmp_path / "wl006-v2-compatibility.xlsx",
        config=_versioned_profile(
            [{"header": "Quantity", "role": "number"}],
            version=2,
            range_ref="A1:B21",
        ),
        rules=(NumericTextRule, LosslessValueNormalizationRule),
    )

    wl006 = next(
        finding
        for finding in scan.findings
        if finding.rule_id == "WL006_NUMERIC_TEXT" and finding.location == "A10"
    )
    numeric_patches = [
        patch
        for patch in scan.patches
        if patch.cell == "A10" and patch.kind == PatchKind.SET_NUMERIC
    ]
    assert len(numeric_patches) == 1
    assert wl006.patch_ids == [numeric_patches[0].id]
    assert numeric_patches[0].kind == PatchKind.SET_NUMERIC
    assert not any(
        finding.rule_id == "WL058_LOSSLESS_VALUE_NORMALIZATION" and finding.location == "A10"
        for finding in scan.findings
    )


def test_copy_number_format_preserves_all_other_target_style_fields(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet.append(["Amount", "Context"])
    for row in range(2, 6):
        worksheet.append([row * 1000, f"R{row}"])
        worksheet[f"A{row}"].number_format = "$#,##0.00"
    worksheet.append(["$1,234.50", "R6"])
    target = worksheet["A6"]
    target.font = Font(name="Arial", size=13, bold=True, color="FF112233")
    target.fill = PatternFill(fill_type="solid", fgColor="FFABCDEF")
    target.border = Border(
        left=Side(style="thick", color="FF010203"),
        right=Side(style="double", color="FF040506"),
    )
    target.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    target.protection = Protection(locked=False, hidden=True)
    source = tmp_path / "copy-number-format-source.xlsx"
    scan = _save_and_scan(
        workbook,
        source,
        config=_versioned_profile(
            [{"header": "Amount", "role": "currency", "repair": "auto"}],
            version=3,
            range_ref="A1:B6",
        ),
        rules=(LosslessValueNormalizationRule,),
    )
    value_patch = next(
        patch
        for patch in scan.patches
        if patch.cell == "A6" and patch.kind == PatchKind.SET_NUMERIC
    )
    format_patch = next(
        patch
        for patch in scan.patches
        if patch.cell == "A6" and patch.kind == PatchKind.COPY_NUMBER_FORMAT
    )
    assert value_patch.atomic_group == format_patch.atomic_group
    assert value_patch.atomic_group is not None

    original = load_workbook(source)
    original_target = original["Data"]["A6"]
    expected_font = copy(original_target.font)
    expected_fill = copy(original_target.fill)
    expected_border = copy(original_target.border)
    expected_alignment = copy(original_target.alignment)
    expected_protection = copy(original_target.protection)
    original.close()

    plan = build_patch_plan(scan)
    output = tmp_path / "copy-number-format-output.xlsx"
    low_level, applied = patch_ooxml_package(
        source,
        plan,
        output,
        selected_ids={value_patch.id, format_patch.id},
        canonical_plan=plan,
    )
    assert {patch.id for patch in applied} == {value_patch.id, format_patch.id}
    assert not low_level.formula_changed

    repaired = load_workbook(output)
    repaired_target = repaired["Data"]["A6"]
    repaired_format_source = repaired["Data"][format_patch.source_cell]
    assert repaired_target.value == 1234.5
    assert repaired_target.number_format == repaired_format_source.number_format
    assert copy(repaired_target.font) == expected_font
    assert copy(repaired_target.fill) == expected_fill
    assert copy(repaired_target.border) == expected_border
    assert copy(repaired_target.alignment) == expected_alignment
    assert copy(repaired_target.protection) == expected_protection
    repaired.close()


def test_profile_rules_reject_hidden_sheets_when_explicitly_configured(tmp_path: Path) -> None:
    workbook = Workbook()
    visible = workbook.active
    assert visible is not None
    visible.title = "Visible"
    visible.append(["ID", "Name"])
    visible.append(["V001", "Valid"])
    hidden = workbook.create_sheet("Data")
    hidden.sheet_state = "hidden"
    hidden.append(["ID", "Name"])
    hidden.append(["H001", None])

    with pytest.raises(UsageError, match="profile worksheet 'Data' is not visible"):
        _save_and_scan(
            workbook,
            tmp_path / "hidden-profile.xlsx",
            config=_profile([{"header": "Name", "role": "text", "required": True}]),
            rules=(RequiredFieldRule,),
        )
