from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import PatternFill

from workbooklens.rules import RuleRegistry
from workbooklens.rules.data_quality import ConditionalFormatSemanticConflictRule
from workbooklens.scanner import ScanResult, scan_workbook


def _scan(tmp_path: Path, workbook: Workbook, name: str) -> ScanResult:
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    return scan_workbook(
        path,
        registry=RuleRegistry([ConditionalFormatSemanticConflictRule()]),
    )


def _quantity_book(
    fill_rgb: str,
    *,
    header: str = "Quantity",
    operator: str = "lessThan",
    values: list[int] | None = None,
) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", header])
    body_values = values or [2, 4, 6, -1, 8, 10, 12, 14]
    for row, value in enumerate(body_values, start=2):
        worksheet.append([f"R{row:03d}", value])
    worksheet.conditional_formatting.add(
        "B2:B9",
        CellIsRule(
            operator=operator,
            formula=["0"],
            fill=PatternFill(fill_type="solid", fgColor=fill_rgb),
        ),
    )
    return workbook


def test_duplicate_success_green_rules_are_semantically_deduplicated(tmp_path: Path) -> None:
    workbook = _quantity_book("FF00B050")
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.conditional_formatting.add(
        "B2:B9",
        CellIsRule(
            operator="lessThan",
            formula=["0"],
            fill=PatternFill(fill_type="solid", fgColor="FF00B050"),
        ),
    )

    scan = _scan(tmp_path, workbook, "duplicate-negative-success-green.xlsx")

    assert [(finding.rule_id, finding.location) for finding in scan.findings] == [
        ("WL051_CONDITIONAL_FORMAT_SEMANTIC_CONFLICT", "B2:B9")
    ]


def test_zero_in_positive_domain_is_not_colored_success_green(tmp_path: Path) -> None:
    scan = _scan(
        tmp_path,
        _quantity_book(
            "FF00B050",
            header="Price",
            operator="lessThanOrEqual",
            values=[2, 4, 6, 0, 8, 10, 12, 14],
        ),
        "zero-success-green.xlsx",
    )

    assert [(finding.rule_id, finding.location) for finding in scan.findings] == [
        ("WL051_CONDITIONAL_FORMAT_SEMANTIC_CONFLICT", "B2:B9")
    ]
    assert scan.findings[0].evidence.observed["violating_cells"] == ["B5"]
    assert scan.findings[0].evidence.observed["zero_cells"] == ["B5"]


def test_prior_stop_if_true_rule_suppresses_later_success_green_rule(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Quantity"])
    for row, value in enumerate([2, 4, 6, -1, 8, 10, 12, 14], start=2):
        worksheet.append([f"R{row:03d}", value])

    green_rule = CellIsRule(
        operator="lessThan",
        formula=["0"],
        fill=PatternFill(fill_type="solid", fgColor="FF00B050"),
    )
    green_rule.priority = 2
    worksheet.conditional_formatting.add("B2:B9", green_rule)
    red_rule = CellIsRule(
        operator="lessThan",
        formula=["0"],
        fill=PatternFill(fill_type="solid", fgColor="FFFF0000"),
        stopIfTrue=True,
    )
    red_rule.priority = 1
    worksheet.conditional_formatting.add("B2:B9", red_rule)

    scan = _scan(tmp_path, workbook, "negative-stop-if-true.xlsx")

    assert not scan.findings


def test_prior_matching_rule_without_stop_if_true_abstains_from_green_conflict(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Quantity"])
    for row, value in enumerate([2, 4, 6, -1, 8, 10, 12, 14], start=2):
        worksheet.append([f"R{row:03d}", value])

    green_rule = CellIsRule(
        operator="lessThan",
        formula=["0"],
        fill=PatternFill(fill_type="solid", fgColor="FF00B050"),
    )
    green_rule.priority = 2
    worksheet.conditional_formatting.add("B2:B9", green_rule)
    red_rule = CellIsRule(
        operator="lessThan",
        formula=["0"],
        fill=PatternFill(fill_type="solid", fgColor="FFFF0000"),
        stopIfTrue=False,
    )
    red_rule.priority = 1
    worksheet.conditional_formatting.add("B2:B9", red_rule)

    scan = _scan(tmp_path, workbook, "negative-prior-without-stop-if-true.xlsx")

    assert not scan.findings


def test_unknown_prior_stop_if_true_rule_abstains_from_green_conflict(tmp_path: Path) -> None:
    workbook = _quantity_book("FF00B050")
    worksheet = workbook.active
    assert worksheet is not None
    green_rule = next(
        rule for conditional in worksheet.conditional_formatting for rule in conditional.rules
    )
    green_rule.priority = 2
    expression = FormulaRule(
        formula=["B2<0"],
        fill=PatternFill(fill_type="solid", fgColor="FFFF0000"),
        stopIfTrue=True,
    )
    expression.priority = 1
    worksheet.conditional_formatting.add("B2:B9", expression)

    scan = _scan(tmp_path, workbook, "unknown-stop-if-true.xlsx")

    assert not scan.findings
