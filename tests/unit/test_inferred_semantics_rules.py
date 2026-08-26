from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook

from workbooklens.rules.inferred_semantics import (
    InferredCategoryVariantRule,
    IsolatedMissingFieldRule,
    _indexed_merged_cell,
    _merged_row_index,
)
from workbooklens.rules.registry import RuleRegistry
from workbooklens.scanner import ScanResult, scan_workbook


def _save_and_scan(
    workbook: Workbook,
    path: Path,
    *,
    config: dict[str, Any] | None = None,
) -> ScanResult:
    workbook.save(path)
    workbook.close()
    return scan_workbook(
        path,
        config=config,
        registry=RuleRegistry(
            [
                InferredCategoryVariantRule(),
                IsolatedMissingFieldRule(),
            ]
        ),
    )


def test_reports_typo_like_categories_placeholder_and_one_isolated_missing_field(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record ID", "Region", "Status", "Product", "Notes"])
    regions = ["East", "West", "North", "South", "Central"]
    statuses = ["Paid", "Pending", "Cancelled", "Refunded"]
    products = ["Model-A", "Model-B", "Service-X"]
    for index in range(1, 23):
        region = "Eest" if index == 7 else regions[(index - 1) % len(regions)]
        if index == 9:
            status = "Piad"
        elif index == 14:
            status = "UnknownStatus"
        else:
            status = statuses[(index - 1) % len(statuses)]
        product = None if index == 11 else products[(index - 1) % len(products)]
        worksheet.append([f"R{index:03d}", region, status, product, f"Reviewed note {index}"])

    scan = _save_and_scan(workbook, tmp_path / "semantic-positive.xlsx")

    variants = {
        finding.location: finding
        for finding in scan.findings
        if finding.rule_id == InferredCategoryVariantRule.rule_id
    }
    assert set(variants) == {"B8", "C10", "C15"}
    assert variants["B8"].evidence.expected == "east"
    assert variants["B8"].evidence.details["reason"] == "single_edit"
    assert variants["C10"].evidence.expected == "paid"
    assert variants["C15"].evidence.details["reason"] == "placeholder_marker"
    missing = [
        finding for finding in scan.findings if finding.rule_id == IsolatedMissingFieldRule.rule_id
    ]
    assert [finding.location for finding in missing] == ["D12"]
    assert all(
        not finding.patch_ids and not finding.safe_patch_available for finding in scan.findings
    )


def test_suppresses_free_text_high_cardinality_sparse_optional_and_normal_new_categories(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(
        [
            "Record ID",
            "Notes",
            "Segment",
            "Optional tag",
            "Department",
            "Model",
        ]
    )
    departments = ["Finance", "Sales", "Legal", "Research"]
    models = ["Model-A", "Model-B"]
    for index in range(1, 25):
        department = "Marketing" if index == 13 else departments[(index - 1) % 4]
        model = "Model-C" if index == 12 else models[(index - 1) % 2]
        optional = f"Tag-{index}" if index in {2, 7, 15, 21} else None
        worksheet.append(
            [
                f"R{index:03d}",
                f"Customer requested wording variant {index}",
                f"Segment-{index:02d}",
                optional,
                department,
                model,
            ]
        )

    scan = _save_and_scan(workbook, tmp_path / "semantic-negative.xlsx")

    assert not scan.findings


def test_suppresses_small_samples_even_when_a_value_is_one_edit_from_a_peer(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record ID", "Region", "Product"])
    for index, region in enumerate(
        ["East", "West", "East", "West", "East", "West", "Eest", "West"],
        start=1,
    ):
        worksheet.append([f"R{index:03d}", region, None if index == 4 else "Core"])

    scan = _save_and_scan(workbook, tmp_path / "semantic-small-sample.xlsx")

    assert not scan.findings


def test_excludes_merged_cells_from_isolated_missing_inference(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record ID", "Product", "Amount"])
    for index in range(1, 23):
        worksheet.append([f"R{index:03d}", f"Product-{index % 3}", index * 10])
    worksheet.merge_cells("B11:B12")

    scan = _save_and_scan(workbook, tmp_path / "semantic-merged.xlsx")

    assert not any(finding.rule_id == IsolatedMissingFieldRule.rule_id for finding in scan.findings)


def test_explicit_allowed_values_take_precedence_over_rare_variant_inference(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record ID", "Region", "Amount"])
    regions = ["East", "North", "South"]
    for index in range(1, 25):
        region = "Eest" if index == 7 else regions[(index - 1) % len(regions)]
        worksheet.append([f"R{index:03d}", region, index * 10])

    scan = _save_and_scan(
        workbook,
        tmp_path / "semantic-explicit-enum.xlsx",
        config={
            "version": 2,
            "profile": {
                "sheets": [
                    {
                        "sheet": "Sheet",
                        "range": "A1:C25",
                        "columns": [
                            {
                                "header": "Region",
                                "allowed_values": ["East", "Eest", "North", "South"],
                            }
                        ],
                    }
                ]
            },
        },
    )

    assert not any(
        finding.rule_id == InferredCategoryVariantRule.rule_id for finding in scan.findings
    )


def test_explicit_optional_column_takes_precedence_over_missing_field_inference(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record ID", "Product", "Amount"])
    for index in range(1, 25):
        worksheet.append(
            [
                f"R{index:03d}",
                None if index == 12 else f"Product-{index % 3}",
                index * 10,
            ]
        )

    scan = _save_and_scan(
        workbook,
        tmp_path / "semantic-explicit-optional.xlsx",
        config={
            "version": 2,
            "profile": {
                "sheets": [
                    {
                        "sheet": "Sheet",
                        "range": "A1:C25",
                        "columns": [{"header": "Product", "required": False}],
                    }
                ]
            },
        },
    )

    assert not any(finding.rule_id == IsolatedMissingFieldRule.rule_id for finding in scan.findings)


def test_legitimate_east_west_categories_are_not_treated_as_spelling_variants(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Record ID", "Region", "Amount"])
    regions = ["East"] * 8 + ["North"] * 7 + ["South"] * 8 + ["West"]
    for index, region in enumerate(regions, start=1):
        worksheet.append([f"R{index:03d}", region, index * 10])

    scan = _save_and_scan(workbook, tmp_path / "semantic-east-west.xlsx")

    assert not any(
        finding.rule_id == InferredCategoryVariantRule.rule_id for finding in scan.findings
    )


def test_merged_row_index_handles_many_ranges_before_repeated_cell_queries() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for row in range(2, 1002, 2):
        worksheet.merge_cells(start_row=row, start_column=5, end_row=row + 1, end_column=6)

    merged_rows = _merged_row_index(worksheet, tuple(range(2, 1002)))

    assert len(merged_rows) == 1000
    for row in range(2, 1002):
        assert _indexed_merged_cell(merged_rows, row, 5)
        assert _indexed_merged_cell(merged_rows, row, 6)
        assert not _indexed_merged_cell(merged_rows, row, 4)
    workbook.close()
