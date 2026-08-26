from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from workbooklens.formulas.ir import (
    WorksheetContentIndex,
    parse_formula_ir,
    worksheet_content_index,
)
from workbooklens.i18n import localize_scan_result
from workbooklens.models import Finding, PatchKind, PatchRisk, Severity
from workbooklens.repair.planning import build_patch_plan, resolve_patch_selection
from workbooklens.rules import RuleRegistry
from workbooklens.rules.builtin import FormulaPatternOutlierRule, TextFormulaInDataRegionRule
from workbooklens.rules.formula_semantics import (
    AggregateRangeCoverageRule,
    CrossSheetReferenceValidityRule,
    DegenerateFormulaRule,
    FormulaFormatRoleMismatchRule,
    FormulaInNotesColumnRule,
    _formula_cells,
)
from workbooklens.scanner import ScanResult, scan_workbook


def _scan(
    tmp_path: Path,
    workbook: Workbook,
    rule: object,
    *,
    name: str = "formula-semantics.xlsx",
) -> ScanResult:
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    registry = rule if isinstance(rule, RuleRegistry) else RuleRegistry([rule])
    return scan_workbook(path, registry=registry)  # type: ignore[arg-type, list-item]


def _amount_table(workbook: Workbook) -> None:
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Sales"
    worksheet.append(["Item", "Amount"])
    for index in range(1, 8):
        worksheet.append([f"Item {index}", index * 10])


def _cross_sheet_book(*, blank_row: int | None = None) -> Workbook:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source Data"
    source.append(["ID", "Amount", "Category"])
    for row in range(2, 9):
        source.append([row - 1, None if row == blank_row else row * 10, "A"])
    dashboard = workbook.create_sheet("Dashboard")
    dashboard["A1"] = "Result"
    return workbook


def test_formula_ir_resolves_quoted_cross_sheet_aggregate() -> None:
    ir = parse_formula_ir(
        "=SUM('Sales Data'!$B$2:$B$8)",
        origin_sheet="Dashboard",
        origin_coordinate="B5",
    )

    assert len(ir.references) == 1
    reference = ir.references[0]
    assert reference.sheet == "Sales Data"
    assert reference.is_cross_sheet
    assert reference.local_range == "B2:B8"
    assert [(item.function, item.reference.raw) for item in ir.aggregate_references] == [
        ("SUM", "'Sales Data'!$B$2:$B$8")
    ]


@pytest.mark.parametrize(
    ("formula", "kind"),
    [
        ("=A1/A1", "self_division"),
        ("=SUM(A1:A3)-SUM(A1:A3)", "self_subtraction"),
        ("=A1*0", "multiply_by_zero"),
        ("=0*A1", "multiply_by_zero"),
    ],
)
def test_formula_ir_identifies_only_top_level_degeneracy(formula: str, kind: str) -> None:
    ir = parse_formula_ir(formula, origin_sheet="Sheet", origin_coordinate="B1")

    assert ir.degeneracy is not None
    assert ir.degeneracy.kind == kind


@pytest.mark.parametrize("formula", ["=IF(A1=0,0,A1/A1)", "=A1/A1+1", '="A1/A1"'])
def test_formula_ir_does_not_flag_nested_or_display_text(formula: str) -> None:
    assert parse_formula_ir(formula, "Sheet", "B1").degeneracy is None


def test_aggregate_range_reports_omitted_tail_rows(tmp_path: Path) -> None:
    workbook = Workbook()
    _amount_table(workbook)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A9"] = "Total"
    worksheet["B9"] = "=SUM(B2:B6)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert [(finding.location, finding.rule_id) for finding in scan.findings] == [
        ("B9", "WL042_AGGREGATE_RANGE_COVERAGE")
    ]
    assert scan.findings[0].evidence.details["excluded_cells"] == ["B7", "B8"]
    assert scan.findings[0].evidence.expected == {"reviewed_reference": "B2:B8"}
    assert len(scan.patches) == 1
    patch = scan.patches[0]
    assert patch.kind == PatchKind.SET_FORMULA
    assert patch.after == "=SUM(B2:B8)"
    assert patch.risk == PatchRisk.FORMULA_DERIVED
    assert not patch.safe
    assert patch.atomic_group is not None
    assert float(patch.confidence) == 0.99
    assert patch.derivation.strategy == "unique_complete_aggregate_range"
    assert patch.derivation.candidate_count == 1
    assert len(patch.derivation.sources) == 2
    assert patch.derivation.requires_recalculation
    assert scan.findings[0].patch_ids == [patch.id]
    localized = localize_scan_result(scan, "zh-CN", strict=True)
    assert localized.patches[0].description == "将聚合范围扩展到唯一确定的完整推断表格主体。"


def test_aggregate_range_can_depend_on_unique_safe_numeric_normalization(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    _amount_table(workbook)
    worksheet = workbook.active
    assert worksheet is not None
    for row in range(2, 8):
        worksheet.cell(row, 2).number_format = "$#,##0"
    worksheet["B8"] = "800"
    worksheet["B8"].number_format = "$#,##0"
    worksheet["A9"] = "Total"
    worksheet["B9"] = "=SUM(B2:B6)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    normalization = next(patch for patch in scan.patches if patch.kind == PatchKind.SET_NUMERIC)
    formula = next(patch for patch in scan.patches if patch.kind == PatchKind.SET_FORMULA)
    assert normalization.cell == "B8"
    assert normalization.safe_only_eligible
    assert formula.after == "=SUM(B2:B8)"
    assert formula.prerequisite_patch_ids == [normalization.id]
    assert any(
        source.startswith("safe_normalization_prerequisite:")
        for source in formula.derivation.sources
    )

    plan = build_patch_plan(scan)
    assert resolve_patch_selection(plan, safe_only=True) == {normalization.id}
    assert resolve_patch_selection(plan, auto_repair=True) == {
        normalization.id,
        formula.id,
    }


def _formula_amount_book() -> Workbook:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source Data"
    source.append(["ID", "Hours", "Rate", "Amount"])
    for row in range(2, 30):
        source.append([row - 1, row, row + 0.5, f"=B{row}*C{row}"])
    dashboard = workbook.create_sheet("Dashboard")
    dashboard["A2"] = "Total amount"
    dashboard["B2"] = "=SUM('Source Data'!D2:D24)"
    return workbook


def test_sum_formula_column_uses_verified_formula_repairs_as_prerequisites(
    tmp_path: Path,
) -> None:
    workbook = _formula_amount_book()
    source = workbook["Source Data"]
    source["D5"] = "=B5+C5"
    source["D6"] = "=B7*C6"
    source["D7"] = "'=B7*C7"

    scan = _scan(
        tmp_path,
        workbook,
        RuleRegistry(
            [
                FormulaPatternOutlierRule(),
                TextFormulaInDataRegionRule(),
                AggregateRangeCoverageRule(),
            ]
        ),
        name="formula-column-prerequisites.xlsx",
    )

    source_repairs = {
        patch.cell: patch
        for patch in scan.patches
        if patch.sheet == "Source Data"
        and patch.kind == PatchKind.SET_FORMULA
        and patch.cell in {"D5", "D6", "D7"}
    }
    assert set(source_repairs) == {"D5", "D6", "D7"}
    aggregate = next(
        patch for patch in scan.patches if patch.sheet == "Dashboard" and patch.cell == "B2"
    )
    assert aggregate.after == "=SUM('Source Data'!D2:D29)"
    assert set(aggregate.prerequisite_patch_ids) == {patch.id for patch in source_repairs.values()}
    assert any(
        source.startswith("formula_repair_prerequisite:") for source in aggregate.derivation.sources
    )
    assert not any(
        source.startswith("safe_normalization_prerequisite:")
        for source in aggregate.derivation.sources
    )

    plan = build_patch_plan(scan)
    selected = resolve_patch_selection(plan, auto_repair=True)
    assert aggregate.id in selected
    assert set(aggregate.prerequisite_patch_ids) <= selected


@pytest.mark.parametrize("blocked", ["mixed_tail_signature", "leading_omission"])
def test_sum_formula_column_requires_unique_signature_and_tail_only_extension(
    tmp_path: Path,
    blocked: str,
) -> None:
    workbook = _formula_amount_book()
    source = workbook["Source Data"]
    dashboard = workbook["Dashboard"]
    if blocked == "mixed_tail_signature":
        source["D28"] = "=B28+C28"
    else:
        dashboard["B2"] = "=SUM('Source Data'!D4:D29)"

    scan = _scan(
        tmp_path,
        workbook,
        AggregateRangeCoverageRule(),
        name=f"formula-column-{blocked}.xlsx",
    )

    assert any(
        finding.rule_id == "WL042_AGGREGATE_RANGE_COVERAGE"
        and finding.sheet == "Dashboard"
        and finding.location == "B2"
        for finding in scan.findings
    )
    assert not any(patch.sheet == "Dashboard" and patch.cell == "B2" for patch in scan.patches)


def test_aggregate_range_accepts_complete_table_body(tmp_path: Path) -> None:
    workbook = Workbook()
    _amount_table(workbook)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A9"] = "Total"
    worksheet["B9"] = "=SUM(B2:B8)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert scan.findings == []


def test_aggregate_range_does_not_guess_an_intentional_window(tmp_path: Path) -> None:
    workbook = Workbook()
    _amount_table(workbook)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A9"] = "Last five periods"
    worksheet["B9"] = "=SUM(B4:B8)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert scan.findings == []


def test_aggregate_range_catches_one_omitted_leading_row(tmp_path: Path) -> None:
    workbook = Workbook()
    _amount_table(workbook)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A9"] = "Grand total"
    worksheet["B9"] = "=SUM(B3:B8)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert len(scan.findings) == 1
    assert scan.findings[0].evidence.details["excluded_cells"] == ["B2"]


def test_aggregate_range_patch_preserves_absolute_anchors(tmp_path: Path) -> None:
    workbook = Workbook()
    _amount_table(workbook)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A9"] = "Total"
    worksheet["B9"] = "=SUM($B$2:$B$6)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert len(scan.patches) == 1
    assert scan.patches[0].after == "=SUM($B$2:$B$8)"


def test_aggregate_range_mismatched_function_label_remains_review_only(tmp_path: Path) -> None:
    workbook = Workbook()
    _amount_table(workbook)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A9"] = "Total"
    worksheet["B9"] = "=AVERAGE(B2:B6)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert len(scan.findings) == 1
    assert scan.patches == []


def test_aggregate_range_candidate_that_creates_cycle_is_withheld(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Cycle Data"
    worksheet.append(["Item", "Amount", "Input"])
    worksheet["C2"] = 10
    for row in range(2, 8):
        worksheet.cell(row, 1, f"Item {row}")
        worksheet.cell(row, 2, "=$C$2")
    worksheet["A8"] = "Item 8"
    worksheet["B8"] = "=B9"
    worksheet["A9"] = "Total"
    worksheet["B9"] = "=SUM(B2:B6)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert len(scan.findings) == 1
    assert scan.findings[0].location == "B9"
    assert scan.patches == []


@pytest.mark.parametrize("blocked", ["protected", "merged", "hidden_row", "hidden_column"])
def test_aggregate_range_high_risk_target_remains_review_only(
    tmp_path: Path,
    blocked: str,
) -> None:
    workbook = Workbook()
    _amount_table(workbook)
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A9"] = "Total"
    worksheet["B9"] = "=SUM(B2:B6)"
    if blocked == "protected":
        worksheet.protection.sheet = True
    elif blocked == "merged":
        worksheet.merge_cells("B9:C9")
    elif blocked == "hidden_row":
        worksheet.row_dimensions[9].hidden = True
    else:
        worksheet.column_dimensions["B"].hidden = True

    scan = _scan(
        tmp_path,
        workbook,
        AggregateRangeCoverageRule(),
        name=f"aggregate-{blocked}.xlsx",
    )

    assert len(scan.findings) == 1
    assert scan.patches == []


def test_cross_sheet_metric_label_reports_truncated_aggregate(tmp_path: Path) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["A2"] = "Gross Sales"
    dashboard["B2"] = "=SUM('Source Data'!B2:B5)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert [(finding.location, finding.rule_id) for finding in scan.findings] == [
        ("B2", "WL042_AGGREGATE_RANGE_COVERAGE")
    ]
    assert scan.findings[0].evidence.details["excluded_cells"] == ["B6", "B7", "B8"]


@pytest.mark.parametrize(
    ("label", "function"),
    [
        ("项目总预算", "SUM"),
        ("项目总支出", "SUM"),
        ("总人工成本", "SUM"),
        ("平均人工成本", "AVERAGE"),
    ],
)
def test_cross_sheet_chinese_amount_metric_reports_truncated_aggregate(
    tmp_path: Path,
    label: str,
    function: str,
) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["A2"] = label
    dashboard["B2"] = f"={function}('Source Data'!B2:B5)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert [(finding.location, finding.rule_id) for finding in scan.findings] == [
        ("B2", "WL042_AGGREGATE_RANGE_COVERAGE")
    ]
    assert scan.findings[0].evidence.details["excluded_cells"] == ["B6", "B7", "B8"]
    assert len(scan.patches) == 1
    patch = scan.patches[0]
    assert patch.risk == PatchRisk.FORMULA_DERIVED
    assert patch.after == f"={function}('Source Data'!B2:B8)"


def _conditional_aggregate_book() -> Workbook:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "报销明细"
    source.append(["ID", "含税金额", "审批状态", "费用类别"])
    for row in range(2, 9):
        source.append(
            [
                row - 1,
                row * 100,
                "已批准" if row % 2 else "待审批",
                "差旅",
            ]
        )
    dashboard = workbook.create_sheet("看板")
    dashboard["A2"] = "已批准报销额"
    return workbook


def test_sumif_reviews_synchronized_cross_sheet_boundaries(tmp_path: Path) -> None:
    workbook = _conditional_aggregate_book()
    workbook["看板"]["B2"] = '=SUMIF(报销明细!C2:C5,"已批准",报销明细!B2:B5)'

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert len(scan.findings) == 1
    finding = scan.findings[0]
    assert finding.location == "B2"
    assert finding.evidence.details["function"] == "SUMIF"
    assert finding.evidence.details["excluded_rows"] == [6, 7, 8]
    assert finding.evidence.expected == {
        "reviewed_references": {
            "criteria_range": "'报销明细'!C2:C8",
            "sum_range": "'报销明细'!B2:B8",
        }
    }
    assert len(scan.patches) == 1
    patch = scan.patches[0]
    assert patch.after == ("=SUMIF('报销明细'!C2:C8,\"已批准\",'报销明细'!B2:B8)")
    assert patch.risk == PatchRisk.FORMULA_DERIVED
    assert patch.derivation.strategy == "unique_synchronized_conditional_aggregate_ranges"
    assert "criteria_and_sum_ranges_share_identical_row_bounds" in patch.derivation.invariants


def test_sumifs_reviews_all_synchronized_cross_sheet_boundaries(tmp_path: Path) -> None:
    workbook = _conditional_aggregate_book()
    workbook["看板"]["B2"] = '=SUMIFS(报销明细!B2:B5,报销明细!C2:C5,"已批准",报销明细!D2:D5,"差旅")'

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert len(scan.findings) == 1
    finding = scan.findings[0]
    assert finding.evidence.details["function"] == "SUMIFS"
    assert finding.evidence.details["excluded_rows"] == [6, 7, 8]
    assert [item["role"] for item in finding.evidence.details["ranges"]] == [
        "sum_range",
        "criteria_range_1",
        "criteria_range_2",
    ]
    assert len(scan.patches) == 1
    patch = scan.patches[0]
    assert patch.after == (
        "=SUMIFS('报销明细'!B2:B8,'报销明细'!C2:C8,\"已批准\",'报销明细'!D2:D8,\"差旅\")"
    )
    assert patch.risk == PatchRisk.FORMULA_DERIVED


def test_sumif_requires_synchronized_observed_boundaries(tmp_path: Path) -> None:
    workbook = _conditional_aggregate_book()
    workbook["看板"]["B2"] = '=SUMIF(报销明细!C2:C5,"已批准",报销明细!B2:B8)'

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert scan.findings == []


def _kpi_block_book(*, proven_count: int = 3, all_proven: bool = False) -> Workbook:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source Data"
    source.append(["ID", "Sales", "Gross", "Salary", "Inventory"])
    for row in range(2, 9):
        source.append([row - 1, row * 10, row * 12, row * 100, row * 4])
    dashboard = workbook.create_sheet("Dashboard")
    dashboard["A4"] = "Metric"
    dashboard["B4"] = "Value"
    labels = ["Total Sales", "Gross Sales", "Avg Salary", "Inventory Value"]
    source_columns = ["B", "C", "D", "E"]
    functions = ["SUM", "SUM", "AVERAGE", "SUM"]
    proven_rows = {5, 6, 8} if proven_count == 3 else {5, 6}
    if all_proven:
        proven_rows = {5, 6, 7, 8}
    for row, label, source_column, function in zip(
        range(5, 9), labels, source_columns, functions, strict=True
    ):
        dashboard.cell(row, 1, label)
        end_row = 5 if row in proven_rows else 8
        dashboard.cell(
            row, 2, f"={function}('Source Data'!{source_column}2:{source_column}{end_row})"
        )
    return workbook


def _aggregate_block_advisories(scan: ScanResult) -> list[Finding]:
    return [
        finding
        for finding in scan.findings
        if finding.evidence.details.get("proof_level") == "advisory"
        and "unproven_formulas" in finding.evidence.details
    ]


def test_aggregate_range_summarizes_only_independently_proven_kpi_omissions(
    tmp_path: Path,
) -> None:
    scan = _scan(tmp_path, _kpi_block_book(), AggregateRangeCoverageRule())

    advisories = _aggregate_block_advisories(scan)
    assert len(advisories) == 1
    advisory = advisories[0]
    assert advisory.location == "B5:B8"
    assert advisory.severity == Severity.INFO
    assert advisory.evidence.observed == {
        "proven_cells": ["B5", "B6", "B8"],
        "unproven_cells": ["B7"],
    }
    assert advisory.patch_ids == []


def test_aggregate_range_does_not_add_block_advisory_without_unproven_formula(
    tmp_path: Path,
) -> None:
    scan = _scan(
        tmp_path,
        _kpi_block_book(all_proven=True),
        AggregateRangeCoverageRule(),
    )

    assert not _aggregate_block_advisories(scan)


def test_aggregate_range_requires_three_independent_block_omissions(tmp_path: Path) -> None:
    scan = _scan(
        tmp_path,
        _kpi_block_book(proven_count=2),
        AggregateRangeCoverageRule(),
    )

    assert not _aggregate_block_advisories(scan)


def test_aggregate_range_counts_one_formula_once_when_it_has_two_truncated_inputs(
    tmp_path: Path,
) -> None:
    workbook = _kpi_block_book(proven_count=2)
    dashboard = workbook["Dashboard"]
    dashboard["B5"] = "=SUM('Source Data'!B2:B5)+SUM('Source Data'!C2:C5)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert not _aggregate_block_advisories(scan)


def test_aggregate_range_does_not_add_block_advisory_for_hidden_dashboard(
    tmp_path: Path,
) -> None:
    workbook = _kpi_block_book()
    workbook["Dashboard"].sheet_state = "hidden"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert not _aggregate_block_advisories(scan)


@pytest.mark.parametrize("label", ["最近项目总预算", "项目预算执行率"])
def test_scoped_or_rate_chinese_metric_does_not_imply_full_range(
    tmp_path: Path,
    label: str,
) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["A2"] = label
    dashboard["B2"] = "=SUM('Source Data'!B2:B5)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert scan.findings == []


@pytest.mark.parametrize("label", ["Q1 Sales", "Rolling Sales", "Last period revenue"])
def test_cross_sheet_window_label_does_not_imply_full_range(tmp_path: Path, label: str) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["A2"] = label
    dashboard["B2"] = "=SUM('Source Data'!B2:B5)"

    scan = _scan(tmp_path, workbook, AggregateRangeCoverageRule())

    assert scan.findings == []


def test_cross_sheet_reference_treats_blank_future_input_as_advisory(tmp_path: Path) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["B1"] = "='Source Data'!F999"

    scan = _scan(tmp_path, workbook, CrossSheetReferenceValidityRule())

    assert len(scan.findings) == 1
    assert scan.findings[0].location == "B1"
    assert scan.findings[0].evidence.details["reason"] == "blank_outside_content"
    assert scan.findings[0].severity == Severity.INFO
    assert float(scan.findings[0].confidence) == 0.7
    assert scan.findings[0].evidence.details["proof_level"] == "advisory"
    assert "planned future input" in scan.findings[0].explanation
    assert scan.patches == []


def test_cross_sheet_missing_worksheet_remains_proven_error(tmp_path: Path) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["B1"] = "='Missing Input'!A1"

    scan = _scan(tmp_path, workbook, CrossSheetReferenceValidityRule())

    assert len(scan.findings) == 1
    finding = scan.findings[0]
    assert finding.evidence.details["reason"] == "missing_worksheet"
    assert finding.severity == Severity.ERROR
    assert float(finding.confidence) == 1.0
    assert finding.evidence.details["proof_level"] == "strong_structural"


def test_sparse_content_index_queries_are_sorted_bounded_and_cached() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = "Header"
    worksheet["C1"] = "Other"
    for row in range(2, 42):
        worksheet.cell(row, 2, row)
    worksheet["D5"] = "Flag"

    index = WorksheetContentIndex.from_worksheet(worksheet)

    assert index.content_bounds == (1, 41, 1, 4)
    assert [cell.coordinate for cell in index.iter_nonblank_row(1)] == ["A1", "C1"]
    assert [cell.coordinate for cell in index.iter_nonblank_column(2, min_row=3, max_row=5)] == [
        "B3",
        "B4",
        "B5",
    ]
    assert [
        cell.coordinate
        for cell in index.iter_nonblank_rectangle(
            min_row=1,
            max_row=41,
            min_column=2,
            max_column=2,
            limit=32,
        )
    ] == [f"B{row}" for row in range(2, 34)]

    context = SimpleNamespace(analysis_cache={})
    assert worksheet_content_index(context, worksheet) is worksheet_content_index(
        context, worksheet
    )


def test_formula_cells_are_sorted_cached_and_do_not_build_content_index(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["C3"] = "=A1"
    worksheet["A1"] = "=1"
    worksheet["B2"] = "=2"
    calls: list[str] = []

    def unexpected_index(
        cls: type[WorksheetContentIndex], sheet: Worksheet
    ) -> WorksheetContentIndex:
        calls.append(sheet.title)
        raise AssertionError("formula collection must not build a content index")

    monkeypatch.setattr(WorksheetContentIndex, "from_worksheet", classmethod(unexpected_index))
    context = SimpleNamespace(analysis_cache={})

    first = _formula_cells(context, worksheet)
    second = _formula_cells(context, worksheet)

    assert first is second
    assert [cell.coordinate for cell in first] == ["A1", "B2", "C3"]
    assert calls == []
    assert context.analysis_cache["formula-semantics-formula-cells-v1"][worksheet.title] is first


@pytest.mark.parametrize("limit", [0, -1])
def test_sparse_iterators_return_no_cells_for_nonpositive_limits(limit: int) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = "first"
    worksheet["B1"] = "row peer"
    worksheet["A2"] = "column peer"
    index = WorksheetContentIndex.from_worksheet(worksheet)

    assert list(index.iter_nonblank_row(1, limit=limit)) == []
    assert list(index.iter_nonblank_column(1, limit=limit)) == []
    assert (
        list(
            index.iter_nonblank_rectangle(
                min_row=1,
                max_row=2,
                min_column=1,
                max_column=2,
                limit=limit,
            )
        )
        == []
    )


def test_sparse_axis_indexes_use_flat_tuples_and_support_tail_queries() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    cells = [worksheet.cell(row, 1, row) for row in range(1, 5001)]

    index = WorksheetContentIndex.from_cells(worksheet, cells)

    assert isinstance(index._rows._cells, tuple)
    assert isinstance(index._columns._cells, tuple)
    assert len(index._rows._cells) == 5000
    assert len(index._columns._cells) == 5000
    assert not hasattr(index, "_row_columns")
    assert not hasattr(index, "_column_rows")
    assert [cell.coordinate for cell in index.iter_nonblank_row(5000)] == ["A5000"]
    assert [
        cell.coordinate for cell in index.iter_nonblank_column(1, min_row=4998, max_row=5000)
    ] == ["A4998", "A4999", "A5000"]
    assert [
        cell.coordinate
        for cell in index.iter_nonblank_rectangle(
            min_row=4998,
            max_row=5000,
            min_column=1,
            max_column=1,
        )
    ] == ["A4998", "A4999", "A5000"]


def test_sparse_rectangle_is_row_major_and_applies_one_global_limit() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for coordinate in ("D1", "B1", "C2", "A2", "B3"):
        worksheet[coordinate] = coordinate
    index = WorksheetContentIndex.from_worksheet(worksheet)

    assert [
        cell.coordinate
        for cell in index.iter_nonblank_rectangle(
            min_row=1,
            max_row=3,
            min_column=1,
            max_column=4,
            limit=4,
        )
    ] == ["B1", "D1", "A2", "C2"]


def test_sparse_index_materializes_one_shot_cell_iterable_once() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    cells = [
        worksheet.cell(1, 3, "C1"),
        worksheet.cell(2, 1, "A2"),
        worksheet.cell(3, 2, "B3"),
    ]
    consumed: list[str] = []

    def cell_stream():
        for cell in cells:
            consumed.append(cell.coordinate)
            yield cell

    index = WorksheetContentIndex.from_cells(worksheet, cell_stream())

    assert consumed == ["C1", "A2", "B3"]
    assert [cell.coordinate for cell in index.iter_nonblank_row(2)] == ["A2"]
    assert [cell.coordinate for cell in index.iter_nonblank_column(2)] == ["B3"]
    assert [
        cell.coordinate
        for cell in index.iter_nonblank_rectangle(
            min_row=1,
            max_row=3,
            min_column=1,
            max_column=3,
        )
    ] == ["C1", "A2", "B3"]


def test_cross_sheet_rule_builds_each_sparse_index_once_per_scan(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["B1"] = "='Source Data'!B2"
    dashboard["B2"] = "='Source Data'!B3"
    calls: list[str] = []
    original = WorksheetContentIndex.from_worksheet

    def counted(cls: type[WorksheetContentIndex], worksheet: Worksheet) -> WorksheetContentIndex:
        calls.append(worksheet.title)
        return original(worksheet)

    monkeypatch.setattr(WorksheetContentIndex, "from_worksheet", classmethod(counted))

    scan = _scan(tmp_path, workbook, CrossSheetReferenceValidityRule())

    assert scan.findings == []
    assert calls == ["Source Data", "Dashboard"]


def test_cross_sheet_reference_reports_blank_inside_dense_table(tmp_path: Path) -> None:
    workbook = _cross_sheet_book(blank_row=4)
    dashboard = workbook["Dashboard"]
    dashboard["B1"] = "='Source Data'!B4"

    scan = _scan(tmp_path, workbook, CrossSheetReferenceValidityRule())

    assert len(scan.findings) == 1
    assert scan.findings[0].evidence.details["reason"] == "blank_inside_dense_table"
    assert scan.findings[0].severity == Severity.WARNING


def test_cross_sheet_reference_accepts_populated_target(tmp_path: Path) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["B1"] = "='Source Data'!B4"

    scan = _scan(tmp_path, workbook, CrossSheetReferenceValidityRule())

    assert scan.findings == []


def test_cross_sheet_sum_reports_incompatible_source_column_roles(tmp_path: Path) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source Data"
    source.append(["ID", "Qty", "Unit Price"])
    for row in range(2, 7):
        source.append([row - 1, row, row * 10])
        source.cell(row, 3).number_format = "$#,##0.00"
    dashboard = workbook.create_sheet("Dashboard")
    dashboard["A1"] = "Check"
    dashboard["B1"] = "=SUM('Source Data'!B2:C2)"

    scan = _scan(tmp_path, workbook, CrossSheetReferenceValidityRule())

    assert len(scan.findings) == 1
    finding = scan.findings[0]
    assert finding.location == "B1"
    assert finding.evidence.details["reason"] == "heterogeneous_aggregate_column_roles"
    assert finding.evidence.details["source_roles"] == ["amount", "quantity"]
    assert finding.evidence.details["source_columns"] == [
        {"column": "B", "header": "Qty", "role": "quantity"},
        {"column": "C", "header": "Unit Price", "role": "amount"},
    ]
    assert finding.severity == Severity.WARNING
    assert scan.patches == []


@pytest.mark.parametrize("function", ["COUNT", "SUM"])
def test_cross_sheet_aggregate_accepts_compatible_or_non_additive_columns(
    tmp_path: Path,
    function: str,
) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source Data"
    if function == "SUM":
        source.append(["ID", "Q1 Sales", "Q2 Sales"])
    else:
        source.append(["ID", "Qty", "Unit Price"])
    for row in range(2, 7):
        source.append([row - 1, row * 10, row * 20])
        source.cell(row, 2).number_format = "$#,##0.00"
        source.cell(row, 3).number_format = "$#,##0.00"
    dashboard = workbook.create_sheet("Dashboard")
    dashboard["A1"] = "Check"
    dashboard["B1"] = f"={function}('Source Data'!B2:C2)"

    scan = _scan(
        tmp_path,
        workbook,
        CrossSheetReferenceValidityRule(),
        name=f"compatible-{function.casefold()}.xlsx",
    )

    assert scan.findings == []


def test_cross_sheet_reference_does_not_overstate_optional_edge_blank(tmp_path: Path) -> None:
    workbook = _cross_sheet_book(blank_row=8)
    dashboard = workbook["Dashboard"]
    dashboard["B1"] = "='Source Data'!B8"

    scan = _scan(tmp_path, workbook, CrossSheetReferenceValidityRule())

    assert scan.findings == []


def test_cross_sheet_reference_reports_obviously_truncated_nonaggregate_range(
    tmp_path: Path,
) -> None:
    workbook = _cross_sheet_book()
    dashboard = workbook["Dashboard"]
    dashboard["A1"] = "Total row count"
    dashboard["B1"] = "=ROWS('Source Data'!B2:B5)"

    scan = _scan(tmp_path, workbook, CrossSheetReferenceValidityRule())

    assert len(scan.findings) == 1
    finding = scan.findings[0]
    assert finding.evidence.details["reason"] == "truncated_cross_sheet_range"
    assert finding.severity == Severity.WARNING
    assert finding.evidence.details["excluded_cells"] == ["B6", "B7", "B8"]


def test_notes_column_reports_isolated_formula(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Notes"])
    for row in range(2, 9):
        worksheet.append([row - 1, f"Reviewed note {row}"])
    worksheet["B5"] = "=SUM(A2:A4)"

    scan = _scan(tmp_path, workbook, FormulaInNotesColumnRule())

    assert len(scan.findings) == 1
    assert scan.findings[0].location == "B5"
    assert scan.findings[0].evidence.details["header"] == "Notes"
    assert scan.patches == []


def test_notes_rule_accepts_formula_in_amount_column(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Amount"])
    for row in range(2, 9):
        worksheet.append([row - 1, f"=A{row}*10"])

    scan = _scan(tmp_path, workbook, FormulaInNotesColumnRule())

    assert scan.findings == []


def test_notes_rule_does_not_match_noted_date_header(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["ID", "Noted Date"])
    for row in range(2, 9):
        worksheet.append([row - 1, f"2026-08-{row:02d}"])
    worksheet["B5"] = "=TODAY()"

    scan = _scan(tmp_path, workbook, FormulaInNotesColumnRule())

    assert scan.findings == []


@pytest.mark.parametrize(
    ("formula", "kind"),
    [
        ("=A1/A1", "self_division"),
        ("=SUM(A1:A3)-SUM(A1:A3)", "self_subtraction"),
        ("=A1*0", "multiply_by_zero"),
    ],
)
def test_degenerate_formula_rule_reports_exact_top_level_forms(
    tmp_path: Path,
    formula: str,
    kind: str,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = 10
    worksheet["A2"] = 5
    worksheet["B1"] = formula

    scan = _scan(tmp_path, workbook, DegenerateFormulaRule(), name=f"{kind}.xlsx")

    assert len(scan.findings) == 1
    assert scan.findings[0].evidence.details["kind"] == kind
    assert scan.findings[0].patch_ids == []
    assert scan.patches == []


@pytest.mark.parametrize("formula", ["=A1/A2", "=A1-A2", "=A1*1"])
def test_degenerate_formula_rule_accepts_non_degenerate_forms(
    tmp_path: Path,
    formula: str,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = 10
    worksheet["A2"] = 5
    worksheet["B1"] = formula

    scan = _scan(tmp_path, workbook, DegenerateFormulaRule())

    assert scan.findings == []


def test_degenerate_formula_rule_ignores_guarded_inner_identity(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = 10
    worksheet["B1"] = "=IF(A1=0,0,A1/A1)"

    scan = _scan(tmp_path, workbook, DegenerateFormulaRule())

    assert scan.findings == []


def _countif_denominator_book(
    *,
    matching_value: bool = False,
    formula_input: bool = False,
) -> Workbook:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source"
    source.append(["Amount"])
    for row in range(2, 9):
        value = 200_000_000 if matching_value and row == 8 else row * 100
        source.cell(row, 1, value)
    if formula_input:
        source["A5"] = "=1+1"
    dashboard = workbook.create_sheet("Dashboard")
    dashboard["B2"] = '=SUM(Source!A2:A8)/COUNTIF(Source!A2:A8,">100000000")'
    return workbook


def test_countif_denominator_reports_proven_zero_from_literals(tmp_path: Path) -> None:
    scan = _scan(
        tmp_path,
        _countif_denominator_book(),
        DegenerateFormulaRule(),
        name="proven-zero-countif-denominator.xlsx",
    )

    assert len(scan.findings) == 1
    finding = scan.findings[0]
    assert finding.location == "B2"
    assert finding.severity == Severity.ERROR
    assert finding.evidence.details["kind"] == "provable_zero_countif_denominator"
    assert finding.evidence.details["literal_value_count"] == 7
    assert float(finding.confidence) == 1.0
    assert scan.patches == []
    localized = localize_scan_result(scan, "zh-CN", strict=True)
    assert localized.findings[0].title == "公式退化为代数恒等式或常量"


@pytest.mark.parametrize("case", ["matching_value", "formula_input"])
def test_countif_denominator_requires_zero_proof_and_literal_inputs(
    tmp_path: Path,
    case: str,
) -> None:
    workbook = _countif_denominator_book(
        matching_value=case == "matching_value",
        formula_input=case == "formula_input",
    )

    scan = _scan(
        tmp_path,
        workbook,
        DegenerateFormulaRule(),
        name=f"countif-denominator-{case}.xlsx",
    )

    assert scan.findings == []


def test_format_role_rule_reports_rate_currency_and_amount_percentage(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = "Conversion Rate"
    worksheet["B1"] = "=C1/D1"
    worksheet["B1"].number_format = "$#,##0.00"
    worksheet["C1"] = 50
    worksheet["D1"] = 100
    worksheet["A2"] = "Revenue"
    worksheet["B2"] = "=C2*D2"
    worksheet["B2"].number_format = "0.0%"
    worksheet["C2"] = 5
    worksheet["D2"] = 100

    scan = _scan(tmp_path, workbook, FormulaFormatRoleMismatchRule())

    assert [finding.location for finding in scan.findings] == ["B1", "B2"]
    assert all(finding.patch_ids == [] for finding in scan.findings)
    assert scan.patches == []


def test_format_role_rule_accepts_compatible_formats(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = "Conversion Rate"
    worksheet["B1"] = "=C1/D1"
    worksheet["B1"].number_format = "0.0%"
    worksheet["C1"] = 50
    worksheet["D1"] = 100
    worksheet["A2"] = "Revenue"
    worksheet["B2"] = "=C2*D2"
    worksheet["B2"].number_format = "$#,##0.00"
    worksheet["C2"] = 5
    worksheet["D2"] = 100

    scan = _scan(tmp_path, workbook, FormulaFormatRoleMismatchRule())

    assert scan.findings == []


@pytest.mark.parametrize("label", ["Hourly Rate", "Exchange Rate", "汇率"])
def test_format_role_rule_skips_ambiguous_currency_rates(tmp_path: Path, label: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = label
    worksheet["B1"] = "=C1/D1"
    worksheet["B1"].number_format = "$#,##0.00"
    worksheet["C1"] = 50
    worksheet["D1"] = 100

    scan = _scan(tmp_path, workbook, FormulaFormatRoleMismatchRule())

    assert scan.findings == []
