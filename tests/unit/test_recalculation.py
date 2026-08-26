from __future__ import annotations

import subprocess
import time
import zipfile
from contextlib import nullcontext
from datetime import date, datetime
from datetime import time as datetime_time
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from workbooklens import conversion
from workbooklens.exceptions import PatchValidationError
from workbooklens.repair import recalculation

_SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _provider(kind: str) -> conversion.ConversionProvider:
    if kind == "excel":
        return conversion.ConversionProvider("excel", "Microsoft Excel", Path("powershell.exe"))
    return conversion.ConversionProvider("libreoffice", "LibreOffice", Path("soffice.exe"))


def _rewrite_zip_member(path: Path, member: str, transform: Any) -> None:
    temporary = path.with_name(f".{path.stem}-rewrite.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as output:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == member:
                payload = transform(payload)
            output.writestr(info, payload)
    temporary.replace(path)


def _formula_error_workbook(
    path: Path,
    errors: tuple[str, ...],
    *,
    dimension: str = "A1",
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Errors"
    for row in range(1, len(errors) + 1):
        worksheet.cell(row=row, column=1, value="=1/0")
    workbook.save(path)
    workbook.close()

    def add_cached_errors(payload: bytes) -> bytes:
        root = etree.fromstring(payload)
        dimension_node = root.find(f"{{{_SPREADSHEET_NS}}}dimension")
        assert dimension_node is not None
        dimension_node.set("ref", dimension)
        cells = root.findall(f".//{{{_SPREADSHEET_NS}}}c")
        assert len(cells) == len(errors)
        for cell, error in zip(cells, errors, strict=True):
            cell.set("t", "e")
            value = cell.find(f"{{{_SPREADSHEET_NS}}}v")
            assert value is not None
            value.text = error
        return etree.tostring(root, encoding="utf-8", xml_declaration=True)

    _rewrite_zip_member(path, "xl/worksheets/sheet1.xml", add_cached_errors)


def _structured_workbook(path: Path, *, formula: str = "=SUM(A1:A2)") -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Data"
    worksheet["A1"] = 2
    worksheet["A2"] = 3
    worksheet["B1"] = formula
    worksheet.merge_cells("C1:D1")
    worksheet["C1"] = "Heading"
    workbook.defined_names.add(DefinedName("InputRange", attr_text="'Data'!$A$1:$A$2"))
    workbook.save(path)
    workbook.close()


def _literal_structure_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Literals"
    worksheet["A1"] = 12.5
    worksheet["A2"] = "alpha"
    worksheet["A3"] = datetime(2026, 8, 26, 14, 30, 15, 125000)
    worksheet["A4"] = date(2026, 8, 26)
    worksheet["A5"] = datetime_time(8, 15, 30, 250000)
    worksheet["A6"] = True
    worksheet["A7"] = "#N/A"
    worksheet["A8"] = "00123"
    worksheet["A8"].quotePrefix = True
    worksheet["B1"] = "=SUM(A1, 1)"
    workbook.save(path)
    workbook.close()


def test_candidate_providers_preserves_preference_and_fidelity_order(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    excel = _provider("excel")
    libreoffice = _provider("libreoffice")
    monkeypatch.setattr(
        conversion,
        "available_conversion_providers",
        lambda: (excel, libreoffice),
    )

    assert recalculation.candidate_providers("auto") == (excel, libreoffice)
    assert recalculation.candidate_providers("excel") == (excel,)
    assert recalculation.candidate_providers("libreoffice") == (libreoffice,)
    assert recalculation.candidate_providers("none") == ()


def test_non_windows_provider_discovery_skips_excel_registry_and_degrades_cleanly(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(conversion.sys, "platform", "linux")
    monkeypatch.setattr(conversion, "_find_powershell", lambda: Path("/usr/bin/pwsh"))
    monkeypatch.setattr(conversion, "_find_libreoffice", lambda: None)
    monkeypatch.setattr(
        conversion.importlib,
        "import_module",
        lambda name: pytest.fail(f"non-Windows discovery imported {name}"),
    )

    assert conversion.available_conversion_providers() == ()
    assert recalculation.candidate_providers("auto") == ()


def test_private_workspace_must_be_a_direct_child(tmp_path: Path) -> None:
    intended_parent = tmp_path / "intended"
    escaped = tmp_path / "escaped"
    intended_parent.mkdir()
    escaped.mkdir()

    with pytest.raises(PatchValidationError, match="escaped its intended parent directory"):
        recalculation._validated_private_workspace(
            escaped,
            intended_parent,
            "Recalculation workspace",
        )


def test_excel_recalculation_rejects_overlong_tool_path_before_process_start(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"source")
    output = tmp_path / "output.xlsx"
    provider = _provider("excel")
    long_workspace = tmp_path / ("x" * 240)
    monkeypatch.setattr(
        recalculation.tempfile,
        "TemporaryDirectory",
        lambda **_kwargs: nullcontext(str(long_workspace)),
    )
    monkeypatch.setattr(
        subprocess,
        "Popen",
        lambda *_args, **_kwargs: pytest.fail("Excel must not start for an overlong tool path"),
    )

    with pytest.raises(PatchValidationError, match="safe Windows path limit"):
        recalculation._run_excel_recalculation(provider, source, output, 30)


def test_excel_tool_path_accepts_limit_and_rejects_one_character_over(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(recalculation.sys, "platform", "win32")
    allowed = Path("C:/" + "x" * (recalculation._WINDOWS_LEGACY_PATH_LIMIT - 3))
    blocked = Path(f"{allowed}x")
    assert len(str(allowed)) == recalculation._WINDOWS_LEGACY_PATH_LIMIT
    assert len(str(blocked)) == recalculation._WINDOWS_LEGACY_PATH_LIMIT + 1

    recalculation._validate_excel_tool_path(allowed, "Allowed path")
    with pytest.raises(PatchValidationError, match="choose a shorter output directory"):
        recalculation._validate_excel_tool_path(blocked, "Blocked path")


def test_excel_recalculation_script_disables_active_content_before_read_only_open() -> None:
    script = recalculation._EXCEL_RECALCULATION_SCRIPT
    open_call = "$excel.Workbooks.Open($inputPath, 0, $true)"
    required_settings = (
        "$excel.DisplayAlerts = $false",
        "$excel.AskToUpdateLinks = $false",
        "$excel.EnableEvents = $false",
        "$excel.AutomationSecurity = 3",
    )

    assert open_call in script
    for setting in required_settings:
        assert setting in script
        assert script.index(setting) < script.index(open_call)


def test_excel_recalculation_script_initializes_calculation_with_temporary_workbook() -> None:
    script = recalculation._EXCEL_RECALCULATION_SCRIPT
    add_workbook = "$calculationWorkbook = $excel.Workbooks.Add()"
    set_calculation = "$excel.Calculation = -4135"
    set_save_calculation = "$excel.CalculateBeforeSave = $false"
    close_workbook = "$calculationWorkbook.Close($false)"
    release_workbook = "Release-ComObject $calculationWorkbook"
    open_source = "$workbook = $excel.Workbooks.Open($inputPath, 0, $true)"

    add_index = script.index(add_workbook)
    set_index = script.index(set_calculation, add_index)
    set_save_index = script.index(set_save_calculation, set_index)
    close_index = script.index(close_workbook, set_save_index)
    release_index = script.index(release_workbook, close_index)
    clear_index = script.index("$calculationWorkbook = $null", release_index)
    assert add_index < set_index < set_save_index < close_index < release_index < clear_index
    assert clear_index < script.index(open_source)

    cleanup_block = """if ($null -ne $calculationWorkbook) {
        try { $calculationWorkbook.Close($false) } catch {}
        Release-ComObject $calculationWorkbook
    }"""
    assert cleanup_block in script
    assert script.index(cleanup_block) < script.index("if ($null -ne $workbook)")


def test_validation_recalculates_source_and_candidate_with_same_provider(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source")
    candidate.write_bytes(b"candidate")
    excel = _provider("excel")
    calls: list[tuple[str, str, Path]] = []
    monkeypatch.setattr(recalculation, "candidate_providers", lambda _preference: (excel,))

    def fake_run(
        provider: conversion.ConversionProvider,
        input_path: Path,
        output_path: Path,
        _timeout_seconds: int,
    ) -> None:
        calls.append((provider.kind, input_path.name, input_path.parent))
        output_path.write_bytes(input_path.read_bytes())

    monkeypatch.setattr(recalculation, "_run_provider", fake_run)
    monkeypatch.setattr(
        recalculation,
        "_validate_recalculated_package",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        recalculation,
        "collect_formula_errors",
        lambda path, **_kwargs: ("Sheet1!B2:#DIV/0!",) if path.name == "s.xlsx" else (),
    )

    result = recalculation.validate_formula_recalculation(source, candidate)

    assert [(kind, name) for kind, name, _parent in calls] == [
        ("excel", "s.xlsx"),
        ("excel", "c.xlsx"),
    ]
    assert calls[0][2] == calls[1][2]
    assert calls[0][2].parent == tmp_path
    assert calls[0][2] != tmp_path
    assert not calls[0][2].exists()
    assert result.provider == excel
    assert result.formula_errors_before == ("Sheet1!B2:#DIV/0!",)
    assert result.formula_errors_after == ()


def test_validation_falls_back_only_by_recalculating_both_copies_with_next_provider(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source")
    candidate.write_bytes(b"candidate")
    excel = _provider("excel")
    libreoffice = _provider("libreoffice")
    calls: list[tuple[str, str]] = []
    monkeypatch.setattr(
        recalculation,
        "candidate_providers",
        lambda _preference: (excel, libreoffice),
    )

    def fake_run(
        provider: conversion.ConversionProvider,
        input_path: Path,
        output_path: Path,
        _timeout_seconds: int,
    ) -> None:
        calls.append((provider.kind, input_path.name))
        if provider.kind == "excel" and input_path.name == "s.xlsx":
            raise PatchValidationError("Excel source recalculation failed")
        output_path.write_bytes(input_path.read_bytes())

    monkeypatch.setattr(recalculation, "_run_provider", fake_run)
    monkeypatch.setattr(
        recalculation,
        "_validate_recalculated_package",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        recalculation,
        "collect_formula_errors",
        lambda _path, **_kwargs: (),
    )

    result = recalculation.validate_formula_recalculation(source, candidate)

    assert calls == [
        ("excel", "s.xlsx"),
        ("libreoffice", "s.xlsx"),
        ("libreoffice", "c.xlsx"),
    ]
    assert result.provider == libreoffice


def test_validation_refuses_fallback_after_provider_validates_source(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source")
    candidate.write_bytes(b"candidate")
    excel = _provider("excel")
    libreoffice = _provider("libreoffice")
    calls: list[tuple[str, str]] = []
    monkeypatch.setattr(
        recalculation,
        "candidate_providers",
        lambda _preference: (excel, libreoffice),
    )

    def fake_run(
        provider: conversion.ConversionProvider,
        input_path: Path,
        output_path: Path,
        _timeout_seconds: int,
    ) -> None:
        calls.append((provider.kind, input_path.name))
        if provider.kind == "excel" and input_path.name == "c.xlsx":
            raise PatchValidationError("Excel candidate recalculation failed")
        output_path.write_bytes(input_path.read_bytes())

    monkeypatch.setattr(recalculation, "_run_provider", fake_run)
    monkeypatch.setattr(
        recalculation,
        "_validate_recalculated_package",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        recalculation,
        "collect_formula_errors",
        lambda _path, **_kwargs: (),
    )

    with pytest.raises(PatchValidationError, match="refusing provider fallback") as caught:
        recalculation.validate_formula_recalculation(source, candidate)

    assert calls == [
        ("excel", "s.xlsx"),
        ("excel", "c.xlsx"),
    ]
    assert caught.value.__dict__["recalculation_provider"] == "excel"


@pytest.mark.parametrize(
    ("hash_argument", "expected_message"),
    [
        ("expected_source_sha256", "Source workbook SHA-256 changed"),
        ("expected_candidate_sha256", "Repair candidate SHA-256 changed"),
    ],
)
def test_validation_rejects_input_hash_mismatch_before_provider_execution(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    hash_argument: str,
    expected_message: str,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source")
    candidate.write_bytes(b"candidate")
    excel = _provider("excel")
    monkeypatch.setattr(recalculation, "candidate_providers", lambda _preference: (excel,))
    monkeypatch.setattr(
        recalculation,
        "_run_provider",
        lambda *_args: pytest.fail("provider must not receive a hash-mismatched input"),
    )

    with pytest.raises(PatchValidationError, match=expected_message):
        recalculation.validate_formula_recalculation(
            source,
            candidate,
            **{hash_argument: "0" * 64},
        )

    assert not tuple(tmp_path.glob(".wl-v-*"))


def test_stage_verified_copy_detects_source_replacement_during_copy(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    staged = tmp_path / "staged.xlsx"
    source.write_bytes(b"stable-original")
    expected = recalculation.sha256_file(source)
    original_copyfileobj = recalculation.shutil.copyfileobj

    def replace_source_during_copy(source_handle: Any, staged_handle: Any, *, length: int) -> None:
        original_copyfileobj(source_handle, staged_handle, length=length)
        source.write_bytes(b"replaced-during-copy")

    monkeypatch.setattr(
        recalculation.shutil,
        "copyfileobj",
        replace_source_during_copy,
    )

    with pytest.raises(PatchValidationError, match="Source workbook SHA-256 changed"):
        recalculation._stage_verified_copy(
            source,
            staged,
            expected_sha256=expected,
            label="Source workbook",
        )

    assert not staged.exists()


def test_validation_rejects_staged_input_drift_without_provider_fallback(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source")
    candidate.write_bytes(b"candidate")
    excel = _provider("excel")
    libreoffice = _provider("libreoffice")
    calls: list[tuple[str, str]] = []
    workspaces: list[Path] = []
    monkeypatch.setattr(
        recalculation,
        "candidate_providers",
        lambda _preference: (excel, libreoffice),
    )

    def mutate_staged_input(
        provider: conversion.ConversionProvider,
        input_path: Path,
        output_path: Path,
        _timeout_seconds: int,
    ) -> None:
        calls.append((provider.kind, input_path.name))
        workspaces.append(input_path.parent)
        output_path.write_bytes(input_path.read_bytes())
        input_path.write_bytes(b"provider-mutated-staged-input")

    monkeypatch.setattr(recalculation, "_run_provider", mutate_staged_input)

    with pytest.raises(PatchValidationError, match="Staged source workbook SHA-256 changed"):
        recalculation.validate_formula_recalculation(source, candidate)

    assert calls == [("excel", "s.xlsx")]
    assert len(workspaces) == 1
    assert not workspaces[0].exists()


def test_validation_rechecks_staged_hash_after_structure_validation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source")
    candidate.write_bytes(b"candidate")
    excel = _provider("excel")
    libreoffice = _provider("libreoffice")
    calls: list[tuple[str, str]] = []
    monkeypatch.setattr(
        recalculation,
        "candidate_providers",
        lambda _preference: (excel, libreoffice),
    )

    def copy_input(
        provider: conversion.ConversionProvider,
        input_path: Path,
        output_path: Path,
        _timeout_seconds: int,
    ) -> None:
        calls.append((provider.kind, input_path.name))
        output_path.write_bytes(input_path.read_bytes())

    def mutate_during_validation(
        _path: Path,
        _limits: Any,
        *,
        expected_input: Path,
    ) -> None:
        if expected_input.name == "c.xlsx":
            expected_input.write_bytes(b"mutated-inside-structure-validation")

    monkeypatch.setattr(recalculation, "_run_provider", copy_input)
    monkeypatch.setattr(
        recalculation,
        "_validate_recalculated_package",
        mutate_during_validation,
    )
    monkeypatch.setattr(recalculation, "collect_formula_errors", lambda *_args, **_kwargs: ())

    with pytest.raises(PatchValidationError, match="Staged repair candidate SHA-256 changed"):
        recalculation.validate_formula_recalculation(source, candidate)

    assert calls == [
        ("excel", "s.xlsx"),
        ("excel", "c.xlsx"),
    ]


def test_validation_uses_fixed_staged_inputs_after_original_paths_change(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source-original")
    candidate.write_bytes(b"candidate-original")
    excel = _provider("excel")
    seen: list[tuple[str, bytes]] = []
    monkeypatch.setattr(recalculation, "candidate_providers", lambda _preference: (excel,))

    def replace_original_paths(
        _provider: conversion.ConversionProvider,
        input_path: Path,
        output_path: Path,
        _timeout_seconds: int,
    ) -> None:
        seen.append((input_path.name, input_path.read_bytes()))
        if input_path.name == "s.xlsx":
            source.write_bytes(b"source-replaced-later")
            candidate.write_bytes(b"candidate-replaced-later")
        output_path.write_bytes(input_path.read_bytes())

    monkeypatch.setattr(recalculation, "_run_provider", replace_original_paths)
    monkeypatch.setattr(
        recalculation,
        "_validate_recalculated_package",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(recalculation, "collect_formula_errors", lambda *_args, **_kwargs: ())

    result = recalculation.validate_formula_recalculation(source, candidate)

    assert result.provider == excel
    assert seen == [
        ("s.xlsx", b"source-original"),
        ("c.xlsx", b"candidate-original"),
    ]
    assert source.read_bytes() == b"source-replaced-later"
    assert candidate.read_bytes() == b"candidate-replaced-later"


def test_validation_rejects_new_formula_errors_without_provider_fallback(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source")
    candidate.write_bytes(b"candidate")
    excel = _provider("excel")
    libreoffice = _provider("libreoffice")
    calls: list[str] = []
    monkeypatch.setattr(
        recalculation,
        "candidate_providers",
        lambda _preference: (excel, libreoffice),
    )

    def fake_run(
        provider: conversion.ConversionProvider,
        input_path: Path,
        output_path: Path,
        _timeout_seconds: int,
    ) -> None:
        calls.append(provider.kind)
        output_path.write_bytes(input_path.read_bytes())

    monkeypatch.setattr(recalculation, "_run_provider", fake_run)
    monkeypatch.setattr(
        recalculation,
        "_validate_recalculated_package",
        lambda *args, **kwargs: None,
    )
    monkeypatch.setattr(
        recalculation,
        "collect_formula_errors",
        lambda path, **_kwargs: () if path.name == "s.xlsx" else ("Sheet1!D4:#VALUE!",),
    )

    with pytest.raises(PatchValidationError, match="introduced new errors"):
        recalculation.validate_formula_recalculation(source, candidate)

    assert calls == ["excel", "excel"]


def test_validation_reports_all_provider_failures_without_raw_process_output(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    source.write_bytes(b"source")
    candidate.write_bytes(b"candidate")
    excel = _provider("excel")
    libreoffice = _provider("libreoffice")
    monkeypatch.setattr(
        recalculation,
        "candidate_providers",
        lambda _preference: (excel, libreoffice),
    )
    monkeypatch.setattr(
        recalculation,
        "_run_provider",
        lambda provider, *args: (_ for _ in ()).throw(
            PatchValidationError(f"{provider.label} unavailable")
        ),
    )

    with pytest.raises(PatchValidationError, match="Every requested recalculation provider failed"):
        recalculation.validate_formula_recalculation(source, candidate)


def test_libreoffice_timeout_terminates_process_tree(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"source")
    output = tmp_path / "output.xlsx"
    provider = _provider("libreoffice")
    terminated: list[Any] = []

    class FakeProcess:
        returncode = None

        def communicate(self, timeout: int) -> tuple[str, str]:
            raise subprocess.TimeoutExpired(["soffice"], timeout)

    process = FakeProcess()
    monkeypatch.setattr(
        subprocess,
        "Popen",
        lambda *args, **kwargs: process,
    )
    monkeypatch.setattr(conversion, "_terminate_process_tree", terminated.append)

    with pytest.raises(PatchValidationError, match="timed out after 7 seconds"):
        recalculation._run_libreoffice_recalculation(provider, source, output, 7)

    assert terminated == [process]
    assert not output.exists()


def test_collect_formula_errors_is_sparse_bounded_and_supports_modern_errors(
    tmp_path: Path,
) -> None:
    path = tmp_path / "errors.xlsx"
    errors = ("#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#UNKNOWN!")
    _formula_error_workbook(path, errors, dimension="A1:XFD1048576")

    result = recalculation.collect_formula_errors(path, max_cells=10, timeout_seconds=5)

    assert result == tuple(f"Errors!A{row}:{error}" for row, error in enumerate(errors, start=1))


def test_collect_formula_errors_rejects_actual_sparse_cell_overflow(tmp_path: Path) -> None:
    path = tmp_path / "too-many-cells.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = 1
    worksheet["A2"] = 2
    worksheet["A3"] = "=SUM(A1:A2)"
    workbook.save(path)
    workbook.close()

    with pytest.raises(PatchValidationError, match="2-cell limit"):
        recalculation.collect_formula_errors(path, max_cells=2)


def test_collect_formula_errors_enforces_time_limit(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    path = tmp_path / "slow.xlsx"
    _formula_error_workbook(path, ("#CALC!",))
    clock = iter((0.0, 2.0))
    monkeypatch.setattr(time, "monotonic", lambda: next(clock))

    with pytest.raises(PatchValidationError, match="time limit"):
        recalculation.collect_formula_errors(path, timeout_seconds=1)


def test_recalculated_package_preserves_formulas_and_structure(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    recalculated = tmp_path / "recalculated.xlsx"
    _structured_workbook(source)
    recalculated.write_bytes(source.read_bytes())

    recalculation._validate_recalculated_package(
        recalculated,
        None,
        expected_input=source,
    )


@pytest.mark.parametrize(
    "mutation",
    [
        "number",
        "text",
        "datetime",
        "date",
        "time",
        "bool",
        "error",
        "quote_prefix",
    ],
)
def test_recalculated_package_rejects_provider_literal_mutation(
    tmp_path: Path,
    mutation: str,
) -> None:
    source = tmp_path / "source.xlsx"
    recalculated = tmp_path / f"recalculated-{mutation}.xlsx"
    _literal_structure_workbook(source)
    recalculated.write_bytes(source.read_bytes())
    workbook = load_workbook(recalculated)
    worksheet = workbook["Literals"]
    if mutation == "number":
        worksheet["A1"] = 99.5
    elif mutation == "text":
        worksheet["A2"] = "beta"
    elif mutation == "datetime":
        worksheet["A3"] = datetime(2027, 1, 1, 9, 0)
    elif mutation == "date":
        worksheet["A4"] = date(2027, 1, 1)
    elif mutation == "time":
        worksheet["A5"] = datetime_time(9, 45)
    elif mutation == "bool":
        worksheet["A6"] = False
    elif mutation == "error":
        worksheet["A7"] = "#DIV/0!"
    else:
        worksheet["A8"].quotePrefix = False
    workbook.save(recalculated)
    workbook.close()

    with pytest.raises(PatchValidationError, match="unexpectedly changed"):
        recalculation._validate_recalculated_package(
            recalculated,
            None,
            expected_input=source,
        )


def test_literal_canonicalization_is_stable_and_rejects_non_finite_values() -> None:
    value = 0.1
    assert recalculation._canonical_literal_value(value, location="Data!A1") == (
        recalculation._canonical_literal_value(
            float.fromhex(value.hex()),
            location="Data!A1",
        )
    )
    assert recalculation._canonical_literal_value(Decimal("1.00"), location="Data!A2") == (
        "decimal:1"
    )

    for value in (
        float("nan"),
        float("inf"),
        float("-inf"),
        Decimal("NaN"),
        Decimal("Infinity"),
        Decimal("-Infinity"),
    ):
        with pytest.raises(PatchValidationError, match="non-finite"):
            recalculation._canonical_literal_value(value, location="Data!A3")


def test_recalculated_package_allows_only_formula_cached_value_change(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    recalculated = tmp_path / "recalculated.xlsx"
    _literal_structure_workbook(source)
    recalculated.write_bytes(source.read_bytes())

    def replace_formula_cache(payload: bytes) -> bytes:
        root = etree.fromstring(payload)
        formula_cell = root.find(f".//{{{_SPREADSHEET_NS}}}c[@r='B1']")
        assert formula_cell is not None
        formula = formula_cell.find(f"{{{_SPREADSHEET_NS}}}f")
        value = formula_cell.find(f"{{{_SPREADSHEET_NS}}}v")
        assert formula is not None
        assert value is not None
        value.text = "13.5"
        return etree.tostring(root, encoding="utf-8", xml_declaration=True)

    _rewrite_zip_member(
        recalculated,
        "xl/worksheets/sheet1.xml",
        replace_formula_cache,
    )

    recalculation._validate_recalculated_package(
        recalculated,
        None,
        expected_input=source,
    )


def test_recalculated_package_accepts_excel_removal_of_optional_sheet_quotes(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    recalculated = tmp_path / "recalculated.xlsx"
    workbook = Workbook()
    data = workbook.active
    assert data is not None
    data.title = "项目预算_练习"
    data["A1"] = 1
    summary = workbook.create_sheet("汇总结果_练习")
    summary["A1"] = "=SUM('项目预算_练习'!A1:A1)"
    workbook.save(source)
    summary["A1"] = "=SUM(项目预算_练习!A1:A1)"
    workbook.save(recalculated)
    workbook.close()

    recalculation._validate_recalculated_package(
        recalculated,
        None,
        expected_input=source,
    )


def test_formula_structure_canonicalization_accepts_ascii_identifier_sheet_quotes() -> None:
    assert recalculation._canonical_formula_text("='Data_2026'!A1") == (
        recalculation._canonical_formula_text("=Data_2026!A1")
    )


@pytest.mark.parametrize(
    ("before", "after"),
    [
        ("='Straße'!A1", "=STRASSE!A1"),
        ("='SheetA'!A1", "=sheeta!A1"),
    ],
)
def test_formula_structure_canonicalization_preserves_exact_sheet_identity(
    before: str,
    after: str,
) -> None:
    assert recalculation._canonical_formula_text(before) != recalculation._canonical_formula_text(
        after
    )


@pytest.mark.parametrize("sheet_name", ["Straße", "SheetA", "Data_2026"])
def test_formula_structure_canonicalization_keeps_same_sheet_quoted_unquoted_equivalent(
    sheet_name: str,
) -> None:
    assert recalculation._canonical_formula_text(f"='{sheet_name}'!A1") == (
        recalculation._canonical_formula_text(f"={sheet_name}!A1")
    )


@pytest.mark.parametrize(
    ("before", "after"),
    [
        ("=\"'Data'!A1\"", '="Data!A1"'),
        ("='Sales Data'!A1", "=Sales Data!A1"),
        ("='O''Brien'!A1", "=O'Brien!A1"),
        ("='[Book.xlsx]Data'!A1", "=[Book.xlsx]Data!A1"),
        ("=SUM(Sheet1:Sheet3!A1)", "=SUM('Sheet1:Sheet3'!A1)"),
        ("='Data'!项目1", "=Data!项目1"),
        ("='Data'!A0", "=Data!A0"),
        ("='A1'!B2", "=A1!B2"),
        ("='R1C1'!A1", "=R1C1!A1"),
    ],
)
def test_formula_structure_canonicalization_keeps_high_risk_quote_changes_distinct(
    before: str,
    after: str,
) -> None:
    assert recalculation._canonical_formula_text(before) != recalculation._canonical_formula_text(
        after
    )


@pytest.mark.parametrize(
    ("before", "after"),
    [
        ("=SUM('Data_2026'!A1:A2)", "=AVERAGE(Data_2026!A1:A2)"),
        ("='Data_2026'!A1", "=Data_2026!B1"),
        ("='Data_2026'!A1+1", "=Data_2026!A1-1"),
        ("=COUNTIF('Data_2026'!A1:A2,\"yes\")", '=COUNTIF(Data_2026!A1:A2,"no")'),
    ],
)
def test_formula_structure_canonicalization_keeps_formula_changes_distinct(
    before: str,
    after: str,
) -> None:
    assert recalculation._canonical_formula_text(before) != recalculation._canonical_formula_text(
        after
    )


@pytest.mark.parametrize("mutation", ["formula", "sheet", "defined_name"])
def test_recalculated_package_rejects_provider_structure_mutation(
    tmp_path: Path,
    mutation: str,
) -> None:
    source = tmp_path / "source.xlsx"
    recalculated = tmp_path / f"recalculated-{mutation}.xlsx"
    _structured_workbook(source)
    recalculated.write_bytes(source.read_bytes())
    workbook = load_workbook(recalculated)
    worksheet = workbook["Data"]
    if mutation == "formula":
        worksheet["B1"] = "=A1-A2"
    elif mutation == "sheet":
        worksheet.title = "Changed"
    else:
        workbook.defined_names["InputRange"].attr_text = "'Data'!$A$1"
    workbook.save(recalculated)
    workbook.close()

    with pytest.raises(PatchValidationError, match="unexpectedly changed"):
        recalculation._validate_recalculated_package(
            recalculated,
            None,
            expected_input=source,
        )


def test_validation_allows_only_formula_changes_already_present_in_candidate(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    _structured_workbook(source, formula="=SUM(A1:A2)")
    _structured_workbook(candidate, formula="=A1+A2")
    excel = _provider("excel")
    monkeypatch.setattr(recalculation, "candidate_providers", lambda _preference: (excel,))
    monkeypatch.setattr(
        recalculation,
        "_run_provider",
        lambda _provider, input_path, output_path, _timeout: output_path.write_bytes(
            input_path.read_bytes()
        ),
    )
    monkeypatch.setattr(
        recalculation,
        "collect_formula_errors",
        lambda _path, **_kwargs: (),
    )

    result = recalculation.validate_formula_recalculation(source, candidate)

    assert result.provider == excel


def test_validation_rejects_provider_formula_mutation(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    candidate = tmp_path / "candidate.xlsx"
    _structured_workbook(source)
    _structured_workbook(candidate, formula="=A1+A2")
    excel = _provider("excel")
    monkeypatch.setattr(recalculation, "candidate_providers", lambda _preference: (excel,))

    def mutate_formula(
        _provider: conversion.ConversionProvider,
        input_path: Path,
        output_path: Path,
        _timeout: int,
    ) -> None:
        output_path.write_bytes(input_path.read_bytes())
        workbook = load_workbook(output_path)
        workbook["Data"]["B1"] = "=A1-A2"
        workbook.save(output_path)
        workbook.close()

    monkeypatch.setattr(recalculation, "_run_provider", mutate_formula)

    with pytest.raises(PatchValidationError, match="unexpectedly changed"):
        recalculation.validate_formula_recalculation(source, candidate)


def test_excel_handshake_cleanup_still_terminates_wrapper_when_excel_cleanup_fails(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"source")
    output = tmp_path / "output.xlsx"
    provider = _provider("excel")
    process = SimpleNamespace()
    actions: list[str] = []
    identity = SimpleNamespace(process_id=73)
    monkeypatch.setattr(conversion, "_snapshot_excel_process_ids", lambda _runner: frozenset())
    monkeypatch.setattr(
        subprocess,
        "Popen",
        lambda *args, **kwargs: process,
    )
    monkeypatch.setattr(
        conversion,
        "_wait_for_owned_excel_process_identity",
        lambda *args: (_ for _ in ()).throw(conversion._ExcelStartupTimeout()),
    )
    monkeypatch.setattr(conversion, "_load_excel_process_identity", lambda _path: identity)

    def fail_excel_cleanup(_identity: Any) -> None:
        actions.append("excel")
        raise conversion._ProviderFailure("cleanup failed")

    monkeypatch.setattr(conversion, "_terminate_excel_identity_if_present", fail_excel_cleanup)
    monkeypatch.setattr(
        conversion,
        "_terminate_process_tree",
        lambda _process: actions.append("wrapper"),
    )

    with pytest.raises(PatchValidationError, match="handshake timed out"):
        recalculation._run_excel_recalculation(provider, source, output, 7)

    assert actions == ["excel", "wrapper"]


def test_excel_timeout_cleanup_still_terminates_wrapper_when_excel_cleanup_fails(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"source")
    output = tmp_path / "output.xlsx"
    provider = _provider("excel")
    identity = SimpleNamespace(process_id=73)
    owned_excel = object()
    actions: list[str] = []

    class FakeProcess:
        returncode = None

        def communicate(self, timeout: int) -> tuple[str, str]:
            raise subprocess.TimeoutExpired(["powershell"], timeout)

    process = FakeProcess()
    monkeypatch.setattr(conversion, "_snapshot_excel_process_ids", lambda _runner: frozenset())
    monkeypatch.setattr(
        subprocess,
        "Popen",
        lambda *args, **kwargs: process,
    )
    monkeypatch.setattr(
        conversion,
        "_wait_for_owned_excel_process_identity",
        lambda *args: identity,
    )
    monkeypatch.setattr(
        conversion, "_open_verified_excel_process", lambda *args, **kwargs: owned_excel
    )
    monkeypatch.setattr(conversion, "_signal_excel_go", lambda _path: None)

    def fail_excel_cleanup(_owned_excel: Any) -> None:
        actions.append("excel")
        raise conversion._ProviderFailure("cleanup failed")

    monkeypatch.setattr(conversion, "_terminate_excel_process", fail_excel_cleanup)
    monkeypatch.setattr(
        conversion,
        "_terminate_process_tree",
        lambda _process: actions.append("wrapper"),
    )
    monkeypatch.setattr(
        conversion,
        "_close_excel_process",
        lambda _process: actions.append("close"),
    )

    with pytest.raises(PatchValidationError, match="timed out after 7 seconds"):
        recalculation._run_excel_recalculation(provider, source, output, 7)

    assert actions == ["excel", "wrapper", "close"]


def test_libreoffice_timeout_cleanup_failure_does_not_mask_timeout(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"source")
    output = tmp_path / "output.xlsx"
    provider = _provider("libreoffice")
    actions: list[str] = []

    class FakeProcess:
        returncode = None

        def communicate(self, timeout: int) -> tuple[str, str]:
            raise subprocess.TimeoutExpired(["soffice"], timeout)

    process = FakeProcess()
    monkeypatch.setattr(
        subprocess,
        "Popen",
        lambda *args, **kwargs: process,
    )

    def fail_cleanup(_process: Any) -> None:
        actions.append("wrapper")
        raise OSError("cleanup failed")

    monkeypatch.setattr(conversion, "_terminate_process_tree", fail_cleanup)

    with pytest.raises(PatchValidationError, match="timed out after 7 seconds"):
        recalculation._run_libreoffice_recalculation(provider, source, output, 7)

    assert actions == ["wrapper"]
