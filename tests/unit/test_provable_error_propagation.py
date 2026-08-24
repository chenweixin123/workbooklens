from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from workbooklens.formulas import error_propagation
from workbooklens.formulas.error_propagation import find_static_formula_errors
from workbooklens.scanner import scan_workbook


def _wl034_by_sheet(path: Path, sheet: str) -> dict[str, object]:
    scan = scan_workbook(path)
    return {
        finding.location: finding
        for finding in scan.findings
        if finding.rule_id == "WL034_PROVABLE_FORMULA_ERROR" and finding.sheet == sheet
    }


def test_provable_formula_errors_propagate_through_exact_references_and_aggregates(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source Data"
    source["I5"] = 10
    source["I6"] = 20
    source["I7"] = 30
    source["I8"] = "=1/0"
    source["I9"] = '=VALUE("abc")'
    source["I10"] = 40
    source["I11"] = "=I8"
    source["I12"] = "#NUM!"

    summary = workbook.create_sheet("Inventory")
    summary["J4"] = "='Source Data'!I8"
    summary["J5"] = "=SUM('Source Data'!I5:I10)"
    summary["J6"] = "=AVERAGE('Source Data'!I9:I10)"
    summary["J7"] = "=MIN('Source Data'!I8:I10)"
    summary["J8"] = "=MAX('Source Data'!I11:I11)"
    summary["J9"] = "=SUM('Source Data'!I12:I12)"
    summary["J10"] = "=J4"
    summary["J11"] = "=SUM(J10:J10)"

    path = tmp_path / "propagated-formula-errors.xlsx"
    workbook.save(path)
    workbook.close()

    findings = _wl034_by_sheet(path, "Inventory")
    assert set(findings) == {"J4", "J5", "J6", "J7", "J8", "J9", "J10", "J11"}
    for finding in findings.values():
        assert float(finding.confidence) == 1.0
        assert finding.evidence.details["proof"] == "propagated_formula_error"
        assert finding.evidence.details["source"]
        assert finding.patch_ids == []


def test_error_propagation_excludes_error_handlers_and_nonexact_expressions(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source"
    source["A1"] = "=1/0"
    source["A2"] = "'#DIV/0!"

    summary = workbook.create_sheet("Summary")
    summary["B1"] = "=IFERROR(Source!A1,0)"
    summary["B2"] = "=IF(FALSE,Source!A1,1)"
    summary["B3"] = "=COUNT(Source!A1:A1)"
    summary["B4"] = "=AGGREGATE(9,6,Source!A1:A1)"
    summary["B5"] = "=SUBTOTAL(9,Source!A1:A1)"
    summary["B6"] = "=SUM(IFERROR(Source!A1,0))"
    summary["B7"] = "=SUM(Source!A1:A1,1)"
    summary["B8"] = "=ABS(Source!A1)+0"
    summary["B9"] = "=IFERROR(SUM(Source!A1:A1),0)"
    summary["B10"] = "=SUM(IF(FALSE,Source!A1,0))"
    summary["B11"] = "=SUM(Source!A2:A2)"
    summary["B12"] = "=Source!A1:A1+0"
    summary["B13"] = "=Source!A1,Source!A1"
    summary["B14"] = "=Source!A1 Source!A1"
    summary["B15"] = "=Source!A1#"
    summary["B16"] = "={1,Source!A1}"
    summary["B17"] = "=@Source!A1+0"
    summary["B18"] = "=Source!A1+NamedValue"
    summary["B19"] = "=Source!A1 +0"
    summary["B20"] = "=Source!A1%%"
    summary["B21"] = "=B21+Source!A1"
    summary["B22"] = "=Missing!A1+Source!A1"
    summary["B23"] = "=Source!A1:A1"

    path = tmp_path / "nonpropagating-or-nonexact-formulas.xlsx"
    workbook.save(path)
    workbook.close()

    assert _wl034_by_sheet(path, "Summary") == {}


def test_direct_reference_and_aggregate_keep_cosmetic_whitespace_compatibility(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source"
    source["A1"] = "=1/0"

    summary = workbook.create_sheet("Summary")
    summary["D1"] = "= Source!A1"
    summary["D2"] = "= SUM(Source!A1:A1)"

    path = tmp_path / "whitespace-compatible-propagation.xlsx"
    workbook.save(path)
    workbook.close()

    findings = _wl034_by_sheet(path, "Summary")
    assert set(findings) == {"D1", "D2"}
    assert findings["D1"].evidence.details["function"] == "DIRECT_REFERENCE"
    assert findings["D2"].evidence.details["function"] == "SUM"
    assert all(finding.evidence.details["error_code_proven"] for finding in findings.values())


def test_provable_formula_errors_propagate_through_pure_operator_expressions(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source"
    source["A1"] = "=1/0"
    source["A2"] = '=VALUE("abc")'
    source["A3"] = "=#REF!"

    summary = workbook.create_sheet("Summary")
    summary["C1"] = "=Source!A1+1"
    summary["C2"] = "=-(Source!A1*2%)"
    summary["C3"] = "=Source!A1=0"
    summary["C4"] = '="prefix"&Source!A2'
    summary["C5"] = "=(Source!A1^2)"
    summary["C6"] = "=+Source!A1"
    summary["C7"] = "=Source!A1%"
    summary["C8"] = "=Source!A3<>0"
    summary["C9"] = "=C1+Source!A2"
    summary["C10"] = "=Source!A1-1"
    summary["C11"] = "=Source!A1/2"
    summary["C12"] = "=Source!A1<0"
    summary["C13"] = "=Source!A1>0"
    summary["C14"] = "=Source!A1<=0"
    summary["C15"] = "=Source!A1>=0"
    summary["C16"] = "=Source!A1+Source!$A$1"

    path = tmp_path / "pure-operator-error-propagation.xlsx"
    workbook.save(path)
    workbook.close()

    findings = _wl034_by_sheet(path, "Summary")
    assert set(findings) == {f"C{row}" for row in range(1, 17)}
    for finding in findings.values():
        assert float(finding.confidence) == 1.0
        assert finding.evidence.details["proof"] == "propagated_formula_error"
        assert finding.evidence.details["function"] == "PURE_OPERATOR_EXPRESSION"
        assert finding.evidence.details["error"] is None
        assert finding.evidence.details["error_code_proven"] is False
        assert finding.patch_ids == []


def test_explicit_broken_reference_seeds_downstream_without_duplicate_wl034(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    source = workbook.active
    assert source is not None
    source.title = "Source"
    source["A1"] = "=#REF!"

    summary = workbook.create_sheet("Summary")
    summary["B1"] = "=Source!A1"
    summary["B2"] = "=SUM(Source!A1:A1)"
    summary["B3"] = "=B1"
    summary["B4"] = "=IFERROR(Source!A1,0)"

    path = tmp_path / "explicit-ref-propagation.xlsx"
    workbook.save(path)
    workbook.close()

    scan = scan_workbook(path)
    wl001 = {
        (finding.sheet, finding.location)
        for finding in scan.findings
        if finding.rule_id == "WL001_BROKEN_REFERENCE"
    }
    assert wl001 == {("Source", "A1")}
    assert _wl034_by_sheet(path, "Source") == {}
    findings = _wl034_by_sheet(path, "Summary")
    assert set(findings) == {"B1", "B2", "B3"}
    for finding in findings.values():
        assert finding.evidence.details["error"] == "#REF!"
        assert finding.evidence.details["error_code_proven"] is True
        assert finding.evidence.details["proof"] == "propagated_formula_error"
    assert "B4" not in findings


def test_error_propagation_skips_formula_ir_parse_failures(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = "=1/0"
    worksheet["B1"] = "=A1"

    original = error_propagation.parse_formula_ir

    def reject_one_formula(formula: str, sheet: str, coordinate: str):  # type: ignore[no-untyped-def]
        if coordinate == "B1":
            raise ValueError("malformed formula")
        return original(formula, sheet, coordinate)

    monkeypatch.setattr(error_propagation, "parse_formula_ir", reject_one_formula)

    proofs = find_static_formula_errors(
        workbook,
        lambda formula: ("#DIV/0!", "constant_division_by_zero") if formula == "=1/0" else None,
    )

    assert set(proofs) == {(worksheet.title, "A1")}
